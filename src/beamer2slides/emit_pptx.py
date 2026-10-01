"""The .pptx that carries a deck's pictures, backgrounds, layouts, template shapes and tables into
Slides; shape requests; a batch of requests.
"""

from __future__ import annotations

import io
from collections.abc import Callable, Mapping, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING

from .emit_metrics import SLIDE_W, xml_text
from .emit_model import (
    JsonMap, PptxTable, SetShape, Template, TemplateKey, set_shape, shape_measures_of, shape_of, template_of,
)
from .ir import Arrow
from .ir_types import Box, MarkedShape, ShapeElement
from .gapi import HttpError, message_of
from .google_types import Outline, SlidesRequest, SlidesService, shape_type
from .gslides import EMU_PER_PT, emu, execute, pt, rgb_color

if TYPE_CHECKING:
    from pptx.oxml.xmlchemy import BaseOxmlElement
    from pptx.presentation import Presentation
    from pptx.slide import Slide, SlideLayout
    from pptx.parts.slide import BaseSlidePart


# Native drop shadows, calibrated against beamer's block shadow (tools/calibrate_shadow.py):
# distance and blur per point of the shadow's width in the PDF, at 45°.
SHADOW_DISTANCE = 0.75
SHADOW_BLUR = 1.0
SHADOW_ALPHA = 0.5
TEMPLATE_PRESETS = {"ROUND_RECTANGLE": "roundRect", "ROUND_2_SAME_RECTANGLE": "round2SameRect", "RECTANGLE": "rect"}
TEMPLATE_LAYOUTS = {"TITLE": 0, "TITLE_ONLY": 5, "BLANK": 6}  # python-pptx default template layout indexes
# what `_add_template_shapes` can put on a source slide
TEMPLATE_KINDS = ("ROUND_RECTANGLE", "ROUND_2_SAME_RECTANGLE", "RECTANGLE", "ELLIPSE", "DIAMOND", "TRIANGLE",
                  "CAN", "CLOUD", "HEXAGON", "OCTAGON", "PENTAGON", "HEPTAGON", "DECAGON", "DODECAGON",
                  "FLOW_CHART_PUNCHED_TAPE")
# presets whose first adjustment a diagram node's template key carries (emit_diagrams.ADJUSTED)
ADJUSTED_KINDS = ("CAN", "HEXAGON", "OCTAGON")


def template_key(el: JsonMap, scale: float) -> TemplateKey | None:
    """`template_key_of` an element dict (sync's base elements, diagrams, element_template_keys)."""
    shape = el["shape"] if el["kind"] == "shape" else None
    if not isinstance(shape, str) or shape not in TEMPLATE_KINDS:
        return None
    bbox, radius, shadow = shape_measures_of(el)
    return template_key_of(shape, bbox, radius, shadow, scale)


def template_key_of(shape: str, bbox: Box, radius: float, shadow: float | None, scale: float) -> TemplateKey | None:
    """Shapes the API can't make exactly: rounded corners of a given radius (the API only
    creates the default rounding) and drop shadows (read-only in the API, `shadow` its size in
    PDF pt). They are duplicated from template shapes that come with the imported .pptx. A kind no
    template is made of has none: createShape refuses it, and the refused element becomes a
    picture (`fallback_pictures`) instead of the whole .pptx failing."""
    if shape not in TEMPLATE_KINDS or (shape == "RECTANGLE" and shadow is None):
        return None
    x0, y0, x1, y1 = bbox
    adj = 0.0
    if shape != "RECTANGLE":
        adj = min(0.5, round(radius / max(min(x1 - x0, y1 - y0), 0.01), 2))
    return shape, adj, None if shadow is None else round(2 * shadow * scale) / 2


ARC = "ARC"
"""An arc template's kind (a diagram's curve, `emit_diagrams.arc_template_key_of`) begins so;
`arc_kind` adds its heads: (kind, |sweep| in degrees, None)."""
ARC_HEADS: dict[Arrow, str] = {"OPEN_ARROW": "arrow", "FILL_ARROW": "triangle", "STEALTH_ARROW": "stealth"}
"""Slides' arrow heads as a .pptx line end's `type`."""


def arc_kind(arrow_from: Arrow | None, arrow_to: Arrow | None) -> str:
    """An arc template's kind: ARC, '<' and the head at its start, '>' and the one at its end
    ('ARC>FILL_ARROW'). The heads come with the template: the API sets none on a shape."""
    return ARC + ("" if arrow_from is None else f"<{arrow_from}") + ("" if arrow_to is None else f">{arrow_to}")


def arc_heads(kind: str) -> tuple[Arrow | None, Arrow | None] | None:
    """The heads (start, end) an arc template's kind names; None for a kind that is no arc's."""
    if not kind.startswith(ARC):
        return None
    rest, to_end, end = kind[len(ARC):].partition(">")
    if rest and not rest.startswith("<"):
        return None

    def head(name: str) -> Arrow | None:
        return next((a for a in ARC_HEADS if a == name), None)
    first, last = head(rest[1:]), head(end)
    if (rest and first is None) or (to_end and last is None):
        return None
    return first, last


def _add_arc(slide: Slide, i: int, sweep: float, heads: tuple[Arrow | None, Arrow | None]) -> None:
    """An arc template: the preset `arc` in a square, from 0° (its right) clockwise through
    `sweep` degrees, unfilled, with its heads (a .pptx brings them in; tools/probe_curves.py)."""
    from lxml import etree

    tree = slide.shapes._spTree
    shape_id = max([int(e.get("id")) for e in tree.iter(f"{{{NS_P}}}cNvPr")] + [1]) + 1
    side, x, y = round(100 * EMU_PER_PT), round((10 + i % 10 * 20) * EMU_PER_PT), round((10 + i // 10 * 20) * EMU_PER_PT)
    start, end = heads
    ends = ("" if start is None else f'<a:headEnd type="{ARC_HEADS[start]}"/>') + \
        ("" if end is None else f'<a:tailEnd type="{ARC_HEADS[end]}"/>')
    tree.append(etree.fromstring(
        f'<p:sp xmlns:p="{NS_P}" xmlns:a="{NS_A}"><p:nvSpPr><p:cNvPr id="{shape_id}" name="Arc {shape_id}"/>'
        f'<p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{side}" cy="{side}"/>'
        f'</a:xfrm><a:prstGeom prst="arc"><a:avLst><a:gd name="adj1" fmla="val 0"/>'
        f'<a:gd name="adj2" fmla="val {round(sweep * 60000)}"/></a:avLst></a:prstGeom><a:noFill/>'
        f'<a:ln w="{round(EMU_PER_PT)}"><a:solidFill><a:srgbClr val="000000"/></a:solidFill>{ends}</a:ln>'
        f'<a:effectLst/></p:spPr><p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:endParaRPr/></a:p></p:txBody></p:sp>'))


NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


PageFill = Mapping[str, object]
"""A page background: {"color": "#rrggbb"} or {"picture": Path} (stretched)."""


def _set_background(part: BaseSlidePart, c_sld: BaseOxmlElement | None, fill: PageFill) -> None:
    """A page background in the .pptx: {"color": "#rrggbb"} or {"picture": Path} (stretched)."""
    from lxml import etree

    if c_sld is None:
        raise ValueError("a page without its cSld")
    if "color" in fill:
        color = fill["color"]
        if not isinstance(color, str):
            raise TypeError(f"a background colour is a hex string, not {color!r}")
        inner = f'<a:solidFill><a:srgbClr val="{color.lstrip("#").upper()}"/></a:solidFill>'
    else:
        _, rid = part.get_or_add_image_part(str(fill["picture"]))  # identical files are stored once
        inner = (f'<a:blipFill dpi="0" rotWithShape="1"><a:blip r:embed="{rid}"/><a:srcRect/>'
                 f'<a:stretch><a:fillRect/></a:stretch></a:blipFill>')
    old = c_sld.find(f"{{{NS_P}}}bg")
    if old is not None:
        c_sld.remove(old)
    c_sld.insert(0, etree.fromstring(
        f'<p:bg xmlns:p="{NS_P}" xmlns:a="{NS_A}" xmlns:r="{NS_R}"><p:bgPr>{inner}<a:effectLst/></p:bgPr></p:bg>'))


def _add_template_shapes(slide: Slide, keys: Sequence[TemplateKey]) -> None:
    from lxml import etree
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.util import Pt

    kinds = dict(zip(TEMPLATE_KINDS, (MSO_SHAPE.ROUNDED_RECTANGLE, MSO_SHAPE.ROUND_2_SAME_RECTANGLE, MSO_SHAPE.RECTANGLE,
                                      MSO_SHAPE.OVAL, MSO_SHAPE.DIAMOND, MSO_SHAPE.ISOSCELES_TRIANGLE,
                                      MSO_SHAPE.CAN, MSO_SHAPE.CLOUD, MSO_SHAPE.HEXAGON, MSO_SHAPE.OCTAGON,
                                      MSO_SHAPE.REGULAR_PENTAGON, MSO_SHAPE.HEPTAGON, MSO_SHAPE.DECAGON,
                                      MSO_SHAPE.DODECAGON, MSO_SHAPE.FLOWCHART_PUNCHED_TAPE)))
    a = NS_A
    for i, (kind, adj, shadow) in enumerate(keys):
        if kind == "BENT_CONNECTOR":
            if adj is None:
                raise ValueError("a bent connector's template key carries where its elbow is")
            line = slide.shapes.add_connector(MSO_CONNECTOR.ELBOW, Pt(10), Pt(10), Pt(110), Pt(110))
            geometry = line._element.spPr.find(f"{{{a}}}prstGeom")
            geometry.set("prst", "bentConnector3")
            for old in geometry.findall(f"{{{a}}}avLst"):
                geometry.remove(old)
            geometry.append(etree.fromstring(
                f'<a:avLst xmlns:a="{a}"><a:gd name="adj1" fmla="val {round(adj * 100000)}"/></a:avLst>'))
            line._element.spPr.append(etree.fromstring(f'<a:effectLst xmlns:a="{a}"/>'))
            continue
        heads = arc_heads(kind)
        if heads is not None:
            if adj is None or not 0 < adj < 360:
                raise ValueError(f"an arc's template key carries its sweep, not {adj!r}")
            _add_arc(slide, i, adj, heads)
            continue
        shape = slide.shapes.add_shape(kinds[kind], Pt(10 + i % 10 * 20), Pt(10 + i // 10 * 20), Pt(100), Pt(100))
        if adj is not None and kind in ("ROUND_RECTANGLE", "ROUND_2_SAME_RECTANGLE"):
            shape.adjustments[0] = adj
            if kind == "ROUND_2_SAME_RECTANGLE":
                shape.adjustments[1] = 0.0
        elif adj is not None and kind in ADJUSTED_KINDS:
            # a can's ellipse height, a hexagon's or an octagon's corner cut, as a share of the
            # shorter side (python-pptx writes the hexagon's vf at its default beside it)
            shape.adjustments[0] = adj
        shape.fill.solid()
        shape.line.fill.background()
        # No text padding (the API can't set it): a diagram label fits a node as tight as TikZ's.
        body_pr = shape.text_frame._txBody.find(f"{{{a}}}bodyPr")
        for side in ("lIns", "tIns", "rIns", "bIns"):
            body_pr.set(side, "0")
        body_pr.set("anchor", "ctr")
        # python-pptx shapes refer to the theme's effect style, which has a shadow: always
        # give an explicit (possibly empty) effect list.
        effects = (f'<a:outerShdw blurRad="{round(SHADOW_BLUR * shadow * EMU_PER_PT)}" '
                   f'dist="{round(SHADOW_DISTANCE * shadow * EMU_PER_PT)}" dir="2700000" algn="tl" rotWithShape="0">'
                   f'<a:srgbClr val="000000"><a:alpha val="{round(SHADOW_ALPHA * 100000)}"/></a:srgbClr>'
                   f'</a:outerShdw>') if shadow else ""
        shape.element.spPr.append(etree.fromstring(f'<a:effectLst xmlns:a="{a}">{effects}</a:effectLst>'))


NO_TABLE_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # PowerPoint's "No Style, No Grid"


def _add_table(slide: Slide, table: PptxTable) -> None:
    """An empty table (pptx_table) on a source slide. Its cell margins are what the API can't
    set: a table made by createTable has 7.2 pt above and below every line, one from a .pptx the
    file's (tools/probe_pptx_table_margins.py); duplicating the slide, inserting rows and columns
    and filling or styling cells through the API all keep them."""
    from lxml import etree
    from pptx.util import Emu

    def e(v: float) -> Emu:
        return Emu(round(v * EMU_PER_PT))

    rows, cols = len(table.heights), len(table.widths)
    frame = slide.shapes.add_table(rows, cols, e(table.x), e(table.y), e(sum(table.widths)), e(sum(table.heights)))
    pr = frame.table._tbl.tblPr
    if pr is None:
        raise ValueError("python-pptx made a table without <a:tblPr>")
    for flag in ("firstRow", "bandRow"):  # (python-pptx's default look: a header row and bands)
        pr.attrib.pop(flag, None)
    style = pr.find(f"{{{NS_A}}}tableStyleId")
    if style is None:
        style = etree.SubElement(pr, f"{{{NS_A}}}tableStyleId")
    style.text = NO_TABLE_STYLE
    for c, w in enumerate(table.widths):
        frame.table.columns[c].width = e(w)
    middle = set(table.middle)
    for r, h in enumerate(table.heights):
        frame.table.rows[r].height = e(h)
        left, top, right, bottom = table.margins[r]
        for c in range(cols):
            cell = frame.table.cell(r, c)
            cell.margin_left, cell.margin_top, cell.margin_right, cell.margin_bottom = \
                e(left), e(0.0 if (r, c) in middle else top), e(right), e(bottom)


VARIANT = "_V"      # layout name suffix: a copy of the layout with another theme decoration (plan_theme)
THEME_VARIANTS = 3


def _clone_layout(prs: Presentation, layout: SlideLayout, name: str) -> SlideLayout:
    """A copy of a layout (placeholders only, no pictures) added to the master, shown as `name`."""
    from copy import deepcopy

    from lxml import etree
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.opc.packuri import PackURI
    from pptx.parts.slide import SlideLayoutPart

    master = prs.slide_master
    taken = {str(p.partname) for p in prs.part.package.iter_parts()}
    k = next(k for k in range(1, 1000) if f"/ppt/slideLayouts/slideLayout{k}.xml" not in taken)
    element = deepcopy(layout.element)
    element.cSld.set("name", name)
    element.set("type", "cust")  # (its own layout name in Slides, never taken for the original's)
    part = SlideLayoutPart(PackURI(f"/ppt/slideLayouts/slideLayout{k}.xml"), layout.part.content_type,
                           layout.part.package, element)
    part.relate_to(master.part, RT.SLIDE_MASTER)
    ids = master.element.get_or_add_sldLayoutIdLst()
    entry = etree.SubElement(ids, f"{{{NS_P}}}sldLayoutId")
    entry.set("id", str(max(int(e.get("id")) for e in ids if e.get("id")) + 1))
    entry.set(f"{{{NS_R}}}id", master.part.relate_to(part, RT.SLIDE_LAYOUT))
    return part.slide_layout


def _add_decoration(layout: SlideLayout, picture: Path, width: int, height: int) -> None:
    """The theme decoration as a full-page picture at the bottom of a layout: above the slide
    background, below everything on the slide."""
    from lxml import etree

    _, rid = layout.part.get_or_add_image_part(str(picture))
    tree = layout.shapes._spTree
    shape_id = max([int(e.get("id")) for e in tree.iter(f"{{{NS_P}}}cNvPr")] + [1]) + 1
    tree.insert(2, etree.fromstring(
        f'<p:pic xmlns:p="{NS_P}" xmlns:a="{NS_A}" xmlns:r="{NS_R}"><p:nvPicPr>'
        f'<p:cNvPr id="{shape_id}" name="Theme" descr="Theme decoration"/><p:cNvPicPr><a:picLocks noGrp="1"/></p:cNvPicPr>'
        f'<p:nvPr userDrawn="1"/></p:nvPicPr><p:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch>'
        f'</p:blipFill><p:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{width}" cy="{height}"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>'))


@dataclass(frozen=True, kw_only=True)
class PptxPicture:
    """A picture on a source slide: its file, its box (slide pt) and its alt text and title (both
    "" for none)."""
    file: str
    bbox: Box
    alt: str
    title: str


@dataclass(frozen=True, kw_only=True)
class PptxPage:
    """A source slide of the .pptx (`build_pptx`, which reads it from its page dict: `pptx_page`)."""
    layout: str
    fill: PageFill | None
    pictures: tuple[PptxPicture, ...]
    tables: tuple[PptxTable, ...]
    templates: bool


def _number(v: object, where: str) -> float:
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        raise TypeError(f"{where}: a number, not {v!r}")
    return v


def _str(v: object, where: str) -> str:
    if not isinstance(v, str):
        raise TypeError(f"{where}: a string, not {v!r}")
    return v


def pptx_page(page: Mapping[str, object]) -> PptxPage:
    """A page dict of `build_pptx` ({"layout", "fill" (None: inherit), "pictures": [{"file",
    "bbox", "alt", "title"}], "tables": [PptxTable] (optional), "templates"}) as its record."""
    layout, fill, pictures = page["layout"], page["fill"], page["pictures"]
    if not isinstance(layout, str):
        raise TypeError(f"a page's layout is a name, not {layout!r}")
    if fill is not None and not isinstance(fill, Mapping):
        raise TypeError(f"a page's fill is a mapping, not {fill!r}")
    if not isinstance(pictures, (list, tuple)):
        raise TypeError(f"a page's pictures are a list, not {pictures!r}")
    pics: list[PptxPicture] = []
    for pic in pictures:
        if not isinstance(pic, Mapping):
            raise TypeError(f"a picture is a mapping, not {pic!r}")
        box = pic["bbox"]
        if not isinstance(box, (list, tuple)) or len(box) != 4:
            raise TypeError(f"a picture's bbox is four numbers, not {box!r}")
        x0, y0, x1, y1 = (_number(v, "a picture's bbox") for v in box)
        alt = pic.get("alt")  # (none, or "": no alt text and no title)
        pics.append(PptxPicture(file=str(pic["file"]), bbox=(x0, y0, x1, y1),
                                alt=_str(alt, "a picture's alt text") if alt else "",
                                title=_str(pic["title"], "a picture's title") if alt else ""))
    tables = page.get("tables", [])
    if not isinstance(tables, (list, tuple)) or not all(isinstance(t, PptxTable) for t in tables):
        raise TypeError(f"a page's tables are PptxTables, not {tables!r}")
    return PptxPage(layout=layout, fill=fill, pictures=tuple(pics),
                    tables=tuple(t for t in tables if isinstance(t, PptxTable)), templates=bool(page["templates"]))


def build_pptx(page_w: float, page_h: float, keys: Sequence[TemplateKey], pages: Sequence[Mapping[str, object]],
               master_fill: PageFill, decorations: Mapping[str, Path | None] | None) -> io.BytesIO:
    """The deck's starting point, imported through Drive. It carries everything the Slides API
    could only insert from a public URL, so no picture ever leaves the user's Drive:

    - the PDF's page size (presentations.create ignores pageSize);
    - the master background (`master_fill`), inherited by the layouts and most slides;
    - the theme decoration (`decorations`, see plan_theme) at the bottom of the layouts; a page
      layout named with the VARIANT suffix is a copy of its layout with that variant's decoration;
    - one source slide per deck slide (`pages`: {"layout", "fill" (None: inherit),
      "pictures": [{"file", "bbox" (slide pt), "alt", "title"}], "tables": [PptxTable],
      "templates" (bool)}), holding its pictures, its tables (empty, with the cell margins the API
      cannot set) and, if it needs any, the template shapes (shadows, exact corner radii).

    emit copies each source slide under our own object IDs and then deletes it."""
    from pptx import Presentation
    from pptx.parts.slide import BaseSlidePart
    from pptx.util import Emu

    parsed = [pptx_page(p) for p in pages]
    prs = Presentation()
    height = SLIDE_W * page_h / page_w
    default_h = prs.slide_height
    if default_h is None:
        raise ValueError("python-pptx's default template has no slide height")
    ratio = height / (default_h / EMU_PER_PT)
    slide_w, slide_h = Emu(round(SLIDE_W * EMU_PER_PT)), Emu(round(height * EMU_PER_PT))
    prs.slide_width, prs.slide_height = slide_w, slide_h
    if abs(ratio - 1) > 1e-3:  # the default template's placeholders are laid out for 4:3
        for page in [prs.slide_master, *prs.slide_layouts]:
            for shape in page.placeholders:
                # A layout placeholder without its own position inherits the master's, rescaled
                # already (setting its top and height would scale it twice and write x and width 0).
                top, shape_h = shape.top, shape.height
                if shape._element.spPr.find(f"{{{NS_A}}}xfrm") is not None and top is not None \
                        and shape_h is not None:
                    shape.top, shape.height = Emu(round(top * ratio)), Emu(round(shape_h * ratio))
    master = prs.slide_master
    master_part = master.part
    if not isinstance(master_part, BaseSlidePart):
        raise TypeError(f"python-pptx's slide master is a {type(master_part).__name__}, not a slide part")
    _set_background(master_part, master.element.find(f"{{{NS_P}}}cSld"), master_fill)
    decorations = decorations or {}
    layouts = {name: prs.slide_layouts[i] for name, i in TEMPLATE_LAYOUTS.items()}
    originals = list(prs.slide_layouts)
    for name in dict.fromkeys(p.layout for p in parsed if VARIANT in p.layout):
        kind, n = name.rsplit(VARIANT, 1)
        picture = decorations.get(f"{'TITLE' if kind == 'TITLE' else '*'}{VARIANT}{n}")
        layouts[name] = _clone_layout(prs, layouts[kind], f"{layouts[kind].name} ({f'theme {int(n) + 1}' if picture else 'no theme'})")
        if picture:
            _add_decoration(layouts[name], picture, slide_w, slide_h)
    for i, layout in enumerate(originals):
        picture = decorations.get("TITLE" if i == TEMPLATE_LAYOUTS["TITLE"] else "*")
        if picture:
            _add_decoration(layout, picture, slide_w, slide_h)
    for page in parsed:
        slide = prs.slides.add_slide(layouts[page.layout])
        if page.fill:
            _set_background(slide.part, slide.element.find(f"{{{NS_P}}}cSld"), page.fill)
        for pic in page.pictures:
            x0, y0, x1, y1 = pic.bbox
            shape = slide.shapes.add_picture(pic.file, Emu(round(x0 * EMU_PER_PT)), Emu(round(y0 * EMU_PER_PT)),
                                             Emu(round((x1 - x0) * EMU_PER_PT)), Emu(round((y1 - y0) * EMU_PER_PT)))
            if pic.alt:  # (text from the PDF: raw Type 3 T1 codes are C0 controls, xml_text)
                nv = shape._element.find(f"{{{NS_P}}}nvPicPr")
                c_nv = nv.find(f"{{{NS_P}}}cNvPr") if nv is not None else None
                if c_nv is None:
                    raise ValueError("python-pptx made a picture without its cNvPr")
                c_nv.set("descr", xml_text(pic.alt))
                c_nv.set("title", xml_text(pic.title))
        for table in page.tables:
            _add_table(slide, table)
        if page.templates:
            _add_template_shapes(slide, keys)
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf


def shape_requests(el: JsonMap, slide_id: str, object_id: str, scale: float,
                   template: JsonMap | None) -> list[SlidesRequest]:
    """`shape_requests_of` a shape dict (a diagram's node, the tests' panels), with a template
    dict ({"id", "w", "h"})."""
    return shape_requests_of(shape_of(el), slide_id, object_id, scale, template_of(template) if template else None)


def shape_element_requests(el: ShapeElement | MarkedShape, slide_id: str, object_id: str, scale: float,
                           template_for: Callable[[TemplateKey], Template]) -> list[SlidesRequest]:
    """The requests of a parsed shape (a panel, rule or marked shape): a duplicate of the slide's
    template shape its key names (`template_for`), where it has one."""
    shape = set_shape(el)
    key = template_key_of(shape.shape, shape.bbox, shape.radius, shape.shadow, scale)
    return shape_requests_of(shape, slide_id, object_id, scale, None if key is None else template_for(key))


def shape_requests_of(el: SetShape, slide_id: str, object_id: str, scale: float,
                      template: Template | None) -> list[SlidesRequest]:
    """A filled shape, outlined only with the frame it carries (`outline`). With a template (a
    template shape on this slide and its size in pt) the shape is a duplicate of it, else a new shape."""
    x0, y0, x1, y1 = (v * scale for v in el.bbox)
    # ROUND_2_SAME_RECTANGLE rounds the top corners; for bottom corners flip both axes
    # (a 180° rotation), which moves the origin to the opposite corner.
    flip = -1 if el.flip else 1
    tx, ty = (x1, y1) if el.flip else (x0, y0)
    reqs: list[SlidesRequest]
    if template is not None:
        reqs = [
            {"duplicateObject": {"objectId": template.id, "objectIds": {template.id: object_id}}},
            {"updatePageElementTransform": {"objectId": object_id, "applyMode": "ABSOLUTE", "transform": {
                "scaleX": flip * (x1 - x0) / template.w, "scaleY": flip * (y1 - y0) / template.h, "unit": "EMU",
                "translateX": round(tx * EMU_PER_PT), "translateY": round(ty * EMU_PER_PT)}}},
            {"updatePageElementsZOrder": {"pageElementObjectIds": [object_id], "operation": "BRING_TO_FRONT"}},
        ]
    else:
        # (a producer's "custom" or "line" is no Slides shape: createShape would be refused with
        # its whole batch, so planning it fails here and `DeckPlan.contain` makes it a picture)
        reqs = [{"createShape": {
            "objectId": object_id, "shapeType": shape_type(el.shape, f"shape {object_id}"),
            "elementProperties": {
                "pageObjectId": slide_id,
                "size": {"width": emu(x1 - x0), "height": emu(y1 - y0)},
                "transform": {"scaleX": flip, "scaleY": flip, "unit": "EMU",
                              "translateX": round(tx * EMU_PER_PT), "translateY": round(ty * EMU_PER_PT)},
            },
        }}]
    # A framed panel (classify.frame_of, framed_panels: \fcolorbox, tcolorbox, a listing's
    # frame=single) carries its frame as the outline, on the frame's centre line; the rules
    # themselves left the background with the panel.
    frame = el.outline
    outline: Outline = {"outlineFill": {"solidFill": {"color": rgb_color(frame.color)}},
                        "weight": pt(round(max(0.25, frame.width * scale), 2)), "propertyState": "RENDERED"} \
        if frame is not None else {"propertyState": "NOT_RENDERED"}
    outline_fields = "outline.outlineFill.solidFill.color,outline.weight,outline.propertyState" if frame is not None \
        else "outline.propertyState"
    return [*reqs,
        {"updateShapeProperties": {
            "objectId": object_id,
            "shapeProperties": {"shapeBackgroundFill": {"solidFill": {"color": rgb_color(el.fill),
                                                                      "alpha": el.opacity}},
                                "outline": outline},
            "fields": "shapeBackgroundFill.solidFill.color,shapeBackgroundFill.solidFill.alpha," + outline_fields,
        }},
    ]


def api_error(e: HttpError) -> str:
    return message_of(e)


def batch(slides: SlidesService, pid: str, reqs: Sequence[SlidesRequest]) -> None:
    execute(slides.presentations().batchUpdate(presentationId=pid, body={"requests": reqs}))
