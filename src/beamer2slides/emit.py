"""Stage 4: build the Google Slides deck from deck.json and the background images."""

import hashlib
import io
import json
import math
import re
import time
import unicodedata
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from importlib import resources
from pathlib import Path

import numpy as np

from . import bidi
from .classify import HOLE_PAD
from .fonts import font_info, google_font
from .gapi import HttpError, media_upload, message_of
from .google_auth import credentials_for_threads, drive_service, fetcher_for_threads, shared_service, slides_service
from .gslides import EMU_PER_PT, emu, execute, per_thread, pt

# Found through the package, never through the checkout: an installed wheel, a zip import and
# a build that stages sources elsewhere all keep the data beside the module, not beside __file__.
CALIBRATION_DIR = resources.files("beamer2slides") / "calibration"
CALIBRATION = CALIBRATION_DIR / "fonts.json"
SLIDE_W = 720.0
BATCH_MAX_REQUESTS = 400  # slides are sent together until a batch reaches this size
# A round trip to Google costs about a second whatever it carries, so the wall clock of a
# conversion is round trips and not work (measured: tools/probe_batch_parallelism.py). Several
# batches may be in flight on one presentation at once - Google takes them and loses nothing -
# and four is where the curve flattens: 8 batches of 200 requests take 9.4 s one at a time,
# 6.1 s two at a time, 3.6 s four at a time and 3.1 s eight at a time.
CONTENT_WORKERS = 4
PPTX_MIME ="application/vnd.openxmlformats-officedocument.presentationml.presentation"

# Slides text box model, measured by tools/calibrate.py and the spacing probes
# (docs/calibration.md).
BASELINE_A = 6.48    # box top -> first baseline = A + ASCENT_EM * size
ASCENT_EM = 0.968
LINE_EM = 1.2        # baseline pitch at lineSpacing 100 (before pixel snapping)
PX_PT = 0.75         # Slides snaps line pitches to whole CSS pixels
DESCENT_EM = LINE_EM - ASCENT_EM
PAD_X = 6.7          # box edge -> text start
BULLET_GAP = 1.9     # bullet glyph's right edge sits this far before indentFirstLine
PPTX_TITLE_DY = 3.9  # title placeholders of pptx-imported decks have a smaller top inset
MIDDLE_BASELINE_EM = ASCENT_EM - LINE_EM / 2  # contentAlignment MIDDLE: baseline below the box middle (tools/probe_middle.py: 0.362)
SOFT_BREAK = chr(11)  # vertical tab: a line break inside a paragraph

FONT_FOR_FAMILY = {"sans": "Lato", "serif": "PT Serif", "mono": "Roboto Mono"}
# Advance widths in Slides text (Lato and its fallback fonts), measured by tools/probe_symbols.py.
SYMBOL_ADVANCE_EM = {
    "=": 0.58, "+": 0.58, "−": 0.58, "<": 0.58, ">": 0.58, "≤": 0.58, "≥": 0.58, "×": 0.58, "·": 0.271,
    "/": 0.313, "∑": 0.682, "∏": 0.682, "∫": 0.397, "∈": 0.984, "∉": 0.492, "⊂": 0.981, "⊆": 0.981,
    "∪": 0.981, "∩": 0.717, "→": 0.998, "←": 0.998, "⇒": 0.981, "⇔": 0.981, "≈": 0.577, "≠": 0.577,
    "±": 0.577, "∞": 0.682, "ℝ": 0.633, "ℕ": 0.65, "ℤ": 0.534, "ℚ": 0.65, "ℂ": 0.65, "α": 0.573,
    "β": 0.57, "γ": 0.496, "δ": 0.552, "ε": 0.443, "θ": 0.552, "λ": 0.496, "μ": 0.573, "π": 0.615,
    "σ": 0.612, "φ": 0.643, "ω": 0.777, "Δ": 0.664, "Σ": 0.615, "Ω": 0.742, "∂": 0.577, "∇": 0.981,
    "′": 0.186, "∀": 0.981, "∃": 0.981, "∧": 0.981, "∨": 0.981, "⊥": 0.981, "∥": 0.981, "∘": 0.489,
    "…": 0.724, " ": 0.19,
}
MATH_SPACE_EM = 0.278  # TeX's \thickmuskip (5 mu) around relations
CMTT_ADVANCE_EM, ROBOTO_MONO_ADVANCE_EM = 0.525, 0.6

BULLET_PRESETS = {
    "number": "NUMBERED_DIGIT_ALPHA_ROMAN",
    "number_parens": "NUMBERED_DIGIT_ALPHA_ROMAN_PARENS",
}
# A preset's three glyphs are those of nesting levels 0, 1, 2; deeper levels repeat ● ○ ■
# whatever the preset. A bullet shape is a preset plus the level showing it (bullet_level).
# Ink height and the gap between ink and indentFirstLine are per em of the bullet's size
# (tools/probe_bullets.py).
BULLET_SHAPES = {  # shape: (preset, level, ink height em, gap em)
    "disc": ("BULLET_DISC_CIRCLE_SQUARE", 0, 0.413, 0.08),
    "circle": ("BULLET_DISC_CIRCLE_SQUARE", 1, 0.43, 0.08),
    "square": ("BULLET_DISC_CIRCLE_SQUARE", 2, 0.45, 0.07),
    "open_square": ("BULLET_CHECKBOX", 0, 0.69, 0.155),  # ❏
    "triangle": ("BULLET_ARROW3D_CIRCLE_SQUARE", 0, 0.525, 0.07),  # ➢: presets have no ▶
    "star": ("BULLET_STAR_CIRCLE_SQUARE", 0, 0.81, 0.07),
    "diamond": ("BULLET_DIAMOND_CIRCLE_SQUARE", 0, 0.81, 0.08),
    "open_diamond": ("BULLET_DIAMONDX_HOLLOWDIAMOND_SQUARE", 1, 0.87, 0.06),
}
GLYPH_SHAPES = {**dict.fromkeys("▶►▸‣", "triangle"), **dict.fromkeys("•●", "disc"), **dict.fromkeys("◦○", "circle"),
                **dict.fromkeys("■▪", "square"), "□": "open_square", **dict.fromkeys("★⋆", "star"),
                **dict.fromkeys("◆♦", "diamond"), **dict.fromkeys("◇⋄", "open_diamond")}


# Width per em of Computer Modern's optical sizes relative to the 10 pt cut, from the glyph
# advances of the Type 1 fonts (cmss8.pfb ... cmss17.pfb) over a sample sentence. EC and
# Latin Modern share these metrics.
DESIGN_WIDTH = {
    "sans": {8: 1.0623, 9: 1.0273, 10: 1.0, 12: 0.9753, 17: 0.9377},
    "serif": {5: 1.3758, 6: 1.2291, 7: 1.1424, 8: 1.0629, 9: 1.0277, 10: 1.0, 12: 0.9786, 17: 0.9136},
    "mono": {8: 1.0114, 9: 1.0, 10: 1.0, 12: 0.979},
}
# A small optical cut is wider per em than the 10 pt one, not taller (cmr5's x-height is cmr10's
# per em), so a substitute sized to its width grows its letters as much: a 5 pt LMRoman5 footline
# (1.38x) came out 1.4 times as tall as the PDF's and filled its bar. The width is matched up to
# the small-caps compromise (SMALL_CAPS_WIDTH) and no further: at 5 pt the line comes out ~18%
# narrower than the PDF's and ~13% taller. 8 and 9 pt cuts (1.03-1.06) are matched in full.
OPTICAL_WIDTH_MAX = 1.13


def optical_width(family: str, design: float) -> float:
    """How much wider per em than its 10 pt cut the size factor takes a run's optical size to be."""
    return min(OPTICAL_WIDTH_MAX, design_width(DESIGN_WIDTH.get(family, DESIGN_WIDTH["sans"]), design))


def u16(text: str) -> int:
    """Length in UTF-16 code units, which is what every Slides text index counts: an astral
    character (𝔼 U+1D53C from amssymb's \\mathbb, 𝛽 U+1D6FD) is two. Counting code points put
    every style range after one a unit early and split the next one's surrogate pair (two tofu)."""
    return len(text.encode("utf-16-le", "surrogatepass")) // 2


# XML 1.0 refuses C0 controls other than tab, newline and return, lone surrogates and U+FFFE/F:
# a Type 3 T1 font's raw glyph codes (quotes, dashes, ligatures at 0x10-0x1d) reach alt texts.
XML_REFUSED = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


def xml_text(text: str) -> str:
    """A string lxml accepts as an attribute value (build_pptx): refused characters dropped."""
    return XML_REFUSED.sub("", text)


# Small caps (tools/probe_text_fit_fonts.py): Slides draws a smallCaps lowercase letter as its
# capital at 0.70 of the size (PT Serif and Lato alike, advance and ink height), where TeX's
# CMCSC10 draws it at 0.755 - and CMCSC is an extended face, its capitals 8% wider than CMR's,
# while PT Serif's are 11% narrower. A CMCSC line came out 0.777 of the PDF's width; at the
# serif size these substitute small caps are 1.27-1.30 times too narrow over any sentence
# (probe advances + CMCSC10.afm), so the size makes up the width, as the size factors do for
# every other face. (Only for Computer Modern / EC / Latin Modern small caps, the ones measured.)
# The full 1.28 made the capitals ~30% taller than CMCSC's; 1.13 (about the square root) is the
# chosen compromise: the line comes out ~12% narrower than the PDF's and ~13% taller.
SMALL_CAPS_WIDTH = {"serif": 1.13}

# Computer Modern's own advance widths, kerning pairs and interword spaces (em; the AFM and TFM
# files, tools/cm_advances.py; EC and Latin Modern share them), per 10 pt face. With Slides'
# advances (ADVANCES) they predict how wide a run comes out in Slides against the PDF - within
# 0.5% of Google's renderer on the calibration rows and the text-fit torture lines.
CM_ADVANCES = json.loads((CALIBRATION_DIR / "cm_advances.json").read_text(encoding="utf-8"))["fonts"]
CM_FACE = {("sans", False, False): "cmss10", ("sans", True, False): "cmssbx10", ("sans", False, True): "cmssi10",
           ("sans", True, True): "cmssbx10",  # beamer's bold italic sans is CMSSBX
           ("serif", False, False): "cmr10", ("serif", True, False): "cmbx10", ("serif", False, True): "cmti10",
           ("serif", True, True): "cmbxti10"}
CM_LIGATURES = (("ffi", "ﬃ"), ("ffl", "ﬄ"), ("ff", "ﬀ"), ("fi", "ﬁ"), ("fl", "ﬂ"))
# TeX's space factor codes (plain/LaTeX \nonfrenchspacing): a space after a factor of 2000 or
# more gets the font's extra space; a capital (999) keeps "A." from ending a sentence.
SPACE_FACTOR = {".": 3000, "?": 3000, "!": 3000, ":": 2000, ";": 1500, ",": 1250}
KEEPS_SPACE_FACTOR = ")]'’”"
# Numbers: Computer Modern's digits are 0.5 em where Slides' Lato draws tabular digits at
# 0.577 em, 13% wider after the size correction, which is calibrated on sentences (and they
# stand at cap height, 7% taller than CM's). A run that is only a number - a table cell, a
# number column, a frame counter - is set at the size that gives it the PDF's width
# (FontMapper.shape_ratio); a number inside a sentence keeps the sentence's size.
DIGITS = "0123456789"
NUMBER_CHARS = DIGITS + ",.:%/-()+"
# Any other run: the calibrated factor makes an average sentence as wide as the PDF's, and a
# text whose letters are unlike a sentence's is off by what its own advances say (serif capitals
# 0.87-0.92, a line of w 1.10). Short texts vary the most - across the 84 test PDFs, a word or a
# label of under 10 letters spreads over 0.92-1.14 - so only a run of at least SHAPE_MIN_CHARS
# counted characters whose width is predicted off by more than SHAPE_TOL is resized: no run of
# 15 characters or more in the test decks is (the widest ordinary one is 6.3% off, "Somewhere
# University"), and text_fit calls a line off at 8%. Ordinary prose keeps the deck-wide size.
SHAPE_MIN_CHARS = 15
SHAPE_TOL = 0.07
# The sentences the size factors were calibrated on (tools/calibrate.py WIDTH_ROWS): a run is
# judged against them in its own face, so bold and italic keep their half correction.
SHAPE_REFERENCE = ("The quick brown fox jumps over the lazy dog",
                   "Another first level item that is long enough to wrap",
                   "Lorem ipsum dolor sit amet, consectetur adipiscing elit")
SHAPE_TITLE_REFERENCE = ("Itemize, nested and frame titles",)  # the title factor's row (CMSS12)
# Dots set apart from words: \dotfill and \dots leaders, \ldots (". . ." in the text layer) and
# "...". TeX sets them as fixed boxes or thin spaces, not as sentence ends; read as periods with
# interword and sentence spaces they made a leader line look 30-40% too wide for Slides, and the
# run was set 1.3-1.7 times too large. Their width is neither the PDF's letters' nor a sentence's,
# so they count on neither side, with the spaces around them (advance_widths).
LEADER = re.compile(r"[   ]*\.(?:[   ]*\.)+[   ]*")
STYLE_KEY = {(False, False): "regular", (True, False): "bold", (False, True): "italic", (True, True): "bold_italic"}
SLANTED = re.compile(r"CMB?X?SL\d|SFSL\d|SFBL\d|LMROMANSLANT")  # slanted roman: upright widths


def cm_face(run: dict) -> str | None:
    """The CM_ADVANCES face a Computer Modern (EC, Latin Modern) run is set in, or None."""
    if font_info(run["font"]).design_size is None:
        return None
    slanted = bool(SLANTED.match(re.sub(r"[^A-Z0-9]", "", run["font"].split("+", 1)[-1].upper())))
    return CM_FACE.get((run["family"], bool(run["bold"]), bool(run["italic"]) and not slanted))


def _advance(table: dict, ch: str) -> float | None:
    """A character's advance in a table, or its base letter's (é -> e: an accent adds no width)."""
    w = table.get(ch)
    if w is None and ch.isalpha():
        w = table.get(unicodedata.normalize("NFD", ch)[0])
    return w


def advance_widths(text: str, cm: dict, slides: dict) -> tuple[float, float, int, int]:
    """(Slides em, PDF em, characters counted, characters skipped) of a text set in a Computer
    Modern face (`cm`, a CM_ADVANCES entry) and in its Slides substitute (`slides`, an ADVANCES
    entry). The PDF side is what TeX set: ligatures, kerning pairs, interword space and the extra
    space after a sentence. A character either table lacks counts on neither side, and so do dot
    leaders and ellipses (LEADER): each of their dots is a skipped character."""
    for seq, lig in CM_LIGATURES:
        if lig in cm["advances"]:
            text = text.replace(seq, lig)
    text = LEADER.sub(lambda m: "\0" * m.group().count("."), text)
    s_em = p_em = 0.0
    counted = skipped = 0
    factor, prev = 1000, None
    for ch in text:
        if ch == "\0":  # a leader's dot
            skipped += 1
            factor, prev = 1000, None
            continue
        if ch in "  ":
            s_em += slides.get(" ", 0.0)
            p_em += cm["space"] + (cm["extra_space"] if factor >= 2000 else 0.0)
            factor, prev = 1000, None
            continue
        parts = next((seq for seq, lig in CM_LIGATURES if lig == ch), ch)  # Slides sets no ligature
        p = _advance(cm["advances"], ch)
        ws = [_advance(slides, c) for c in parts]
        if p is None or None in ws:
            skipped += len(parts)
            prev = None
            continue
        p_em += p + (cm["kerns"].get(prev + ch, 0.0) if prev else 0.0)
        s_em += sum(ws)
        counted += len(parts)
        prev = ch
        if ch.isupper():
            factor = 999
        elif ch in SPACE_FACTOR:
            factor = SPACE_FACTOR[ch] if factor >= 1000 else 1000
        elif ch not in KEEPS_SPACE_FACTOR:
            factor = 1000
    return s_em, p_em, counted, skipped


def design_width(table: dict[int, float], design: float) -> float:
    keys = sorted(table)
    if design <= keys[0]:
        return table[keys[0]]
    if design >= keys[-1]:
        return table[keys[-1]]
    hi = next(k for k in keys if k >= design)
    lo = max(k for k in keys if k <= design)
    if hi == lo:
        return table[lo]
    t = (design - lo) / (hi - lo)
    return table[lo] + t * (table[hi] - table[lo])


class FontMapper:
    """TeX font + size -> Slides font family + size with calibrated width correction."""

    def __init__(self):
        self.factors = {}  # family -> (running text factor, title factor)
        self.style = {}    # family -> {"bold": ratio, "italic": ratio} relative to running text
        self._reference = {}  # reference_ratio's answers
        for family, path in (("sans", CALIBRATION), ("serif", CALIBRATION_DIR / "fonts_serif.json")):
            cal = json.loads(path.read_text(encoding="utf-8"))["fonts"]
            ratios = cal[FONT_FOR_FAMILY[family]]["width_ratio"]
            self.factors[family] = (ratios["text_mean"], ratios["by_row"]["title"])
            self.style[family] = {k: ratios["relative_to_text"][k] for k in ("bold", "italic")}

    def text_style(self, run: dict, scale: float) -> tuple[dict, list[str]]:
        """Font part of a Slides TextStyle. Google fonts used by the PDF itself keep their
        family and weight (e.g. Fira Sans Light); TeX fonts get a calibrated substitute."""
        family, size = self(run, scale)
        google = google_font(run["font"])
        if google:
            return ({"weightedFontFamily": {"fontFamily": google[0], "weight": google[1]},
                     "fontSize": pt(size), "italic": google[2] or run["italic"]},
                    ["weightedFontFamily", "fontSize", "italic"])
        return ({"fontFamily": family, "fontSize": pt(size), "bold": run["bold"], "italic": run["italic"]},
                ["fontFamily", "fontSize", "bold", "italic"])

    def width_ratio(self, font: str, family: str, bold: bool, italic: bool) -> float:
        """Expected Slides width / PDF width of a run after the size correction: bold and
        italic are only half corrected (see __call__)."""
        if google_font(font) or family not in self.style:
            return 1.0
        ratio = 1.0
        for key, on in (("bold", bold), ("italic", italic)):
            if on:
                rel = self.style[family][key]
                ratio *= rel / (1 + (rel - 1) / 2)
        return ratio

    def __call__(self, run: dict, scale: float) -> tuple[str, float]:
        google = google_font(run["font"])
        if google:  # same font in Slides: no width correction
            return google[0], round(run["size"] * scale, 1)
        info = font_info(run["font"])
        family = FONT_FOR_FAMILY.get(run["family"], "Lato")
        design = info.design_size or 10
        if run["family"] == "mono":
            factor = ROBOTO_MONO_ADVANCE_EM / (CMTT_ADVANCE_EM * design_width(DESIGN_WIDTH["mono"], design))
        else:
            text, title = self.factors.get(run["family"], self.factors["sans"])
            if run["family"] != "serif" and 11.5 <= design < 14:
                factor = title  # calibrated directly on CMSS12 titles
            else:
                # Other optical sizes: CM's small cuts are wider per em (up to OPTICAL_WIDTH_MAX),
                # its large ones narrower.
                factor = text / optical_width(run["family"], design)
            # Bold and italic substitutes run 4-8% narrower than CM's; correct half of that, so
            # widths come closer without emphasised words looking visibly larger.
            style = self.style.get(run["family"], self.style["sans"])
            for key in ("bold", "italic"):
                if run[key]:
                    factor *= 1 + (style[key] - 1) / 2
            if info.design_size is not None:  # Computer Modern metrics (CM, EC, Latin Modern)
                if run.get("smallcaps"):
                    factor /= SMALL_CAPS_WIDTH.get(run["family"], 1.0)
                else:
                    factor *= self.shape_ratio(run, family, factor, design)
        return family, round(run["size"] * scale / factor, 1)

    def shape_ratio(self, run: dict, family: str, factor: float, design: float) -> float:
        """How much wider than the PDF's Slides sets this run for its letters, beyond what the
        size factor corrects; 1.0 when that is within tolerance (ordinary prose, a word), not
        known (a character neither table has) or not this run's to fix (scripts, holes).

        A run that is only a number (a digit, nothing but digits and a number's punctuation)
        gets the whole ratio at the size `factor` gives it, when it comes out wider. Any other
        run of SHAPE_MIN_CHARS counted characters or more is judged against the calibration
        sentences in its own face and gets the ratio when it is off by more than SHAPE_TOL, unless
        it shares its paragraph with other runs (`in_sentence`: sized like them)."""
        text = run.get("text", "")
        cm = CM_ADVANCES.get(cm_face(run) or "")
        # (a table cell's run keeps the table's size: its column is made as wide as Slides sets
        # it (fit_columns), and a number set smaller rode high in its top-anchored cell)
        if cm is None or not text.strip() or run.get("script") or run.get("hole") or run.get("cell") or \
                family not in ADVANCES:
            return 1.0
        slides = ADVANCES[family][STYLE_KEY[(bool(run["bold"]), bool(run["italic"]))]]
        number = "".join(text.split())
        if any(c in DIGITS for c in number) and all(c in NUMBER_CHARS for c in number):
            s_em = sum(slides.get(c, UNMEASURED_ADVANCE_EM) for c in number)
            p_em = sum(cm["advances"][c] for c in number)
            return max(1.0, s_em / factor / (p_em * optical_width(run["family"], design)))
        if run.get("in_sentence"):  # (a run among others keeps their size: in_sentence)
            return 1.0
        s_em, p_em, counted, skipped = advance_widths(text, cm, slides)
        if counted < SHAPE_MIN_CHARS or skipped > 0.1 * counted or p_em <= 0:
            return 1.0
        title = run["family"] != "serif" and 11.5 <= design < 14  # sized by the title factor
        ratio = s_em / p_em / self.reference_ratio(cm_face(run), slides, title)
        return ratio if abs(ratio - 1) > SHAPE_TOL else 1.0

    def reference_ratio(self, face: str, slides: dict, title: bool) -> float:
        """Slides em / PDF em of the calibration sentences in a face: where its size factor
        puts the widths of ordinary text."""
        key = (face, id(slides), title)  # `slides` is one of ADVANCES' tables, which live as long
        if key not in self._reference:
            s_em = p_em = 0.0
            for sentence in SHAPE_TITLE_REFERENCE if title else SHAPE_REFERENCE:
                s, p, _, _ = advance_widths(sentence, CM_ADVANCES[face], slides)
                s_em, p_em = s_em + s, p_em + p
            self._reference[key] = s_em / p_em
        return self._reference[key]


def bullet_shape(bullet: dict) -> str | None:
    """The BULLET_SHAPES entry for a bullet; None for numbers. A glyph's is its character's,
    except a bullet character the PDF draws as a filled square (LM Sans's \\textbullet)."""
    text = bullet.get("text", "")
    if bullet["kind"] == "number" or (bullet["kind"] == "image" and text.isdigit()):
        return None
    if bullet["kind"] == "glyph":
        shape = GLYPH_SHAPES.get(text, "disc")
        return "square" if shape == "disc" and inked_square(bullet) else shape
    return bullet.get("shape") if bullet.get("shape") in BULLET_SHAPES else "disc"


def inked_square(bullet: dict) -> bool:
    ink, fill = bullet.get("ink"), bullet.get("fill") or 0.0
    if not ink or fill < 0.9:  # (a disc fills 0.79 of its box)
        return False
    w, h = ink[2] - ink[0], ink[3] - ink[1]
    return 0.8 <= w / h <= 1.25 if h > 0 else False


# A glyph bullet whose ink (render.glyph_ink) would come out this much smaller than a bullet at
# the label's size is sized by its ink, as vector bullets are. Beamer's own glyphs (MSAM's ▶
# 0.58 em, CMSY's • 0.39 em against the disc's 0.41) keep the label's size.
INK_SIZED = 0.75


def ink_sized(bullet: dict, size: float, scale: float) -> float | None:
    """The size that gives a glyph bullet its PDF ink height, when it is to be used."""
    if bullet["kind"] != "glyph" or not bullet.get("ink"):
        return None
    shape = bullet_shape(bullet)
    label = bullet.get("label") or {}
    full = min(size, label.get("size", size / scale) * scale)
    height = (bullet["ink"][3] - bullet["ink"][1]) * scale
    inked = max(0.3 * size, min(size, height / BULLET_SHAPES[shape][2]))
    return inked if inked < INK_SIZED * full or shape != GLYPH_SHAPES.get(bullet.get("text", ""), "disc") else None


def bullet_preset(bullet: dict) -> str:
    shape = bullet_shape(bullet)
    if shape is None:
        return BULLET_PRESETS["number_parens" if ")" in bullet.get("text", "") else "number"]
    return BULLET_SHAPES[shape][0]


def bullet_level(bullet: dict, level: int) -> int:
    """Slides nesting level: numbers count by depth (1., a., i.), glyphs pick their shape.
    ● ○ ■ keep the depth (they repeat every 3 levels); other glyphs exist at one level only.
    (Indents are set explicitly, so the level only decides the glyph and what Tab does.)"""
    shape = bullet_shape(bullet)
    if shape is None:
        return level
    preset, first = BULLET_SHAPES[shape][:2]
    return 3 * min(level, 2) + first if preset == "BULLET_DISC_CIRCLE_SQUARE" else first


def bullet_size(bullet: dict, size: float, scale: float) -> float:
    """Font size giving the bullet its PDF height (at most the text's: a larger bullet would
    push the line down). Glyph and number boxes are font boxes: their size is the font's, or
    their ink's where that is much smaller (`ink_sized`)."""
    shape = bullet_shape(bullet)
    inked = ink_sized(bullet, size, scale)
    if inked is not None:
        return round(inked, 1)
    if bullet["kind"] in ("glyph", "number") or shape is None:
        label = bullet.get("label") or {}
        return round(min(size, label.get("size", size / scale) * scale), 1)
    height = (bullet["bbox"][3] - bullet["bbox"][1]) * scale
    return round(max(0.3 * size, min(size, height / BULLET_SHAPES[shape][2])), 1)


def bullet_gap(bullet: dict, size: float) -> float:
    """Distance from the bullet box's right edge to indentFirstLine."""
    shape = bullet_shape(bullet)
    if bullet["kind"] in ("glyph", "number") or shape is None:
        return BULLET_GAP
    return BULLET_SHAPES[shape][3] * size


def bullet_extent(bullet: dict, size: float, scale: float) -> tuple[float, float, float]:
    """(PDF x0, PDF x1, gap after it in Slides pt) of what the Slides bullet stands for: the
    ink of a glyph sized by its ink (then placed as a vector bullet is), else the bullet's box."""
    z = bullet_size(bullet, size, scale)
    if ink_sized(bullet, size, scale) is not None:
        return bullet["ink"][0], bullet["ink"][2], BULLET_SHAPES[bullet_shape(bullet)][3] * z
    return bullet["bbox"][0], bullet["bbox"][2], bullet_gap(bullet, z)


def rgb(hex_color: str) -> dict:
    h = hex_color.lstrip("#")
    return {"opaqueColor": {"rgbColor": {k: int(h[i:i + 2], 16) / 255 for k, i in
                                         (("red", 0), ("green", 2), ("blue", 4))}}}


# ---------------------------------------------------------------- text boxes

def line_size(run: dict, z: float) -> float:
    """The size Slides lays out a line holding this run at, when the run is its largest: its
    own - except a smallCaps run of lowercase letters only, which Slides draws wholly in the
    small font, at SMALL_CAPS_SIZE. One space, capital or comma in the run and the line takes
    the full size (tools/probe_text_fit_fonts.py, `line_size_pt`: PT Serif 26 small caps in a
    Lato 20 line - 'mm' lays out at 20.3, 'mm mm' and 'Mm' at 26.3; 34 pt 'mm' at 24.0)."""
    text = run.get("text", "")
    if run.get("smallcaps") and text and all(c.islower() for c in text):
        return z * SMALL_CAPS_SIZE
    return z


def body_size(runs: list[dict], sizes: list[float]) -> float | None:
    """The Slides size most of a paragraph's characters are set at (its scripts and holes aside)."""
    count: dict[float, int] = {}
    for run, z in zip(runs, sizes):
        if not run.get("script") and not run.get("hole") and not run.get("hole_size") and run["text"].strip():
            count[z] = count.get(z, 0) + len(run["text"].strip())
    return max(count, key=lambda z: (count[z], z)) if count else None


def run_sizes(runs: list[dict], scale: float, fonts: "FontMapper") -> list[float]:
    """The Slides size of each run. A subscript is set no larger than the text around it
    (`body_size`): Slides lowers a SUBSCRIPT run 0.371 em of its own size, where TeX lowers one
    0.15 em (0.25 beside a superscript), so every point it grows reaches further into the line
    below; and FontMapper would give it more than the text (the size is the text's, the font a
    small optical cut, which FontMapper reads as wider per em - 23.4 pt against 21.2). At the
    text's size Slides draws its digits as tall as TeX's (0.665 x 0.71 em against cmss8's) and
    it is what a person typing a subscript gets (tools/probe_subscripts.py)."""
    sizes = [fonts(r, scale)[1] for r in runs]
    body = body_size(runs, sizes)
    if body is None:
        return sizes
    return [min(z, body) if r.get("script") == "sub" else z for r, z in zip(runs, sizes)]


def run_width(run: dict, z: float, scale: float, fonts: "FontMapper") -> float:
    """About how wide Slides sets a run (pt): measured advances where there are some."""
    if run.get("hole_size"):
        return len(run["text"]) * HOLE_SPACE_EM * run["hole_size"]
    width = slides_width([run], scale, fonts)
    return width if width is not None else len(run["text"]) * 0.5 * z * (SCRIPT_SIZE if run.get("script") else 1.0)


def line_sizes(p: dict, sizes: list[float], scale: float, fonts: "FontMapper") -> list[float]:
    """The size Slides lays out each of a paragraph's lines at: the largest run *on that line*
    (`line_size`). A wrapped paragraph whose runs differ in size - an inline formula's italic
    letters, a larger word - does not have that size on every line, and Slides' pitch from one
    line to the next is the first one's descent and the next one's ascent (0.227 and 0.968 of
    each's own size, tools/probe_subscripts.py `crowding`). Which runs land on which line is
    estimated from the PDF's line widths, the runs laid end to end at their Slides widths."""
    n = len(p["lines"])
    sized = [line_size(r, z) for r, z in zip(p["runs"], sizes)]
    if not sized:
        return [p["size"] * scale] * n
    if n == 1 or len(set(sized)) == 1:
        return [max(sized)] * n
    widths = [run_width(r, z, scale, fonts) for r, z in zip(p["runs"], sizes)]
    total = sum(widths)
    spans = [max(1e-6, (ln.get("x1") or 0.0) - (ln.get("x0") or 0.0)) for ln in p["lines"]]
    bounds = [0.0]
    for w in spans:
        bounds.append(bounds[-1] + w * total / sum(spans))
    out = [0.0] * n
    pos = 0.0
    for s, w in zip(sized, widths):
        a, b = pos, pos + w
        pos = b
        mid = (a + b) / 2
        for j in range(n):
            lo, hi = bounds[j], bounds[j + 1]
            # a run lies on the line holding its middle, and on any line it covers half an em of
            if lo <= mid < hi or (j == n - 1 and mid >= hi) or min(b, hi) - max(a, lo) > 0.5 * s:
                out[j] = max(out[j], s)
    common = body_size(p["runs"], sized) or max(sized)
    return [z or common for z in out]


def extra_below(r: float, z: float) -> float:
    """Extra space lineSpacing r adds under a line of size z (negative when r < 1)."""
    return (r - 1) * LINE_EM * z if r >= 1 else -(1 - r) * 0.25 * LINE_EM * z


def extra_above(r: float, z: float) -> float:
    """lineSpacing >= 100% never moves a line down; below 100% it pulls the baseline up."""
    return 0.0 if r >= 1 else -(1 - r) * 0.75 * LINE_EM * z


def snap(v: float) -> float:
    """Slides lays lines out on whole CSS pixels (0.75 pt)."""
    return round(v / PX_PT) * PX_PT


def line_pitch(z: float, r: float, z2: float | None = None) -> float:
    """Baseline distance between wrapped lines of one paragraph (from a line of size z to one of
    size z2, by default the same)."""
    return snap(inner_pitch(z, r, z if z2 is None else z2))


def inner_pitch(z1: float, r: float, z2: float) -> float:
    """Unsnapped baseline distance between two wrapped lines of one paragraph, of sizes z1 and z2."""
    if z1 == z2:
        return LINE_EM * z1 * r
    return DESCENT_EM * z1 + ASCENT_EM * z2 + extra_below(r, z1) + extra_above(r, z2)


def pitch_between(z1: float, r1: float, z2: float, r2: float) -> float:
    """Baseline distance from the last line of one paragraph to the first line of the next."""
    return snap(DESCENT_EM * z1 + ASCENT_EM * z2 + extra_below(r1, z1) + extra_above(r2, z2))


def solve_increasing(f, target: float, lo: float = 0.5, hi: float = 3.0) -> float:
    for _ in range(40):
        mid = (lo + hi) / 2
        lo, hi = (mid, hi) if f(mid) < target else (lo, mid)
    return (lo + hi) / 2


HOLE_FONT, HOLE_SPACE_EM = "Roboto Mono", 0.6  # monospaced: a space is exactly 0.6 em


def in_sentence(runs: list[dict]) -> list[dict]:
    """A paragraph's runs, each marked `in_sentence` when another run with words or a formula
    shares the paragraph (a table cell, a node's line): FontMapper then leaves its letters' shape
    alone (`shape_ratio`). Sized to its own PDF width, a reference's author list ("J. Park, W.
    Zhang, et al. ", 0.905) or a journal abbreviation came out 10% larger than the title beside
    it, and a bullet with it: one size across a line matters more than one run's width. Copies,
    never the IR (sync diffs it)."""
    if sum(1 for r in runs if r.get("text", "").strip() or r.get("hole")) < 2:
        return runs
    return [r if r.get("in_sentence") else {**r, "in_sentence": True} for r in runs]


def hole_run(run: dict, scale: float, fonts: FontMapper) -> dict:
    """The gap under an inline formula picture (and the word space after it): no-break spaces
    in a monospaced font, sized so they are exactly as wide as the formula and no taller than
    the line."""
    z = fonts(run, scale)[1]
    width = run["hole"] * scale
    n = max(1, math.ceil(width / (HOLE_SPACE_EM * z)))
    return {**run, "text": " " * n, "hole_size": round(width / (HOLE_SPACE_EM * n), 2)}


def vertical_layout(paras: list[dict], baselines: list[list[float]], sizes: list):
    """lineSpacing ratio and spaceAbove per paragraph so Slides baselines land on the PDF's.
    `sizes` holds a paragraph's size, or the size of each of its lines (`line_sizes`).

    Slides ignores spaceAbove/spaceBelow between items of a bulleted list, so there the gap
    to the next item has to come from the item's own lineSpacing; for a wrapped item one
    ratio covers its inner lines plus that gap, spreading the difference evenly.

    Pitches snap to whole pixels, so each paragraph aims at the original position measured
    from where Slides will actually have put the previous one: rounding errors don't add up."""
    lines = [list(s) if isinstance(s, (list, tuple)) else [s] * len(bl) for s, bl in zip(sizes, baselines)]
    estimate = None
    for _ in range(4):  # a paragraph's ratio depends on the next one's (below 100% it moves up)
        ratios, space_above = _vertical_pass(paras, baselines, lines, estimate)
        if ratios == estimate:
            break
        estimate = ratios
    return ratios, space_above


def _vertical_pass(paras, baselines, lines, estimate):
    ratios: list[float] = []
    space_above = [0.0] * len(paras)
    pulled: dict[int, float] = {}  # paragraph -> lineSpacing < 1 that pulls it up to its target
    first = baselines[0][0]  # predicted Slides baseline of the current paragraph's first line
    for i, (p, bl, zs) in enumerate(zip(paras, baselines, lines)):
        n = len(bl)
        z = zs[-1]  # the last line's size: what the next paragraph is spaced from
        uniform = len(set(zs)) == 1

        def inner(r: float) -> float:  # first to last baseline of this paragraph, unsnapped
            return (n - 1) * LINE_EM * z * r if uniform else sum(inner_pitch(a, r, b) for a, b in zip(zs, zs[1:]))

        has_next = i + 1 < len(paras)
        list_link = has_next and p["bullet"] and paras[i + 1]["bullet"]
        next_r = estimate[i + 1] if estimate and has_next else 1.0
        if list_link:
            target = baselines[i + 1][0] - first
            zn = lines[i + 1][0]
            r = solve_increasing(lambda r: inner(r) + DESCENT_EM * z + ASCENT_EM * zn +
                                 extra_below(r, z) + extra_above(next_r, zn), target)
        elif n > 1:
            r = (bl[-1] - bl[0]) / (n - 1) / (LINE_EM * z) if uniform else solve_increasing(inner, bl[-1] - bl[0])
        else:
            r = pulled.get(i, 1.0)
        r = round(min(3.0, max(0.5, r)), 4)
        ratios.append(r)
        last = first + ((n - 1) * line_pitch(z, r) if uniform else sum(line_pitch(a, r, b) for a, b in zip(zs, zs[1:])))
        if has_next:
            zn = lines[i + 1][0]
            natural = pitch_between(z, r, zn, next_r)
            if not list_link:
                gap = baselines[i + 1][0] - last - pitch_between(z, r, zn, 1.0)
                nxt = paras[i + 1]
                free = len(baselines[i + 1]) == 1 and not (nxt["bullet"] and i + 2 < len(paras) and paras[i + 2]["bullet"])
                if gap < -PX_PT and free:
                    # Tighter than Slides' natural pitch (block title right above its body):
                    # a lineSpacing below 100% moves the next single line up.
                    rn = max(0.5, 1 + gap / (0.75 * LINE_EM * zn))
                    pulled[i + 1] = rn
                    natural = pitch_between(z, r, zn, rn)
                space_above[i + 1] = max(0.0, baselines[i + 1][0] - last - natural)
            first = last + natural + space_above[i + 1]
    return ratios, space_above


def hugs(p: dict) -> str:
    """Which page edge the paragraph's lines are drawn against, which `align` says for a
    left-to-right paragraph and understates for a right-to-left one: a Hebrew paragraph
    whose lines all end together hugs the **right**, and `align` calls that "left" because
    its lines also start together (justified prose) or because it has only one line, where
    nothing was measured at all. Slides is told an alignment relative to the reading
    direction, so it needs the edge, not the name."""
    if p.get("direction") == "rtl" and p["align"] == "left" and \
            max(l["x1"] for l in p["lines"]) - min(l["x1"] for l in p["lines"]) <= 1:
        return "right"
    return p["align"]


def box_lines(paras: list[dict], edges: list[str], scale: float, fonts: FontMapper) -> tuple[float, float] | None:
    """(right edge of the widest line, the least right edge at which a line would take its next
    word) in Slides pt over a left-aligned box's paragraphs as Slides sets their PDF lines
    (slides_lines); None when a paragraph cannot be measured."""
    if set(edges) != {"left"} or any(p.get("direction") == "rtl" or not p["runs"] for p in paras):
        return None
    got = [slides_lines(p, scale, fonts) for p in paras]
    if None in got:
        return None
    return max(g[0] for g in got), min(g[1] for g in got)


def text_box_requests(el: dict, slide_id: str, object_id: str, scale: float, fonts: FontMapper,
                      placeholder: dict | None = None, page_slide: dict[int, str] | None = None,
                      bar: list[float] | None = None, right_limit: float | None = None,
                      marks: list[str] | None = None) -> list[dict]:
    """A text box for a text element. With `bar` (the PDF box of a block's title bar that this
    one-line text sits on) the box fills the bar and centres its text vertically, so the
    title stays in the middle of the bar when the block is resized. `right_limit` (PDF x) is
    how far a box of unwrapped left-aligned text may extend. `marks` highlights the hole runs,
    one colour each (measure_places)."""
    marks = list(marks or [])
    paras = [{**p, "runs": [hole_run(r, scale, fonts) if r.get("hole") else r for r in in_sentence(p["runs"])]}
             for p in el["paragraphs"]]
    # A line is as tall as its largest run, as Slides lays it out (line_size: small caps), and a
    # subscript is no larger than its text (run_sizes).
    sized = [run_sizes(p["runs"], scale, fonts) for p in paras]
    base_sizes = [max(zs) if p["runs"] else p["size"] * scale for p, zs in zip(paras, sized)]
    # A bullet is no larger than its item's text (`body_size`), not its largest run: one {\Large}
    # word or a superscript's optical cut grew that item's bullet over its neighbours'.
    bullet_caps = [body_size(p["runs"], zs) or base for p, zs, base in zip(paras, sized, base_sizes)]
    per_line = [line_sizes(p, zs, scale, fonts) for p, zs in zip(paras, sized)]
    sizes = [max(ls) for ls in per_line]

    edges = [hugs(p) for p in paras]
    # (a centred or right-aligned paragraph's longest line can start left of its first line)
    left_pdf = min(p["bullet"]["bbox"][0] if p["bullet"] else
                   min([p["text_x0"]] + ([l["x0"] for l in p["lines"]] if e != "left" else []))
                   for p, e in zip(paras, edges))
    # (a right-to-left paragraph's bullet hangs right of its text, as a left-to-right one's
    # hangs left of it, so it is that side's edge)
    right_pdf = max([line["x1"] for p in paras for line in p["lines"]] +
                    [p["bullet"]["bbox"][2] for p in paras if p["bullet"] and p.get("direction") == "rtl"])
    first_baseline = paras[0]["lines"][0]["baseline"] * scale
    last_baseline = paras[-1]["lines"][-1]["baseline"] * scale

    baselines = [[line["baseline"] * scale for line in p["lines"]] for p in paras]
    ratios, space_above = vertical_layout(paras, baselines, per_line)

    inner_w = (right_pdf - left_pdf) * scale
    # Titles carry their line breaks as soft breaks (SOFT_BREAK) and need no tight width.
    multiline = any(len(p["lines"]) > 1 and not any(SOFT_BREAK in r["text"] for r in p["runs"]) for p in paras)
    # Wrapped paragraphs need a width that breaks where TeX did: wide enough for the longest
    # line, narrower than where the next line's first word would fit. The middle of that range
    # tolerates the substitute font being a little wider or narrower. Single lines get room so
    # that a slightly wider font never wraps them.
    limits = [p["wrap_limit"] for p in paras if p.get("wrap_limit") and not any(SOFT_BREAK in r["text"] for r in p["runs"])]
    room = (min(limits) - right_pdf) * scale if limits else 0.0
    measured = box_lines(paras, edges, scale, fonts) if multiline else None
    if not multiline:
        slack = max(0.15 * inner_w, 2 * max(sizes))
    elif measured:
        # Each PDF line as Slides sets its words: a TeX-full line in a narrow column comes out a
        # few points wider in Lato, more than the room the PDF leaves before the next word.
        widest, joins = measured
        inner_w = widest - left_pdf * scale
        room = joins - widest
        slack = room / 2 if room > 2 * WRAP_MARGIN else WRAP_MARGIN
        if all(p.get("justified") for p in paras if len(p["lines"]) > 1):
            # JUSTIFIED sets every line but the last out to the box's edge: room there would
            # move the column's right edge, not just leave the next word out.
            slack = min(slack, 2 * WRAP_MARGIN)
    elif room > 4:
        slack = room / 2
    else:
        slack = 2 + 0.01 * inner_w
    x = left_pdf * scale - PAD_X
    aligns = set(edges)
    if aligns == {"center"}:
        x -= slack / 2
    elif aligns == {"right"}:
        x -= slack
    z_first, z_last = per_line[0][0], per_line[-1][-1]
    y = first_baseline - (BASELINE_A + ASCENT_EM * z_first + extra_above(ratios[0], z_first))
    w = inner_w + 2 * PAD_X + slack
    h = last_baseline - y + DESCENT_EM * z_last + extra_below(ratios[-1], z_last) + 4
    if right_limit and not multiline and aligns == {"left"} and not any(SOFT_BREAK in r["text"] for p in paras for r in p["runs"]):
        # Room up to the block edge or the next element: text typed later wraps where a user
        # expects, not a few points after the converted words.
        w = max(w, right_limit * scale - x)
    middle = bool(bar) and not placeholder and len(paras) == 1 and len(paras[0]["lines"]) == 1 and ratios[0] == 1
    if middle:
        h = (bar[3] - bar[1]) * scale
        y = first_baseline - MIDDLE_BASELINE_EM * sizes[0] - h / 2
        if aligns == {"left"}:
            w = max(w, (bar[2] - 1) * scale - x)  # to the bar's end: the title wraps with the block

    if placeholder:
        # An existing layout placeholder (the slide title): its size is fixed at creation, so
        # it is resized through the transform's scale. Text is not scaled by that.
        y += placeholder["dy"]
        reqs = [
            {"updatePageElementTransform": {"objectId": object_id, "applyMode": "ABSOLUTE", "transform": {
                "scaleX": w / placeholder["base_w"], "scaleY": h / placeholder["base_h"], "unit": "EMU",
                "translateX": round(x * EMU_PER_PT), "translateY": round(y * EMU_PER_PT)}}},
            {"updateShapeProperties": {"objectId": object_id, "fields": "contentAlignment,autofit.autofitType",
                                       "shapeProperties": {"contentAlignment": "TOP",
                                                           "autofit": {"autofitType": "NONE"}}}},
        ]
    else:
        transform = {"scaleX": 1, "scaleY": 1, "unit": "EMU",
                     "translateX": round(x * EMU_PER_PT), "translateY": round(y * EMU_PER_PT)}
        if el.get("rotation"):
            # Laid out in the text's own frame (classify.rotated_texts): turn the box onto the page.
            turn = 1 if el["rotation"] > 0 else -1
            transform = {"scaleX": 0, "scaleY": 0, "shearX": -turn, "shearY": turn, "unit": "EMU",
                         "translateX": round(-turn * y * EMU_PER_PT), "translateY": round(turn * x * EMU_PER_PT)}
        reqs = [{"createShape": {
            "objectId": object_id, "shapeType": "TEXT_BOX",
            "elementProperties": {
                "pageObjectId": slide_id,
                "size": {"width": emu(w), "height": emu(h)},
                "transform": transform,
            },
        }}]
        if middle:
            reqs.append({"updateShapeProperties": {"objectId": object_id, "fields": "contentAlignment",
                                                   "shapeProperties": {"contentAlignment": "MIDDLE"}}})

    texts =["".join(r["text"] for r in p["runs"]) for p in paras]
    # Bullets: contiguous ranges with the same preset. createParagraphBullets consumes the
    # leading tabs and sets nesting levels relative to the range's shallowest paragraph, so a
    # range whose levels start above 0 begins with a dummy paragraph, deleted right after.
    levels = [bullet_level(p["bullet"], p["level"]) if p["bullet"] else 0 for p in paras]
    ranges = []
    for i, p in enumerate(paras):
        preset = bullet_preset(p["bullet"]) if p["bullet"] else None
        if preset and ranges and ranges[-1][2] == preset and ranges[-1][1] == i - 1:
            ranges[-1][1] = i
        elif preset:
            ranges.append([i, i, preset])
    dummies = {first for first, last, _ in ranges if min(levels[first:last + 1]) > 0}
    starts_tabbed, parts, pos = [], [], 0
    for i, (p, t) in enumerate(zip(paras, texts)):
        if i in dummies:
            parts.append("-")
            pos += 2
        starts_tabbed.append(pos)
        parts.append("\t" * levels[i] + t)
        pos += levels[i] + u16(t) + 1  # (every index below counts UTF-16 units, as Slides does: u16)
    reqs.append({"insertText": {"objectId": object_id, "text": "\n".join(parts), "insertionIndex": 0}})
    # A bullet keeps the text style it was created with, unless a later style request covers
    # its whole paragraph. So every paragraph first gets its base family and size, bulleted
    # ones the bullet's size and colour, and the runs are styled below in parts.
    for i, (p, start, level, size, cap) in enumerate(zip(paras, starts_tabbed, levels, base_sizes, bullet_caps)):
        family = fonts(p["runs"][0], scale)[0] if p["runs"] else "Lato"
        length = level + u16("".join(r["text"] for r in p["runs"]))
        style = {"fontFamily": family, "fontSize": pt(size)}
        if p["bullet"]:
            style["fontSize"] = pt(bullet_size(p["bullet"], cap, scale))
            color = p["bullet"].get("color") or (p["runs"][0]["color"] if p["runs"] else None)
            if color:
                style["foregroundColor"] = rgb(color)
        if length:
            reqs.append({"updateTextStyle": {
                "objectId": object_id, "fields": ",".join(style), "style": style,
                "textRange": {"type": "FIXED_RANGE", "startIndex": start - (2 if i in dummies else 0),
                              "endIndex": start + length},
            }})
    for first, last, preset in reversed(ranges):
        start = starts_tabbed[first] - (2 if first in dummies else 0)
        end = starts_tabbed[last] + levels[last] + u16(texts[last])
        reqs.append({"createParagraphBullets": {
            "objectId": object_id, "bulletPreset": preset,
            "textRange": {"type": "FIXED_RANGE", "startIndex": start, "endIndex": end},
        }})
        if first in dummies:
            reqs.append({"deleteText": {"objectId": object_id,
                                        "textRange": {"type": "FIXED_RANGE", "startIndex": start, "endIndex": start + 2}}})

    # From here on indices refer to the final text, without tabs.
    pos = 0
    for p, t, ratio, above, cap, edge, zs in zip(paras, texts, ratios, space_above, bullet_caps, edges, sized):
        p_start, p_end = pos, pos + u16(t)
        pos = p_end + 1
        start = p_start
        for run, z in zip(p["runs"], zs):
            if not run["text"]:
                continue
            style, fields = fonts.text_style(run, scale)
            if "fontSize" in style:
                style["fontSize"] = pt(z)  # (a subscript no larger than its text: run_sizes)
            if run.get("hole_size"):
                style, fields = {"fontFamily": HOLE_FONT, "fontSize": pt(run["hole_size"]), "bold": False,
                                 "italic": False}, ["fontFamily", "fontSize", "bold", "italic"]
            style.update({"smallCaps": run["smallcaps"], "foregroundColor": rgb(run["color"]),
                          "underline": bool(run.get("underline")), "strikethrough": bool(run.get("strike")),
                          "baselineOffset": {"super": "SUPERSCRIPT", "sub": "SUBSCRIPT"}.get(run.get("script"), "NONE")})
            fields = fields + ["smallCaps", "foregroundColor", "underline", "strikethrough", "baselineOffset"]
            if run.get("highlight"):
                style["backgroundColor"] = rgb(run["highlight"])
                fields.append("backgroundColor")
            if run.get("hole_size") and marks:
                style["backgroundColor"] = rgb(marks.pop(0))
                fields.append("backgroundColor")
            fields = ",".join(dict.fromkeys(fields))
            if run["link"] and run["link"].startswith("#page="):
                target = page_slide.get(int(run["link"][6:])) if page_slide else None
                if target:
                    style["link"] = {"pageObjectId": target}  # TOC entries jump to their slide
                    fields += ",link"
            elif run["link"]:
                style["link"] = {"url": run["link"]}
                fields += ",link"
            end = start + u16(run["text"])
            # (one request over the whole paragraph would restyle its bullet too; the cut before
            # the last character, never inside its surrogate pair)
            cuts = [start, end - u16(run["text"][-1]), end] \
                if p["bullet"] and start == p_start and end == p_end and len(run["text"]) > 1 else [start, end]
            for c0, c1 in zip(cuts, cuts[1:]):
                reqs.append({"updateTextStyle": {
                    "objectId": object_id, "style": style, "fields": fields,
                    "textRange": {"type": "FIXED_RANGE", "startIndex": c0, "endIndex": c1},
                }})
            start = end

        # Slides measures a paragraph from where it *starts*, which is the right edge of a
        # right-to-left one (classify.Paragraph.direction): START is that edge, and so is
        # indentStart. What classify measured is the page's own left and right, so both are
        # mirrored here, around the edge the paragraph hugs (`hugs`).
        rtl = p.get("direction") == "rtl"
        # Code lines carry their indentation as leading spaces already.
        # A paragraph that hugs the other edge, or the middle, places itself: an indent would
        # only offset it.
        start_edge = "right" if rtl else "left"
        room = (right_pdf - max(l["x1"] for l in p["lines"])) if rtl else (p["text_x0"] - left_pdf)
        text_indent = 0.0 if el.get("code") or edge != start_edge else room * scale
        if p["bullet"]:
            # Slides ends the bullet glyph a little before indentFirstLine.
            b_x0, b_x1, gap = bullet_extent(p["bullet"], cap, scale)
            side = (right_pdf - b_x0) if rtl else (b_x1 - left_pdf)
            first_indent = side * scale + gap
        elif p.get("tab_x0") and not el.get("code") and not rtl:
            # "label<TAB>content": a tab after the hanging label jumps to indentStart.
            # (a right-to-left label's tab lands where nothing in the PDF says: no hang)
            first_indent, text_indent = text_indent, (p["tab_x0"] - left_pdf) * scale
        else:
            first_indent = text_indent
            if edge == start_edge and not rtl and not el.get("code") and len(p["lines"]) > 1 and \
                    p["lines"][0]["x0"] > p["text_x0"] + 0.2 * p["size"]:
                # a first line set in by \parindent (classify.Paragraph.indent)
                first_indent += (p["lines"][0]["x0"] - p["text_x0"]) * scale
        # Justified prose stays justified (classify.PageClassifier.is_justified); Slides leaves
        # the last line ragged, as TeX does.
        justify = p.get("justified") and edge == start_edge
        reqs.append({"updateParagraphStyle": {
            "objectId": object_id,
            "textRange": {"type": "FIXED_RANGE", "startIndex": p_start, "endIndex": max(p_end, p_start + 1)},
            "style": {
                "alignment": "JUSTIFIED" if justify else "CENTER" if edge == "center" else
                             "START" if (edge == "right") == rtl else "END",
                "lineSpacing": round(100 * ratio, 1),
                "spaceAbove": pt(round(above, 2)), "spaceBelow": pt(0),
                "indentStart": pt(round(text_indent, 2)), "indentFirstLine": pt(round(first_indent, 2)),
                **({"direction": "RIGHT_TO_LEFT"} if rtl else {}),
            },
            "fields": "alignment,lineSpacing,spaceAbove,spaceBelow,indentStart,indentFirstLine" +
                      (",direction" if rtl else ""),
        }})
    return reqs


def number_box_requests(number: dict, slide_id: str, object_id: str, scale: float, fonts: FontMapper) -> list[dict]:
    """A literal list number centred on its ball picture (classify.literal_list_numbers): a
    box around the ball's centre with centred text and contentAlignment MIDDLE. Lato digits
    are 0.72 em tall, so a baseline 0.362 em below the middle puts them in the middle too."""
    run = {**number, "smallcaps": False}
    style, fields = fonts.text_style(run, scale)
    size = fonts(run, scale)[1]
    cx, cy = number["center"][0] * scale, number["center"][1] * scale
    w = number["height"] * scale + 2 * PAD_X + len(number["text"]) * size  # never wraps "(iv)"
    h = max(number["height"] * scale, LINE_EM * size + 2)
    style["foregroundColor"] = rgb(number["color"])
    return [
        {"createShape": {"objectId": object_id, "shapeType": "TEXT_BOX", "elementProperties": {
            "pageObjectId": slide_id, "size": {"width": emu(w), "height": emu(h)},
            "transform": {"scaleX": 1, "scaleY": 1, "unit": "EMU",
                          "translateX": round((cx - w / 2) * EMU_PER_PT), "translateY": round((cy - h / 2) * EMU_PER_PT)}}}},
        {"updateShapeProperties": {"objectId": object_id, "fields": "contentAlignment,autofit.autofitType",
                                   "shapeProperties": {"contentAlignment": "MIDDLE", "autofit": {"autofitType": "NONE"}}}},
        {"insertText": {"objectId": object_id, "text": number["text"], "insertionIndex": 0}},
        {"updateTextStyle": {"objectId": object_id, "style": style, "fields": ",".join(fields + ["foregroundColor"]),
                             "textRange": {"type": "ALL"}}},
        {"updateParagraphStyle": {"objectId": object_id, "textRange": {"type": "ALL"},
                                  "style": {"alignment": "CENTER", "lineSpacing": 100, "spaceAbove": pt(0), "spaceBelow": pt(0),
                                            "indentStart": pt(0), "indentFirstLine": pt(0)},
                                  "fields": "alignment,lineSpacing,spaceAbove,spaceBelow,indentStart,indentFirstLine"}},
    ]


def merge_blocks(elements: list[dict]) -> list[dict]:
    """A block body (see classify.blocks) reaches up under its title bar, with the outline
    of the whole block: resizing the block as a group can then never open a gap between
    the two, and the body's shadow falls behind the whole block."""
    heads = {e["block"]: e for e in elements if e["kind"] == "shape" and e.get("block") is not None and not e.get("title_bar")}
    out = []
    for el in elements:
        head = heads.get(el.get("block")) if el.get("title_bar") else None
        if head is None:
            out.append(el)
            continue
        top = el["title_bar"][1]
        top_round = head["shape"] == "ROUND_RECTANGLE" or (head["shape"] == "ROUND_2_SAME_RECTANGLE" and not head["flip"])
        bottom_round = el["shape"] == "ROUND_RECTANGLE" or (el["shape"] == "ROUND_2_SAME_RECTANGLE" and el["flip"])
        shape, flip = {(True, True): ("ROUND_RECTANGLE", False), (False, False): ("RECTANGLE", False),
                       (True, False): ("ROUND_2_SAME_RECTANGLE", False),
                       (False, True): ("ROUND_2_SAME_RECTANGLE", True)}[(top_round, bottom_round)]
        out.append({**el, "bbox": [el["bbox"][0], top, el["bbox"][2], el["bbox"][3]], "shape": shape, "flip": flip,
                    "radius": max(el["radius"], head["radius"])})
    # Creation order is z-order: title bars go above every body (shapes come first, then the rest).
    bar = lambda e: e["kind"] == "shape" and e.get("block") is not None and not e.get("title_bar")
    return sorted(out, key=lambda e: 0 if e["kind"] == "shape" and not bar(e) else 1 if bar(e) else 2)


# Native drop shadows, calibrated against beamer's block shadow (tools/calibrate_shadow.py):
# distance and blur per point of the shadow's width in the PDF, at 45°.
SHADOW_DISTANCE = 0.75
SHADOW_BLUR = 1.0
SHADOW_ALPHA = 0.5
TEMPLATE_PRESETS = {"ROUND_RECTANGLE": "roundRect", "ROUND_2_SAME_RECTANGLE": "round2SameRect", "RECTANGLE": "rect"}
TEMPLATE_LAYOUTS = {"TITLE": 0, "TITLE_ONLY": 5, "BLANK": 6}  # python-pptx default template layout indexes


def template_key(el: dict, scale: float) -> tuple | None:
    """Shapes the API can't make exactly: rounded corners of a given radius (the API only
    creates the default rounding) and drop shadows (read-only in the API). They are
    duplicated from template shapes that come with the imported .pptx."""
    if el["kind"] != "shape" or (el["shape"] == "RECTANGLE" and not el.get("shadow")):
        return None
    x0, y0, x1, y1 = el["bbox"]
    adj = 0.0
    if el["shape"] != "RECTANGLE":
        adj = min(0.5, round(el.get("radius", 0.0) / max(min(x1 - x0, y1 - y0), 0.01), 2))
    shadow = round(2 * el["shadow"]["size"] * scale) / 2 if el.get("shadow") else None
    return el["shape"], adj, shadow


NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _set_background(part, c_sld, fill: dict) -> None:
    """A page background in the .pptx: {"color": "#rrggbb"} or {"picture": Path} (stretched)."""
    from lxml import etree

    if "color" in fill:
        inner = f'<a:solidFill><a:srgbClr val="{fill["color"].lstrip("#").upper()}"/></a:solidFill>'
    else:
        _, rid = part.get_or_add_image_part(str(fill["picture"]))  # identical files are stored once
        inner = (f'<a:blipFill dpi="0" rotWithShape="1"><a:blip r:embed="{rid}"/><a:srcRect/>'
                 f'<a:stretch><a:fillRect/></a:stretch></a:blipFill>')
    old = c_sld.find(f"{{{NS_P}}}bg")
    if old is not None:
        c_sld.remove(old)
    c_sld.insert(0, etree.fromstring(
        f'<p:bg xmlns:p="{NS_P}" xmlns:a="{NS_A}" xmlns:r="{NS_R}"><p:bgPr>{inner}<a:effectLst/></p:bgPr></p:bg>'))


def _add_template_shapes(slide, keys: list[tuple]) -> None:
    from lxml import etree
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.util import Pt

    kinds = {"ROUND_RECTANGLE": MSO_SHAPE.ROUNDED_RECTANGLE, "ROUND_2_SAME_RECTANGLE": MSO_SHAPE.ROUND_2_SAME_RECTANGLE,
             "RECTANGLE": MSO_SHAPE.RECTANGLE, "ELLIPSE": MSO_SHAPE.OVAL, "DIAMOND": MSO_SHAPE.DIAMOND,
             "TRIANGLE": MSO_SHAPE.ISOSCELES_TRIANGLE}
    a = NS_A
    for i, (kind, adj, shadow) in enumerate(keys):
        if kind == "BENT_CONNECTOR":
            line = slide.shapes.add_connector(MSO_CONNECTOR.ELBOW, Pt(10), Pt(10), Pt(110), Pt(110))
            geometry = line._element.spPr.find(f"{{{a}}}prstGeom")
            geometry.set("prst", "bentConnector3")
            for old in geometry.findall(f"{{{a}}}avLst"):
                geometry.remove(old)
            geometry.append(etree.fromstring(
                f'<a:avLst xmlns:a="{a}"><a:gd name="adj1" fmla="val {round(adj * 100000)}"/></a:avLst>'))
            line._element.spPr.append(etree.fromstring(f'<a:effectLst xmlns:a="{a}"/>'))
            continue
        shape = slide.shapes.add_shape(kinds[kind], Pt(10 + i % 10 * 20), Pt(10 + i // 10 * 20), Pt(100), Pt(100))
        if adj is not None and kind in ("ROUND_RECTANGLE", "ROUND_2_SAME_RECTANGLE"):
            shape.adjustments[0] = adj
            if kind == "ROUND_2_SAME_RECTANGLE":
                shape.adjustments[1] = 0.0
        shape.fill.solid()
        shape.line.fill.background()
        # No text padding (the API can't set it): a diagram label fits a node as tight as TikZ's.
        body_pr = shape.text_frame._txBody.find(f"{{{a}}}bodyPr")
        for side in ("lIns", "tIns", "rIns", "bIns"):
            body_pr.set(side, "0")
        body_pr.set("anchor", "ctr")
        # python-pptx shapes refer to the theme's effect style, which has a shadow: always
        # give an explicit (possibly empty) effect list.
        effects = (f'<a:outerShdw blurRad="{round(SHADOW_BLUR * shadow * EMU_PER_PT)}" '
                   f'dist="{round(SHADOW_DISTANCE * shadow * EMU_PER_PT)}" dir="2700000" algn="tl" rotWithShape="0">'
                   f'<a:srgbClr val="000000"><a:alpha val="{round(SHADOW_ALPHA * 100000)}"/></a:srgbClr>'
                   f'</a:outerShdw>') if shadow else ""
        shape.element.spPr.append(etree.fromstring(f'<a:effectLst xmlns:a="{a}">{effects}</a:effectLst>'))


NO_TABLE_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # PowerPoint's "No Style, No Grid"


def _add_table(slide, table: dict) -> None:
    """An empty table (pptx_table) on a source slide. Its cell margins are what the API can't
    set: a table made by createTable has 7.2 pt above and below every line, one from a .pptx the
    file's (tools/probe_pptx_table_margins.py); duplicating the slide, inserting rows and columns
    and filling or styling cells through the API all keep them."""
    from lxml import etree
    from pptx.util import Emu

    def e(v: float) -> Emu:
        return Emu(round(v * EMU_PER_PT))

    rows, cols = len(table["heights"]), len(table["widths"])
    frame = slide.shapes.add_table(rows, cols, e(table["x"]), e(table["y"]), e(sum(table["widths"])),
                                   e(sum(table["heights"])))
    pr = frame._element.graphic.graphicData.tbl.tblPr
    for flag in ("firstRow", "bandRow"):  # (python-pptx's default look: a header row and bands)
        pr.attrib.pop(flag, None)
    style = pr.find(f"{{{NS_A}}}tableStyleId")
    if style is None:
        style = etree.SubElement(pr, f"{{{NS_A}}}tableStyleId")
    style.text = NO_TABLE_STYLE
    for c, w in enumerate(table["widths"]):
        frame.table.columns[c].width = e(w)
    for r, h in enumerate(table["heights"]):
        frame.table.rows[r].height = e(h)
        left, top, right, bottom = table["margins"][r]
        for c in range(cols):
            cell = frame.table.cell(r, c)
            cell.margin_left, cell.margin_top, cell.margin_right, cell.margin_bottom = e(left), e(top), e(right), e(bottom)


VARIANT = "_V"      # layout name suffix: a copy of the layout with another theme decoration (plan_theme)
THEME_VARIANTS = 3


def _clone_layout(prs, layout, name: str):
    """A copy of a layout (placeholders only, no pictures) added to the master, shown as `name`."""
    from copy import deepcopy

    from lxml import etree
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.opc.packuri import PackURI
    from pptx.parts.slide import SlideLayoutPart

    master = prs.slide_master
    taken = {str(p.partname) for p in prs.part.package.iter_parts()}
    k = next(k for k in range(1, 1000) if f"/ppt/slideLayouts/slideLayout{k}.xml" not in taken)
    element = deepcopy(layout.element)
    element.cSld.set("name", name)
    element.set("type", "cust")  # (its own layout name in Slides, never taken for the original's)
    part = SlideLayoutPart(PackURI(f"/ppt/slideLayouts/slideLayout{k}.xml"), layout.part.content_type,
                           layout.part.package, element)
    part.relate_to(master.part, RT.SLIDE_MASTER)
    ids = master.element.get_or_add_sldLayoutIdLst()
    entry = etree.SubElement(ids, f"{{{NS_P}}}sldLayoutId")
    entry.set("id", str(max(int(e.get("id")) for e in ids if e.get("id")) + 1))
    entry.set(f"{{{NS_R}}}id", master.part.relate_to(part, RT.SLIDE_LAYOUT))
    return part.slide_layout


def _add_decoration(layout, picture: Path, width: int, height: int) -> None:
    """The theme decoration as a full-page picture at the bottom of a layout: above the slide
    background, below everything on the slide."""
    from lxml import etree

    _, rid = layout.part.get_or_add_image_part(str(picture))
    tree = layout.shapes._spTree
    shape_id = max([int(e.get("id")) for e in tree.iter(f"{{{NS_P}}}cNvPr")] + [1]) + 1
    tree.insert(2, etree.fromstring(
        f'<p:pic xmlns:p="{NS_P}" xmlns:a="{NS_A}" xmlns:r="{NS_R}"><p:nvPicPr>'
        f'<p:cNvPr id="{shape_id}" name="Theme" descr="Theme decoration"/><p:cNvPicPr><a:picLocks noGrp="1"/></p:cNvPicPr>'
        f'<p:nvPr userDrawn="1"/></p:nvPicPr><p:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch>'
        f'</p:blipFill><p:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{width}" cy="{height}"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>'))


def build_pptx(page_w: float, page_h: float, keys: list[tuple], pages: list[dict], master_fill: dict,
               decorations: dict | None = None) -> io.BytesIO:
    """The deck's starting point, imported through Drive. It carries everything the Slides API
    could only insert from a public URL, so no picture ever leaves the user's Drive:

    - the PDF's page size (presentations.create ignores pageSize);
    - the master background (`master_fill`), inherited by the layouts and most slides;
    - the theme decoration (`decorations`, see plan_theme) at the bottom of the layouts; a page
      layout named with the VARIANT suffix is a copy of its layout with that variant's decoration;
    - one source slide per deck slide (`pages`: {"layout", "fill" (None: inherit),
      "pictures": [{"file", "bbox" (slide pt), "alt", "title"}], "tables": [pptx_table(...)],
      "templates" (bool)}), holding its pictures, its tables (empty, with the cell margins the API
      cannot set) and, if it needs any, the template shapes (shadows, exact corner radii).

    emit copies each source slide under our own object IDs and then deletes it."""
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation()
    height = SLIDE_W * page_h / page_w
    ratio = height / (prs.slide_height / EMU_PER_PT)
    prs.slide_width, prs.slide_height = Emu(round(SLIDE_W * EMU_PER_PT)), Emu(round(height * EMU_PER_PT))
    if abs(ratio - 1) > 1e-3:  # the default template's placeholders are laid out for 4:3
        for page in [prs.slide_master, *prs.slide_layouts]:
            for shape in page.placeholders:
                # A layout placeholder without its own position inherits the master's, rescaled
                # already (setting its top and height would scale it twice and write x and width 0).
                if shape._element.spPr.find(f"{{{NS_A}}}xfrm") is not None and shape.height is not None:
                    shape.top, shape.height = Emu(round(shape.top * ratio)), Emu(round(shape.height * ratio))
    master = prs.slide_master
    _set_background(master.part, master.element.find(f"{{{NS_P}}}cSld"), master_fill)
    decorations = decorations or {}
    layouts = {name: prs.slide_layouts[i] for name, i in TEMPLATE_LAYOUTS.items()}
    originals = list(prs.slide_layouts)
    for name in dict.fromkeys(p["layout"] for p in pages if VARIANT in p["layout"]):
        kind, n = name.rsplit(VARIANT, 1)
        picture = decorations.get(f"{'TITLE' if kind == 'TITLE' else '*'}{VARIANT}{n}")
        layouts[name] = _clone_layout(prs, layouts[kind], f"{layouts[kind].name} ({f'theme {int(n) + 1}' if picture else 'no theme'})")
        if picture:
            _add_decoration(layouts[name], picture, prs.slide_width, prs.slide_height)
    for i, layout in enumerate(originals):
        picture = decorations.get("TITLE" if i == TEMPLATE_LAYOUTS["TITLE"] else "*")
        if picture:
            _add_decoration(layout, picture, prs.slide_width, prs.slide_height)
    for page in pages:
        slide = prs.slides.add_slide(layouts[page["layout"]])
        if page["fill"]:
            _set_background(slide.part, slide.element.find(f"{{{NS_P}}}cSld"), page["fill"])
        for pic in page["pictures"]:
            x0, y0, x1, y1 = pic["bbox"]
            shape = slide.shapes.add_picture(str(pic["file"]), Emu(round(x0 * EMU_PER_PT)), Emu(round(y0 * EMU_PER_PT)),
                                             Emu(round((x1 - x0) * EMU_PER_PT)), Emu(round((y1 - y0) * EMU_PER_PT)))
            if pic.get("alt"):  # (text from the PDF: raw Type 3 T1 codes are C0 controls, xml_text)
                shape._element.nvPicPr.cNvPr.set("descr", xml_text(pic["alt"]))
                shape._element.nvPicPr.cNvPr.set("title", xml_text(pic["title"]))
        for table in page.get("tables", []):
            _add_table(slide, table)
        if page["templates"]:
            _add_template_shapes(slide, keys)
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf


def shape_requests(el: dict, slide_id: str, object_id: str, scale: float, template: dict | None = None) -> list[dict]:
    """A filled shape without outline. With a template ({"id", "w", "h"}: a template shape
    on this slide and its size in pt) the shape is a duplicate of it, else a new shape."""
    x0, y0, x1, y1 = (v * scale for v in el["bbox"])
    # ROUND_2_SAME_RECTANGLE rounds the top corners; for bottom corners flip both axes
    # (a 180° rotation), which moves the origin to the opposite corner.
    flip = -1 if el["flip"] else 1
    tx, ty = (x1, y1) if el["flip"] else (x0, y0)
    if template:
        reqs = [
            {"duplicateObject": {"objectId": template["id"], "objectIds": {template["id"]: object_id}}},
            {"updatePageElementTransform": {"objectId": object_id, "applyMode": "ABSOLUTE", "transform": {
                "scaleX": flip * (x1 - x0) / template["w"], "scaleY": flip * (y1 - y0) / template["h"], "unit": "EMU",
                "translateX": round(tx * EMU_PER_PT), "translateY": round(ty * EMU_PER_PT)}}},
            {"updatePageElementsZOrder": {"pageElementObjectIds": [object_id], "operation": "BRING_TO_FRONT"}},
        ]
    else:
        reqs = [{"createShape": {
            "objectId": object_id, "shapeType": el["shape"],
            "elementProperties": {
                "pageObjectId": slide_id,
                "size": {"width": emu(x1 - x0), "height": emu(y1 - y0)},
                "transform": {"scaleX": flip, "scaleY": flip, "unit": "EMU",
                              "translateX": round(tx * EMU_PER_PT), "translateY": round(ty * EMU_PER_PT)},
            },
        }}]
    return reqs + [
        {"updateShapeProperties": {
            "objectId": object_id,
            "shapeProperties": {"shapeBackgroundFill": {"solidFill": {"color": rgb(el["fill"])["opaqueColor"],
                                                                      "alpha": el.get("opacity", 1.0)}},
                                "outline": {"propertyState": "NOT_RENDERED"}},
            "fields": "shapeBackgroundFill.solidFill.color,shapeBackgroundFill.solidFill.alpha,outline.propertyState",
        }},
    ]


TABLE_MIN_COLUMN_PT = 32.0  # the API refuses narrower columns
# A table made by createTable has 7.2 pt of cell padding above and below that the API cannot
# change; a table the .pptx brings keeps the file's `a:tcPr` margins, down to 0
# (tools/probe_pptx_table_margins.py). emit's own tables come with the .pptx (`build_pptx`).
TABLE_ROW_PAD = 14.4        # an API-made table's padding above and below
TABLE_ROW_EM = 1.195        # a row of one line is at least its insets + 1.195·z·lineSpacing
TABLE_MIN_SPACING = 0.5
TABLE_TEXT_TOP = BASELINE_A - TABLE_ROW_PAD / 2  # cell top inset -> first baseline, less ASCENT_EM·z


def table_line_spacing(pitch: float, z: float, pad: float = TABLE_ROW_PAD) -> float:
    """lineSpacing ratio at which a row of text size z fits into the given row pitch."""
    return min(1.0, max(TABLE_MIN_SPACING, (pitch - pad) / (TABLE_ROW_EM * z)))


TABLE_CELL_PAD = 7.2  # cell padding left and right
# Advance widths (em) of text characters on Slides' own renderer, per substitute font and
# style (tools/probe_advances.py). Slides' Lato is not the Lato on google/fonts (its space is
# 0.19 em, its slash 0.31), so these are measured, not read out of a font file.
ADVANCES = json.loads((CALIBRATION_DIR / "advances.json").read_text(encoding="utf-8"))["fonts"]
UNMEASURED_ADVANCE_EM = 0.6  # a character the probe did not measure: as wide as the widest digits
SCRIPT_SIZE = 2 / 3          # super- and subscripts in Slides (measured 0.665: tools/probe_text_fit_fonts.py)
WRAP_MARGIN = 1.0            # Slides pt kept free in a cell so kerning or rounding cannot wrap it
SMALL_CAPS_SIZE = 0.70       # Slides draws a small capital at 70% of its capital (tools/probe_text_fit_fonts.py)


def slides_width(runs: list[dict], scale: float, fonts: "FontMapper") -> float | None:
    """Advance width (Slides pt) of one line of these runs as Slides sets them, or None when a
    run is in a font the probe did not measure (a Google font the PDF itself uses).

    The calibrated size factors make a *sentence* as wide in Slides as in the PDF; a number is
    not a sentence - Lato's digits are tabular at 0.58 em where Computer Modern's are 0.5 - so a
    number column comes out ~14% wider than the PDF's, more than a column's slack."""
    total = 0.0
    for run in runs:
        family, size = fonts(run, scale)
        style = {(False, False): "regular", (True, False): "bold", (False, True): "italic",
                 (True, True): "bold_italic"}[(bool(run["bold"]), bool(run["italic"]))]
        if family == FONT_FOR_FAMILY["mono"]:
            table = {}
            unmeasured = ROBOTO_MONO_ADVANCE_EM
        elif family in ADVANCES:
            table, unmeasured = ADVANCES[family][style], UNMEASURED_ADVANCE_EM
        else:
            return None
        if run.get("script"):
            size *= SCRIPT_SIZE
        for ch in run["text"]:
            if run.get("smallcaps") and ch.islower():
                total += table.get(ch.upper(), unmeasured) * size * SMALL_CAPS_SIZE
            else:
                total += table.get(ch, wide_advance(ch, unmeasured)) * size
    return total


def wide_advance(ch: str, unmeasured: float) -> float:
    """The advance (em) of a character the probe did not measure: a CJK ideograph, kana or
    full-width form is a whole em in every fallback font (unicodedata's East Asian Width W / F).
    At 0.6 em a Japanese header came out 40% narrower than Slides sets it and wrapped its cell."""
    return 1.0 if unicodedata.east_asian_width(ch) in "WF" else unmeasured


def runs_between(runs: list[dict], a: int, b: int) -> list[dict]:
    """The runs' characters a to b (indices into their joined text)."""
    out, at = [], 0
    for run in runs:
        text = run["text"]
        lo, hi = max(a, at), min(b, at + len(text))
        if lo < hi:
            out.append({**run, "text": text[lo - at:hi - at]})
        at += len(text)
    return out


def wrapped_width(runs: list[dict], starts: list[int], scale: float, fonts: "FontMapper") -> float | None:
    """Slides width of the widest of a wrapped cell's PDF lines (`starts`: where each line after
    the first begins), a word TeX hyphenated at a line's end taken whole: Slides does not
    hyphenate, and a column that holds each of the PDF's lines so wraps the cell in as many lines
    or fewer - never more, which would grow its row."""
    text = "".join(r["text"] for r in runs)
    bounds = [0, *starts, len(text)]
    widths = []
    for a, b in zip(bounds, bounds[1:]):
        while b < len(text) and not text[b - 1].isspace() and not text[b].isspace():
            b += 1  # (the line ended inside a word)
        while b > a and text[b - 1].isspace():
            b -= 1
        widths.append(slides_width(runs_between(runs, a, b), scale, fonts))
    return None if None in widths else max(widths)


def wrap_joins(runs: list[dict], starts: list[int], scale: float, fonts: "FontMapper") -> float | None:
    """Slides width of the narrowest of a wrapped cell's PDF lines with the next line's first
    word joined to it: a column whose text room is less than that breaks the cell where TeX did
    (table_columns). None when a run is in a font the probe did not measure."""
    text = "".join(r["text"] for r in runs)
    bounds = [len(text) - len(text.lstrip()), *starts]
    joins = []
    for a, s in zip(bounds, bounds[1:]):
        b = s
        while b < len(text) and text[b].isspace():
            b += 1
        while b < len(text) and not text[b].isspace():
            b += 1  # (through the next word; a word TeX hyphenated is taken whole by wrapped_width)
        joins.append(slides_width(runs_between(runs, a, b), scale, fonts))
    return None if not joins or None in joins else min(joins)


def pdf_width(runs: list[dict]) -> float | None:
    """Width (PDF pt) TeX sets these Computer Modern runs at, from CM's own advances (within
    0.3% on one-line paragraphs: test_computer_modern_advances_give_the_pdf_its_own_line_widths),
    or None for a run in another font or with characters the tables lack."""
    total = 0.0
    for run in runs:
        face, info = cm_face(run), font_info(run["font"])
        if face is None or run.get("hole") or run.get("smallcaps"):
            return None
        slides = ADVANCES.get(FONT_FOR_FAMILY.get(run["family"], ""), ADVANCES["Lato"])["regular"]
        _, em, _, skipped = advance_widths(run["text"], CM_ADVANCES[face], slides)
        if skipped:
            return None
        total += em * design_width(DESIGN_WIDTH.get(run["family"], DESIGN_WIDTH["sans"]), info.design_size) * run["size"]
    return total


LINE_FIT_TOL = 0.06  # a PDF line may be this much wider than its words (a justified line's spaces)


def pdf_line_breaks(p: dict) -> list[int] | None:
    """Where each of a paragraph's PDF lines after the first starts (indices into its runs'
    joined text, at a word), found by setting its words with CM's advances into the lines'
    extents; None when the paragraph's words cannot be measured or do not come out as its lines
    (a hyphenated line end, a font without CM metrics)."""
    runs, lines = p["runs"], p["lines"]
    text = "".join(r["text"] for r in runs)
    if any(SOFT_BREAK in r["text"] or "\t" in r["text"] for r in runs):
        return None
    spaces = [i for i, ch in enumerate(text) if ch == " "]
    starts, a = [], 0
    while text[a:a + 1] == " ":
        a += 1
    for k, line in enumerate(lines):
        extent = line["x1"] - line["x0"]
        if k == len(lines) - 1:
            end = len(text.rstrip())
        else:
            end = None
            for b in (s for s in spaces if s > a):
                w = pdf_width(runs_between(runs, a, b))
                if w is None:
                    return None
                if w > extent * (1 + 0.01) + 0.5:
                    break
                end = b
            if end is None:
                return None
        w = pdf_width(runs_between(runs, a, end))
        if w is None or not extent * (1 - LINE_FIT_TOL) - 0.5 <= w <= extent * 1.01 + 0.5:
            return None
        if k < len(lines) - 1:
            a = end
            while text[a:a + 1] == " ":
                a += 1
            starts.append(a)
    return starts


def slides_lines(p: dict, scale: float, fonts: "FontMapper") -> tuple[float, float] | None:
    """(right edge of the widest line, the least right edge at which a line's next word would
    join it) in Slides pt, of a left-aligned paragraph's PDF lines as Slides sets their words
    (measured advances: slides_width); None when that is not known. A text box whose text ends
    between the two breaks the paragraph where TeX did."""
    runs, lines = p["runs"], p["lines"]
    text = "".join(r["text"] for r in runs)
    if any(r.get("hole") or SOFT_BREAK in r["text"] or "\t" in r["text"] for r in runs) or not text.strip():
        return None
    starts = [] if len(lines) == 1 else pdf_line_breaks(p)
    if starts is None:
        return None
    end = len(text.rstrip())
    bounds = [len(text) - len(text.lstrip()), *starts, end]
    widest, joins = 0.0, math.inf
    for k, (a, b) in enumerate(zip(bounds, bounds[1:])):
        x0 = lines[k]["x0"] * scale
        b = a + len(text[a:b].rstrip())  # (Slides lets a line's last space hang past the edge)
        w = slides_width(runs_between(runs, a, b), scale, fonts)
        if w is None:
            return None
        widest = max(widest, x0 + w)
        if k + 1 < len(lines):
            nxt = text.find(" ", bounds[k + 1])
            nxt = end if nxt < 0 or nxt > end else nxt
            j = slides_width(runs_between(runs, a, nxt), scale, fonts)
            if j is None:
                return None
            joins = min(joins, x0 + j)
    return widest, joins


def fit_columns(bounds: list[float], cols: list[dict], scale: float, need: list[float | None] | None = None,
                tight: bool = False, cap: list[float | None] | None = None) -> list[float]:
    """Column boundaries (PDF pt) moved just enough that every column's text fits inside the
    Slides cell padding, with room for the substitute font. Tables typeset with @{} have text
    touching the frame, which would otherwise wrap in Slides.

    need[i] is the width (PDF pt) column i's widest one-column cell takes in Slides
    (slides_width), when it is known. A cell that wraps in Slides doubles its row and pushes the
    table down over whatever stands under it - a caption - so where the PDF left less room
    between two columns than their text needs, the table grows sideways instead: the columns
    after move right.

    `tight`, for a table that would otherwise run off the page (table_layout): a measured column
    is as wide as its text in Slides and WRAP_MARGIN, lined up on its alignment edge - narrower
    than the PDF's when the text is set smaller.

    cap[i], for a left-aligned column holding a wrapped cell (table_columns): the text width
    (PDF pt) at which a line of that cell would take the next line's first word. The column's
    text room stays under it where its own lines allow, so the cell breaks where TeX did: joined
    into fewer lines than the PDF's, it left its row - which keeps the PDF's pitch - half empty."""
    pad = TABLE_CELL_PAD / scale
    need = need or [None] * len(cols)
    cap = cap or [None] * len(cols)
    if tight:
        def extent(c: dict, n: float | None) -> dict:
            if n is None:
                return c
            x0 = c["x0"] if c["align"] == "left" else c["x1"] - n if c["align"] == "right" else (c["x0"] + c["x1"] - n) / 2
            return {**c, "x0": x0, "x1": x0 + n}
        cols = [extent(c, n) for c, n in zip(cols, need)]
    # Text grows away from its alignment edge: room on the right of left-aligned columns, on
    # the left of right-aligned ones, half on each side of centred ones.
    width = [c["x1"] - c["x0"] for c in cols]
    room = [(1 + WRAP_MARGIN) / scale if tight and n is not None else
            max(0.08 * w + 1 / scale, (n - w + (1 + WRAP_MARGIN) / scale) if n is not None else 0.0)
            for w, n in zip(width, need)]
    least_room = [((n - w) if n is not None else 0.0) + (1 + WRAP_MARGIN) / scale for w, n in zip(width, need)]
    capped = [k is not None and c["align"] == "left" and k - w >= lo
              for c, w, k, lo in zip(cols, width, cap, least_room)]
    room = [min(r, k - w) if ok else r for r, w, k, ok in zip(room, width, cap, capped)]
    right =[r if c["align"] == "left" else r / 2 if c["align"] == "center" else 0.0 for c, r in zip(cols, room)]
    left = [r if c["align"] == "right" else r / 2 if c["align"] == "center" else 0.0 for c, r in zip(cols, room)]
    out = list(bounds)
    out[0] = min(out[0], cols[0]["x0"] - pad - left[0])
    out[-1] = max(out[-1], cols[-1]["x1"] + pad + right[-1])
    for i in range(1, len(cols)):
        lo, hi = cols[i - 1]["x1"] + pad + right[i - 1], cols[i]["x0"] - pad - left[i]
        out[i] = min(max(out[i], lo), hi) if lo <= hi else (lo + hi) / 2
    for i, ok in enumerate(capped):
        if ok:  # (the PDF's boundary can leave more room than the cap)
            out[i + 1] = min(out[i + 1], cols[i]["x1"] + pad + right[i])
    # Crowded columns: a column narrower than its text, its room and both paddings pushes every
    # boundary after it along (table_requests then caps the cell's alignment indent).
    for i in range(len(cols)):
        least = width[i] + room[i] + 2 * pad
        if out[i + 1] - out[i] < least:
            shift = least - (out[i + 1] - out[i])
            out[i + 1:] = [x + shift for x in out[i + 1:]]
    return out


def table_rows(el: dict, z: float, scale: float, imported: bool = False
               ) -> tuple[float, list[float], list[float], list[float]]:
    """Table top, row heights, per-row lineSpacing and per-row top cell inset (Slides pt).

    A row is at least its top and bottom insets + 1.195·z·lineSpacing tall (+ LINE_EM·z·lineSpacing
    per further line of a wrapped cell, `row_lines`). Row boundaries sit on the PDF's rules where
    there are any (booktabs puts extra space around them) and else just above the next row's text.

    An API-made table (`imported` False) has 7.2 pt insets above and below, which TeX's rows are
    too tight for (tools/probe_table_rows.py): the line spacing is tightened until rows keep the
    original pitch, and each row's lineSpacing then moves its baseline to the PDF's. A table the
    .pptx brings (`build_pptx`) has no inset below and a top inset of its own per row: the text
    keeps its natural line spacing and the inset moves the baseline down to the PDF's (booktabs'
    space under a rule), so a row keeps the PDF's pitch down to 1.195 em."""
    baselines = [b * scale for b in el["row_baselines"]]
    n = len(baselines)
    pitches = [h * scale for h in el["row_heights"]]
    lines = el.get("row_lines") or [1] * n
    pad_top = pad_bottom = 0.0 if imported else TABLE_ROW_PAD / 2
    default = 1.0 if imported else table_line_spacing(min(pitches), z)

    def offset(r: float) -> float:  # row top -> baseline, with the fixed top inset
        return pad_top + TABLE_TEXT_TOP + ASCENT_EM * z + extra_above(r, z)

    def body(i: int) -> float:  # the height of row i's text at lineSpacing 100
        return (TABLE_ROW_EM + (lines[i] - 1) * LINE_EM) * z

    ruled: dict[int, float] = {}
    for rule in el.get("rules", []) + [b for b in el.get("borders", []) if b["position"] in ("TOP", "BOTTOM")]:
        if "y" in rule:
            ruled.setdefault(rule["row"] + (rule["position"] == "BOTTOM"), rule["y"] * scale)
    def target(i: int) -> float:  # where row i should start
        if i in ruled:
            return ruled[i]
        if i < n:
            return baselines[i] - offset(default)
        return baselines[-1] - offset(default) + pitches[-1]

    def clamp(r: float) -> float:
        return min(1.0, max(TABLE_MIN_SPACING, r))

    # \multirow heads are centred in their rows (contentAlignment MIDDLE): no inset of their own.
    middle = {m["row"] for m in el.get("merges", []) if m["rows"] > 1}
    # Row by row from the actual top: a row's baseline offset and its minimum height both grow
    # with lineSpacing, so when the room above the text (from a rule) asks for more height than
    # the row has, the error is split between this baseline and the rows below.
    top = target(0)
    y, heights, ratios, insets = top, [], [], []
    for i in range(n):
        room = baselines[i] - y  # offset(r) = offset(1) - (1 - r)·0.9·z
        h_target = target(i + 1) - y
        inset = pad_top
        if imported and i not in middle:
            # The inset takes the room above the text, as far as the row's height allows.
            inset = max(0.0, min(room - offset(1.0), h_target - pad_bottom - body(i)))
        r_room = clamp(1.0 if room >= offset(1.0) + inset else 1 - (offset(1.0) + inset - room) / (0.75 * LINE_EM * z))
        r_fit = clamp((h_target - inset - pad_bottom) / body(i))
        r = r_room if r_room <= r_fit else (r_room + r_fit) / 2
        h = max(h_target, inset + pad_bottom + body(i) * r)
        ratios.append(r)
        heights.append(h)
        insets.append(inset)
        y += h
    return top, heights, ratios, insets


def table_columns(el: dict, cells: list[list[list[dict]]], scale: float, fonts: FontMapper,
                  tight: bool = False) -> tuple[list[float], dict]:
    """Column boundaries (PDF pt) of a table whose cells hold `cells`, and each cell's Slides
    width where it is known (slides_width, wrapped_width). `tight`: see fit_columns."""
    cols = el["columns"]
    fx0, _, fx1, _ = el["frame"]
    bounds = el.get("bounds") or [fx0] + [(a["x1"] + b["x0"]) / 2 for a, b in zip(cols, cols[1:])] + [fx1]
    spanned = {(m["row"], m["col"]) for m in el.get("merges", []) if m["cols"] > 1}
    # A cell set in a paragraph column (p{3cm}) wraps in Slides as in the PDF (classify
    # `wrapped`): it takes as much room as its widest line.
    wrapped = {(r, c): starts for r, c, starts in el.get("wrapped", [])}
    cell_width = {(r, c): wrapped_width(runs, wrapped[(r, c)], scale, fonts) if (r, c) in wrapped else
                  slides_width(runs, scale, fonts)
                  for r, row in enumerate(cells) for c, runs in enumerate(row) if runs}
    need: list[float | None] = []
    cap: list[float | None] = []
    for c, col in enumerate(cols):
        ws = [w for (r, cc), w in cell_width.items() if cc == c and (r, c) not in spanned]
        need.append(None if not ws or None in ws else max(ws) / scale)
        # A wrapped cell breaks where TeX did while no line has room for the next word (wrap_joins).
        # A font the probe did not measure is the PDF's own (a Google font it uses): TeX broke
        # each line because the next word overflowed its p{} width, at least the widest line.
        joins = [wrap_joins(cells[r][cc], starts, scale, fonts) for (r, cc), starts in wrapped.items()
                 if cc == c and (r, cc) not in spanned]
        w = col["x1"] - col["x0"]
        cap.append(None if not joins else
                   min(joins) / scale - 0.5 / scale if None not in joins else w + max((1 + WRAP_MARGIN) / scale, 0.02 * w))
    bounds = fit_columns(bounds, cols, scale, need, tight, cap)
    # A cell spanning columns wraps as readily as one that does not: the columns it spans grow.
    for m in el.get("merges", []):
        w = cell_width.get((m["row"], m["col"]))
        end = m["col"] + m["cols"]
        if m["cols"] > 1 and w is not None:
            short = (w + 2 * TABLE_CELL_PAD + 1 + WRAP_MARGIN) / scale - (bounds[end] - bounds[m["col"]])
            if short > 0:
                bounds[end:] = [x + short for x in bounds[end:]]
    return bounds, cell_width


# A table fit_columns widens past the page (eleven \scriptsize columns, each given room for the
# substitute and both cell paddings; a full-width tabularx) lost its last column off the slide.
# It may reach into the right margin by TABLE_MARGIN of the room left of it; past that, it first
# keeps only the room its measured text needs, then its text is set smaller - down to
# TABLE_MIN_SHRINK of its size - until it ends there (or, if the margin is out of reach, at the
# page edge). A table still on the page that would shrink by less than 2% keeps its size.
TABLE_MARGIN = 0.5
TABLE_MIN_SHRINK = 0.75
TABLE_KEEP_SIZE = 0.98


def table_layout(el: dict, scale: float, fonts: FontMapper, imported: bool = False,
                 page_w: float | None = None) -> dict:
    """Where a table goes in Slides (Slides pt): {"x", "y", "widths", "heights", "ratios" (per-row
    lineSpacing), "insets" (per-row top cell inset), "bounds" (column boundaries, PDF pt),
    "cell_width" ((row, col) -> the text's Slides width where known), "z", "first_run", "cells"
    (the runs as they are written: in_sentence, shrunk), "shrink" (the share of its size the text
    is set at)}. `imported`: the table comes with the .pptx (table_rows). `page_w`: the PDF page's
    width, by default that of a deck SLIDE_W wide (what every Slides page size is)."""
    page_w = SLIDE_W / scale if page_w is None else page_w
    # (`cell`: set at the table's size, not shaped to the PDF's width: FontMapper.shape_ratio)
    cells = [[[{**r, "cell": True} for r in in_sentence(runs)] for runs in row] for row in el["cells"]]
    bounds, cell_width = table_columns(el, cells, scale, fonts)
    shrink = 1.0
    # Into the right margin at most half as far as the table stands from the left edge (TABLE_MARGIN),
    # unless the PDF's own table reaches further.
    limit = min(page_w, max(page_w - TABLE_MARGIN * max(0.0, bounds[0]), el["frame"][2], (el.get("bounds") or [0.0])[-1]))
    if bounds[-1] > limit + 0.01:
        roomy = bounds, cell_width
        tight = table_columns(el, cells, scale, fonts, tight=True)

        def shrunk(s: float) -> list[list[list[dict]]]:
            return [[[{**r, "size": r["size"] * s} for r in runs] for runs in row] for row in cells]

        least = table_columns(el, shrunk(TABLE_MIN_SHRINK), scale, fonts, tight=True)
        # The margin if the smallest size reaches it, else the page edge; a table that does not
        # fit even then keeps its size (smaller words would not bring its last column back).
        goal = next((g for g in (limit, page_w) if least[0][-1] <= g + 0.01), None)
        best = None
        if tight[0][-1] > limit + 0.01 and goal is not None:
            lo, hi, best = TABLE_MIN_SHRINK, 1.0, (TABLE_MIN_SHRINK, least)
            for _ in range(12):
                mid = (lo + hi) / 2
                got = table_columns(el, shrunk(mid), scale, fonts, tight=True)
                if got[0][-1] <= goal + 0.01:
                    lo, best = mid, (mid, got)
                else:
                    hi = mid
        if tight[0][-1] <= limit + 0.01:
            bounds, cell_width = tight
        elif best and (best[0] < TABLE_KEEP_SIZE or roomy[0][-1] > page_w + 0.01):
            shrink, (bounds, cell_width) = best[0], best[1]
            cells = shrunk(shrink)
        else:
            # A table on the page that a smaller size would pull back only a little from the
            # margin keeps its size and its room (every size step is a residual for pull).
            bounds, cell_width = roomy if roomy[0][-1] <= page_w + 0.01 else tight
    widths = [max(TABLE_MIN_COLUMN_PT, (b - a) * scale) for a, b in zip(bounds, bounds[1:])]
    first_run = next((r for row in cells for cell in row for r in cell), None)
    z = fonts(first_run, scale)[1] if first_run else el["size"] * scale * shrink
    y, heights, ratios, insets = table_rows(el, z, scale, imported)
    return {"x": bounds[0] * scale, "y": y, "widths": widths, "heights": heights, "ratios": ratios, "insets": insets,
            "bounds": bounds, "cell_width": cell_width, "z": z, "first_run": first_run, "cells": cells,
            "shrink": round(shrink, 3)}


def pptx_table(el: dict, scale: float, fonts: FontMapper, page_w: float | None = None) -> dict:
    """The empty table the .pptx carries for a table element (build_pptx): its box, grid and
    per-row cell margins (left, top, right, bottom; Slides pt). The API fills it in (table_requests)."""
    lay = table_layout(el, scale, fonts, imported=True, page_w=page_w)
    return {"x": lay["x"], "y": lay["y"], "widths": lay["widths"], "heights": lay["heights"],
            "margins": [(TABLE_CELL_PAD, round(t, 2), TABLE_CELL_PAD, 0.0) for t in lay["insets"]]}


def table_requests(el: dict, slide_id: str, object_id: str, scale: float, fonts: FontMapper,
                   imported: bool = False, page_w: float | None = None) -> list[dict]:
    """A table filled in through the API. `imported`: the table (`object_id`) came with the .pptx,
    empty and with the cell margins `pptx_table` gave it; else it is made here by createTable."""
    cols = el["columns"]
    lay = table_layout(el, scale, fonts, imported, page_w)
    x, y, widths, heights, row_ratio = lay["x"], lay["y"], lay["widths"], lay["heights"], lay["ratios"]
    bounds, cell_width, first_run = lay["bounds"], lay["cell_width"], lay["first_run"]
    n_rows, n_cols = len(el["cells"]), len(cols)

    reqs: list[dict] = [
        # (it lies where the .pptx put it, below what the slide's elements before it made)
        {"updatePageElementsZOrder": {"pageElementObjectIds": [object_id], "operation": "BRING_TO_FRONT"}}
    ] if imported else [
        {"createTable": {"objectId": object_id, "rows": n_rows, "columns": n_cols, "elementProperties": {
            "pageObjectId": slide_id,
            "size": {"width": emu(sum(widths)), "height": emu(sum(heights))},
            "transform": {"scaleX": 1, "scaleY": 1, "unit": "EMU",
                          "translateX": round(x * EMU_PER_PT), "translateY": round(y * EMU_PER_PT)}}}},
    ]
    reqs += [
        # No grid: only the rules of the original are drawn.
        {"updateTableBorderProperties": {
            "objectId": object_id, "borderPosition": "ALL",
            "tableRange": {"location": {"rowIndex": 0, "columnIndex": 0}, "rowSpan": n_rows, "columnSpan": n_cols},
            "tableBorderProperties": {"tableBorderFill": {"solidFill": {"color": {"rgbColor": {}}, "alpha": 0}}},
            "fields": "tableBorderFill.solidFill.alpha"}},
    ]
    for i, w in enumerate(widths):
        reqs.append({"updateTableColumnProperties": {"objectId": object_id, "columnIndices": [i],
                                                     "tableColumnProperties": {"columnWidth": emu(w)},
                                                     "fields": "columnWidth"}})
    for i, h in enumerate(heights):
        reqs.append({"updateTableRowProperties": {"objectId": object_id, "rowIndices": [i],
                                                  "tableRowProperties": {"minRowHeight": emu(h)},
                                                  "fields": "minRowHeight"}})
    for rule in el["rules"]:
        reqs.append({"updateTableBorderProperties": {
            "objectId": object_id, "borderPosition": rule["position"],
            "tableRange": {"location": {"rowIndex": rule["row"], "columnIndex": 0}, "rowSpan": 1, "columnSpan": n_cols},
            "tableBorderProperties": {
                "tableBorderFill": {"solidFill": {"color": rgb(rule["color"])["opaqueColor"], "alpha": 1}},
                "weight": pt(round(max(0.5, rule["weight"] * scale), 2))},
            "fields": "tableBorderFill.solidFill.color,tableBorderFill.solidFill.alpha,weight"}})
    for b in el.get("borders", []):
        reqs.append({"updateTableBorderProperties": {
            "objectId": object_id, "borderPosition": b["position"],
            "tableRange": {"location": {"rowIndex": b["row"], "columnIndex": b["col"]}, "rowSpan": 1, "columnSpan": 1},
            "tableBorderProperties": {
                "tableBorderFill": {"solidFill": {"color": rgb(b["color"])["opaqueColor"], "alpha": 1}},
                "weight": pt(round(max(0.5, b["weight"] * scale), 2))},
            "fields": "tableBorderFill.solidFill.color,tableBorderFill.solidFill.alpha,weight"}})
    for f in el.get("fills", []):
        reqs.append({"updateTableCellProperties": {
            "objectId": object_id, "tableRange": {"location": {"rowIndex": f["row"], "columnIndex": f["col"]},
                                                  "rowSpan": 1, "columnSpan": 1},
            "tableCellProperties": {"tableCellBackgroundFill": {"solidFill": {"color": rgb(f["color"])["opaqueColor"]}}},
            "fields": "tableCellBackgroundFill.solidFill.color"}})
    merged = {(m["row"], m["col"]): m for m in el.get("merges", [])}
    for m in merged.values():
        reqs.append({"mergeTableCells": {"objectId": object_id, "tableRange": {
            "location": {"rowIndex": m["row"], "columnIndex": m["col"]}, "rowSpan": m["rows"], "columnSpan": m["cols"]}}})
        if m["rows"] > 1:  # \multirow centres its text vertically
            reqs.append({"updateTableCellProperties": {
                "objectId": object_id, "tableRange": {"location": {"rowIndex": m["row"], "columnIndex": m["col"]},
                                                      "rowSpan": m["rows"], "columnSpan": m["cols"]},
                "tableCellProperties": {"contentAlignment": "MIDDLE"}, "fields": "contentAlignment"}})

    hidden = {(m["row"] + i, m["col"] + j) for m in merged.values()
              for i in range(m["rows"]) for j in range(m["cols"])} - set(merged)
    # The indent puts the text where the PDF has it, but never so far that a cell's text no
    # longer fits on one line (a wrapped cell doubles its row): at most what the column's widest
    # cell leaves, the same for every cell of the column so that they stay aligned. Capped cell
    # by cell, a tight right-aligned column of signed numbers came out left-aligned. A width
    # nothing measured (a font the probe did not measure: Arial, Calibri) is the PDF's.
    spare = []
    for c, col in enumerate(cols):
        # (a head set its own way leaves the indent to its body: classify `head`, `body`)
        ex0, ex1 = col.get("body") or (col["x0"], col["x1"])
        ws = [(cell_width.get((r, c)) if cell_width.get((r, c)) is not None else (ex1 - ex0) * scale * lay["shrink"])
              for r, row in enumerate(el["cells"]) if c < len(row) and row[c] and (r, c) not in merged
              and (r, c) not in hidden and not (r == 0 and "head" in col)]
        spare.append(max(0.0, widths[c] - 2 * TABLE_CELL_PAD - WRAP_MARGIN - max(ws)) if ws else None)
    for r, row in enumerate(lay["cells"]):
        for c, runs in enumerate(row):
            text = "".join(run["text"] for run in runs).strip()
            loc = {"rowIndex": r, "columnIndex": c}
            if not text:
                if (r, c) in hidden or first_run is None:
                    continue
                # An empty cell still has a line of the default font, which would set the row's
                # minimum height: give it a space in the table's font and line spacing.
                style, fields = fonts.text_style(first_run, scale)
                reqs += [
                    {"insertText": {"objectId": object_id, "cellLocation": loc, "text": " "}},
                    {"updateTextStyle": {"objectId": object_id, "cellLocation": loc, "textRange": {"type": "ALL"},
                                         "style": style, "fields": ",".join(fields)}},
                    {"updateParagraphStyle": {"objectId": object_id, "cellLocation": loc, "textRange": {"type": "ALL"},
                                              "style": {"lineSpacing": round(100 * row_ratio[r], 1), "spaceAbove": pt(0),
                                                        "spaceBelow": pt(0)},
                                              "fields": "lineSpacing,spaceAbove,spaceBelow"}},
                ]
                continue
            reqs.append({"insertText": {"objectId": object_id, "cellLocation": loc, "text": text}})
            start = 0
            for run, z in zip(runs, run_sizes(runs, scale, fonts)):
                piece = run["text"].strip() if len(runs) == 1 else run["text"]
                if start == 0:
                    piece = piece.lstrip()
                if not piece:
                    continue
                style, fields = fonts.text_style(run, scale)
                if "fontSize" in style:
                    style["fontSize"] = pt(z)  # (a subscript no larger than its text: run_sizes)
                style.update({"smallCaps": run["smallcaps"], "foregroundColor": rgb(run["color"]),
                              "baselineOffset": {"super": "SUPERSCRIPT", "sub": "SUBSCRIPT"}.get(run.get("script"), "NONE")})
                reqs.append({"updateTextStyle": {
                    "objectId": object_id, "cellLocation": loc,  # (UTF-16 units: u16)
                    "textRange": {"type": "FIXED_RANGE", "startIndex": start, "endIndex": min(u16(text), start + u16(piece))},
                    "style": style, "fields": ",".join(fields + ["smallCaps", "foregroundColor", "baselineOffset"])}})
                start += u16(piece)
            col = cols[c]
            # A head set otherwise than its column's body (classify `head`): the head row its own
            # way, the body to the body's own edges.
            head = r == 0 and "head" in col
            align = col["head"] if head else col["align"]
            x0, x1 = (col["x0"], col["x1"]) if head or "body" not in col else col["body"]
            # Line the text up with the original inside the (contiguous) Slides columns.
            left_pad = max(0.0, (x0 - bounds[c]) * scale - PAD_X) if align == "left" else 0.0
            right_pad = max(0.0, (bounds[c + 1] - x1) * scale - PAD_X) if align == "right" else 0.0
            if (r, c) in merged and (merged[(r, c)]["cols"] > 1 or merged[(r, c)].get("align") == "center"):
                # (a \multirow or \makecell head centred over a flush column stays centred)
                align, left_pad, right_pad = merged[(r, c)]["align"], 0.0, 0.0
            if spare[c] is not None and (r, c) not in merged and not head:  # (the column's: see `spare` above)
                left_pad, right_pad = min(left_pad, spare[c]), min(right_pad, spare[c])
            # A cell that reads right to left starts at its right edge, so its alignment and
            # its two indents are mirrored (the text element's rule, one cell wide).
            rtl = bidi.reads_rtl(text)
            indent_start, indent_end = (right_pad, left_pad) if rtl else (left_pad, right_pad)
            reqs.append({"updateParagraphStyle": {
                "objectId": object_id, "cellLocation": loc, "textRange": {"type": "ALL"},
                "style": {"alignment": ({"left": "END", "center": "CENTER", "right": "START"} if rtl else
                                        {"left": "START", "center": "CENTER", "right": "END"})[align],
                          "lineSpacing": round(100 * row_ratio[r], 1), "spaceAbove": pt(0), "spaceBelow": pt(0),
                          "indentStart": pt(round(indent_start, 2)), "indentFirstLine": pt(round(indent_start, 2)),
                          "indentEnd": pt(round(indent_end, 2)),
                          **({"direction": "RIGHT_TO_LEFT"} if rtl else {})},
                "fields": "alignment,lineSpacing,spaceAbove,spaceBelow,indentStart,indentFirstLine,indentEnd" +
                          (",direction" if rtl else "")}})
    return reqs


# Share of a preset shape's width its text may use: Slides lays text out in the shape's .pptx
# text rectangle (an ellipse's is its inscribed square, a diamond's half its width).
TEXT_RECT_WIDTH = {"RECTANGLE": 1.0, "ROUND_RECTANGLE": 0.9, "ELLIPSE": 0.707, "DIAMOND": 0.5}
# Connection sites of preset shapes in their .pptx order (fractions of the box): a line connected
# to a site follows the shape when it is moved in Slides.
CONNECTION_SITES = {
    "RECTANGLE": [(0.5, 0), (0, 0.5), (0.5, 1), (1, 0.5)],
    "ROUND_RECTANGLE": [(0.5, 0), (0, 0.5), (0.5, 1), (1, 0.5)],
    "DIAMOND": [(0.5, 0), (0, 0.5), (0.5, 1), (1, 0.5)],
    "ELLIPSE": [(0.5, 0), (0.1464, 0.1464), (0, 0.5), (0.1464, 0.8536), (0.5, 1), (0.8536, 0.8536), (1, 0.5), (0.8536, 0.1464)],
    "TRIANGLE": [(0.5, 0), (0.25, 0.5), (0, 1), (0.5, 1), (1, 1), (0.75, 0.5)],
}
LABEL_ROOM = 1.08  # the substitute font may run this much wider


def label_inside(node: dict) -> bool:
    """A node's label goes into the node shape itself (it then moves and resizes with it) when
    it fits the shape's text rectangle without wrapping."""
    x0, _, x1, _ = node["bbox"]
    text = "".join(r["text"] for runs in node["paragraphs"] for r in runs).strip()
    return bool(text) and node["shape"] in TEXT_RECT_WIDTH and \
        node.get("label_w", 0.0) * LABEL_ROOM <= (x1 - x0) * TEXT_RECT_WIDTH[node["shape"]] - 0.5


def node_template_key(node: dict) -> tuple:
    return node["shape"], None, None


def bend_template_key(line: dict) -> tuple:
    """An elbow connector (bentConnector3) turning at its start (|-, adj 0) or its end (-|)."""
    return "BENT_CONNECTOR", 0.0 if line["bend"] == "vh" else 1.0, None


def element_template_keys(el: dict, scale: float) -> list[tuple]:
    if el["kind"] == "diagram":
        return [node_template_key(n) for n in el["nodes"] if n["shape"] and label_inside(n)] + \
               [bend_template_key(ln) for ln in el["lines"] if ln.get("bend")]
    key = template_key(el, scale)
    return [key] if key else []


def connection(point: list[float], nodes: list[dict], oids: list[str]) -> dict | None:
    """The node connection site a line end sits on (PDF pt, within 1.5 pt), if any."""
    best = None
    for node, oid in zip(nodes, oids):
        if not node["shape"] or node["shape"] not in CONNECTION_SITES:
            continue
        x0, y0, x1, y1 = node["bbox"]
        for index, (fx, fy) in enumerate(CONNECTION_SITES[node["shape"]]):
            d = math.hypot(point[0] - (x0 + fx * (x1 - x0)), point[1] - (y0 + fy * (y1 - y0)))
            if d <= 1.5 and (best is None or d < best[0]):
                best = (d, {"connectedObjectId": oid, "connectionSiteIndex": index})
    return best[1] if best else None


def arrow_style(arrow) -> str:
    if not arrow:
        return "NONE"
    return arrow if isinstance(arrow, str) else "OPEN_ARROW"


def diagram_requests(el: dict, slide_id: str, object_id: str, scale: float, fonts: FontMapper,
                     template=None) -> list[dict]:
    """Nodes become shapes, edges become lines with arrow heads; the parts are grouped so the
    diagram moves as one piece but stays editable. A label that fits goes inside its node (a
    template shape without text padding, see label_inside); one that doesn't gets a text box
    grouped with its node. Line ends on a node's connection site are connected to it, so edges
    follow nodes moved in Slides. `template(key)` gives this slide's template shape for a key."""
    reqs: list[dict] = []
    children: list[str] = []
    node_oids = [f"{object_id}_n{j}" for j in range(len(el["nodes"]))]
    segments = []  # (object id, from, to, line): elbows without a template fall back to two lines
    for j, ln in enumerate(el["lines"]):
        if ln.get("bend") and template is None:
            segments += [(f"{object_id}_l{j}", ln["from"], ln["via"], {**ln, "arrow_to": None}),
                         (f"{object_id}_l{j}b", ln["via"], ln["to"], {**ln, "arrow_from": None})]
        else:
            segments.append((f"{object_id}_l{j}", ln["from"], ln["to"], ln))
    for oid, (x1, y1), (x2, y2), ln in segments:
        dx, dy = (x2 - x1) * scale, (y2 - y1) * scale
        if ln.get("bend") and template is not None:
            tpl = template(bend_template_key(ln))
            reqs += [
                {"duplicateObject": {"objectId": tpl["id"], "objectIds": {tpl["id"]: oid}}},
                # Like a straight line: from the transform origin along +size, flipped by negative scales.
                {"updatePageElementTransform": {"objectId": oid, "applyMode": "ABSOLUTE", "transform": {
                    "scaleX": dx / tpl["w"], "scaleY": dy / tpl["h"], "unit": "EMU",
                    "translateX": round(x1 * scale * EMU_PER_PT), "translateY": round(y1 * scale * EMU_PER_PT)}}},
                {"updatePageElementsZOrder": {"pageElementObjectIds": [oid], "operation": "BRING_TO_FRONT"}},
            ]
        else:
            reqs.append({"createLine": {"objectId": oid, "lineCategory": "STRAIGHT", "elementProperties": {
                "pageObjectId": slide_id,
                "size": {"width": emu(abs(dx)), "height": emu(abs(dy))},
                # The line runs from the transform origin along +size, flipped by negative scales.
                "transform": {"scaleX": -1 if dx < 0 else 1, "scaleY": -1 if dy < 0 else 1, "unit": "EMU",
                              "translateX": round(x1 * scale * EMU_PER_PT), "translateY": round(y1 * scale * EMU_PER_PT)}}}})
        reqs.append({"updateLineProperties": {"objectId": oid, "fields": "lineFill.solidFill.color,weight,startArrow,endArrow",
                                              "lineProperties": {
                                                  "lineFill": {"solidFill": {"color": rgb(ln["stroke"])["opaqueColor"]}},
                                                  "weight": pt(round(max(0.5, ln["width"] * scale), 2)),
                                                  "startArrow": arrow_style(ln["arrow_from"]),
                                                  "endArrow": arrow_style(ln["arrow_to"])}}})
        children.append(oid)
    for j, node in enumerate(el["nodes"]):
        oid = node_oids[j]
        x0, y0, x1, y1 = (v * scale for v in node["bbox"])
        text = "\n".join("".join(r["text"] for r in runs).strip() for runs in node["paragraphs"])
        card = node.get("text")  # a card's text: a text box on the PDF baselines (classify.card_text)
        inside = not card and bool(node["shape"]) and label_inside(node) and template is not None
        props = None if node["shape"] is None else {"contentAlignment": "MIDDLE", "autofit": {"autofitType": "NONE"},
                 "shapeBackgroundFill": ({"solidFill": {"color": rgb(node["fill"])["opaqueColor"]}} if node["fill"]
                                         else {"propertyState": "NOT_RENDERED"}),
                 "outline": ({"outlineFill": {"solidFill": {"color": rgb(node["stroke"])["opaqueColor"]}},
                              "weight": pt(round(max(0.5, (node["width"] or 0.4) * scale), 2))}
                             if node["stroke"] else {"propertyState": "NOT_RENDERED"})}
        members = []
        if props:  # free labels (edge labels, captions) have no shape, only the text box below
            fields = ["contentAlignment", "autofit.autofitType", "shapeBackgroundFill"]
            fields += ["outline.outlineFill.solidFill.color", "outline.weight"] if node["stroke"] else ["outline.propertyState"]
            if not node["fill"]:
                fields[fields.index("shapeBackgroundFill")] = "shapeBackgroundFill.propertyState"
            else:
                fields[fields.index("shapeBackgroundFill")] = "shapeBackgroundFill.solidFill.color"
            if inside:
                tpl = template(node_template_key(node))
                reqs += [
                    {"duplicateObject": {"objectId": tpl["id"], "objectIds": {tpl["id"]: oid}}},
                    {"updatePageElementTransform": {"objectId": oid, "applyMode": "ABSOLUTE", "transform": {
                        "scaleX": (x1 - x0) / tpl["w"], "scaleY": (y1 - y0) / tpl["h"], "unit": "EMU",
                        "translateX": round(x0 * EMU_PER_PT), "translateY": round(y0 * EMU_PER_PT)}}},
                    {"updatePageElementsZOrder": {"pageElementObjectIds": [oid], "operation": "BRING_TO_FRONT"}},
                ]
            else:
                reqs.append({"createShape": {"objectId": oid, "shapeType": node["shape"], "elementProperties": {
                    "pageObjectId": slide_id, "size": {"width": emu(x1 - x0), "height": emu(y1 - y0)},
                    "transform": {"scaleX": 1, "scaleY": 1, "unit": "EMU",
                                  "translateX": round(x0 * EMU_PER_PT), "translateY": round(y0 * EMU_PER_PT)}}}})
            reqs.append({"updateShapeProperties": {"objectId": oid, "shapeProperties": props, "fields": ",".join(fields)}})
            members.append(oid)
        if card:
            for k, box in enumerate(card):
                label = f"{object_id}_x{j}" + (f"_{k}" if k else "")
                reqs += text_box_requests(box, slide_id, label, scale, fonts)
                members.append(label)
        elif text:
            if not inside:
                # A label wider than the node's text rectangle would wrap inside the shape: it
                # gets its own wider text box, centred on the node and grouped with it.
                label = f"{object_id}_x{j}"
                cx, w = (x0 + x1) / 2, (x1 - x0) + 2 * PAD_X + 40
                reqs += [
                    {"createShape": {"objectId": label, "shapeType": "TEXT_BOX", "elementProperties": {
                        "pageObjectId": slide_id, "size": {"width": emu(w), "height": emu(y1 - y0)},
                        "transform": {"scaleX": 1, "scaleY": 1, "unit": "EMU",
                                      "translateX": round((cx - w / 2) * EMU_PER_PT), "translateY": round(y0 * EMU_PER_PT)}}}},
                    {"updateShapeProperties": {"objectId": label, "fields": "contentAlignment,autofit.autofitType",
                                               "shapeProperties": {"contentAlignment": "MIDDLE",
                                                                   "autofit": {"autofitType": "NONE"}}}},
                ]
                members.append(label)
            target = oid if inside else label
            reqs.append({"insertText": {"objectId": target, "text": text}})
            start = 0
            for runs in map(in_sentence, node["paragraphs"]):
                line_text = "".join(r["text"] for r in runs).strip()
                offset = 0
                for run in runs:
                    piece = run["text"].strip() if len(runs) == 1 else run["text"]
                    if offset == 0:
                        piece = piece.lstrip()
                    if not piece:
                        continue
                    style, sfields = fonts.text_style(run, scale)
                    style["foregroundColor"] = rgb(run["color"])
                    reqs.append({"updateTextStyle": {  # (UTF-16 units: u16)
                        "objectId": target, "style": style, "fields": ",".join(sfields + ["foregroundColor"]),
                        "textRange": {"type": "FIXED_RANGE", "startIndex": start + offset,
                                      "endIndex": min(start + u16(line_text), start + offset + u16(piece))}}})
                    offset += u16(piece)
                start += u16(line_text) + 1
            reqs.append({"updateParagraphStyle": {
                "objectId": target, "textRange": {"type": "ALL"}, "fields": "alignment,lineSpacing,spaceAbove,spaceBelow",
                "style": {"alignment": "CENTER", "lineSpacing": 100, "spaceAbove": pt(0), "spaceBelow": pt(0)}}})
        if len(members) >= 2:  # a node with its label (or card texts) outside: they move together
            reqs.append({"groupObjects": {"groupObjectId": f"{object_id}_g{j}", "childrenObjectIds": members}})
            members = [f"{object_id}_g{j}"]
        children += members
    # Edges follow the nodes they start or end on.
    for oid, start, end, ln in segments:
        ends = {"startConnection": connection(start, el["nodes"], node_oids) if start == ln["from"] else None,
                "endConnection": connection(end, el["nodes"], node_oids) if end == ln["to"] else None}
        ends = {k: v for k, v in ends.items() if v}
        if ends:
            reqs.append({"updateLineProperties": {"objectId": oid, "fields": ",".join(ends), "lineProperties": ends}})
    if len(children) >= 2:
        reqs.append({"groupObjects": {"groupObjectId": object_id, "childrenObjectIds": children}})
    return reqs


def block_groups(elements: list[dict], object_ids: list[str], title_oid: str | None) -> list[list[str]]:
    """Object ids per block: its panel shapes (title bar and body, see classify.blocks), plus
    the text and pictures lying on them."""
    blocks: dict[int, list] = {}  # block -> [x0, y0, x1, y1, [oids]]
    for el, oid in zip(elements, object_ids):
        if el["kind"] == "shape" and el.get("block") is not None:
            x0, y0, x1, y1 = el["bbox"]
            b = blocks.setdefault(el["block"], [x0, y0, x1, y1, []])
            b[:4] = [min(b[0], x0), min(b[1], y0), max(b[2], x1), max(b[3], y1)]
            b[4].append(oid)
    out = []
    for x0, y0, x1, y1, members in blocks.values():
        if len(members) < 2:
            continue  # a lone panel is not recognisably a block
        for el, oid in zip(elements, object_ids):
            # Tables can't be grouped in Slides: a table in a block stays on its own.
            if el["kind"] in ("text", "image") and oid != title_oid and not el.get("anchor"):
                ex0, ey0, ex1, ey1 = el["bbox"]
                cx, cy = (ex0 + ex1) / 2, (ey0 + ey1) / 2
                if x0 <= cx <= x1 and y0 <= cy <= y1:
                    members.append(oid)
        out.append(members)
    return out


def rule_groups(elements: list[dict], object_ids: list[str]) -> list[list[str]]:
    """Rules lying on one another (a progress bar on its track) move as one."""
    rules = [(el["bbox"], oid) for el, oid in zip(elements, object_ids) if el["kind"] == "shape" and el.get("role") == "rule"]
    out: list[list] = []  # [bbox, [oids]]
    for (x0, y0, x1, y1), oid in rules:
        for g in out:
            gx0, gy0, gx1, gy1 = g[0]
            if x0 < gx1 and gx0 < x1 and y0 < gy1 and gy0 < y1:
                g[0] = [min(x0, gx0), min(y0, gy0), max(x1, gx1), max(y1, gy1)]
                g[1].append(oid)
                break
        else:
            out.append([[x0, y0, x1, y1], [oid]])
    return [g[1] for g in out if len(g[1]) >= 2]


def text_right_limit(el: dict, slide: dict) -> float | None:
    """How far right (PDF x) a text element's box may reach: inside a panel (a block body),
    as far from the panel's right edge as the text is from its left edge; elsewhere the
    mirrored left margin of the page, stopping short of anything to the right on the same
    lines (the other column, a picture)."""
    if el.get("role") not in ("body", "title", None) or not el["paragraphs"] or el.get("rotation"):
        return None
    x0, y0, x1, y1 = el["bbox"]
    cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
    panels = [e["bbox"] for e in slide["elements"] if e["kind"] == "shape" and e.get("role") == "panel"
              and e["bbox"][0] <= cx <= e["bbox"][2] and e["bbox"][1] <= cy <= e["bbox"][3]]
    if panels:
        px0, _, px1, _ = min(panels, key=lambda b: (b[2] - b[0]) * (b[3] - b[1]))
        limit = px1 - max(x0 - px0, 2.0)
    else:
        margin = min((e["bbox"][0] for e in slide["elements"] if e["kind"] == "text" and e.get("role") in ("body", None)
                      and not e.get("rotation")),
                     default=x0)
        limit = slide["size"][0] - margin
        size = el["paragraphs"][0]["size"]
        for o in slide["elements"]:
            if o is el or o["kind"] == "shape":
                continue
            ox0, oy0, ox1, oy1 = o["bbox"]
            if ox0 >= x1 - 1 and oy0 < y1 and oy1 > y0:
                limit = min(limit, ox0 - size)
    return limit if limit > x1 else None


def earlier_holes(p: dict, run: dict) -> list[tuple[float, float]]:
    """(x0, width) of the holes before `run` on its line (PDF pt)."""
    words = [b[6] for b in run.get("before", []) if len(b) >= 7]
    if not words:
        return []
    i = next(k for k, r in enumerate(p["runs"]) if r is run)
    return [(r["hole_x0"], r["hole"]) for r in p["runs"][:i]
            if r.get("hole") and min(words) <= r["hole_x0"] < run["hole_x0"]]


def space_shift(run: dict, em: float, holes: list[tuple[float, float]] = ()) -> float:
    """What word spaces add to a formula gap's position in Slides. The substitute's size
    calibration makes its glyphs wider and its spaces narrower than TeX's, evening out at word
    ends; a formula starts after a space, so it lands one space deficit early. TeX also
    stretches some spaces (after a colon, around math): Slides sets a plain space there.
    A gap holding an earlier hole (`holes`) becomes a space and that hole's no-break spaces."""
    if google_font(run["font"]):
        return 0.0  # the PDF's own font: its spaces too
    words = sorted((b[6], b[6] + b[0]) for b in run.get("before", []) if len(b) >= 7)
    if not words:
        return 0.0
    size, space = run["size"], SYMBOL_ADVANCE_EM[" "] * em
    ends = [(a[1], b[0]) for a, b in zip(words, words[1:])] + [(words[-1][1], run["hole_x0"])]
    shift, gaps = 0.0, []
    for a, b in ends:
        inside = [w for x, w in holes if a - 0.5 <= x < b]
        shift += sum(space + w for w in inside) - (b - a) if inside else 0.0
        gaps.append(None if inside else b - a)
    spaces = [g for g in gaps if g is not None and g > 0.15 * size]
    nominal = min(sorted(spaces)[len(spaces) // 2], 0.4 * size) if spaces else 0.0
    shift += sum(nominal - g for g in spaces)
    if gaps[-1] is None:
        shift += space  # (the gap's own PDF space went with the hole before)
    elif gaps[-1] > 0.15 * size:
        shift += space - nominal
    return shift


def slide_holes(slide: dict) -> list[tuple[dict, dict, dict, dict | None]]:
    """(text element, paragraph, hole run, its picture) for every hole on a slide, in text order.
    A picture covers most of its hole (it usually starts where the hole does, but a radical's
    sign can reach further left) on one of the paragraph's lines; holes on two lines can have
    the same x: each takes the nearest picture not taken yet."""
    pictures = [e for e in slide["elements"] if e["kind"] == "image" and e.get("anchor")]
    out, taken = [], set()
    for el in slide["elements"]:
        if el["kind"] != "text":
            continue
        for p in el["paragraphs"]:
            for run in p["runs"]:
                if not run.get("hole"):
                    continue
                a, b = run["hole_x0"] - HOLE_PAD, run["hole_x0"] - HOLE_PAD + run["hole"]
                near = [e for e in pictures if e["anchor"] == el["id"] and e["id"] not in taken
                        and min(b, e["bbox"][2]) - max(a, e["bbox"][0]) > 0.5 * min(b - a, e["bbox"][2] - e["bbox"][0])]
                pic = min(near, default=None, key=lambda e: (min(
                    abs((e["bbox"][1] + e["bbox"][3]) / 2 - line["baseline"] + 0.35 * run["size"]) for line in p["lines"]) // 4,
                    abs(e["bbox"][0] + HOLE_PAD - run["hole_x0"])))
                if pic:
                    taken.add(pic["id"])
                out.append((el, p, run, pic))
    return out


def hole_neighbours(p: dict, run: dict) -> tuple[float | None, float | None]:
    """PDF x where the word before a hole ends, if a space separates them (Slides has a word
    space there too), and where the word after it starts, if nothing separates them in the
    text (the hole reaches up to it)."""
    i = next(k for k, r in enumerate(p["runs"]) if r is run)
    words = [b[6] + b[0] for b in run.get("before", []) if len(b) >= 7]
    prev_end = None
    if i > 0 and p["runs"][i - 1]["text"].endswith(" ") and words and \
            not any(x >= max(words) for x, _ in earlier_holes(p, run)):
        prev_end = max(words)
    nxt = p["runs"][i + 1]["text"] if i + 1 < len(p["runs"]) else ""
    return prev_end, (run.get("next_x0") if not nxt[:1].isspace() else None)


def fit_holes(slide: dict, scale: float, fonts: FontMapper) -> dict:
    """The slide with each hole as wide as the PDF's room between the words around it, less the
    Slides word space before it, so the words after it keep their place. Never narrower than
    the picture."""
    widths = {}
    for _, p, run, pic in slide_holes(slide):
        prev_end, next_x0 = hole_neighbours(p, run)
        if pic is None:
            continue
        if next_x0 is None:
            # (render grows a formula picture to its glyphs' ink: an integral's overhang)
            if pic["bbox"][2] - pic["bbox"][0] > run["hole"]:
                widths[id(run)] = round(pic["bbox"][2] - pic["bbox"][0], 2)
            continue
        start = prev_end + SYMBOL_ADVANCE_EM[" "] * fonts(run, scale)[1] / scale if prev_end is not None else pic["bbox"][0]
        widths[id(run)] = round(max(pic["bbox"][2] - pic["bbox"][0], next_x0 - start), 2)
    if not widths:
        return slide
    return {**slide, "elements": [
        {**el, "paragraphs": [{**p, "runs": [{**r, "hole": widths[id(r)]} if id(r) in widths else r for r in p["runs"]]}
                              for p in el["paragraphs"]]} if el["kind"] == "text" else el
        for el in slide["elements"]]}


def hole_offset(p: dict, run: dict, pic: dict, scale: float, z: float) -> float:
    """Where a picture sits in its gap (slide pt from the gap's start): the Slides word space
    before the gap and the gap's room after the picture are shared in the PDF's proportion."""
    prev_end, next_x0 = hole_neighbours(p, run)
    if prev_end is None:
        return 0.0
    x0, _, x1, _ = pic["bbox"]
    room = max(0.0, run["hole"] - (x1 - x0)) * scale
    left = max(0.0, x0 - prev_end) * scale
    right = max(0.0, next_x0 - x1) * scale if next_x0 is not None else room
    space = SYMBOL_ADVANCE_EM[" "] * z
    return (space + room) * left / (left + right) - space if left + right > 0 else 0.0


def formula_shifts(slide: dict, scale: float, fonts: FontMapper) -> dict[str, float]:
    """PDF-point x offsets for inline formula pictures, so each sits over the gap where Slides
    will put it: the words before it on its line come out a little narrower or wider."""
    out = {}
    for el, p, run, pic in slide_holes(slide):
        if p["align"] != "left" or pic is None:
            continue
        em = fonts(run, scale)[1] / scale  # the line's Slides font size, in PDF points
        shift = HOLE_PAD  # the gap has room for the picture's padding on its left too
        for w, font, family, bold, italic, *text in run.get("before", []):
            if family == "math" and text:
                # Math symbols come from Lato or Slides' fallback fonts, at their own widths.
                shift += sum(SYMBOL_ADVANCE_EM.get(ch, 0.55) for ch in text[0]) * em - w
            else:
                spaces = len(text[0]) - len(text[0].strip()) if text else 0
                # A math space inside the span (" 2," after ≥): TeX's thick space is wider.
                pdf_spaces = spaces * MATH_SPACE_EM * run["size"]
                shift += (w - pdf_spaces) * (fonts.width_ratio(font, family, bold, italic) - 1) + \
                    spaces * SYMBOL_ADVANCE_EM[" "] * em - pdf_spaces
        shift += space_shift(run, em, earlier_holes(p, run)) + hole_offset(p, run, pic, scale, em * scale) / scale
        if abs(shift) >= 0.2:
            out[pic["id"]] = shift
    return out


OVERLAY_STRETCH = 0.1  # how much a graphic between words may widen or narrow with predicted words
OVERLAY_STRETCH_MEASURED = 0.3  # ... with measured ones (a brace spans its phrase)


def mark_drifts(el: dict, scale: float, fonts: FontMapper) -> list[float]:
    """Predicted drift (PDF pt, Slides minus PDF) of each of an overlay's marks: a formula gap
    starting at the mark (formula_shifts)."""
    drifts = []
    for m in el["marks"]:
        run = {"text": " ", "font": m["font"], "family": m["family"], "size": m["size"], "bold": m["bold"],
               "italic": m["italic"], "hole": 1.0, "hole_x0": m["hole_x0"], "before": m["before"]}
        probe = {"elements": [{"kind": "text", "id": "mark", "paragraphs": [
            {"align": "left", "runs": [run], "lines": [{"baseline": 0.35 * m["size"]}]}]},
                              {"kind": "image", "id": "gap", "anchor": "mark",
                               "bbox": [m["hole_x0"] - HOLE_PAD, 0, m["hole_x0"], 0]}]}
        # (the gap's picture starts HOLE_PAD before the gap; shifts under 0.2 pt are not reported)
        drifts.append(formula_shifts(probe, scale, fonts).get("gap", HOLE_PAD) - HOLE_PAD + m.get("pads", 0))
    return drifts


def fit_overlay(bbox: list[float], drifts: list[tuple[float, float]], cap: float) -> tuple[float, float]:
    """PDF x extent of a picture following its marks' drifts ((x, drift)): moved by the mean
    drift and stretched by the least-squares slope (at most `cap`) when the marks are 10 pt apart or more."""
    xs = [x for x, _ in drifts]
    mx, md = sum(xs) / len(xs), sum(d for _, d in drifts) / len(drifts)
    spread = sum((x - mx) ** 2 for x in xs)
    slope = sum((x - mx) * (d - md) for x, d in drifts) / spread if max(xs) - min(xs) >= 10 else 0.0
    slope = max(-cap, min(cap, slope))
    x0, _, x1, _ = bbox
    return x0 + md + slope * (x0 - mx), x1 + md + slope * (x1 - mx)


def overlay_boxes(slide: dict, scale: float, fonts: FontMapper) -> dict[str, tuple[float, float]]:
    """PDF x extents for graphics drawn over or at words (classify.overlay: arrows, braces,
    callouts), following where Slides sets those words, as predicted (mark_drifts); measure_places
    corrects them on the thumbnail."""
    return {el["id"]: fit_overlay(el["bbox"], list(zip([m["x"] for m in el["marks"]], mark_drifts(el, scale, fonts))),
                                  OVERLAY_STRETCH)
            for el in slide["elements"] if el.get("marks")}


# ---------------------------------------------------------------- measured picture positions
#
# The prediction above is off by several points now and then (fallback fonts, kerning, wraps).
# So each slide with holes or overlays gets a scratch copy of those text boxes on a white slide,
# every hole run and every word an overlay's marks lie on highlighted in a mark colour and all
# text black; Google's thumbnail of it shows where the gaps and words really are, and the
# pictures are moved (overlays also stretched) there before they are grouped.

HOLE_MARKS = ["#ff00ff", "#00ffff", "#ffff00", "#00ff00"]  # 0/255 channels only (mark_alpha)
MARK_CORE = 0.9  # a pixel at least this much covered by a mark is inside it


def mark_alpha(img: np.ndarray, color: str) -> np.ndarray:
    """How much of each pixel of an RGB thumbnail a highlight in `color` covers (white page).
    Dark glyph pixels have the colour's full channels low too: they count as uncovered."""
    c = [int(color.lstrip("#")[i:i + 2], 16) for i in (0, 2, 4)]
    img = img.astype(np.float32)
    full = np.min([img[..., i] for i in range(3) if c[i] == 255], axis=0)
    alpha = np.mean([(255 - img[..., i]) / 255 for i in range(3) if c[i] == 0], axis=0)
    return np.where(full >= 200, np.clip(alpha, 0.0, 1.0), 0.0)


def _runs(mask: np.ndarray) -> list[tuple[int, int]]:
    """Index ranges [a, b] of consecutive True values."""
    idx = np.flatnonzero(mask)
    if idx.size == 0:
        return []
    cuts = np.flatnonzero(np.diff(idx) > 1)
    return [(int(g[0]), int(g[-1])) for g in np.split(idx, cuts + 1)]


def _edges(profile: np.ndarray, a: int, b: int) -> tuple[float, float]:
    """Sub-pixel extent of a covered range [a, b]: partly covered neighbours add their share."""
    lo = a - (profile[a - 1] if a > 0 else 0.0)
    hi = b + 1 + (profile[b + 1] if b + 1 < profile.size else 0.0)
    return float(lo), float(hi)


def find_marks(alpha: np.ndarray, px_per_pt: float) -> list[tuple[float, float, float, float]]:
    """Highlighted rectangles (x0, y0, x1, y1 in pt) in a mark_alpha map. A glyph lying over a
    highlight hides it only in some rows, so columns count as covered if any row of the band is."""
    out = []
    core = alpha >= MARK_CORE
    for r0, r1 in _runs(core.sum(axis=1) >= 3):
        cols = alpha[r0:r1 + 1].max(axis=0)
        for c0, c1 in _runs(cols >= MARK_CORE):
            if c1 - c0 < 2:
                continue
            rows = alpha[:, c0:c1 + 1].max(axis=1)
            x0, x1 = _edges(cols, c0, c1)
            y0, y1 = _edges(rows, r0, r1)
            out.append((x0 / px_per_pt, y0 / px_per_pt, x1 / px_per_pt, y1 / px_per_pt))
    return out


def pick_gap(marks: list[tuple[float, float, float, float]], x0: float, cy: float, width: float,
             pitch: float) -> tuple[float, float] | None:
    """(dx, dy) from the predicted gap start `x0` and line middle `cy` (slide pt) to the mark as
    wide as the hole that lies nearest; dy is whole line pitches (a line Slides wrapped
    differently), 0 on the predicted line."""
    fits = [m for m in marks if abs((m[2] - m[0]) - width) <= max(1.5, 0.12 * width)]
    if not fits:
        return None
    m = min(fits, key=lambda m: abs(m[0] - x0) + abs((m[1] + m[3]) / 2 - cy))
    lines = round(((m[1] + m[3]) / 2 - cy) / pitch)
    return m[0] - x0, lines * pitch


def ink_end(img: np.ndarray, px_per_pt: float, y0: float, y1: float, x0: float, x1: float) -> float | None:
    """Right end (pt) of the dark text pixels between rows y0..y1 and columns x0..x1 (pt)."""
    h, w = img.shape[:2]
    a, b = max(0, int(x0 * px_per_pt)), min(w, int(math.ceil(x1 * px_per_pt)))
    band = img[max(0, int(y0 * px_per_pt)):min(h, int(math.ceil(y1 * px_per_pt))), a:b]
    cols = np.flatnonzero((band.max(axis=2) < 110).any(axis=0)) if band.size else []
    return (a + cols[-1] + 1) / px_per_pt if len(cols) else None


def slides_texts(el: dict, scale: float, fonts: FontMapper) -> list[str]:
    """Each paragraph's text as its text box ends up holding it (holes as their no-break spaces)."""
    return ["".join((hole_run(r, scale, fonts) if r.get("hole") else r)["text"] for r in p["runs"]) for p in el["paragraphs"]]


def mark_words(overlay: dict, text: dict, scale: float, fonts: FontMapper) -> list[dict]:
    """The words an overlay's marks lie on, in the final text of its anchor's text box:
    {"range": (start, end), "marks": [(mark index, 0 left edge / 1 right edge)], "width": PDF
    width or None}. A right edge ends the last word before it (the whole run of words before it
    on its line finds it); a left edge starts the word a right edge on the same line ends, or else
    the word after the ones before it. Marks whose words aren't found are left out."""
    paras = slides_texts(text, scale, fonts)
    # (ranges in UTF-16 units, as Slides counts: u16)
    starts = [sum(u16(t) + 1 for t in paras[:i]) for i in range(len(paras))]
    tokens = [[(u16(t[:m.start()]), u16(t[:m.end()]), m.group()) for m in re.finditer(r"\S+", t)] for t in paras]

    def locate(want: list[str]) -> tuple[int, int] | None:
        """(paragraph, index of the last token) of a run of words; shorter tails if unique."""
        for k in range(len(want), 0, -1):
            hits = [(pi, j) for pi, toks in enumerate(tokens) for j in range(k - 1, len(toks))
                    if [t[2] for t in toks[j - k + 1:j + 1]] == want[-k:]]
            if hits and (k == len(want) or len(hits) == 1):
                return hits[0]
        return None

    words: dict[tuple, dict] = {}

    def add(key: tuple[int, int, int], mark: tuple[int, int], width: float | None) -> None:
        pi, j0, j1 = key
        w = words.setdefault(key, {"range": (starts[pi] + tokens[pi][j0][0], starts[pi] + tokens[pi][j1][1]),
                                   "marks": [], "width": width})
        w["marks"].append(mark)
        w["width"] = w["width"] or width

    ends, lefts = [], []  # (mark index, texts before the word, its closed x0, word key)
    for i, m in enumerate(overlay["marks"]):
        before = m["before"]
        if not (before and abs(before[-1][6] + before[-1][0] - m["hole_x0"]) <= 0.3):
            lefts.append(i)
            continue
        n = len(before[-1][5].split())
        hit = locate(" ".join(b[5] for b in before).split())
        if hit and n and hit[1] - n + 1 >= 0:
            key = (hit[0], hit[1] - n + 1, hit[1])
            ends.append((i, tuple(b[5] for b in before[:-1]), before[-1][6], key))
            add(key, (i, 1), before[-1][0])
    for i in lefts:
        m = overlay["marks"][i]
        texts = tuple(b[5] for b in m["before"])
        # (the same words before and x0 on two lines, "Stochastic" and "but" at the left margin:
        # classify lists a word's left edge right before its right edge)
        key = min(((abs(j - i - 1), k) for j, t, x0, k in ends if t == texts and abs(x0 - m["hole_x0"]) <= 0.3),
                  default=(0, None))[1]
        if key is None and texts:
            hit = locate(" ".join(texts).split())
            if hit and hit[1] + 1 < len(tokens[hit[0]]):
                key = (hit[0], hit[1] + 1, hit[1] + 1)
        if key is not None:
            add(key, (i, 0), None)
    return list(words.values())


def pick_word(marks: list[tuple[float, float, float, float]], x0: float, width: float | None,
              top: float, bottom: float) -> tuple[float, float, float, float] | None:
    """The highlighted word (slide pt) starting nearest the predicted `x0`, within its text box's
    rows top..bottom, about as wide as the PDF word if `width` is known, at most 25 pt away."""
    fits = [m for m in marks if top <= (m[1] + m[3]) / 2 <= bottom and abs(m[0] - x0) <= 25
            and (width is None or abs((m[2] - m[0]) - width) <= max(3.0, 0.3 * width))]
    return min(fits, key=lambda m: abs(m[0] - x0)) if fits else None


def title_bar_under(el: dict, slide: dict) -> list[float] | None:
    """The PDF box of the block title bar a text element sits on, if any."""
    cx, cy = (el["bbox"][0] + el["bbox"][2]) / 2, (el["bbox"][1] + el["bbox"][3]) / 2
    return next((e["bbox"] for e in slide["elements"] if e["kind"] == "shape" and e.get("block") is not None
                 and not e.get("title_bar") and e["bbox"][0] <= cx <= e["bbox"][2]
                 and e["bbox"][1] <= cy <= e["bbox"][3]), None)


def measure_jobs(deck: dict, scale: float, fonts: FontMapper, placed, page_slide: dict) -> tuple[list[dict], list[tuple]]:
    """The scratch slides of measure_places: their requests, and per slide (scratch slide id, page,
    [hole to find: picture, expected gap start x0 and line middle cy, width, pitch, colour, hang],
    [overlay: picture, predicted PDF extent, marks' predicted drifts, words (mark_words with their
    colour, predicted slide x0 and text box rows)])."""
    reqs, jobs = [], []
    for slide in deck["slides"]:
        n = slide["page"]
        holes = [h for h in slide_holes(slide) if h[3] is not None]
        texts = {e["id"]: e for e in slide["elements"] if e["kind"] == "text"}
        overlays = [e for e in slide["elements"] if e.get("marks") and e.get("anchor") in texts]
        if not holes and not overlays:
            continue
        sid = f"b2s_m{n:03}"
        reqs += [{"createSlide": {"objectId": sid, "slideLayoutReference": {"predefinedLayout": "BLANK"}}},
                 {"updatePageProperties": {"objectId": sid, "fields": "pageBackgroundFill.solidFill.color",
                                           "pageProperties": {"pageBackgroundFill": {"solidFill": {
                                               "color": {"rgbColor": {"red": 1, "green": 1, "blue": 1}}}}}}}]
        found, measured, taken = [], [], []  # taken: (colour, x0, x1, y0, y1) of every highlight, slide pt
        for i, el in enumerate(slide["elements"]):
            mine = [h for h in holes if h[0] is el]
            marked = [o for o in overlays if o["anchor"] == el["id"]]
            if not mine and not marked:
                continue
            colours = [HOLE_MARKS[(len(found) + k) % len(HOLE_MARKS)] for k in range(len(mine))]
            oid = f"{sid}_t{i}"
            plain = {**el, "paragraphs": [{**p, "runs": [{**r, "highlight": None} for r in p["runs"]]} for p in el["paragraphs"]]}
            reqs += text_box_requests(plain, sid, oid, scale, fonts, None, page_slide, title_bar_under(el, slide),
                                      text_right_limit(el, slide), colours)
            reqs.append({"updateTextStyle": {"objectId": oid, "textRange": {"type": "ALL"}, "fields": "foregroundColor",
                                             "style": {"foregroundColor": rgb("#000000")}}})
            top, bottom = el["bbox"][1] * scale - 4, el["bbox"][3] * scale + 4
            for (_, _, _, pic), colour in zip(mine, colours):
                x0, y0, x1, y1 = (v * scale for v in placed(pic, n)["bbox"])
                taken.append((colour, x0, x1, y0, y1))
            for o in marked:
                bx0, _, bx1, _ = o["bbox"]
                a0, _, a1, _ = placed(o, n)["bbox"]
                words = []
                for w in mark_words(o, el, scale, fonts):
                    m = o["marks"][w["marks"][0][0]]
                    x = m["x"] - (w["width"] or 0.0) * w["marks"][0][1]  # the word's PDF x0
                    x0 = (a0 + (x - bx0) * (a1 - a0) / (bx1 - bx0)) * scale
                    x1 = x0 + (w["width"] or 0.0) * scale

                    def clash(c):  # how close the nearest highlight of colour c on these rows is
                        return min([max(0.0, t[1] - x1, x0 - t[2]) for t in taken
                                    if t[0] == c and t[3] < bottom and t[4] > top] + [math.inf])
                    colour = max(HOLE_MARKS, key=lambda c: (clash(c), -sum(t[0] == c for t in taken)))
                    taken.append((colour, x0, x1, top, bottom))
                    words.append({**w, "colour": colour, "x0": x0, "top": top, "bottom": bottom})
                    reqs.append({"updateTextStyle": {"objectId": oid, "fields": "backgroundColor",
                                                     "style": {"backgroundColor": rgb(colour)},
                                                     "textRange": {"type": "FIXED_RANGE", "startIndex": w["range"][0],
                                                                   "endIndex": w["range"][1]}}})
                measured.append({"pic": o, "box": (a0, a1), "drifts": mark_drifts(o, scale, fonts), "words": words})
            for (_, p, run, pic), colour in zip(mine, colours):
                x0, y0, _, y1 = placed(pic, n)["bbox"]
                z = fonts(run, scale)[1]
                bl = [line["baseline"] for line in p["lines"]]
                pitch = (bl[-1] - bl[0]) / (len(bl) - 1) * scale if len(bl) > 1 else LINE_EM * z
                prev_end, next_x0 = hole_neighbours(p, run)
                # The middle of the hole's text line, not of the picture: a brace's label below
                # or above the formula would put that a line off.
                line = min(p["lines"], key=lambda l: (not y0 - 0.5 <= l["baseline"] <= y1 + 0.5,
                                                      abs((y0 + y1) / 2 - l["baseline"] + 0.35 * run["size"])))
                found.append({"pic": pic, "x0": x0 * scale - hole_offset(p, run, pic, scale, z),  # expected gap start
                              "cy": (line["baseline"] - 0.35 * run["size"]) * scale, "width": run["hole"] * scale, "pitch": pitch, "colour": colour,
                              # last on its PDF line: a hole Slides lets hang past the box edge isn't drawn
                              "hang": None if run.get("next_x0") is not None else
                              (SYMBOL_ADVANCE_EM[" "] * z if prev_end is not None else 0.0,
                               el["bbox"][0] * scale - PAD_X, el["bbox"][2] * scale + 2 * PAD_X)})
        jobs.append((sid, n, found, measured))
    return reqs, jobs


def overlay_move(o: dict, marks: dict[str, list], scale: float) -> tuple[tuple[float, float, float] | None, list]:
    """(dx, dy, scaleX) taking an overlay picture from its predicted extent to the one its
    measured words give (fit_overlay; marks not measured keep their predicted drift), or None if
    no word was found; and [x, predicted drift, measured drift or None] per mark (PDF pt)."""
    got = {}
    for w in o["words"]:
        rect = pick_word(marks[w["colour"]], w["x0"], w["width"] and w["width"] * scale, w["top"], w["bottom"])
        for i, side in w["marks"] if rect else []:
            got[i] = rect[2 * side] / scale - o["pic"]["marks"][i]["x"]
    xs = [m["x"] for m in o["pic"]["marks"]]
    table = [[x, round(d, 2), round(got[i], 2) if i in got else None] for i, (x, d) in enumerate(zip(xs, o["drifts"]))]
    if not got:
        return None, table
    b0, b1 = fit_overlay(o["pic"]["bbox"], [(x, got.get(i, d)) for i, (x, d) in enumerate(zip(xs, o["drifts"]))],
                         OVERLAY_STRETCH_MEASURED)
    a0, a1 = o["box"]
    return ((b0 - a0) * scale, 0.0, (b1 - b0) / (a1 - a0)), table


def measure_places(slides, pid: str, deck: dict, scale: float, fonts: FontMapper, placed, page_slide: dict,
                   out: Path, page_width: float = SLIDE_W) -> tuple[dict[str, tuple], list[str]]:
    """Moves for hole pictures ((dx, dy) in slide pt) and overlay pictures ((dx, dy, scaleX): dx
    moves the left edge), measured on scratch slides (see above), and the scratch slides to
    delete. What can't be found keeps its predicted place.

    `page_width`: how wide the deck being measured is, in slide pt. A thumbnail is a fixed number
    of pixels wide whatever the page is, so it alone says how many pixels a point is."""
    from PIL import Image
    from .gslides import save_thumbnail

    started = time.monotonic()
    reqs, jobs = measure_jobs(deck, scale, fonts, placed, page_slide)
    if not jobs:
        return {}, []
    try:
        batch(slides, pid, reqs)
    except HttpError as e:
        print(f"warning: could not measure the picture places ({api_error(e)}); keeping the predicted places")
        return {}, []  # (a refused batch created nothing)

    # One client per worker thread, not one per slide: `build(...)` fetches a discovery document
    # every time it is called. Where a caller handed its own client over there is only that one,
    # which is not thread-safe, so the thumbnails are fetched one at a time.
    creds, fetch = credentials_for_threads(), fetcher_for_threads()
    client = per_thread(lambda: slides_service(creds))
    workers = 1 if shared_service("slides", "v1") else 6
    tables = {}

    def measure(job):
        sid, n, found, overlays = job
        path = out / "holes" / f"marks-{n + 1:03}.png"
        try:
            save_thumbnail(client(), pid, sid, path, fetch)
        except Exception as e:  # noqa: BLE001 - HttpError, OSError, or a harness fetcher's own
            print(f"warning: slide {n + 1}: no thumbnail to measure the picture places ({e})")
            return {}
        img = np.asarray(Image.open(path).convert("RGB"))
        px_per_pt = img.shape[1] / page_width
        marks = {c: find_marks(mark_alpha(img, c), px_per_pt)
                 for c in {f["colour"] for f in found} | {w["colour"] for o in overlays for w in o["words"]}}
        moves = {}
        for o in overlays:
            move, tables[o["pic"]["id"]] = overlay_move(o, marks, scale)
            if move is None:
                print(f"warning: slide {n + 1}: words of {o['pic']['id']} not found; keeping its predicted place")
            elif abs(move[0]) >= 0.2 or abs(move[2] - 1) >= 0.002:
                moves[o["pic"]["id"]] = move
        for f in found:
            move = pick_gap(marks[f["colour"]], f["x0"], f["cy"], f["width"], f["pitch"])
            if move is None and f["hang"]:
                space, left, right = f["hang"]
                end = ink_end(img, px_per_pt, f["cy"] - 0.2 * f["pitch"], f["cy"] + 0.2 * f["pitch"], left, right)
                if end is not None and abs(end + space - f["x0"]) < 0.5 * f["width"] + 10:
                    move = (end + space - f["x0"], 0.0)
            if move is None:
                print(f"warning: slide {n + 1}: gap of {f['pic']['id']} not found; keeping its predicted place")
            elif abs(move[0]) >= 0.2 or move[1]:
                moves[f["pic"]["id"]] = move
        return moves

    moves = {}
    with ThreadPoolExecutor(max_workers=workers) as pool:
        for result in pool.map(measure, jobs):
            moves.update(result)
    (out / "holes" / "moves.json").write_text(json.dumps({k: [round(v, 4) for v in m] for k, m in moves.items()},
                                                         indent=1), encoding="utf-8")
    # (per overlay mark: PDF x, predicted drift, measured drift)
    (out / "holes" / "overlays.json").write_text(json.dumps(tables), encoding="utf-8")
    n_holes, n_overlays = sum(len(job[2]) for job in jobs), sum(len(job[3]) for job in jobs)
    print(f"  picture places: {n_holes} formula gaps and {n_overlays} overlays measured on {len(jobs)} slides, "
          f"{len(moves)} pictures moved ({time.monotonic() - started:.1f} s)")
    return moves, [job[0] for job in jobs]


def subtitle_element(slide: dict, title_idx: int) -> int | None:
    """On the title page, the biggest plain text box below the title (authors, institute,
    date) goes into the TITLE layout's subtitle placeholder."""
    if not slide.get("title_page"):
        return None
    title_bottom = slide["elements"][title_idx]["bbox"][3]
    below = [(sum(len(r["text"]) for p in e["paragraphs"] for r in p["runs"]), i)
             for i, e in enumerate(slide["elements"])
             if e["kind"] == "text" and e["role"] == "body" and e["bbox"][1] > title_bottom
             and not any(p["bullet"] for p in e["paragraphs"])]
    return max(below)[1] if below else None


def title_element(slide: dict) -> int | None:
    """Index of the element that becomes the slide's title placeholder."""
    for i, el in enumerate(slide["elements"]):
        if el["kind"] == "text" and el["role"] == "title":
            return i
    return None


PICTURE_TITLES = {"math": "Formula", "icon": "Icon", "fallback": "Picture"}


# ---------------------------------------------------------------- presentation

def import_presentation(slides, drive, title: str, page_w: float, page_h: float, pptx: io.BytesIO,
                        existing: str | None) -> dict:
    """A new deck from the .pptx, or an existing one (same URL) with its content replaced by it."""
    media = media_upload(pptx, PPTX_MIME)
    if existing:
        execute(drive.files().update(fileId=existing, media_body=media, fields="id"))
        pid = existing
    else:
        pid = execute(drive.files().create(
            body={"name": title, "mimeType": "application/vnd.google-apps.presentation"},
            media_body=media, fields="id"))["id"]
    pres = execute(slides.presentations().get(presentationId=pid))
    got = pres["pageSize"]["height"]["magnitude"] / pres["pageSize"]["width"]["magnitude"]
    if abs(got - page_h / page_w) > 0.003:
        raise RuntimeError(f"page aspect {got:.4f} != PDF aspect {page_h / page_w:.4f}")
    return pres


def background_key(slide: dict, out: Path) -> tuple:
    if slide.get("background_color"):
        return ("color", slide["background_color"].lower())
    return ("png", hashlib.sha1((out / slide["background"]).read_bytes()).hexdigest())


def plan_theme(deck: dict, out: Path, bg_key: dict) -> dict | None:
    """Theme decoration on the layouts (render.theme_decoration), so a background colour set in
    Slides changes only the page ground and new slides get the decoration too. Title pages (TITLE
    layout) and the other slides get one each. Backgrounds that don't show it (a standout frame, a
    closing page without the headline) get the decoration they share among themselves on a copy of
    their layout (VARIANT suffix; up to THEME_VARIANTS copies, the last without decoration). Slide
    backgrounds stay as they are: the decoration drawn over a background showing it changes nothing.

    Returns None without decoration, else {"ground": "#rrggbb", "decorations": {"TITLE" (title layout),
    "*" (every other layout), "TITLE_V1", "*_V1", ... (copies): Path or None}, "layouts": {page:
    layout name, e.g. "TITLE_ONLY_V1"}, "exact": {"TITLE", "*": background key that is exactly the
    ground plus the decoration, or None}}."""
    from PIL import Image

    from .render import page_ground, save_png, theme_decoration

    counts = Counter(bg_key.values())
    groups = {"TITLE": [s for s in deck["slides"] if slide_layout(s)[0] == "TITLE"]}
    groups["*"] = [s for s in deck["slides"] if slide_layout(s)[0] != "TITLE"]
    load = lambda s: np.asarray(Image.open(out / s["background"]).convert("RGB"))
    main = groups["*"] or groups["TITLE"]
    if not main:
        return None
    for old in (out / "backgrounds").glob("theme-*.png"):
        old.unlink()
    ground = page_ground(load(max(main, key=lambda s: counts[bg_key[s["page"]]])))
    theme = {"ground": "#" + "".join(f"{int(v):02x}" for v in ground), "decorations": {},
             "layouts": {s["page"]: slide_layout(s)[0] for s in deck["slides"]}, "exact": {}}
    for name, members in groups.items():
        for variant in range(THEME_VARIANTS + 1):
            if not members:
                break
            key = name if variant == 0 else f"{name}{VARIANT}{variant}"
            keys = sorted(dict.fromkeys(bg_key[s["page"]] for s in members), key=lambda k: -counts[k])
            first = {k: next(s for s in members if bg_key[s["page"]] == k) for k in keys}
            picture, inside = None, [True] * len(keys)
            if variant < THEME_VARIANTS:
                picture, inside, exact = theme_decoration((load(first[k]) for k in keys), ground)
                inside = inside if picture is not None else [True] * len(keys)
                if variant == 0:
                    theme["exact"][name] = keys[0] if picture is not None and exact else None
            if picture is not None:
                path = out / "backgrounds" / f"theme-{'title' if name == 'TITLE' else 'main'}-{variant}.png"
                save_png(picture, path)
            theme["decorations"][key] = path if picture is not None else None
            if variant:
                for s in members:
                    if inside[keys.index(bg_key[s["page"]])]:
                        theme["layouts"][s["page"]] += f"{VARIANT}{variant}"
            members = [s for s in members if not inside[keys.index(bg_key[s["page"]])]]
    if not any(theme["decorations"].values()):
        return None
    if not groups["TITLE"]:
        theme["decorations"]["TITLE"] = theme["decorations"].get("*")  # (for title slides added in Slides)
    return theme


def master_ground(shared: tuple | None, files: dict, page_w: float):
    """bbox (PDF pt) -> the master background's median colour there (see background_key)."""
    if shared is None or shared[0] == "color":
        colour = shared[1] if shared else "#ffffff"
        return lambda bbox: colour
    from PIL import Image

    img = np.asarray(Image.open(files[shared]).convert("RGB"))
    k =img.shape[1] / page_w

    def ground(bbox: list[float]) -> str:
        x0, y0, x1, y1 = (max(0, int(round(v * k))) for v in bbox)
        area = img[y0:max(y1, y0 + 1), x0:max(x1, x0 + 1)].reshape(-1, 3)
        return "#" + "".join(f"{int(v):02x}" for v in np.median(area, axis=0)) if len(area) else "#ffffff"
    return ground


LAYOUT_TEXT_PREFIX = "b2s_L"


def write_layout_texts(slides, pid: str, texts: list[dict], scale: float, fonts: FontMapper) -> None:
    """Header/footer text shared by every slide goes onto the layouts our slides use, so
    it is edited once for the whole deck. Previous runs' layout texts are replaced."""
    pres = execute(slides.presentations().get(
        presentationId=pid, fields="layouts(objectId,layoutProperties,pageElements(objectId))"))
    layouts = pres.get("layouts", [])  # all of them: slides added later in Slides get the footer too
    # All old ones first: object IDs are unique across the whole presentation.
    reqs = [{"deleteObject": {"objectId": e["objectId"]}} for l in pres.get("layouts", [])
            for e in l.get("pageElements", []) if e["objectId"].startswith(LAYOUT_TEXT_PREFIX)]
    for li, layout in enumerate(layouts):
        for ti, el in enumerate(texts):
            reqs += text_box_requests(el, layout["objectId"], f"{LAYOUT_TEXT_PREFIX}{li}_{ti}", scale, fonts)
    if reqs:
        execute(slides.presentations().batchUpdate(presentationId=pid, body={"requests": reqs}))


def contrast(a: str, b: str) -> float:
    """WCAG contrast ratio of two #rrggbb colours."""
    def luminance(h: str) -> float:
        c = [int(h.lstrip("#")[i:i + 2], 16) / 255 for i in (0, 2, 4)]
        c = [v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4 for v in c]
        return 0.2126 * c[0] + 0.7152 * c[1] + 0.0722 * c[2]
    la, lb = sorted((luminance(a), luminance(b)), reverse=True)
    return (la + 0.05) / (lb + 0.05)


MIN_CONTRAST = 2.0


def readable_run(run: dict, slide: dict, bbox: list[float], ground) -> dict:
    """The run in a colour readable on a new slide: that is on the master background under `bbox`
    (`ground(bbox)`: its colour there), without the shapes of the converted slide. White title
    text on a native title panel becomes the panel's colour (else black or white)."""
    under = ground(bbox) if ground else None
    if under is None or contrast(run["color"], under) >= MIN_CONTRAST:
        return run
    cx, cy = (bbox[0] + bbox[2]) / 2, (bbox[1] + bbox[3]) / 2
    panels = [e["fill"] for e in reversed(slide["elements"]) if e["kind"] == "shape" and e.get("fill")
              and e["bbox"][0] <= cx <= e["bbox"][2] and e["bbox"][1] <= cy <= e["bbox"][3]]
    colour = next((c for c in panels if contrast(c, under) >= MIN_CONTRAST), None) or \
        max(("#000000", "#ffffff"), key=lambda c: contrast(c, under))
    return {**run, "color": colour}


def style_layout_placeholders(slides, pid: str, deck: dict, scale: float, fonts: FontMapper, dy: float,
                              ground=None) -> None:
    """Title and body placeholders of every layout take the deck's own look (font, size,
    colour, title position), so slides added later in Slides match the converted ones.
    `ground(bbox)`: the master background's colour under a PDF box, to keep text readable there."""
    texts = [(s, e) for s in deck["slides"] for e in s["elements"] if e["kind"] == "text" and e["paragraphs"][0]["runs"]]
    frame = next(((s, e) for s, e in texts if e["role"] == "title" and not s.get("title_page")), None)
    page = next(((s, e) for s, e in texts if e["role"] == "title" and s.get("title_page")), None) or frame
    frame_title, page_title = (frame or (None, None))[1], (page or (None, None))[1]
    title_runs = {id(e): readable_run(e["paragraphs"][0]["runs"][0], s, e["bbox"], ground)
                  for s, e in (pair for pair in (frame, page) if pair)}
    body_runs = Counter((r["font"], r["size"], r["color"], r["family"]) for s, e in texts if e["role"] == "body"
                        for p in e["paragraphs"] for r in p["runs"] for _ in range(len(r["text"])))
    body = None
    if body_runs:
        font, size, color, family = body_runs.most_common(1)[0][0]
        body = {"font": font, "size": size, "color": color, "family": family, "bold": False, "italic": False}
        w, h = deck["slides"][0]["size"]
        body = readable_run(body, {"elements": []}, [0.1 * w, 0.3 * h, 0.9 * w, 0.8 * h], ground)
    page_fields = "objectId,pageElements(objectId,size,transform,shape(placeholder/type,text/textElements))"
    pres = execute(slides.presentations().get(presentationId=pid, fields=f"masters({page_fields}),layouts({page_fields})"))
    reqs = []
    # The master too: layout placeholders inherit whatever style they don't set themselves.
    for layout in pres.get("masters", []) + pres.get("layouts", []):
        for pe in layout.get("pageElements", []):
            kind = pe.get("shape", {}).get("placeholder", {}).get("type")
            if kind in ("TITLE", "CENTERED_TITLE"):
                el = page_title if kind == "CENTERED_TITLE" else frame_title
                if el is None:
                    continue
                p = el["paragraphs"][0]
                run = title_runs[id(el)]
                z = fonts(run, scale)[1]
                x = p["text_x0"] * scale - PAD_X
                if p["align"] == "center":
                    x = min(x, 0.1 * SLIDE_W)
                w = SLIDE_W - 2 * max(x, 10)
                h = 2 * LINE_EM * z + 2 * BASELINE_A
                y = p["lines"][0]["baseline"] * scale - (BASELINE_A + ASCENT_EM * z) + dy
                reqs.append({"updatePageElementTransform": {"objectId": pe["objectId"], "applyMode": "ABSOLUTE", "transform": {
                    "scaleX": w / (pe["size"]["width"]["magnitude"] / EMU_PER_PT),
                    "scaleY": h / (pe["size"]["height"]["magnitude"] / EMU_PER_PT), "unit": "EMU",
                    "translateX": round(max(x, 10) * EMU_PER_PT), "translateY": round(max(0.0, y) * EMU_PER_PT)}}})
                reqs.append({"updateShapeProperties": {"objectId": pe["objectId"], "fields": "contentAlignment",
                                                       "shapeProperties": {"contentAlignment": "TOP"}}})
                align = {"left": "START", "center": "CENTER", "right": "END"}[p["align"]]
            elif kind == "BODY" and body:
                run, align = body, "START"
            else:
                continue
            if not pe["shape"].get("text", {}).get("textElements"):
                # Without any text (not even the imported "\n" per list level) a placeholder can't
                # be styled, and the API refuses to put text into layout placeholders.
                continue
            style, fields = fonts.text_style(run, scale)
            style["foregroundColor"] = rgb(run["color"])
            if "bold" not in fields:  # a weighted family: the layout's own bold (section header) would add to it
                style["bold"], fields = False, fields + ["bold"]
            reqs += [
                {"updateTextStyle": {"objectId": pe["objectId"], "textRange": {"type": "ALL"}, "style": style,
                                     "fields": ",".join(fields + ["foregroundColor"])}},
                {"updateParagraphStyle": {"objectId": pe["objectId"], "textRange": {"type": "ALL"},
                                          "style": {"alignment": align}, "fields": "alignment"}},
            ]
    for i in range(0, len(reqs), 200):
        try:
            execute(slides.presentations().batchUpdate(presentationId=pid, body={"requests": reqs[i:i + 200]}))
        except HttpError as e:  # the deck's look for new slides is a nicety, never a reason to fail
            print(f"warning: could not style the layouts ({api_error(e)})")


def write_layouts(client, pid: str, deck: dict, scale: float, fonts: FontMapper, ground) -> None:
    """The whole of the layout and master work, as one job: what a deck's own look gives a slide
    somebody adds later in Slides. It reads and writes layout and master pages only, which is
    what lets it run on a thread of its own (`client()` gives that thread its own Slides client,
    `gslides.per_thread`) - but writing a layout is not the same as leaving the slides alone: a
    slide's TITLE placeholder inherits its layout parent's box until a content batch gives it one
    of its own, so this pass and that batch write one value and the later commit wins. It is
    joined before the first content batch goes out (`build_deck`, `tools/probe_layout_race.py`).

    Inside itself it stays serial, which is the one place in `build_deck` where that was measured
    rather than assumed: a layout write reaches every slide inheriting from it and Google charges
    for that, so more of this pass in the air is *slower*. Interleaved A/B on a 10-slide deck, the
    placeholder requests cut into four batches at once against one batch: 10.10 s against 4.95 s;
    the texts beside the placeholders rather than after them: 5.24 s against 3.96 s (18.6 s of
    conversion against 16.4 s). `tools/probe_batch_parallelism.py`'s finding - that several
    `batchUpdate`s may be in flight on one presentation and nothing is lost - is about slide
    content and does not carry here. Nor is there anything to win by reading less: merging the
    two passes' reads into one saves a round trip inside the pass and nothing at all outside it
    (17.17 s against 16.93 s over four interleaved pairs, which the arms split two each), because
    what is left of the pass hides behind phase 1, the placeholder read and `measure_places`."""
    slides = client()
    write_layout_texts(slides, pid, deck.get("layout_texts", []), scale, fonts)
    style_layout_placeholders(slides, pid, deck, scale, fonts, PPTX_TITLE_DY, ground)


# ---------------------------------------------------------------- the same, as data (theme_sync)

def master_plan(deck: dict, out: Path, theme="plan") -> dict:
    """`build_deck`'s decisions about the master and the layouts, as data, for a sync that has to
    compare them with what convert wrote (`theme_sync`): {"bg_key": page -> background key,
    "bg_file": key -> file, "shared": the key most slides share (None: none shared), "theme":
    plan_theme's answer, "fill": the key on the master, "ground": master_ground}. `theme="plan"`
    plans the theme (and so rewrites out/backgrounds/theme-*.png, as plan_theme does); a theme
    passed in is taken as it is, None as no theme."""
    bg_key = {s["page"]: background_key(s, out) for s in deck["slides"]}
    bg_file = {bg_key[s["page"]]: out / s["background"] for s in deck["slides"] if not s.get("background_color")}
    counts = Counter(bg_key.values())
    shared = counts.most_common(1)[0][0] if counts and counts.most_common(1)[0][1] >= 2 else None
    if theme == "plan":
        theme = plan_theme(deck, out, bg_key)
    fill = shared or ("color", "#ffffff")
    if theme and theme.get("exact") is not None:
        group = lambda s: "TITLE" if slide_layout(s)[0] == "TITLE" else "*"
        if shared is None or all(theme["exact"].get(group(s)) == shared for s in deck["slides"] if bg_key[s["page"]] == shared):
            fill = ("color", theme["ground"])
    page_w = deck["slides"][0]["size"][0] if deck["slides"] else SLIDE_W
    return {"bg_key": bg_key, "bg_file": bg_file, "shared": shared, "theme": theme, "fill": fill,
            "ground": master_ground(shared, bg_file, page_w)}


def layout_style_spec(deck: dict, scale: float, fonts: FontMapper, dy: float, ground=None) -> dict:
    """What `style_layout_placeholders` writes into each kind of placeholder, as data a base can
    hold: {"TITLE" / "CENTERED_TITLE" / "BODY": {"box": [x, y, w, h] (slide pt; None for the
    body, whose box is left alone), "style", "fields", "align"}, or None where nothing is written}.
    `layout_placeholder_requests` turns one entry into the requests for one placeholder;
    tests/test_theme_sync.py holds the two to the requests `style_layout_placeholders` sends."""
    texts = [(s, e) for s in deck["slides"] for e in s["elements"] if e["kind"] == "text" and e["paragraphs"][0]["runs"]]
    frame = next(((s, e) for s, e in texts if e["role"] == "title" and not s.get("title_page")), None)
    page = next(((s, e) for s, e in texts if e["role"] == "title" and s.get("title_page")), None) or frame
    frame_title, page_title = (frame or (None, None))[1], (page or (None, None))[1]
    title_runs = {id(e): readable_run(e["paragraphs"][0]["runs"][0], s, e["bbox"], ground)
                  for s, e in (pair for pair in (frame, page) if pair)}
    body_runs = Counter((r["font"], r["size"], r["color"], r["family"]) for s, e in texts if e["role"] == "body"
                        for p in e["paragraphs"] for r in p["runs"] for _ in range(len(r["text"])))
    body = None
    if body_runs:
        font, size, color, family = body_runs.most_common(1)[0][0]
        body = {"font": font, "size": size, "color": color, "family": family, "bold": False, "italic": False}
        w, h = deck["slides"][0]["size"]
        body = readable_run(body, {"elements": []}, [0.1 * w, 0.3 * h, 0.9 * w, 0.8 * h], ground)

    def styled(run: dict) -> tuple[dict, str]:
        style, fields = fonts.text_style(run, scale)
        style["foregroundColor"] = rgb(run["color"])
        if "bold" not in fields:
            style["bold"], fields = False, fields + ["bold"]
        return style, ",".join(fields + ["foregroundColor"])

    spec: dict = {}
    for kind, el in (("TITLE", frame_title), ("CENTERED_TITLE", page_title)):
        if el is None:
            spec[kind] = None
            continue
        p = el["paragraphs"][0]
        run = title_runs[id(el)]
        z = fonts(run, scale)[1]
        x = p["text_x0"] * scale - PAD_X
        if p["align"] == "center":
            x = min(x, 0.1 * SLIDE_W)
        w = SLIDE_W - 2 * max(x, 10)
        h = 2 * LINE_EM * z + 2 * BASELINE_A
        y = p["lines"][0]["baseline"] * scale - (BASELINE_A + ASCENT_EM * z) + dy
        style, fields = styled(run)
        spec[kind] = {"box": [max(x, 10), max(0.0, y), w, h], "style": style, "fields": fields,
                      "align": {"left": "START", "center": "CENTER", "right": "END"}[p["align"]]}
    if body:
        style, fields = styled(body)
        spec["BODY"] = {"box": None, "style": style, "fields": fields, "align": "START"}
    else:
        spec["BODY"] = None
    return json.loads(json.dumps(spec))  # (plain data: nothing shared with the caller's runs)


def layout_placeholder_requests(entry: dict, pe: dict) -> list[dict]:
    """The requests `style_layout_placeholders` sends for one placeholder `pe` (a layout or master
    page element as presentations.get gives it) from its `layout_style_spec` entry."""
    reqs = []
    if entry.get("box") is not None:
        x, y, w, h = entry["box"]
        reqs.append({"updatePageElementTransform": {"objectId": pe["objectId"], "applyMode": "ABSOLUTE", "transform": {
            "scaleX": w / (pe["size"]["width"]["magnitude"] / EMU_PER_PT),
            "scaleY": h / (pe["size"]["height"]["magnitude"] / EMU_PER_PT), "unit": "EMU",
            "translateX": round(x * EMU_PER_PT), "translateY": round(y * EMU_PER_PT)}}})
        reqs.append({"updateShapeProperties": {"objectId": pe["objectId"], "fields": "contentAlignment",
                                               "shapeProperties": {"contentAlignment": "TOP"}}})
    if pe.get("shape", {}).get("text", {}).get("textElements"):
        reqs += [
            {"updateTextStyle": {"objectId": pe["objectId"], "textRange": {"type": "ALL"},
                                 "style": json.loads(json.dumps(entry["style"])), "fields": entry["fields"]}},
            {"updateParagraphStyle": {"objectId": pe["objectId"], "textRange": {"type": "ALL"},
                                      "style": {"alignment": entry["align"]}, "fields": "alignment"}},
        ]
    return reqs


def api_error(e: HttpError) -> str:
    return message_of(e)


def batch(slides, pid: str, reqs: list[dict]) -> None:
    execute(slides.presentations().batchUpdate(presentationId=pid, body={"requests": reqs}))


# ---------------------------------------------------------------- main entry

def existing_presentation(drive, out: Path) -> str | None:
    """The deck from a previous run of this output folder, if it still exists (not trashed)."""
    from .guard import previous_deck
    previous = previous_deck(drive, out)
    return previous["presentationId"] if previous and previous["state"] == "live" else None


def size_pt(element: dict) -> tuple[float, float]:
    size = element["size"]
    return tuple(size[k]["magnitude"] / (EMU_PER_PT if size[k]["unit"] == "EMU" else 1) for k in ("width", "height"))


def fallback_pictures(deck: dict, refused: list[tuple[int, str]], out: Path) -> dict:
    """The deck with every element the API refused ((PDF page, element id)) replaced by a
    picture of its region, cropped from the PDF the deck was built from."""
    from .render import crop_region

    source_pdf = out / "slides.pdf" if (out / "slides.pdf").exists() else Path(deck["source"]["pdf"])
    if not source_pdf.exists():
        # The one step of a conversion that needs the PDF itself rather than what was classified
        # out of it, and the only reason `agent.deck_tools.deck_upload` asks for one at all. A
        # folder that travelled without its source says so here rather than inside `crop_region`.
        raise FileNotFoundError(
            f"the API refused {len(refused)} element(s) and the region of each has to be cropped "
            f"from the page, but the PDF this deck was built from is not at {source_pdf}. Put it "
            f"back beside the folder (or pass it in) and build the deck again.")
    new_slides = []
    for slide in deck["slides"]:
        ids = {eid for page, eid in refused if page == slide["page"]}
        elements = []
        for el in slide["elements"]:
            if el["id"] in ids and el["kind"] != "image":
                ids.discard(el["id"])
                x0, y0, x1, y1 = el["bbox"]
                bbox = [x0 - 2, y0 - 2, x1 + 2, y1 + 2]
                path = out / "figures" / f"fallback-{el['id']}.png"
                crop_region(source_pdf, slide["page"], bbox, path, 6.0)
                el = {"kind": "image", "id": el["id"], "role": "fallback", "bbox": bbox,
                      "file": str(path.relative_to(out)).replace("\\", "/")}
            elements.append(el)
        new_slides.append({**slide, "elements": elements})
    return {**deck, "slides": new_slides}


def preflight_rebuild(out: Path, source_pdf: Path | None, new_deck: bool = False, force_rebuild: bool = False,
                      slides=None, drive=None) -> dict | None:
    """The guard's question (guard.check_rebuild) before the conversion work starts, so a refusal
    comes in a second instead of after extract, classify and render. `emit` asks again - and backs
    the deck up - immediately before the write, in case the deck is edited in between.

    Returns what it found ({"presentationId", "found"}), which that second ask confirms with one
    field of one read instead of asking the whole question again (`plan_rebuild`'s `checked`);
    None where there was nothing to ask."""
    from . import guard

    if new_deck or force_rebuild or not (out / "emit.json").exists():
        return None
    drive = drive or drive_service()
    previous = guard.previous_deck(drive, out)
    if not previous or previous["state"] != "live":
        return None
    pid = previous["presentationId"]
    return {"presentationId": pid,
            "found": guard.check_rebuild(slides or slides_service(), drive, pid, out, source_pdf, False)}


def preflight_in_background(out: Path, source_pdf: Path | None, new_deck: bool = False,
                            force_rebuild: bool = False):
    """`preflight_rebuild` on a thread of its own. Returns the function that asks for its answer:
    it raises whatever the check raised, and gives back what it found (`emit`'s `checked`).

    The check is three Drive reads and a whole `presentations.get` - three seconds that need
    nothing the conversion produces and answer a question only the first write to Drive really
    asks - so it is made while the PDF is being converted and collected just before `emit`. The
    price is that a refusal now comes after the local conversion instead of in a second, and what
    that writes is the output folder's own files: the deck itself is still never touched.
    """
    if new_deck or force_rebuild or not (out / "emit.json").exists() \
            or shared_service("slides", "v1") or shared_service("drive", "v3"):
        # Nothing to ask, or a caller's own clients, which are that caller's one thread's.
        found = preflight_rebuild(out, source_pdf, new_deck, force_rebuild)
        return lambda: found
    creds = credentials_for_threads()  # here: a worker thread inherits no context (google_auth)
    pool = ThreadPoolExecutor(1, thread_name_prefix="b2s-preflight")
    work = pool.submit(lambda: preflight_rebuild(out, source_pdf, new_deck, force_rebuild,
                                                 slides_service(creds), drive_service(creds)))
    pool.shutdown(wait=False)

    def answer() -> dict | None:
        return work.result()
    return answer


def look_again(slides, drive, out: Path, checked: dict | None) -> tuple[dict | None, dict | None]:
    """What the output folder points at, and the preflight's finding where it still stands.

    Two reads that need nothing of each other - the deck's place in Drive and its revision - so
    they are made at once, one client per thread. `guard.recheck` is the cheap half of the second
    ask; a deck that has moved since the preflight (or one the folder no longer points at, or one
    now in the trash) gets the whole question again, in `plan_rebuild`."""
    from . import guard

    pid = (checked or {}).get("presentationId")
    if not pid or shared_service("slides", "v1"):
        return guard.previous_deck(drive, out), None
    creds = credentials_for_threads()  # here: a worker thread inherits no context (google_auth)
    with ThreadPoolExecutor(1, thread_name_prefix="b2s-recheck") as pool:
        again = pool.submit(lambda: guard.recheck(slides_service(creds), pid, checked["found"]))
        previous = guard.previous_deck(drive, out)
        found = again.result()
    if previous and previous["presentationId"] == pid and previous["state"] == "live":
        return previous, found
    return previous, None


def plan_rebuild(slides, drive, out: Path, new_deck: bool, force_rebuild: bool, backup: str,
                 source_pdf: Path | None, checked: dict | None = None) -> tuple[str | None, dict | None]:
    """Decide what happens to the deck this output folder already has: rebuild it in place (the id
    is returned), or leave it alone and make a new one. Nothing destructive happens before this:
    `guard.check_rebuild` raises `guard.RebuildRefused` when the deck was edited in Slides, and a
    forced rebuild keeps a backup and records the deck's revision first
    (`<out>/backups/backups.json`, printed too) - and is refused in turn when that backup could
    not be kept (`guard.demand_way_back`), because then nothing could bring the deck back. The
    second value goes into emit.json as "previous".

    `checked`: what the preflight found a few seconds ago (`preflight_rebuild`). The question is
    asked again here because the deck may have been edited in between - but only what could have
    changed since is really in question, so where the deck is still at the revision the preflight
    read, that finding stands and the deck and the base are not read again."""
    from . import guard

    previous, found = look_again(slides, drive, out, None if new_deck or force_rebuild else checked)
    if previous is None:
        return None, None
    pid, url = previous["presentationId"], guard.deck_url(previous["presentationId"])
    if previous["state"] != "live":
        where = {"trashed": "is in the Drive trash", "gone": "is gone (deleted, or not this app's file any more)",
                 "other": "is not a presentation any more"}[previous["state"]]
        print(f"the deck of the previous run ({pid}) {where}: making a new one, that deck is left as it is")
        return None, {"presentationId": pid, "state": previous["state"], "action": "new deck", "url": url}
    if new_deck:
        print(f"--new-deck: the previous deck is left as it is at {url}\n"
              f"  (this folder tracks the new deck from now on; the old one is only reachable by that link)")
        return None, {"presentationId": pid, "state": "kept", "action": "new deck", "url": url}
    found = found or guard.check_rebuild(slides, drive, pid, out, source_pdf, force_rebuild)
    mode = backup if backup != "auto" else ("file" if found["reason"] else "none")
    entry = {"presentationId": pid, "url": url, "action": "rebuilt in place", "revisionId": found.get("revisionId"),
             "modifiedTime": previous.get("modifiedTime"), "out": str(out),  # Drive's clock, and where to restore from
             "checked": found.get("checked"), "reason": found.get("reason") or "no deck edits",
             "summary": guard.summary_line(found) if found.get("edited") else "no deck edits",
             "examples": found.get("examples", []), "base_from": found.get("base_from")}
    if found["reason"]:
        print(f"WARNING: rebuilding a deck that {'was edited in Slides' if found['reason'] == 'edited' else found['reason']} "
              f"(--force-rebuild): {entry['summary']}")
    entry["backup"] = guard.backup_deck(drive, pid, out, mode, entry["reason"])
    guard.record(out, entry)  # the attempt belongs in the log even when it failed, and what follows
    if found["reason"]:
        guard.demand_way_back(pid, out, source_pdf, entry, mode)  # no backup, no forced rebuild
    print(f"updating existing deck {pid} (revision {found.get('revisionId')})")
    for line in guard.restore_hint(entry) if found["reason"] else []:
        print(line)
    return pid, entry


def emit(deck: dict, out: Path, title: str, new_deck: bool = False, measure: bool = True,
         force_rebuild: bool = False, backup: str = "auto", source_pdf: Path | None = None,
         checked: dict | None = None) -> dict:
    """Build the deck. An output folder that already has a deck is rebuilt in place unless
    `new_deck`; that replaces the deck's whole content, so `guard.check_rebuild` refuses when
    the deck was edited in Slides (`force_rebuild` goes ahead, after a backup). `checked`: what
    the preflight found (`preflight_rebuild`), which saves the second ask a read of the deck."""
    from . import guard

    slides, drive = slides_service(), drive_service()
    deck = {**deck, "slides": [{**s, "elements": merge_blocks(s["elements"])} for s in deck["slides"]]}
    existing, previous_entry = plan_rebuild(slides, drive, out, new_deck, force_rebuild, backup, source_pdf, checked)
    state, refused = build_deck(slides, drive, deck, out, title, existing, measure)
    if refused:
        # A picture can only come with the imported .pptx (the API inserts images from public
        # URLs only), so the deck is built once more with the refused elements as pictures.
        print(f"rebuilding the deck with {len(refused)} refused element(s) as pictures")
        state, again = build_deck(slides, drive, fallback_pictures(deck, refused, out), out, title,
                                  state["presentationId"], measure)
        for page, eid in again:
            print(f"warning: slide {page + 1}: {eid} was refused again and is missing")
    if previous_entry:
        state["previous"] = previous_entry  # what this run replaced, and how to get it back
    (out / "emit.json").write_text(json.dumps({k: v for k, v in state.items() if k != "deck"}, indent=1), encoding="utf-8")
    return state


def build_deck(slides, drive, deck: dict, out: Path, title: str, existing: str | None,
               measure: bool = True) -> tuple[dict, list[tuple[int, str]]]:
    """Import the .pptx and fill in the content. Returns the state for emit.json and the
    elements the API refused ((PDF page, element id)). `measure`: hole and overlay pictures go
    where a thumbnail shows their gaps and words (measure_places), not only where they are predicted."""
    page_w, page_h = deck["slides"][0]["size"]
    plan = DeckPlan(deck, pptx_tables=True)
    deck, scale, fonts = plan.deck, plan.scale, plan.fonts

    # Backgrounds: the most common one becomes the master's (the deck's theme): layouts and
    # slides inherit it, and slides added later too. Identical pictures are stored once.
    bg_key = {s["page"]: background_key(s, out) for s in deck["slides"]}
    bg_file = {bg_key[s["page"]]: out / s["background"] for s in deck["slides"] if not s.get("background_color")}
    counts = Counter(bg_key.values())
    shared = counts.most_common(1)[0][0] if counts and counts.most_common(1)[0][1] >= 2 else None

    def fill(key: tuple) -> dict:
        return {"color": key[1]} if key[0] == "color" else {"picture": bg_file[key]}

    theme = plan_theme(deck, out, bg_key)
    master_fill = fill(shared or ("color", "#ffffff"))
    if theme:
        group = lambda s: "TITLE" if slide_layout(s)[0] == "TITLE" else "*"
        # The master's ground colour where the shared background is nothing but ground and decoration.
        if shared is None or all(theme["exact"].get(group(s)) == shared for s in deck["slides"] if bg_key[s["page"]] == shared):
            master_fill = {"color": theme["ground"]}
    pages = [{
        "layout": theme["layouts"][s["page"]] if theme else slide_layout(s)[0],
        "fill": None if bg_key[s["page"]] == shared else fill(bg_key[s["page"]]),
        "pictures": [{"file": out / e["file"], "bbox": bbox, "alt": e.get("alt"),
                      "title": PICTURE_TITLES.get(e.get("role"), "Figure")} for e, bbox in plan.pictures(s)],
        "tables": plan.tables(s),
        "templates": plan.uses_templates[s["page"]],
    } for s in deck["slides"]]
    pptx = build_pptx(page_w, page_h, plan.keys, pages, master_fill, theme and theme["decorations"])
    pres = import_presentation(slides, drive, title, page_w, page_h, pptx, existing)
    pid = pres["presentationId"]
    sources = pres.get("slides", [])
    if len(sources) != len(deck["slides"]):
        raise RuntimeError(f"the import brought {len(sources)} slides, expected {len(deck['slides'])}")

    # Phase 1: every source slide is copied under our object IDs (slide, title and subtitle
    # placeholders, pictures, template shapes); the sources are deleted at the end.
    template_sizes: list[tuple[float, float]] = []
    reqs = []
    for slide, source in zip(deck["slides"], sources):
        request, sizes = plan.copy_request(slide, source)
        template_sizes = template_sizes or sizes
        reqs.append(request)
    batch(slides, pid, reqs)

    # A round trip to Google costs about a second whatever it carries, so what a conversion
    # spends is round trips and not work, and two of them may be in the air at once. The layouts
    # and the master are nobody else's business - neither pass reads the slides being filled in
    # beside them - so they go on a thread of their own, and the content batches below go out
    # CONTENT_WORKERS at a time, each thread with its own client. A caller that handed one ready
    # Slides client over (google_auth.use_services with a mapping) keeps the old serial order:
    # that client is its own, and a service object is not thread-safe.
    threaded = not shared_service("slides", "v1")
    creds = credentials_for_threads() if threaded else None  # here: a worker inherits no context
    client = per_thread(lambda: slides_service(creds)) if threaded else (lambda: slides)
    pool = ThreadPoolExecutor(CONTENT_WORKERS, thread_name_prefix="b2s-content") if threaded else None
    ground = master_ground(shared, bg_file, page_w)
    layout_pool, layout_work = None, None
    if threaded:
        layout_pool = ThreadPoolExecutor(1, thread_name_prefix="b2s-layout")
        layout_work = layout_pool.submit(write_layouts, client, pid, deck, scale, fonts, ground)
    else:
        write_layouts(client, pid, deck, scale, fonts, ground)

    state = {"presentationId": pid, "url": f"https://docs.google.com/presentation/d/{pid}/edit",
             "scale": scale, "slides": []}
    if theme:
        state["theme"] = {"ground": theme["ground"], "master": master_fill.get("color"),
                          "decorations": {k: str(p.relative_to(out)).replace("\\", "/") if p else None
                                          for k, p in theme["decorations"].items()},
                          "layouts": {str(k): v for k, v in theme["layouts"].items()}}
    # Placeholder sizes (needed to resize them) and any extra layout placeholders.
    created = execute(slides.presentations().get(
        presentationId=pid,
        fields="slides(objectId,pageElements(objectId,size),slideProperties/notesPage/notesProperties)"))
    page_elements = {s["objectId"]: s.get("pageElements", []) for s in created["slides"]}
    speaker_notes = {s["objectId"]: s.get("slideProperties", {}).get("notesPage", {})
                     .get("notesProperties", {}).get("speakerNotesObjectId") for s in created["slides"]}
    moves, scratch = measure_places(slides, pid, deck, scale, fonts, plan.placed, plan.page_slide, out) \
        if measure else ({}, [])

    # Phase 2: content, batched over slides. Each slide's requests come in parts (one per
    # element) so that a rejected batch can be narrowed down to the element at fault.
    refused: list[tuple[int, str]] = []

    def send(items: list[tuple[str, int, list[tuple[dict | None, list[dict]]]]]) -> None:
        reqs = [r for _, _, parts in items for _, rs in parts for r in rs]
        if not reqs:
            return
        try:
            batch(client(), pid, reqs)
            return
        except HttpError as e:
            if len(items) > 1:
                for item in items:
                    send([item])
                return
            print(f"warning: {items[0][0]}: batch rejected ({api_error(e)}); retrying element by element")
        slide_id, page, parts = items[0]
        for el, rs in parts:
            try:
                if rs:
                    batch(client(), pid, rs)
            except HttpError as e:
                print(f"warning: {slide_id}: {el['kind'] + ' ' + el['id'] if el else 'request'} rejected "
                      f"({api_error(e)})" + ("; using a picture of it instead" if el and el["kind"] != "image" else ""))
                if el and el["kind"] != "image":
                    refused.append((page, el["id"]))

    # A full batch is sent while the next slides are still being planned, several at a time.
    sent = []

    def dispatch(items: list) -> None:
        if pool:
            sent.append(pool.submit(send, items))
        else:
            send(items)

    pending: list[tuple[str, int, list]] = []
    pending_size = 0
    try:
        # The layouts first, and never beside the slides: a slide's title placeholder inherits
        # the layout's box until we give it one of its own, and while the layout batch is in the
        # air together with the content batch that does that, the one Google commits LAST wins -
        # a layout batch landing second takes every title's own box away again and the deck's
        # titles all sit at the layout's, silently. Measured with three conversions at once
        # (`tools/probe_layout_race.py`): the deck whose layout batch landed after its first
        # content batch lost all ten titles, the two that landed first kept theirs. So the
        # layout pass overlaps the read and `measure_places` above it, and nothing below.
        if layout_work is not None:
            layout_work.result()
        for slide in deck["slides"]:
            n = slide["page"]
            slide_id = f"b2s_s{n:03}"
            parts, element_ids = plan.slide_parts(slide, page_elements, speaker_notes, moves, template_sizes)
            size = sum(len(rs) for _, rs in parts)
            # Several slides per round trip; a slide's requests are never split across batches.
            if pending and pending_size + size > BATCH_MAX_REQUESTS:
                dispatch(pending)
                pending, pending_size = [], 0
            pending.append((slide_id, n, parts))
            pending_size += size
            objects, groups = element_objects(parts, element_ids)
            state["slides"].append({"page": n, "objectId": slide_id, "elements": element_ids, "objects": objects,
                                    "groups": groups})
            if plan.pptx_tables:  # the base records them: a sync refills such a table in place (sync.table_refill)
                state["slides"][-1]["table_margins"] = {str(i): pptx_table(el, plan.scale, plan.fonts)["margins"]
                                                        for i, el in enumerate(slide["elements"]) if el["kind"] == "table"}
            kinds = [el["kind"] for el in slide["elements"]]
            print(f"  slide {n + 1}: {kinds.count('text')} text boxes, {kinds.count('image')} pictures, "
                  f"{kinds.count('shape')} shapes, {kinds.count('table')} tables")
        if pending:
            dispatch(pending)
        for job in sent:             # every content batch has landed
            job.result()
    finally:
        for p in (pool, layout_pool):
            if p:
                p.shutdown()
    refused.sort()                   # several threads appended to it
    batch(slides, pid, [{"deleteObject": {"objectId": oid}} for oid in [s["objectId"] for s in sources] + scratch])
    state["deck"] = deck  # (what was built, for the sync snapshot; not written to emit.json)
    return state, refused


def created_ids(reqs: list[dict]) -> list[str]:
    """Object ids a list of requests creates."""
    out = []
    for r in reqs:
        (kind, body), = r.items()
        if kind in ("createShape", "createLine", "createTable", "createImage", "createSlide"):
            out.append(body["objectId"])
        elif kind == "duplicateObject":
            out += list(body.get("objectIds", {}).values())
        elif kind == "groupObjects":
            out.append(body["groupObjectId"])
    return out


def element_objects(parts: list[tuple[dict | None, list[dict]]], element_ids: list[str]) -> tuple[list[list[str]], list[str]]:
    """Per element (in slide_parts order) every object id created for it: its main object first,
    then what its requests create and its group with anchored pictures ({oid}_g); and the slide's
    other groups (blocks, rules)."""
    objects = [[oid] for oid in element_ids]
    index = {oid: i for i, oid in enumerate(element_ids)}
    groups = []
    k = 0
    for el, reqs in parts:
        if el is not None:
            objects[k] += [o for o in created_ids(reqs) if o != element_ids[k]]
            k += 1
            continue
        for oid in created_ids(reqs):
            owner = index.get(oid[:-2]) if oid.endswith("_g") else None
            if owner is not None:
                objects[owner].append(oid)
            else:
                groups.append(oid)
    return [list(dict.fromkeys(o)) for o in objects], groups


class DeckPlan:
    """The requests build_deck sends, apart from what only Google knows (the imported slides'
    object IDs, placeholder and template sizes, measured hole moves): pure, so tests can check
    them offline (plan_offline)."""

    def __init__(self, deck: dict, page_width: float = SLIDE_W, pptx_tables: bool = False):
        # `page_width`: the width of the deck this plan is for, in slide pt. A deck `convert` makes
        # is always SLIDE_W wide (it uploads the .pptx that says so), but `sync` may be writing into
        # a deck a person built at any size (`adopt_sync`), and every box, font size and hole width
        # below is this converter's PDF pt times `scale`.
        # `pptx_tables`: tables come with the imported .pptx, empty and with their cell margins
        # (`tables`, build_pptx), and are filled in; else (sync) they are made by createTable.
        self.page_width = page_width
        self.pptx_tables = pptx_tables
        self.scale = scale = page_width / deck["slides"][0]["size"][0]
        self.fonts = fonts = FontMapper()
        self.deck = deck = {**deck, "slides": [fit_holes(s, scale, fonts) for s in deck["slides"]]}
        self.keys = list(dict.fromkeys(k for s in deck["slides"] for e in s["elements"] for k in element_template_keys(e, scale)))
        self.uses_templates = {s["page"]: any(element_template_keys(e, scale) for e in s["elements"]) for s in deck["slides"]}
        self.shifts = {s["page"]: formula_shifts(s, scale, fonts) for s in deck["slides"]}
        self.overlays = {s["page"]: overlay_boxes(s, scale, fonts) for s in deck["slides"]}
        # Internal link targets: PDF page -> slide. A skipped overlay step maps to the kept
        # (last) step of its frame, which comes right after it.
        kept = sorted(s["page"] for s in deck["slides"])
        self.page_slide = {}
        for page in range(kept[-1] + 1):
            target = next(k for k in kept if k >= page)
            self.page_slide[page] = f"b2s_s{target:03}"

    def placed(self, el: dict, n: int) -> dict:
        """Inline formula pictures sit over the gap Slides leaves for them (formula_shifts),
        graphics drawn at words over those words (overlay_boxes)."""
        overlays, shifts = self.overlays[n], self.shifts[n]
        if el["id"] in overlays:
            return {**el, "bbox": [overlays[el["id"]][0], el["bbox"][1], overlays[el["id"]][1], el["bbox"][3]]}
        dx = shifts.get(el["id"])
        return el if dx is None else {**el, "bbox": [el["bbox"][0] + dx, el["bbox"][1], el["bbox"][2] + dx, el["bbox"][3]]}

    def pictures(self, slide: dict) -> list[tuple[dict, list[float]]]:
        """The slide's pictures with their boxes in the .pptx (slide pt)."""
        return [(e, [v * self.scale for v in self.placed(e, slide["page"])["bbox"]])
                for e in slide["elements"] if e["kind"] == "image"]

    def tables(self, slide: dict) -> list[dict]:
        """The slide's tables as the .pptx carries them (pptx_table), when it does."""
        if not self.pptx_tables:
            return []
        return [pptx_table(e, self.scale, self.fonts) for e in slide["elements"] if e["kind"] == "table"]

    def copy_request(self, slide: dict, source: dict) -> tuple[dict, list[tuple[float, float]]]:
        """Phase 1: the duplicateObject copying a slide's imported source under our object IDs,
        and the sizes of the template shapes on the source."""
        n = slide["page"]  # PDF page index; slides may skip pages (overlays)
        slide_id = f"b2s_s{n:03}"
        keys, uses_templates = self.keys, self.uses_templates
        els = source.get("pageElements", [])
        placeholders = {e["shape"]["placeholder"]["type"]: e["objectId"] for e in els if "placeholder" in e.get("shape", {})}
        pictures = [e["objectId"] for e in els if "image" in e]
        tables = [e["objectId"] for e in els if "table" in e]
        shapes = [e for e in els if "image" not in e and "table" not in e and "placeholder" not in e.get("shape", {})]
        picture_idx = [i for i, e in enumerate(slide["elements"]) if e["kind"] == "image"]
        table_idx = [i for i, e in enumerate(slide["elements"]) if e["kind"] == "table"] if self.pptx_tables else []
        if len(pictures) != len(picture_idx) or len(tables) != len(table_idx) or \
                len(shapes) != (len(keys) if uses_templates[n] else 0):
            raise RuntimeError(f"slide {n + 1}: the import brought {len(pictures)} pictures, {len(tables)} tables and "
                               f"{len(shapes)} template shapes, expected {len(picture_idx)}, {len(table_idx)} and "
                               f"{len(keys) if uses_templates[n] else 0}")
        ids = {source["objectId"]: slide_id}
        ids.update({oid: f"{slide_id}_f{i}" for oid, i in zip(pictures, picture_idx)})
        ids.update({oid: f"{slide_id}_tab{i}" for oid, i in zip(tables, table_idx)})
        ids.update({e["objectId"]: f"{slide_id}_k{j}" for j, e in enumerate(shapes)})
        title_idx = title_element(slide)
        if title_idx is not None:
            ids[placeholders[slide_layout(slide)[1]]] = f"{slide_id}_t{title_idx}"
            sub_idx = subtitle_element(slide, title_idx)
            if sub_idx is not None and "SUBTITLE" in placeholders:
                ids[placeholders["SUBTITLE"]] = f"{slide_id}_t{sub_idx}"
        return {"duplicateObject": {"objectId": source["objectId"], "objectIds": ids}}, [size_pt(e) for e in shapes]

    def slide_parts(self, slide: dict, page_elements: dict[str, list[dict]], speaker_notes: dict[str, str | None],
                    moves: dict[str, tuple[float, float]], template_sizes: list[tuple[float, float]]
                    ) -> tuple[list[tuple[dict | None, list[dict]]], list[str]]:
        """Phase 2 for one slide after its copy: requests in parts ((element, requests), so a
        rejected batch can be narrowed down to the element at fault) and the element object IDs.
        `page_elements` and `speaker_notes` describe the copied slides (slide id -> elements with
        objectId and size, speaker notes object id), `moves` are measure_places' results."""
        scale, fonts, keys = self.scale, self.fonts, self.keys
        placed, page_slide, uses_templates = self.placed, self.page_slide, self.uses_templates
        placeholder_dy = PPTX_TITLE_DY

        def template_on_slide(slide_id: str, key: tuple) -> dict:
            """The slide's copy of a template shape ({"id", "w", "h"}: its unscaled size in pt)."""
            j = keys.index(key)
            w, h = template_sizes[j]
            return {"id": f"{slide_id}_k{j}", "w": w, "h": h}

        n = slide["page"]
        slide_id = f"b2s_s{n:03}"
        title_idx = title_element(slide)
        title_oid = f"{slide_id}_t{title_idx}" if title_idx is not None else None
        sub_idx = subtitle_element(slide, title_idx) if title_idx is not None else None
        subtitle_oid = f"{slide_id}_t{sub_idx}" if sub_idx is not None else None
        ours = (f"{slide_id}_k", f"{slide_id}_f", f"{slide_id}_tab")  # template shapes, pictures, tables from the .pptx
        parts: list[tuple[dict | None, list[dict]]] = [(None, [
            {"deleteObject": {"objectId": e["objectId"]}}
            for e in page_elements.get(slide_id, [])
            if e["objectId"] not in (title_oid, subtitle_oid) and not e["objectId"].startswith(ours)])]
        element_ids = []
        for i, el in enumerate(slide["elements"]):  # shapes, then pictures, then text on top
            el = placed(el, n)
            if el["kind"] == "shape":
                oid = f"{slide_id}_s{i}"
                key = template_key(el, scale)
                reqs = shape_requests(el, slide_id, oid, scale, template_on_slide(slide_id, key) if key else None)
            elif el["kind"] == "table":
                oid = f"{slide_id}_tab{i}"
                reqs = table_requests(el, slide_id, oid, scale, fonts, self.pptx_tables)
            elif el["kind"] == "diagram":
                oid = f"{slide_id}_dg{i}"
                reqs = diagram_requests(el, slide_id, oid, scale, fonts,
                                        (lambda key, s=slide_id: template_on_slide(s, key)) if keys else None)
            elif el["kind"] == "image":
                # The picture came with the slide: move it to its place in the z-order.
                oid = f"{slide_id}_f{i}"
                reqs = [{"updatePageElementsZOrder": {"pageElementObjectIds": [oid], "operation": "BRING_TO_FRONT"}}]
                if el["id"] in moves:  # to the gap or words measured for it (measure_places)
                    dx, dy, sx = (*moves[el["id"]], 1.0)[:3]
                    # (a relative transform scales about the page origin: the left edge keeps its dx)
                    reqs.insert(0, {"updatePageElementTransform": {"objectId": oid, "applyMode": "RELATIVE", "transform": {
                        "scaleX": sx, "scaleY": 1, "unit": "EMU",
                        "translateX": round((dx + (1 - sx) * el["bbox"][0] * scale) * EMU_PER_PT),
                        "translateY": round(dy * EMU_PER_PT)}}})
                if el.get("number"):
                    reqs += number_box_requests(el["number"], slide_id, f"{oid}n", scale, fonts)
            else:
                oid = f"{slide_id}_t{i}"
                placeholder = None
                if oid in (title_oid, subtitle_oid):
                    size = next(e["size"] for e in page_elements[slide_id] if e["objectId"] == oid)
                    placeholder = {"base_w": size["width"]["magnitude"] / EMU_PER_PT,
                                   "base_h": size["height"]["magnitude"] / EMU_PER_PT, "dy": placeholder_dy}
                reqs = text_box_requests(el, slide_id, oid, scale, fonts, placeholder, page_slide,
                                         title_bar_under(el, slide), text_right_limit(el, slide))
            parts.append((el, reqs))
            element_ids.append(oid)
        # The slide's copies of the template shapes have been duplicated from: remove them.
        extra = [{"deleteObject": {"objectId": f"{slide_id}_k{j}"}} for j in range(len(keys)) if uses_templates[n]]
        # Inline formula pictures move with their text: group them (placeholders can't be grouped).
        by_id = {el["id"]: oid for el, oid in zip(slide["elements"], element_ids)}
        anchored: dict[str, list[str]] = {}
        for el, oid in zip(slide["elements"], element_ids):
            if el.get("anchor") in by_id and by_id[el["anchor"]] not in (title_oid, subtitle_oid):
                anchored.setdefault(by_id[el["anchor"]], []).extend([oid, f"{oid}n"] if el.get("number") else [oid])
        grouped = set()
        for text_oid, pictures in anchored.items():
            extra.append({"groupObjects": {"groupObjectId": f"{text_oid}_g", "childrenObjectIds": [text_oid] + pictures}})
            grouped |= {text_oid, *pictures}
        # A beamer block (title bar and body shapes plus everything on them) moves as one.
        for bi, members in enumerate(block_groups(slide["elements"], element_ids, title_oid)):
            children = [f"{m}_g" if m in anchored else m for m in members if m not in grouped or m in anchored]
            if len(children) >= 2:
                extra.append({"groupObjects": {"groupObjectId": f"{slide_id}_blk{bi}", "childrenObjectIds": children}})
                # A group takes the place of its topmost member, above a table lying on the
                # block (tables can't join the group): blocks are backdrops, send them back.
                extra.append({"updatePageElementsZOrder": {"pageElementObjectIds": [f"{slide_id}_blk{bi}"],
                                                           "operation": "SEND_TO_BACK"}})
        for ri, members in enumerate(rule_groups(slide["elements"], element_ids)):
            extra.append({"groupObjects": {"groupObjectId": f"{slide_id}_rules{ri}", "childrenObjectIds": members}})
        if slide.get("notes") and speaker_notes.get(slide_id):
            extra.append({"insertText": {"objectId": speaker_notes[slide_id], "text": slide["notes"]}})
        if title_oid and len(slide["elements"]) > 1:
            # The placeholder was created with the slide, below everything added since.
            extra.append({"updatePageElementsZOrder": {"pageElementObjectIds": [o for o in (title_oid, subtitle_oid) if o],
                                                       "operation": "BRING_TO_FRONT"}})
        parts += [(None, [r]) for r in extra]
        return parts, element_ids


def slide_layout(slide: dict) -> tuple[str, str | None]:
    """The title page uses the TITLE layout (centered title), frames with a title TITLE_ONLY."""
    if title_element(slide) is None:
        return "BLANK", None
    return ("TITLE", "CENTERED_TITLE") if slide.get("title_page") else ("TITLE_ONLY", "TITLE")


LAYOUT_PLACEHOLDERS = {"TITLE": ["CENTERED_TITLE", "SUBTITLE"], "TITLE_ONLY": ["TITLE"], "BLANK": []}


def plan_offline(deck: dict, placeholder_size: tuple[float, float] = (612.0, 90.0),
                 template_size: tuple[float, float] = (100.0, 100.0)) -> dict:
    """What emit would send for a classified deck, without Google: the imported slides are made
    up as the .pptx brings them (layout placeholders, pictures, template shapes) and hole
    pictures keep their predicted places. {"plan": DeckPlan, "pictures": {page: [(element, .pptx
    box)]}, "copies": phase 1 requests, "page_elements" and "speaker_notes": the copied slides,
    "measure": measure_places' scratch slide requests, "slides": [(slide id, page, parts, element ids)]}."""
    plan = DeckPlan({**deck, "slides": [{**s, "elements": merge_blocks(s["elements"])} for s in deck["slides"]]},
                    pptx_tables=True)

    def size(w: float, h: float) -> dict:
        return {"width": emu(w), "height": emu(h)}

    copies, page_elements, template_sizes = [], {}, []
    for slide in plan.deck["slides"]:
        n, source = slide["page"], f"src{slide['page']:03}"
        els = [{"objectId": f"{source}_{kind}", "size": size(*placeholder_size), "shape": {"placeholder": {"type": kind}}}
               for kind in LAYOUT_PLACEHOLDERS[slide_layout(slide)[0]]]
        els += [{"objectId": f"{source}_p{i}", "size": size(1, 1), "image": {}} for i, _ in enumerate(plan.pictures(slide))]
        els += [{"objectId": f"{source}_tb{i}", "size": size(sum(t["widths"]), sum(t["heights"])), "table": {}}
                for i, t in enumerate(plan.tables(slide))]
        els += [{"objectId": f"{source}_k{j}", "size": size(*template_size), "shape": {}}
                for j in range(len(plan.keys) if plan.uses_templates[n] else 0)]
        request, sizes = plan.copy_request(slide, {"objectId": source, "pageElements": els})
        copies.append(request)
        template_sizes = template_sizes or sizes
        ids = request["duplicateObject"]["objectIds"]
        page_elements[ids[source]] = [{"objectId": ids.get(e["objectId"], f"{e['objectId']}_copy"), "size": e["size"]}
                                      for e in els]
    speaker_notes = {slide_id: f"{slide_id}_notes" for slide_id in page_elements}
    slides = []
    for slide in plan.deck["slides"]:
        parts, element_ids = plan.slide_parts(slide, page_elements, speaker_notes, {}, template_sizes)
        slides.append((f"b2s_s{slide['page']:03}", slide["page"], parts, element_ids))
    return {"plan": plan, "pictures": {s["page"]: plan.pictures(s) for s in plan.deck["slides"]}, "copies": copies,
            "page_elements": page_elements, "speaker_notes": speaker_notes,
            "measure": measure_jobs(plan.deck, plan.scale, plan.fonts, plan.placed, plan.page_slide)[0], "slides": slides}


class _OnePage(dict):
    """`page_slide` for `slide_emission`: every internal link goes to one stand-in slide. Which
    slide a link names is the element's own IR (`identity.ir_fields` keys it), not its neighbours'."""

    def __bool__(self) -> bool:
        return True

    def get(self, page, default=None):
        return "b2s_link"


def slide_emission(slide: dict, scale: float, fonts: FontMapper, placeholder_size: tuple[float, float] = (612.0, 90.0),
                   template_size: tuple[float, float] = (100.0, 100.0)) -> dict:
    """What emit writes for one slide of a DeckPlan (blocks merged, holes fitted: `DeckPlan.deck`),
    worked out from that slide alone, so the same slide gives the same answer in whichever deck it
    stands: {"slide_id", "parts" and "element_ids" (`DeckPlan.slide_parts`), "boxes" (each
    picture's predicted place in slide pt, None for other kinds), "title" and "subtitle" (element
    indices of the layout placeholders' texts), "templates" ({object id of the slide's copy: its
    template key})}. As in `plan_offline` the layout's placeholders and the template shapes have
    made-up sizes; unlike it, links to other slides all go to one page and measure_places' moves
    are left out (only Google's renderer knows them), so pictures keep their predicted places.
    For sync.mark_emitted, which compares a base slide's emission with the new one's."""
    n = slide["page"]
    slide_id = f"b2s_s{n:03}"
    keys = list(dict.fromkeys(k for e in slide["elements"] for k in element_template_keys(e, scale)))
    plan = DeckPlan.__new__(DeckPlan)  # (one slide: none of the deck-wide work __init__ does)
    plan.page_width, plan.pptx_tables, plan.scale, plan.fonts = slide["size"][0] * scale, False, scale, fonts
    plan.deck = {"slides": [slide]}
    plan.keys, plan.uses_templates = keys, {n: bool(keys)}
    plan.shifts = {n: formula_shifts(slide, scale, fonts)}
    plan.overlays = {n: overlay_boxes(slide, scale, fonts)}
    plan.page_slide = _OnePage()
    title = title_element(slide)
    subtitle = subtitle_element(slide, title) if title is not None else None
    w, h = placeholder_size
    page_elements = {slide_id: [{"objectId": f"{slide_id}_t{i}", "size": {"width": emu(w), "height": emu(h)}}
                                for i in (title, subtitle) if i is not None]}
    parts, element_ids = plan.slide_parts(slide, page_elements, {}, {}, [template_size] * len(keys))
    boxes = [[v * scale for v in plan.placed(e, n)["bbox"]] if e["kind"] == "image" else None for e in slide["elements"]]
    return {"slide_id": slide_id, "parts": parts, "element_ids": element_ids, "boxes": boxes, "title": title,
            "subtitle": subtitle, "templates": {f"{slide_id}_k{j}": k for j, k in enumerate(keys)}}
