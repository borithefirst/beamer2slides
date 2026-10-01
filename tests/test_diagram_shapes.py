"""Node shapes beyond the plain ones (`classify_shapes`): each recogniser on a synthetic path drawn
the way TikZ draws the shape, a near miss that must stay a picture, and what emit makes of the
preset (its template key, text rectangle and connection sites, the .pptx template's adjustment)."""

import math
from collections.abc import Sequence

from beamer2slides.classify_model import Rect
from beamer2slides.classify_shapes import (
    Preset, can_adjust, cloud_puffs, polygon_preset, preset_shape, punched_tape, split_rectangle, stadium_radius,
    turned_angle,
)
from beamer2slides.emit_diagrams import connection, connection_sites, label_inside, node_template_key, text_rect_width
from beamer2slides.json_types import Json, JsonObject
from beamer2slides.raw_types import PathItem

Pt = tuple[float, float]


def arc(cx: float, cy: float, rx: float, ry: float, a0: float, a1: float, turn: float) -> PathItem:
    """One cubic of an elliptic arc from angle a0 to a1 (radians, y down), the ellipse turned
    `turn` radians (y down: a negative turn is counterclockwise on the page)."""
    k = 4 / 3 * math.tan((a1 - a0) / 4)
    c, s = math.cos(turn), math.sin(turn)

    def at(x: float, y: float) -> list[float]:
        return [cx + c * x - s * y, cy + s * x + c * y]

    def along(x: float, y: float) -> tuple[float, float]:
        return c * x - s * y, s * x + c * y
    p0, p3 = at(rx * math.cos(a0), ry * math.sin(a0)), at(rx * math.cos(a1), ry * math.sin(a1))
    d0, d1 = along(-rx * math.sin(a0), ry * math.cos(a0)), along(-rx * math.sin(a1), ry * math.cos(a1))
    return ("c", [p0, [p0[0] + k * d0[0], p0[1] + k * d0[1]], [p3[0] - k * d1[0], p3[1] - k * d1[1]], p3])


def line(a: Pt, b: Pt) -> PathItem:
    return ("l", [[a[0], a[1]], [b[0], b[1]]])


def polygon(vs: Sequence[Pt]) -> list[PathItem]:
    return [line(vs[k], vs[(k + 1) % len(vs)]) for k in range(len(vs))]


def box_of(path: list[PathItem]) -> Rect:
    """The path's drawn box, as extract gives it (curves sampled)."""
    xs: list[float] = []
    ys: list[float] = []
    for op, pts in path:
        if op == "c":
            for i in range(21):
                t, u = i / 20, 1 - i / 20
                xs.append(u ** 3 * pts[0][0] + 3 * u * u * t * pts[1][0] + 3 * u * t * t * pts[2][0] + t ** 3 * pts[3][0])
                ys.append(u ** 3 * pts[0][1] + 3 * u * u * t * pts[1][1] + 3 * u * t * t * pts[2][1] + t ** 3 * pts[3][1])
        else:
            xs += [p[0] for p in pts]
            ys += [p[1] for p in pts]
    return Rect(min(xs), min(ys), max(xs), max(ys))


HALF = math.pi / 2


def cylinder(ry_top: float, ry_bottom: float) -> list[PathItem]:
    """TikZ's `cylinder, shape border rotate=90` on a 40 x 40 box: the bottom's front half, the
    right side, the top's back half, the left side, the top's front half."""
    top, bottom = ry_top, 40 - ry_bottom
    return [arc(20, bottom, 20, ry_bottom, math.pi, HALF, 0), arc(20, bottom, 20, ry_bottom, HALF, 0, 0),
            line((40, bottom), (40, top)),
            arc(20, top, 20, ry_top, 0, -HALF, 0), arc(20, top, 20, ry_top, -HALF, -math.pi, 0),
            line((0, top), (0, bottom)),
            arc(20, top, 20, ry_top, math.pi, HALF, 0), arc(20, top, 20, ry_top, HALF, 0, 0)]


def test_a_cylinder_is_a_can_at_its_ellipse_depth() -> None:
    path = cylinder(4, 4)
    assert box_of(path) == Rect(0, 0, 40, 40)
    assert can_adjust(path, Rect(0, 0, 40, 40)) == 0.2, "the ellipse is 8 pt of the 40 pt side"
    assert preset_shape(path, Rect(0, 0, 40, 40), True) == Preset(shape="CAN", adjust=0.2, radius=None)


def test_a_lopsided_cylinder_stays_a_picture() -> None:
    path = cylinder(8, 3)
    assert can_adjust(path, box_of(path)) is None, "no CAN draws two ellipses of different depths"
    assert preset_shape(path, box_of(path), True) is None


def test_a_hexagon_is_one_at_its_corner_cut() -> None:
    path = polygon([(0, 21), (12, 0), (36, 0), (48, 21), (36, 42), (12, 42)])
    assert preset_shape(path, Rect(0, 0, 48, 42), True) == Preset(shape="HEXAGON", adjust=round(12 / 42, 3), radius=None)


def test_an_irregular_hexagon_stays_a_picture() -> None:
    path = polygon([(0, 21), (12, 0), (40, 3), (48, 21), (36, 42), (12, 42)])
    assert preset_shape(path, box_of(path), True) is None


def test_a_chamfered_rectangle_is_an_octagon() -> None:
    path = polygon([(6, 0), (54, 0), (60, 6), (60, 18), (54, 24), (6, 24), (0, 18), (0, 6)])
    assert polygon_preset([p for _, pts in path for p in pts], Rect(0, 0, 60, 24)) == \
        Preset(shape="OCTAGON", adjust=0.25, radius=None)
    uneven = polygon([(10, 0), (50, 0), (60, 4), (60, 20), (50, 24), (10, 24), (0, 20), (0, 4)])
    assert preset_shape(uneven, Rect(0, 0, 60, 24), True) is None, "cut 10 along, 4 down: no octagon draws that"


def regular(n: int, start: float) -> list[PathItem]:
    """A regular polygon of radius 20 around (30, 30), its first corner at `start` degrees (y down)."""
    return polygon([(30 + 20 * math.cos(math.radians(start + 360 * k / n)), 30 + 20 * math.sin(math.radians(start + 360 * k / n)))
                    for k in range(n)])


def test_regular_polygons_standing_on_a_side() -> None:
    for n, kind in ((5, "PENTAGON"), (7, "HEPTAGON"), (10, "DECAGON"), (12, "DODECAGON")):
        path = regular(n, 90 + 180 / n)
        got = preset_shape(path, box_of(path), True)
        assert got is not None and got.shape == kind, n
    tilted = regular(5, 100 + 180 / 5)
    assert preset_shape(tilted, box_of(tilted), True) is None, "a pentagon on its corner: the preset stands on a side"
    nine = regular(9, 90 + 180 / 9)
    assert preset_shape(nine, box_of(nine), True) is None, "no preset has nine sides"


def test_an_open_polyline_is_no_polygon() -> None:
    path = polygon([(0, 21), (12, 0), (36, 0), (48, 21), (36, 42), (12, 42)])[:-1]
    assert preset_shape(path, Rect(0, 0, 48, 42), False) is None, "stroked and not closed: lines"


def cloud(puffs: int) -> list[PathItem]:
    """Puffs around an ellipse: valleys at 0.8 of the radius, peaks at 1, cusps between."""
    def at(deg: float, share: float) -> list[float]:
        return [40 + 30 * share * math.cos(math.radians(deg)), 30 + 20 * share * math.sin(math.radians(deg))]
    step = 360 / puffs
    out: list[PathItem] = []
    for k in range(puffs):
        a = k * step
        out.append(("c", [at(a, 0.8), at(a, 1.15), at(a + 0.35 * step, 1.0), at(a + step / 2, 1.0)]))
        out.append(("c", [at(a + step / 2, 1.0), at(a + 0.65 * step, 1.0), at(a + step, 1.15), at(a + step, 0.8)]))
    return out


def test_a_cloud_is_puffs_meeting_in_cusps() -> None:
    path = cloud(10)
    assert cloud_puffs(path, box_of(path)) == 10
    assert preset_shape(path, box_of(path), True) == Preset(shape="CLOUD", adjust=None, radius=None)


def test_a_smooth_outline_of_many_curves_is_no_cloud() -> None:
    path = [arc(40, 30, 30, 20, math.radians(30 * k), math.radians(30 * (k + 1)), 0) for k in range(12)]
    assert cloud_puffs(path, box_of(path)) is None, "no cusps: an ellipse drawn in twelve pieces"


def tape(top_left_in: bool) -> list[PathItem]:
    """TikZ's `tape` on 40 x 30: waves 3 pt deep, 3 pt in from the top and the bottom."""
    down, up = (7.0, -1.0) if top_left_in else (-1.0, 7.0)
    return [("c", [[0, 3], [7, down], [13, down], [20, 3]]), ("c", [[20, 3], [27, up], [33, up], [40, 3]]),
            line((40, 3), (40, 27)),
            ("c", [[40, 27], [33, 23], [27, 23], [20, 27]]), ("c", [[20, 27], [13, 31], [7, 31], [0, 27]]),
            line((0, 27), (0, 3))]


def test_a_tape_is_punched_tape() -> None:
    assert box_of(tape(True)) == Rect(0, 0, 40, 30)
    assert punched_tape(tape(True), Rect(0, 0, 40, 30))
    assert preset_shape(tape(True), Rect(0, 0, 40, 30), True) == \
        Preset(shape="FLOW_CHART_PUNCHED_TAPE", adjust=None, radius=None)
    assert not punched_tape(tape(False), Rect(0, 0, 40, 30)), "its top dipping in on the right: the preset's mirror"


def pill(r_left: float) -> list[PathItem]:
    return [line((10, 0), (50, 0)), arc(50, 10, 10, 10, -HALF, 0, 0), arc(50, 10, 10, 10, 0, HALF, 0),
            line((50, 20), (10, 20)), arc(10, 10, r_left, 10, HALF, math.pi, 0), arc(10, 10, r_left, 10, math.pi, 3 * HALF, 0)]


def test_a_pill_is_a_rounded_rectangle_at_its_largest_rounding() -> None:
    assert stadium_radius(pill(10), Rect(0, 0, 60, 20)) == 10
    assert preset_shape(pill(10), Rect(0, 0, 60, 20), True) == Preset(shape="ROUND_RECTANGLE", adjust=None, radius=10)
    flat = pill(3)
    assert preset_shape(flat, box_of(flat), True) is None, "a flattened end is no half circle"
    key = node_template_key({"bbox": [0, 0, 60, 20], "shape": "ROUND_RECTANGLE", "radius": 10.0})
    assert key == ("ROUND_RECTANGLE", 0.5, None)


def test_a_split_rectangle_is_its_parts() -> None:
    rect: PathItem = ("re", [[0, 0], [100, 60]])
    parts = split_rectangle([rect, line((0, 20), (100, 20)), line((0, 45), (100, 45))], Rect(0, 0, 100, 60))
    assert parts == [Rect(0, 0, 100, 20), Rect(0, 20, 100, 45), Rect(0, 45, 100, 60)]
    assert split_rectangle([rect, line((0, 20), (100, 20)), line((50, 0), (50, 60))], Rect(0, 0, 100, 60)) is None, \
        "a grid is a table's"
    assert split_rectangle([rect, line((0, 20), (60, 20))], Rect(0, 0, 100, 60)) is None, "a line splitting nothing"


def test_turned_nodes_say_so() -> None:
    turned = [arc(40, 30, 30, 15, HALF * k, HALF * (k + 1), math.radians(-30)) for k in range(4)]
    angle = turned_angle(turned, box_of(turned))
    assert angle is not None and abs(angle - 30) < 0.5
    upright = [arc(40, 30, 30, 15, HALF * k, HALF * (k + 1), 0) for k in range(4)]
    assert turned_angle(upright, box_of(upright)) is None
    c, s = math.cos(math.radians(20)), math.sin(math.radians(20))
    corners = [(30 + c * x - s * y, 30 + s * x + c * y) for x, y in ((-20, -10), (20, -10), (20, 10), (-20, 10))]
    angle = turned_angle(polygon(corners), box_of(polygon(corners)))
    assert angle is not None and abs(angle + 20) < 0.5, "turned clockwise on the page"
    square = polygon([(0, 0), (40, 0), (40, 20), (0, 20)])
    assert turned_angle(square, Rect(0, 0, 40, 20)) is None


def node(bbox: Sequence[float], shape: str, adjust: float | None, label_w: float) -> JsonObject:
    box: list[Json] = [v for v in bbox]
    out: JsonObject = {"bbox": box, "shape": shape, "label_w": label_w, "paragraphs": [[{"text": "Label"}]]}
    if adjust is not None:
        out["adjust"] = adjust
    return out


def test_adjusted_presets_carry_their_adjustment() -> None:
    assert node_template_key(node([0, 0, 48, 42], "HEXAGON", 0.286, 10)) == ("HEXAGON", 0.29, None)
    assert node_template_key(node([0, 0, 40, 40], "CAN", 0.2, 10)) == ("CAN", 0.2, None)
    assert node_template_key(node([0, 0, 60, 40], "CLOUD", None, 10)) == ("CLOUD", None, None)


def test_preset_text_rectangles() -> None:
    assert text_rect_width("OCTAGON", (0, 0, 60, 24), 0.25) == 1 - 0.25 * 24 / 60
    hexagon = text_rect_width("HEXAGON", (0, 0, 42, 42), 0.25)
    assert hexagon is not None and abs(hexagon - 2 / 3) < 1e-9, "the default's corners on the slanted sides"
    assert text_rect_width("PENTAGON", (0, 0, 42, 40), None) is None, "a pentagon's label stays beside it"
    assert label_inside(node([0, 0, 60, 40], "CLOUD", None, 30))
    assert not label_inside(node([0, 0, 60, 40], "CLOUD", None, 40)), "a cloud's text rectangle is 0.65 of it"
    assert not label_inside(node([0, 0, 42, 40], "PENTAGON", None, 10))


def test_preset_connection_sites() -> None:
    assert connection_sites("CAN", (0, 0, 40, 40), 0.2)[0] == (0.5, 0.2), "the top site on the near rim"
    # Google's can, measured live (a line connected to each index, read back): five sites, the
    # bottom at 3 - a line to the bottom connected at OOXML's 2 was moved to the left side.
    assert connection_sites("CAN", (0, 0, 40, 40), 0.2)[1:] == [(0.5, 0), (0, 0.5), (0.5, 1), (1, 0.5)]
    nodes = [node([0, 0, 48, 42], "HEXAGON", 0.25, 10), node([100, 0, 160, 24], "OCTAGON", 0.25, 10)]
    assert connection([48, 21], nodes, ["h", "o"]) == {"connectedObjectId": "h", "connectionSiteIndex": 0}
    assert connection([106, 0], nodes, ["h", "o"]) == {"connectedObjectId": "o", "connectionSiteIndex": 6}, \
        "the top side's left end, 6 pt (0.25 x 24) in"
    assert connection([100, 12], nodes, ["h", "o"]) is None, "an octagon has no site mid-side"


def test_template_shapes_carry_the_adjustment() -> None:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.shapes.autoshape import Shape

    from beamer2slides.emit_pptx import _add_template_shapes
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_template_shapes(slide, [("CAN", 0.2, None), ("HEXAGON", 0.29, None), ("OCTAGON", 0.28, None),
                                 ("CLOUD", None, None), ("FLOW_CHART_PUNCHED_TAPE", None, None), ("PENTAGON", None, None)])
    shapes = [s for s in slide.shapes if isinstance(s, Shape)]
    assert [s.auto_shape_type for s in shapes] == [MSO_SHAPE.CAN, MSO_SHAPE.HEXAGON, MSO_SHAPE.OCTAGON, MSO_SHAPE.CLOUD,
                                                   MSO_SHAPE.FLOWCHART_PUNCHED_TAPE, MSO_SHAPE.REGULAR_PENTAGON]
    can, hexagon, octagon = shapes[:3]
    assert abs(can.adjustments[0] - 0.2) < 1e-4 and abs(octagon.adjustments[0] - 0.28) < 1e-4
    assert abs(hexagon.adjustments[0] - 0.29) < 1e-4
    assert abs(hexagon.adjustments[1] - 1.1547) < 1e-4, "the hexagon's vf stays its default: it fills the box's height"
