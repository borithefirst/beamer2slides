"""Tables on synthetic pages, from the visual hunt: shaded grids with no rules, a heatmap under one
\\hline, white stripes on an off-white theme, \\multirow heads, double vertical rules, wrapped
p{} cells (justified, several columns at once), list continuations beside each other that are no
table, and in emit a wrapped cell's column that keeps the PDF's breaks and number cells kept at
the table's size."""

import pytest

from beamer2slides import emit as E
from beamer2slides.classify import PageClassifier, Rect, Span, classify
from beamer2slides.fonts import font_info

W, H = 453.54, 255.12
SANS = "LMSans10-Regular"
SIZE = 10.0
CHAR = 0.46  # em per character: Latin Modern Sans' average


class Page:
    def __init__(self):
        self.spans, self.drawings = [], []

    def text(self, text: str, x0: float, baseline: float, size: float = SIZE, color: str = "#000000") -> dict:
        w = len(text) * CHAR * size
        s = {"id": f"p0s{len(self.spans)}", "text": text, "font": SANS, "size": size, "color": color,
             "alpha": 255, "origin": [x0, baseline], "bbox": [x0, baseline - 0.75 * size, x0 + w, baseline + 0.25 * size],
             "dir": [1.0, -0.0], "smallcaps": False}
        self.spans.append(s)
        return s

    def words(self, text: str, x0: float, baseline: float, size: float = SIZE, color: str = "#000000") -> float:
        """A line of words a word space apart, one span each; returns where it ends."""
        for word in text.split():
            x0 = self.text(word, x0, baseline, size, color)["bbox"][2] + 0.33 * size
        return x0 - 0.33 * size

    def fill(self, x0, y0, x1, y1, color: str) -> None:
        self._draw("f", [["re", [[x0, y0], [x1, y1]]]], fill=color)

    def hline(self, x0, x1, y, width: float = 0.4) -> None:
        self._draw("s", [["l", [[x0, y], [x1, y]]]], stroke="#000000", width=width)

    def vline(self, x, y0, y1, width: float = 0.4) -> None:
        self._draw("s", [["l", [[x, y1], [x, y0]]]], stroke="#000000", width=width)

    def _draw(self, type: str, path: list, fill=None, stroke=None, width=None) -> None:
        pts = [p for _, ps in path for p in ps]
        self.drawings.append({
            "id": f"p0d{len(self.drawings)}", "type": type, "items": "".join(op for op, _ in path),
            "bbox": [min(x for x, _ in pts), min(y for _, y in pts), max(x for x, _ in pts), max(y for _, y in pts)],
            "fill": fill, "stroke": stroke, "width": width, "fill_opacity": 1.0, "corners": {}, "path": path})

    def raw(self) -> dict:
        return {"index": 0, "label": "1", "size": [W, H], "spans": self.spans, "images": [],
                "drawings": self.drawings, "links": [], "frame_label": None}

    def elements(self) -> list[dict]:
        self.words("Body text that sets the size of the deck", 30, 245)
        deck = classify({"version": 1, "source": {"title": ""}, "pages": [self.raw()]})
        return deck["slides"][0]["elements"]


def tables(els: list[dict]) -> list[dict]:
    return [e for e in els if e["kind"] == "table"]


def cell_texts(t: dict) -> list[list[str]]:
    return [[" ".join("".join(r["text"] for r in c).split()) for c in row] for row in t["cells"]]


# ---------------------------------------------------------------- shading and fills

def shaded_grid(page: Page, colors: list[str], x0: float = 60, col_w: float = 110, top: float = 50,
                pitch: float = 18, rows: list[list[str]] | None = None) -> None:
    rows = rows or [["Region", "Share", "Trend"], ["North", "31", "up"], ["South", "27", "flat"], ["East", "42", "down"]]
    for i, row in enumerate(rows):
        y = top + i * pitch
        for j, word in enumerate(row):
            page.fill(x0 + j * col_w, y, x0 + (j + 1) * col_w, y + pitch, colors[i % len(colors)])
            page.words(word, x0 + j * col_w + 5, y + 13)


def test_cells_shaded_edge_to_edge_are_a_table_not_a_diagram_of_boxes():
    """\\rowcolors with no rules: every cell a filled box touching its neighbours. diagram_from
    took the boxes for nodes, the words for their labels (r2_tables_v1 slide 4)."""
    page = Page()
    shaded_grid(page, ["#c8d4e8", "#eef2f8"])
    els = page.elements()
    assert not [e for e in els if e["kind"] == "diagram"]
    (t,) = tables(els)
    assert cell_texts(t) == [["Region", "Share", "Trend"], ["North", "31", "up"], ["South", "27", "flat"], ["East", "42", "down"]]
    assert {f["color"] for f in t["fills"]} == {"#c8d4e8", "#eef2f8"}


def test_a_heatmap_under_one_hline_keeps_its_header_row():
    """A heatmap's shaded body under a single \\hline, the column heads above it unshaded: one
    wide rule is no group of table rules, and it was dropped as theme decoration."""
    page = Page()
    heads = ["Model", "A", "B", "C", "D", "E"]
    for j, word in enumerate(heads):
        page.words(word, 45 + j * 60 + 5, 60)
    page.hline(45, 45 + 6 * 60, 65)
    shades = ["#fde0dd", "#fa9fb5", "#c51b8a", "#7a0177"]
    for i, name in enumerate(["M1", "M2", "M3"]):
        y = 65 + i * 18
        page.words(name, 50, y + 13)
        for j in range(1, 6):
            page.fill(45 + j * 60, y, 45 + (j + 1) * 60, y + 18, shades[(i + j) % 4])
            page.words(f"0.{i}{j}", 45 + j * 60 + 5, y + 13)
    (t,) = tables(page.elements())
    assert cell_texts(t)[0] == heads
    assert len(t["cells"]) == 4


def test_white_stripes_on_an_off_white_page_are_kept():
    """metropolis' page is #fafafa: a \\rowcolors stripe in white is a fill of its own, not the
    page showing through (r2_tables_v3 slide 6)."""
    page = Page()
    page.fill(0, 0, W, H, "#fafafa")
    shaded_grid(page, ["#ffffff", "#e0e0e0"])
    (t,) = tables(page.elements())
    assert "#ffffff" in {f["color"] for f in t["fills"]}


# ---------------------------------------------------------------- rows and rules

def ruled_table(page: Page, rows: list[list[tuple[float, str]]], baselines: list[float], x0=50, x1=400) -> None:
    page.hline(x0, x1, baselines[0] - 12, 0.8)
    page.hline(x0, x1, baselines[0] + 4, 0.5)
    page.hline(x0, x1, baselines[-1] + 5, 0.8)
    for row, b in zip(rows, baselines):
        for x, words in row:
            page.words(words, x, b)


def test_a_multirow_head_halfway_between_widely_spaced_rows_is_merged():
    """\\multirow{2}{*}{9:00} in rows 1.6 em apart: the head sits 0.8 em from each of its rows,
    more than the 0.75 em the halfway test allowed, and came out as a row of its own."""
    page = Page()
    rows = [[(55, "Time"), (150, "Talk"), (300, "Room")],
            [(150, "Opening words"), (300, "Hall A")],
            [(55, "9:00")],
            [(150, "Keynote on reefs"), (300, "Hall A")],
            [(55, "10:30"), (150, "Coffee"), (300, "Foyer")]]
    ruled_table(page, rows, [70, 90, 98, 106, 122])
    (t,) = tables(page.elements())
    assert len(t["cells"]) == 4
    assert any(m["row"] == 1 and m["col"] == 0 and m["rows"] == 2 for m in t["merges"])
    assert cell_texts(t)[1][0] == "9:00"


def test_a_double_vertical_rule_is_one_border():
    """{l||r|r}: two vertical rules 2 pt apart at one column boundary; the second was taken for a
    rule in the middle of a cell and the table refused (r2_tables_v3 slide 4)."""
    page = Page()
    xs, top, pitch = [50, 150, 250, 350], 50, 18
    rows = [["Team", "Won", "Lost"], ["Owls", "12", "3"], ["Hawks", "9", "6"], ["Crows", "4", "11"]]
    for i, row in enumerate(rows):
        for j, word in enumerate(row):
            page.words(word, xs[j] + 6, top + i * pitch + 13)
    bottom = top + len(rows) * pitch
    for i in range(len(rows) + 1):
        page.hline(xs[0], xs[-1], top + i * pitch)
    for x in xs:
        page.vline(x, top, bottom)
    page.vline(xs[1] - 2, top, bottom)
    els = page.elements()
    (t,) = tables(els)
    assert cell_texts(t) == rows


# ---------------------------------------------------------------- heads and alignment

def alignments(t: dict) -> dict[tuple[int, int], tuple[str, float, float]]:
    """(row, col) -> (alignment, indent start, indent end) as table_requests writes them."""
    reqs = E.table_requests(t, "s", "o", E.SLIDE_W / W, E.FontMapper(), True, W)
    out = {}
    for r in reqs:
        u = r.get("updateParagraphStyle")
        if u and "alignment" in u["style"]:
            loc = u["cellLocation"]
            out[(loc["rowIndex"], loc["columnIndex"])] = (u["style"]["alignment"], u["style"]["indentStart"]["magnitude"],
                                                          u["style"]["indentEnd"]["magnitude"])
    return out


def test_an_s_columns_numbers_stay_flush_right_under_a_wider_centred_head():
    """siunitx: numbers of different widths set flush right under a head centred over them and
    wider. The column read 'center' from all its cells, and every number came out centred - the
    shorter one out of line (r3_univ_v3 slide 4)."""
    page = Page()
    rows = [[(55, "Name"), (154.2, "Salary [EUR]")], [(55, "Ciara"), (168, "48 200")], [(55, "Zeynep"), (168, "61 750")],
            [(55, "Nadia"), (172.6, "7 450")]]
    ruled_table(page, rows, [70, 86, 100, 114], x1=230)
    (t,) = tables(page.elements())
    col = t["columns"][1]
    assert (col["align"], col["head"]) == ("right", "center")
    got = alignments(t)
    assert got[(0, 1)][0] == "CENTER" and {got[(r, 1)][0] for r in (1, 2, 3)} == {"END"}
    assert len({got[(r, 1)][2] for r in (1, 2, 3)}) == 1


def test_a_multirow_head_centred_over_a_left_column_stays_centred():
    """\\thead{Deliverable} beside a two-line \\makecell: a multirow head centred over the
    body's left-aligned words took the column's left edge (r2_tables_v3 slide 8). A multirow
    word flush with its column's edge keeps that edge, however near the middle of its cell."""
    page = Page()
    page.hline(50, 330, 54, 0.8)
    page.words("No.", 57.3, 71)
    page.words("Deliverable", 150.95, 71)
    page.words("Due", 300, 65)
    page.words("month", 300, 77)
    page.hline(50, 330, 80, 0.5)
    for b, (no, what, due) in zip((92, 106, 120), (("D3.1", "Data harmonisation report", "M6"),
                                                   ("D3.2", "Open source aggregation library", "M12"),
                                                   ("D3.3", "Privacy accounting paper", "M18"))):
        page.words(no, 55, b)
        page.words(what, 100, b)
        page.words(due, 300, b)
    page.hline(50, 330, 125, 0.8)
    (t,) = tables(page.elements())
    merges = {(m["row"], m["col"]): m for m in t["merges"]}
    assert merges[(0, 1)]["align"] == "center"
    got = alignments(t)
    assert got[(0, 1)][0] == "CENTER" and got[(0, 0)][0] == "CENTER"
    assert got[(2, 1)][0] == "START"


def test_the_cells_beside_a_multirow_keep_their_rows_top_inset():
    """{9:00} over two rows with 'WP1 status' beside it: the row holding a \\multirow got no top
    inset at all, and 'WP1 status' sat against the rule above it (r2_tables_v3 slide 7). The
    row keeps its inset; the \\multirow cell alone, centred in its rows, has none."""
    from pptx import Presentation
    page = Page()
    rows = [[(55, "Time"), (150, "Track"), (300, "Room")],
            [(150, "WP1 status"), (300, "WP2 status")],
            [(55, "9:00")],
            [(150, "WP3 status"), (300, "WP4 status")],
            [(55, "10:30"), (150, "Coffee"), (300, "Foyer")]]
    ruled_table(page, rows, [70, 90, 98, 106, 122])
    (t,) = tables(page.elements())
    assert any(m["row"] == 1 and m["col"] == 0 and m["rows"] == 2 for m in t["merges"])
    scale = E.SLIDE_W / W
    lay = E.table_layout(t, scale, E.FontMapper(), imported=True, page_w=W)
    assert lay["insets"][1] > 1.0  # (the space under the rule above it, as every row has)
    table = E.pptx_table(t, scale, E.FontMapper(), page_w=W)
    assert table["middle"] == [[1, 0]]
    prs = Presentation(E.build_pptx(W, H, [], [{"layout": "BLANK", "fill": None, "pictures": [], "tables": [table],
                                                "templates": False}], {"color": "#ffffff"}))
    grid = next(s for s in prs.slides[0].shapes if s.has_table).table
    assert grid.cell(1, 0).margin_top == 0 and grid.cell(1, 1).margin_top == round(table["margins"][1][1] * E.EMU_PER_PT)


def test_a_multirow_word_flush_left_in_a_narrow_column_is_not_centred():
    """'Sparse' over two rows of a narrow left column sits near its cell's middle; it is flush
    with the column's words, so it stays flush left (r2_tables_v1 slide 3)."""
    page = Page()
    rows = [[(55, "Family"), (100, "Model"), (250, "Score")],
            [(100, "BM25"), (250, "18.7")],
            [(55, "Sparse")],
            [(100, "SPLADE"), (250, "38.0")],
            [(55, "Dense"), (100, "ANCE"), (250, "33.0")]]
    ruled_table(page, rows, [70, 86, 93, 100, 114], x1=280)
    (t,) = tables(page.elements())
    merges = {(m["row"], m["col"]): m for m in t["merges"]}
    assert merges[(1, 0)]["align"] == "left"
    assert alignments(t)[(1, 0)][0] == "START"

def test_a_justified_cell_at_a_loose_pitch_is_one_cell_of_two_lines():
    """A justified p{} cell whose lines are 1.45 em apart (\\arraystretch): the second line, lower
    case, under a line that reaches the column's edge, continues the cell (r2_tables_v3 slide 1)
    - it was read as a row of its own, its words split into columns."""
    page = Page()
    rows = [[(55, "Term"), (140, "Meaning")],
            [(55, "Reef"), (140, "a ridge of rock and coral near")],
            [(140, "the surface of the sea")],
            [(55, "Atoll"), (140, "a ring shaped reef")]]
    ruled_table(page, rows, [70, 88.5, 103, 121.5], x1=320)
    # (justified: the first line of the cell reaches the column's right edge)
    (t,) = tables(page.elements())
    assert cell_texts(t) == [["Term", "Meaning"], ["Reef", "a ridge of rock and coral near the surface of the sea"],
                             ["Atoll", "a ring shaped reef"]]
    assert t["row_lines"] == [1, 2, 1]


def test_cells_wrapping_in_two_columns_at_once_stay_their_rows():
    """Two p{} cells of one row that both wrap: their second lines make a line of two phrases
    with nothing in the first column (r2_tables_v1 slide 2), each lower case under a phrase of
    words - not a row of its own."""
    page = Page()
    rows = [[(55, "Method"), (130, "Key idea"), (270, "Main limit")],
            [(55, "BM25"), (130, "term counts with a"), (270, "no notion of synonyms")],
            [(130, "length norm"), (270, "at all")],
            [(55, "DPR"), (130, "dual encoders"), (270, "needs labels")]]
    ruled_table(page, rows, [70, 86, 98, 114])
    (t,) = tables(page.elements())
    assert cell_texts(t) == [["Method", "Key idea", "Main limit"],
                             ["BM25", "term counts with a length norm", "no notion of synonyms at all"],
                             ["DPR", "dual encoders", "needs labels"]]
    assert t["row_lines"] == [1, 2, 1]


# ---------------------------------------------------------------- no table

def test_wrapped_items_of_side_by_side_lists_are_no_plain_table():
    """Two columns of bullet lists whose items wrap at the same heights: the continuation lines,
    one short phrase in each column, lined up like the rows of a rule-less table
    (r3_layout_v1 slide 8) and were taken out of their items."""
    page = Page()
    columns = [("Hardware", [["SD cards fail first"], ["Thermal throttling", "above 70 C"]]),
               ("Network", [["Wi-Fi backhaul is", "bursty"], ["NAT timeouts kill", "idle TCP"]]),
               ("Software", [["Clock skew breaks", "windows"], ["Containers too", "heavy on Pi"]]),
               ("Process", [["Test failover weekly"], ["Log everything", "centrally"]])]
    images = []
    for k, (head, items) in enumerate(columns):
        x = 32.7 + 109.4 * k
        page.words(head, x - 22, 102.7, 10.9)
        b = 117.6
        for lines in items:
            images.append({"id": f"p0i{len(images)}", "bbox": [x - 10.7, b - 5.6, x - 4.7, b + 0.4], "px": [6, 6]})
            for line in lines:
                page.words(line, x, b)
                b += 12.0
            b += 2.9
    raw = page.raw
    page.raw = lambda: {**raw(), "images": images}
    els = page.elements()
    assert not tables(els)
    words = " ".join("".join(r["text"] for r in p["runs"]) for e in els if e["kind"] == "text" for p in e["paragraphs"])
    assert "Wi-Fi backhaul is bursty" in " ".join(words.split())


def test_a_node_split_by_a_vertical_line_is_no_diagram_node():
    """A record node (or a ruled table row diagram_from took) with words on both sides of a
    vertical line crossing it: one label would run them together."""
    def sp(text, x0, baseline=100.0):
        return Span(id=text, text=text, font=SANS, size=SIZE, color="#000000",
                    rect=Rect(x0, baseline - 7.5, x0 + len(text) * 5, baseline + 2.5),
                    baseline=baseline, horizontal=True, info=font_info(SANS))

    node = {"bbox": [50, 90, 250, 106], "shape": "rect"}
    line = {"from": [150, 90], "to": [150, 106]}
    words = [sp("left", 60), sp("right", 160)]
    assert PageClassifier.splits_cells({"nodes": [node], "lines": [line]}, words)
    assert not PageClassifier.splits_cells({"nodes": [node], "lines": [line]}, [sp("left", 60), sp("more", 90)])
    assert not PageClassifier.splits_cells({"nodes": [node], "lines": [{**line, "via": [[150, 80]]}]}, words)


# ---------------------------------------------------------------- emit

def run(text: str, font: str = "LMRoman10-Regular", family: str = "serif", **kw) -> dict:
    return {"text": text, "font": font, "family": family, "size": SIZE, "color": "#000000", "bold": False,
            "italic": False, "smallcaps": False, "link": None, "script": None, **kw}


def wrapped_table(font: str, family: str) -> dict:
    """Two columns; the second's middle cell wraps after 'near' in the PDF."""
    cell = "a ridge of rock and coral near the surface"
    starts = [cell.index("the")]
    return {"id": "t", "kind": "table", "frame": [50, 60, 300, 110], "size": SIZE,
            "row_baselines": [70, 85, 110], "row_heights": [15, 25, 15], "row_lines": [1, 2, 1],
            "wrapped": [[1, 1, starts]],
            "columns": [{"x0": 52, "x1": 90, "align": "left"}, {"x0": 100, "x1": 240, "align": "left"}],
            "bounds": [50, 95, 300],
            "cells": [[[run("Term", font, family)], [run("Meaning", font, family)]],
                      [[run("Reef", font, family)], [run(cell, font, family)]],
                      [[run("Atoll", font, family)], [run("a ring", font, family)]]],
            "merges": [], "rules": [], "borders": [], "fills": []}


def test_a_wrapped_cells_column_does_not_take_the_next_lines_first_word():
    """A wrapped p{} cell keeps its PDF pitch: its row stays two lines tall. The column had
    room for the first line and 'the' too, so Slides set the cell in one line under an empty
    one (r1_lang_v3 slide 3). The column's text room now stops short of that."""
    fonts = E.FontMapper()
    scale = E.SLIDE_W / W
    for font, family in (("LMRoman10-Regular", "serif"), ("Cambria", "serif")):  # (PT Serif, Caladea)
        el = wrapped_table(font, family)
        lay = E.table_layout(el, scale, fonts, imported=True, page_w=W)
        _, start, end = alignments(el)[(1, 1)]
        room = lay["widths"][1] - 2 * E.TABLE_CELL_PAD - start - end
        runs = lay["cells"][1][1]
        joined = E.wrap_joins(runs, el["wrapped"][0][2], scale, fonts)
        if joined is not None:  # measured: the first line fits, the first line and 'the' do not
            assert E.wrapped_width(runs, el["wrapped"][0][2], scale, fonts) < room < joined
        else:  # a font the PDF itself uses: little more than the PDF's own column
            assert room <= (240 - 100) * scale * 1.03


def test_a_number_cell_is_set_at_the_tables_size():
    """shape_ratio shrank a number-only run by Lato's digits against Computer Modern's; in a
    top-anchored cell the smaller number rode high beside its row's words."""
    fonts = E.FontMapper()
    scale = E.SLIDE_W / W
    el = wrapped_table("LMSans10-Regular", "sans")
    el["cells"][2][1] = [run("123456789012345678", "LMSans10-Regular", "sans")]
    lay = E.table_layout(el, scale, fonts, imported=True, page_w=W)
    sizes = {fonts(r, scale)[1] for row in lay["cells"] for c in row for r in c}
    assert len(sizes) == 1


# ---------------------------------------------------------------- wave 3: wrapped cells and width

def three_columns(cell: str, second_line: str, col1: tuple[float, float] = (100, 240),
                  frame: tuple[float, float] = (50, 305)) -> dict:
    """Term | Meaning | Note, the middle cell of row 1 wrapped where `second_line` starts (in the
    middle of a word TeX hyphenated, or at a word)."""
    starts = [cell.index(second_line)]
    fx0, fx1 = frame
    return {"id": "t", "kind": "table", "frame": [fx0, 60, fx1, 110], "size": SIZE,
            "row_baselines": [70, 85, 110], "row_heights": [15, 25, 15], "row_lines": [1, 2, 1],
            "wrapped": [[1, 1, starts]],
            "columns": [{"x0": fx0 + 2, "x1": fx0 + 30, "align": "left"},
                        {"x0": col1[0], "x1": col1[1], "align": "left"},
                        {"x0": col1[1] + 32, "x1": fx1 - 5, "align": "left"}],
            "bounds": [fx0, (fx0 + 30 + col1[0]) / 2, col1[1] + 16, fx1],
            "cells": [[[run("Term")], [run("Meaning")], [run("Note")]],
                      [[run("Reef")], [run(cell)], [run("x")]],
                      [[run("Atoll")], [run("a ring")], [run("y")]]],
            "merges": [], "rules": [], "borders": [], "fills": []}


def lines_at(runs: list[dict], room: float, scale: float, fonts) -> int:
    """How many lines Slides breaks these runs into at this text room."""
    return next(n for n in range(1, 10) if E.wrap_window(runs, n, scale, fonts)[0] <= room)


def cell_room(el: dict, lay: dict, r: int, c: int) -> float:
    _, start, end = alignments(el)[(r, c)]
    return lay["widths"][c] - 2 * E.TABLE_CELL_PAD - start - end


def test_a_word_hyphenated_at_a_line_end_goes_down_a_line_instead_of_widening_its_column():
    """'No notion of synonyms or para-/phrase': the column was made wide enough for the first
    line with 'paraphrase' whole, 30 pt past the PDF's, and the table ran past its frame with its
    text set 13% smaller (r2_tables_v1 slide 2). Slides breaks before 'paraphrase' in the PDF's
    column and keeps the cell's two lines."""
    fonts, scale = E.FontMapper(), E.SLIDE_W / W
    el = three_columns("No notion of synonyms or paraphrase at all", "phrase")
    lay = E.table_layout(el, scale, fonts, imported=True, page_w=W)
    assert lay["shrink"] == 1.0
    assert lay["bounds"][2] <= el["bounds"][2] + 0.5  # (the rule after the column where the PDF has it)
    assert lines_at(lay["cells"][1][1], cell_room(el, lay, 1, 1), scale, fonts) == 2


def test_a_wrapped_cells_cap_keeps_the_rule_after_its_column_on_the_pdfs_boundary():
    """A justified line's words take less room in Slides than the PDF's stretched line: the column
    is held short of taking the next line's first word. Held by moving its right boundary left,
    the rule after it and the next column's words left the PDF's place (r2_tables_v1 slide 2,
    ~17 pt). The cells' indentEnd holds the room now."""
    fonts, scale = E.FontMapper(), E.SLIDE_W / W
    el = three_columns("a ridge of rock and coral near the surface", "coral")
    lay = E.table_layout(el, scale, fonts, imported=True, page_w=W)
    assert abs(lay["bounds"][2] - el["bounds"][2]) <= 0.5
    runs = lay["cells"][1][1]
    starts = el["wrapped"][0][2]
    room = cell_room(el, lay, 1, 1)
    assert E.wrapped_width(runs, starts, scale, fonts) < room < E.wrap_joins(runs, starts, scale, fonts)
    # (the next column's words where the PDF has them)
    assert alignments(el)[(1, 2)][1] == alignments(el)[(0, 2)][1] > 0


def test_a_centred_table_grows_into_both_margins_alike():
    """A centred table whose wrapped cells need more room in Slides than the PDF's columns (a
    hyphenated word whole on no line fits): grown to the right only, it reached the slide's edge
    with its left margin whole (r2_tables_v2 slide 10). It stays centred, its words with it."""
    fonts, scale = E.FontMapper(), E.SLIDE_W / W
    el = three_columns("Needs hand-made descriptors; poor", "tors", col1=(100, 215), frame=(50, W - 50))
    lay = E.table_layout(el, scale, fonts, imported=True, page_w=W)
    width = sum(lay["widths"])
    assert width > (el["frame"][2] - el["frame"][0]) * scale + 5  # (it did grow)
    left, right = lay["x"], E.SLIDE_W - lay["x"] - width
    assert abs(left - right) <= 0.1
    # The first column's words keep their place in the table: moved with it.
    _, start, _ = alignments(el)[(1, 0)]
    text_x = lay["x"] + E.TABLE_CELL_PAD + start
    assert text_x == pytest.approx((el["columns"][0]["x0"] + lay["dx"]) * scale + E.TABLE_CELL_PAD - E.PAD_X, abs=0.05)
    assert lines_at(lay["cells"][1][1], cell_room(el, lay, 1, 1), scale, fonts) == 2
