"""Defects the visual hunt found in what emit writes (out/hunt/verified), pinned offline: dot
leaders and runs inside a sentence sized like their neighbours, small optical cuts no taller than
the small-caps compromise, text ranges in UTF-16 units, XML-safe alt texts, table indents per
column and tables kept on the page."""

import io

import numpy as np
import pytest

from beamer2slides import emit
from beamer2slides.emit import EMU_PER_PT, SLIDE_W, table_requests

from .test_emit_requests import FONTS, SCALE, column_widths, pt_of, run_of


def text_element(*paragraphs: list[dict], bullet: dict | None = None) -> dict:
    """A text element of one-line paragraphs (runs each), 15 pt apart."""
    paras = []
    for i, runs in enumerate(paragraphs):
        b = 60.0 + 15 * i
        paras.append({"align": "left", "level": 0, "bullet": bullet and {**bullet, "bbox": [60.0, b - 6, 64.0, b - 2]},
                      "size": 10.91, "text_x0": 70.0, "tab_x0": None, "wrap_limit": None, "runs": runs,
                      "lines": [{"baseline": b, "x0": 70.0, "x1": 300.0}]})
    return {"id": "p0t0", "kind": "text", "role": "body", "code": False,
            "bbox": [60.0, 50.0, 300.0, 65.0 + 15 * len(paras)], "paragraphs": paras}


def styled(reqs: list[dict]) -> list[tuple[str, dict]]:
    """(the text each updateTextStyle range covers, its style), ranges read as Slides reads them:
    UTF-16 code units of the inserted text."""
    text = next(r["insertText"]["text"] for r in reqs if "insertText" in r)
    units = text.encode("utf-16-le", "surrogatepass")
    out = []
    for r in reqs:
        st = r.get("updateTextStyle")
        if st and st["textRange"]["type"] == "FIXED_RANGE" and "baselineOffset" in st["style"]:
            a, b = st["textRange"]["startIndex"], st["textRange"]["endIndex"]
            out.append((units[2 * a:2 * b].decode("utf-16-le", "surrogatepass"), st["style"]))
    return out


# ---------------------------------------------------------------- sizes

def test_dot_leaders_and_ellipses_keep_the_size_factor():
    # textfx v3 s2 / v2 s10 / v3 s3: \dotfill and \ldots reach the text layer as '. . . .', read
    # as sentence ends with TeX's extra space: the runs were set 1.3-1.7 times too large.
    prose = FONTS(run_of("Monitor"), SCALE)[1]
    for text in ("Welcome and coffee . . . . . . . . . . . . . . . . . . . . . . . . . .",
                 "Comments: . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . ",
                 "A: “Well . . . on the ones we tried, yes. On the others . . . we don’t",
                 "Things to bring: laptop, adapter, clicker, . . . , and a backup",
                 "A pause . . . or three dots... or an ellipsis at the end. . ."):
        assert FONTS(run_of(text), SCALE)[1] == prose, text


def test_a_run_among_others_keeps_their_size():
    # ml v3 s10, control c2 s12, sci v3 s9: an author list or a journal abbreviation of 15+
    # characters, sized to its own PDF width, came out 10% larger than the title beside it.
    authors = run_of("J. Park, W. Zhang, et al. ")
    title = run_of("Graph networks for crystal property prediction", font="CMSSI10", italic=True)
    assert FONTS(authors, SCALE)[1] != FONTS(run_of("Monitor"), SCALE)[1]  # alone: sized to its width
    el = text_element([authors, title], [run_of("J. Park, W. Zhang, et al. ")])
    sizes = [(t, pt_of(s["fontSize"])) for t, s in styled(emit.text_box_requests(el, "b2s_s001", "b2s_s001_t0", SCALE, FONTS))]
    prose = FONTS(run_of("Monitor"), SCALE)[1]
    assert sizes[0] == ("J. Park, W. Zhang, et al. ", prose)
    assert sizes[-1] == ("J. Park, W. Zhang, et al. ", FONTS(authors, SCALE)[1])  # its own paragraph: as before
    assert el["paragraphs"][0]["runs"][0] is authors and "in_sentence" not in authors  # the IR is left alone
    # In a table cell too.
    table = {"id": "p0b0", "kind": "table", "frame": [100.0, 50.0, 400.0, 70.0], "size": 10.91,
             "columns": [{"x0": 106.0, "x1": 390.0, "align": "left"}], "cells": [[[authors, {**title, "text": "Graph nets"}]]],
             "row_baselines": [60.0], "row_heights": [15.0], "rules": []}
    cell = [pt_of(r["updateTextStyle"]["style"]["fontSize"]) for r in table_requests(table, "s", "b2s_s001_b0", SCALE, FONTS)
            if "updateTextStyle" in r]
    assert cell[0] == prose


def test_a_math_letter_among_words_in_a_google_font_takes_that_font():
    # lang v2 s6 (r6): Calibri moved to the sans family (fonts.font_info), and so the β between
    # Carlito words came out in Lato Italic, heavier than its words.
    words = dict(font="ABCDEF+Calibri", family="sans")
    el = text_element([run_of("The rate ", **words), run_of("β", font="CMMI10", family="sans", italic=True),
                       run_of(" is fitted per cohort", **words)])
    got = styled(emit.text_box_requests(el, "b2s_s001", "b2s_s001_t0", SCALE, FONTS))
    beta = next(style for text, style in got if text == "β")
    assert beta["weightedFontFamily"]["fontFamily"] == "Carlito" and beta["italic"] is True
    # Among TeX words it keeps the substitute, as before.
    el = text_element([run_of("The rate "), run_of("β", font="CMMI10", family="sans", italic=True), run_of(" is fitted")])
    beta = next(style for text, style in styled(emit.text_box_requests(el, "b2s_s001", "b2s_s001_t0", SCALE, FONTS)) if text == "β")
    assert beta["fontFamily"] == "Lato"


def test_a_small_optical_cut_is_set_no_taller_than_the_small_caps_compromise():
    # dense v2 s1: a Boadilla footline in LMRoman5 at 4.98 pt. CM's 5 pt cut is 1.38 times as wide
    # per em as the 10 pt one; sized to that width, PT Serif's letters filled the footline bar.
    footline = run_of("Econometrics II", 4.98, font="LMRoman5-Regular", family="serif")
    ten = run_of("Econometrics II", 4.98, font="LMRoman10-Regular", family="serif")
    assert FONTS(footline, SCALE)[1] / FONTS(ten, SCALE)[1] == pytest.approx(emit.OPTICAL_WIDTH_MAX, abs=0.02)
    # 8 and 9 pt cuts are matched in full, as before.
    eight = run_of("Econometrics II", 7.97, font="CMR8", family="serif")
    assert FONTS(eight, SCALE)[1] / FONTS({**eight, "font": "CMR10"}, SCALE)[1] == \
        pytest.approx(emit.DESIGN_WIDTH["serif"][8], abs=0.02)


# ---------------------------------------------------------------- wrapped lines

def column_list() -> dict:
    """themes v4 s3: two triangle items in a 0.48 textwidth column (Goettingen, a 4:3 page)."""
    def item(text: str, top: float, x1s: list[float], limit: float) -> dict:
        return {"align": "left", "level": 0, "size": 10.91, "text_x0": 34.97, "tab_x0": None, "wrap_limit": limit,
                "bullet": {"kind": "glyph", "text": "▶", "color": "#3333b3", "bbox": [21.03, top, 29.51, top + 10.91]},
                "runs": [run_of(text)],
                "lines": [{"baseline": top + 12.16 + 13.54 * i, "x0": 34.97, "x1": x1} for i, x1 in enumerate(x1s)]}
    return {"id": "p2t1", "kind": "text", "role": "body", "code": False, "bbox": [21.03, 70.06, 144.54, 155.0],
            "paragraphs": [item("Mass bleaching in 2016, 2017, 2020, 2022 and 2024", 70.06, [144.54, 133.87, 56.79], 158.42),
                           item("Recovery between events now takes longer than the gap between", 113.7,
                                [116.29, 143.5, 133.92], 147.48)]}


def test_a_wrapped_line_has_room_for_its_words_as_slides_sets_them():
    # themes v4 s3, v2 s5, v3 s8: the box ended 2.9 pt past the widest PDF line (half the room to
    # the next word); Lato set 'events now takes longer' 7.5 pt wider than TeX did, the item
    # wrapped one word early and grew a line, over the paragraph below.
    scale = SLIDE_W / 362.83
    el = column_list()
    reqs = emit.text_box_requests(el, "b2s_s003", "b2s_s003_t1", scale, FONTS)
    props = next(r["createShape"] for r in reqs if "createShape" in r)["elementProperties"]
    right = props["transform"]["translateX"] / EMU_PER_PT + pt_of(props["size"]["width"]) - emit.PAD_X
    assert [emit.pdf_line_breaks(p) for p in el["paragraphs"]] == [[24, 45], [17, 41]]
    for p in el["paragraphs"]:
        text = p["runs"][0]["text"]
        cuts = [0, *emit.pdf_line_breaks(p), len(text)]
        for k, (a, b) in enumerate(zip(cuts, cuts[1:])):
            line = text[a:b].rstrip()
            x0 = p["lines"][k]["x0"] * scale
            assert x0 + emit.slides_width([run_of(line)], scale, FONTS) <= right - emit.WRAP_MARGIN, line
            if b < len(text):  # and the next word still goes onto the next line
                joined = text[a:text.find(" ", b) if " " in text[b:] else len(text)]
                assert x0 + emit.slides_width([run_of(joined)], scale, FONTS) > right, joined
    # A paragraph whose words do not come out as its lines (a hyphenated line end) is not measured.
    hyphen = column_list()["paragraphs"][1]
    hyphen["lines"][0]["x1"] = 140.0
    assert emit.pdf_line_breaks(hyphen) is None


def box_right(el: dict, scale: float) -> float:
    """Where the box's text ends (Slides pt): its right edge less the inset."""
    props = next(r["createShape"] for r in emit.text_box_requests(el, "b2s_s003", "b2s_s003_t1", scale, FONTS)
                 if "createShape" in r)["elementProperties"]
    return props["transform"]["translateX"] / EMU_PER_PT + pt_of(props["size"]["width"]) - emit.PAD_X


def test_a_full_line_keeps_its_margin_when_a_neighbours_next_word_is_close():
    # dense v1 s3, themes v4 s4 (r6): another paragraph's next word would join 1 pt past the
    # widest line, so the box ended 1.0 pt past it (WRAP_MARGIN) - and Slides, which sets a line
    # up to 0.6 pt wider than slides_width, wrapped "if" over "any." and a TOC entry's "Design".
    scale = SLIDE_W / 362.83
    el = column_list()
    widest, joins = zip(*(emit.slides_lines(p, scale, FONTS) for p in el["paragraphs"]))
    line = "Mass bleaching in 2016"  # a one-line paragraph ending 1 pt short of the join
    x0 = (min(joins) - 1.0 - emit.slides_width([run_of(line)], scale, FONTS)) / scale
    el["paragraphs"].append({**el["paragraphs"][0], "bullet": None, "text_x0": x0, "wrap_limit": None, "runs": [run_of(line)],
                             "lines": [{"baseline": 170.0, "x0": x0, "x1": x0 + 100.0}]})
    need = max(max(widest), min(joins) - 1.0)
    assert box_right(el, scale) >= need + emit.LINE_MARGIN - 0.01


def test_an_unmeasured_paragraph_leaves_the_box_as_wide_as_the_measured_lines_need():
    # figures v1 s2 (r6): one item could not be measured, so the box came from the PDF's
    # extents and ended 0.01 pt short of a one-line item's words as Slides sets them: it wrapped.
    scale = SLIDE_W / 362.83
    el = column_list()
    digits = "2016, 2017, 2020, 2022 and 2024"  # Lato's digits are wider than CM's
    wide = run_of(digits)
    extent = emit.pdf_width([wide])
    el["paragraphs"][0]["runs"] = [{**run_of(el["paragraphs"][0]["runs"][0]["text"]), "font": "ArialMT"}]
    assert emit.slides_lines(el["paragraphs"][0], scale, FONTS) is None
    el["paragraphs"].append({**el["paragraphs"][1], "bullet": None, "wrap_limit": None, "runs": [wide],
                             "lines": [{"baseline": 170.0, "x0": 34.97, "x1": 34.97 + extent}]})
    need = 34.97 * scale + emit.slides_width([wide], scale, FONTS)
    assert need > (el["paragraphs"][0]["lines"][0]["x1"]) * scale  # wider than any PDF line
    assert box_right(el, scale) >= need + emit.LINE_MARGIN - 0.01


def test_a_thin_space_and_a_math_symbol_leave_a_paragraph_measurable():
    # figures v1 s2: "1\,mM" reaches the text as "1 mM", a word space where TeX put a thin one,
    # and the line came out 1.8 pt wider than its extent. econ v3 s6: "≈", "×" and a math-italic
    # "." have no CM advances. Either left the whole paragraph unmeasured, sized from the PDF.
    gfp = {"lines": [{"x0": 255.1, "x1": 436.88, "baseline": 50.02}, {"x0": 255.1, "x1": 308.89, "baseline": 63.57}],
           "runs": [run_of("GFP (green): lac reporter, induced with 1 mM IPTG", font="LMSans10-Regular")]}
    assert emit.pdf_line_breaks(gfp) == [40]
    sans, oblique = dict(font="LMSans10-Regular", size=9.96), dict(font="LMSans10-Oblique", size=9.96, italic=True)
    runs = [run_of("ITT; LATE via ", **sans), run_of("D̄", **oblique), run_of("v", **{**oblique, "font": "LMSans8-Oblique"}, script="sub"),
            run_of(" instrument is ", **sans), run_of("≈", size=9.96, font="LMMathSymbols10-Regular"), run_of(" 1", **sans),
            run_of(".", size=9.96, font="LMMathItalic10-Regular"), run_of("2", **sans), run_of("× ", size=9.96, font="LMMathSymbols10-Regular"),
            run_of("larger", **sans)]
    econ = {"lines": [{"x0": 207.58, "x1": 332.04, "baseline": 180.45}, {"x0": 207.58, "x1": 274.45, "baseline": 192.41}], "runs": runs}
    scale = SLIDE_W / 362.83
    assert emit.pdf_line_breaks(econ) is None  # CM advances alone cannot say
    assert emit.pdf_line_breaks(econ, scale, FONTS) == [29]  # "... instrument " | "is ≈ 1.2× larger"
    assert emit.slides_lines(econ, scale, FONTS) is not None
    # (a combining macron has no advance of its own in Slides either)
    assert emit.slides_width([run_of("D̄", **oblique)], scale, FONTS) == emit.slides_width([run_of("D", **oblique)], scale, FONTS)


# ---------------------------------------------------------------- UTF-16 ranges

def test_text_ranges_count_utf16_units():
    # fonts newtx s4, scripts el s3: 𝛽 (U+1D6FD) and \mathbb's 𝔼 (U+1D53C) are two UTF-16 units;
    # counted as one, every later style range started a unit early and split the next pair.
    runs = [run_of("For "), run_of("𝔼", font="MSBM10", family="math"), run_of("[aX + b] = a "),
            run_of("𝔼", font="MSBM10", family="math"), run_of("[X] with "), run_of("𝛽", font="CMMI10", italic=True),
            run_of("⊤", font="CMSY7", script="super"), run_of(" upright.")]
    el = text_element(runs, [run_of("ends on "), run_of("𝔼", font="MSBM10", family="math")],
                      bullet={"kind": "glyph", "text": "•", "color": "#000000"})
    reqs = emit.text_box_requests(el, "b2s_s001", "b2s_s001_t0", SCALE, FONTS)
    got = styled(reqs)
    want = [r["text"] for r in runs] + ["ends on ", "𝔼"]
    assert [t for t, _ in got] == want
    assert [s["baselineOffset"] for t, s in got if t == "⊤"] == ["SUPERSCRIPT"]
    for r in reqs:  # every range ends inside the text, never inside a surrogate pair
        rng = (r.get("updateParagraphStyle") or r.get("createParagraphBullets") or {}).get("textRange")
        if rng:
            text = next(q["insertText"]["text"] for q in reqs if "insertText" in q)
            assert rng["endIndex"] <= emit.u16(text) + 1


def test_table_cell_ranges_count_utf16_units():
    table = {"id": "p0b0", "kind": "table", "frame": [100.0, 50.0, 300.0, 70.0], "size": 10.91,
             "columns": [{"x0": 106.0, "x1": 290.0, "align": "left"}],
             "cells": [[[run_of("𝔼", font="MSBM10", family="math"), run_of("[X]", italic=True, font="CMMI10")]]],
             "row_baselines": [60.0], "row_heights": [15.0], "rules": []}
    ranges = [(r["updateTextStyle"]["textRange"]["startIndex"], r["updateTextStyle"]["textRange"]["endIndex"])
              for r in table_requests(table, "s", "b2s_s001_b0", SCALE, FONTS) if "updateTextStyle" in r]
    assert ranges == [(0, 2), (2, 5)]


def test_the_emit_models_count_utf16_units():
    from .slides_sim import joined, units
    assert units("a𝔼b") == ["a", "\ud835", "\udd3c", "b"] and joined(units("a𝔼b")) == "a𝔼b"


# ---------------------------------------------------------------- the .pptx

def test_control_characters_in_an_alt_text_do_not_break_the_pptx(tmp_path):
    # econ v1: a Type 3 T1 font without ToUnicode gives raw codes (0x10 and 0x11 are quotes, 0x1d
    # a ligature); a figure's alt text holding one made lxml refuse the .pptx: no deck at all.
    from PIL import Image
    from pptx import Presentation
    png = tmp_path / "f.png"
    Image.fromarray(np.zeros((4, 4, 3), dtype=np.uint8)).save(png)
    alt = "\x10A living wage\x11 is a \x1door\x00 \ud835 plan"
    page = {"layout": "BLANK", "fill": None, "templates": False, "tables": [],
            "pictures": [{"file": png, "bbox": [10, 10, 50, 50], "alt": alt, "title": "b2s:\x1a1/f0"}]}
    prs = Presentation(emit.build_pptx(453.54, 340.16, [], [page], {"color": "#ffffff"}))
    pic = next(s for s in prs.slides[0].shapes if s.shape_type == 13)
    assert pic._element.nvPicPr.cNvPr.get("descr") == "A living wage is a oor  plan"
    assert pic._element.nvPicPr.cNvPr.get("title") == "b2s:1/f0"
    assert emit.xml_text("tab\tnew\nline\r") == "tab\tnew\nline\r"


# ---------------------------------------------------------------- tables

def signed_table() -> dict:
    """econ v2 s9: a tight right-aligned column of signed estimates under a wider header."""
    rows = [["Variable", "Estimate"], ["Wage", "-0.021"], ["Tenure", "0.034"], ["Union", "-12.345"]]
    return {"id": "p0b0", "kind": "table", "frame": [100.0, 50.0, 196.0, 110.0], "size": 10.91,
            "columns": [{"x0": 100.0, "x1": 140.0, "align": "left"}, {"x0": 150.0, "x1": 190.0, "align": "right"}],
            "bounds": [100.0, 145.0, 196.0],
            "cells": [[[run_of(t)] for t in row] for row in rows],
            "row_baselines": [60.0, 75.0, 90.0, 105.0], "row_heights": [15.0] * 4, "rules": []}


def indents(reqs: list[dict], col: int) -> list[tuple[str, float, float]]:
    return [(p["style"]["alignment"], p["style"]["indentStart"]["magnitude"], p["style"]["indentEnd"]["magnitude"])
            for r in reqs if (p := r.get("updateParagraphStyle")) and p["cellLocation"]["columnIndex"] == col]


def test_a_right_aligned_column_keeps_one_indent():
    # Capped cell by cell, every cell of a tight column reached its own cap and started at the
    # same x: right alignment became left alignment.
    reqs = table_requests(signed_table(), "s", "b2s_s001_b0", SCALE, FONTS)
    got = indents(reqs, 1)
    assert len(set(got)) == 1 and got[0][0] == "END" and got[0][2] > 0
    widths = column_widths(reqs)
    widest = max(emit.slides_width([run_of(t)], SCALE, FONTS) for t in ("Estimate", "-0.021", "0.034", "-12.345"))
    assert widths[1] - 2 * emit.TABLE_CELL_PAD - got[0][2] >= widest  # and the widest still fits


def test_an_unmeasured_font_cell_keeps_room_for_its_words():
    # scripts ruxe s5, lang v2 s6: Arial and Calibri have no measured advances; the alignment
    # indent was not capped at all, and the widest header wrapped.
    table = signed_table()
    table["cells"] = [[[{**r, "font": "ArialMT"} for r in cell] for cell in row] for row in table["cells"]]
    reqs = table_requests(table, "s", "b2s_s001_b0", SCALE, FONTS)
    got = indents(reqs, 1)
    widths = column_widths(reqs)
    assert len(set(got)) == 1
    assert widths[1] - 2 * emit.TABLE_CELL_PAD - emit.WRAP_MARGIN - got[0][2] >= (190.0 - 150.0) * SCALE - 0.01


def test_wide_characters_are_an_em():
    # scripts ja s5: 'マクロ F1' - kana measured at 0.6 em, as an unmeasured Latin letter - wrapped.
    z = FONTS(run_of("マクロ"), SCALE)[1]
    assert emit.slides_width([run_of("マクロ")], SCALE, FONTS) == pytest.approx(3 * z)


def eleven_columns() -> dict:
    """tables v1 s10 (its first rows): eleven 8 pt columns with 3 pt \\tabcolsep, the table just
    inside a 4:3 page (x 15-348 of 363)."""
    edges = [(18.25, 49.79), (55.79, 81.31), (87.31, 111.57), (117.54, 139.65), (145.62, 167.72),
             (173.7, 195.8), (201.85, 227.41), (233.33, 255.44), (261.5, 283.96), (289.96, 317.0), (323.0, 344.56)]
    cols = [{"x0": a, "x1": b, "align": "center" if i else "left"} for i, (a, b) in enumerate(edges)]

    def cell(t: str) -> list[dict]:
        run = lambda s, f="LMSans8-Regular": run_of(s, 7.97, font=f)  # noqa: E731
        return [run("−", "LMMathSymbols8-Regular"), run(t[1:])] if t.startswith("−") else [run(t)] if t else []
    rows = [["Model", "TREC-", "NF-", "Fi", "Argu-", "Sci-", "Touché", "DB-", "SCI-", "Climate", "Quora"],
            ["", "COVID", "Corpus", "QA", "Ana", "Fact", "2020", "Pedia", "DOCS", "FEVER", ""],
            ["DPR", "−67%", "−30%", "−6%", "−58%", "−65%", "−37%", "−24%", "−47%", "+24%", "−0%"]]
    return {"id": "p9tab0", "kind": "table", "frame": [15.26, 95.31, 347.58, 134.0], "size": 7.97, "columns": cols,
            "bounds": [15.26, 52.79, 84.31, 114.56, 142.63, 170.71, 198.82, 230.37, 258.47, 286.96, 320.0, 347.58],
            "cells": [[cell(t) for t in row] for row in rows], "row_baselines": [105.52, 114.99, 130.09],
            "row_heights": [9.47, 15.1, 9.46], "rules": [],
            "merges": [{"row": 0, "col": 0, "rows": 2, "cols": 1, "align": "left"},
                       {"row": 0, "col": 10, "rows": 2, "cols": 1, "align": "center"}]}


def test_a_wide_table_stays_on_the_page():
    # tables v1 s10, v2 s3: room for the substitute and both cell paddings in every column added
    # up past the page edge, and the last column was lost off the slide.
    page_scale = SLIDE_W / 362.83
    el = eleven_columns()
    assert el["frame"][2] < 362.83
    lay = emit.table_layout(el, page_scale, FONTS, imported=True)
    assert lay["bounds"][-1] <= 362.83 - emit.TABLE_MARGIN * lay["bounds"][0] + 0.02
    assert emit.TABLE_MIN_SHRINK <= lay["shrink"] < 1
    # Every cell still fits its column on one line, at the size it is written in.
    for (r, c), w in lay["cell_width"].items():
        assert w + 2 * emit.TABLE_CELL_PAD <= lay["widths"][c] + 0.01, (r, c)
    assert lay["cells"][0][0][0]["size"] == pytest.approx(7.97 * lay["shrink"], rel=1e-3)
    assert el["cells"][0][0][0]["size"] == 7.97  # the IR is left alone
    reqs = table_requests(el, "s", "b2s_s001_b0", page_scale, FONTS, imported=True)
    written = {pt_of(r["updateTextStyle"]["style"]["fontSize"]) for r in reqs
               if "updateTextStyle" in r and r["updateTextStyle"]["textRange"]["type"] == "FIXED_RANGE"}
    full = {FONTS(r, page_scale)[1] for row in el["cells"] for c in row for r in c}
    assert max(written) < max(full) and min(written) < min(full)
    assert emit.pptx_table(el, page_scale, FONTS)["widths"] == lay["widths"]
    # A table that fits keeps its size and its room.
    ok = signed_table()
    assert emit.table_layout(ok, SCALE, FONTS)["shrink"] == 1.0
