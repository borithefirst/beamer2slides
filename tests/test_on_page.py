"""Nothing emit writes reaches past its slide's page where what it shows does not (`emit_text.on_page`,
`flush_right`, the number box's and a diagram label's narrowing, an arc's upright or elliptical copies): one family per
test, on synthetic decks, each planned twice - on its page and on a page too large to matter, as
emit planned before - so a test says both that the box came in and that the words did not move
or wrap by our own line model (`text_layout`, on the read-back `slides_sim` makes of the plan)."""

import math

import pytest

from pptx import Presentation

from beamer2slides import curves, emit, ir_types, text_layout
from beamer2slides.emit_diagrams import ELLIPSE_SPAN, ELLIPSE_TOLERANCE, KINK, LABEL_ROOM, MAX_ELLIPSES, arc_plan, \
    diagram_requests_of, element_template_keys, element_template_keys_on, ellipse_axes, ellipse_pieces, fine_step, \
    label_width_on, upright_at
from beamer2slides.emit_model import Template, TemplateKey, number_box_of
from beamer2slides.emit_pptx import NS_A, _add_template_shapes, arc_heads
from beamer2slides.emit_text import LINE_MARGIN, PAD_X, Page, number_box_requests_of
from beamer2slides.google_types import presentation, slides_json
from beamer2slides.gslides import EMU_PER_PT
from beamer2slides.ir_types import DiagramElement
from beamer2slides.json_types import Json, JsonObject
from beamer2slides.snapshot import read_presentation

from . import slides_sim
from .json_reads import jat, jnum, jobj, jobjs, jstr
from .test_emit_requests import FONTS, OFF_PAGE_TOLERANCE, Emitted, box_of, off_page, pt_of
from .test_sync import readback, run, text_ir

SCALE = 2.0                  # a 360 x 270 pt PDF page on a 720 pt slide
PAGE = Page(width=720.0, height=540.0)
Box4 = tuple[float, float, float, float]
Lines = dict[str, tuple[tuple[int, int, float], ...]]


def deck_of(*elements: JsonObject) -> JsonObject:
    listed: list[Json] = [e for e in elements]
    return {"source": {"pdf": "nowhere.pdf"}, "slides": [{"page": 0, "size": [360.0, 270.0], "elements": listed}]}


def huge(*, width: float, height: float) -> Page:
    """A page no box reaches past: emit's plan as it was before boxes were kept on theirs."""
    return Page(width=1e6, height=1e6)


def laid_out(deck: JsonObject) -> Lines:
    """Each text box's lines as `text_layout` sets them: (start, end, x where its words start)."""
    pres = read_presentation(presentation(slides_sim.presentation_of(deck), "the simulated deck"))
    out: Lines = {}
    for slide in jobjs(pres, "slides"):
        for oid, rb in jobj(slide, "objects").items():
            lay = text_layout.layout(jobj(rb))
            if lay is not None:
                out[oid] = tuple((ln.start, ln.end, ln.x) for ln in lay.lines)
    return out


def both_ways(deck: JsonObject, monkeypatch: pytest.MonkeyPatch) -> tuple[Emitted, Emitted, Lines, Lines]:
    """(on its page, unbounded) replays and line layouts of a deck."""
    bounded, lines = Emitted("on its page", deck), laid_out(deck)
    with monkeypatch.context() as m:
        m.setattr(emit, "Page", huge)
        free, free_lines = Emitted("unbounded", deck), laid_out(deck)
    return bounded, free, lines, free_lines


def frame(d: Emitted, oid: str) -> Box4:
    got = d.frame(oid)
    assert got is not None, f"{oid} has no frame"
    return got


def assert_same_lines(lines: Lines, free_lines: Lines, oid: str, moved: float) -> None:
    """The same line breaks, each line's words starting within `moved` pt of where they did."""
    a, b = lines[oid], free_lines[oid]
    assert [ln[:2] for ln in a] == [ln[:2] for ln in b], "rewrapped"
    assert all(abs(x - y) <= moved for (_, _, x), (_, _, y) in zip(a, b)), (a, b)


def test_a_footline_number_at_the_edge_is_set_right_aligned_inside_the_page(monkeypatch: pytest.MonkeyPatch) -> None:
    """'1 / 12' ending 8 pt from the page's edge: a left-aligned box must keep LINE_MARGIN past its
    words and a single line's slack, which ran 40 pt off the page (12,259 footers live). Its
    lines end together and Slides sets them as wide as the PDF does, so it is written
    right-aligned, the box ending PAD_X past the words: they end where they did."""
    words = run("1 / 12")
    w = emit.slides_width([words], SCALE, FONTS)
    assert w is not None
    right = 355.8
    foot = text_ir("1 / 12", [right - w / SCALE, 255.0, right, 265.0], "foot", role="footer")
    for p in jobjs(foot, "paragraphs"):
        del p["tab_x0"], p["wrap_limit"]  # (a footer is the theme's: its paragraphs have neither)
    deck = deck_of(foot)
    d, free, lines, free_lines = both_ways(deck, monkeypatch)
    oid = "b2s_s000_t0"
    assert frame(free, oid)[2] > PAGE.width + 10, "the old box ran off the page"
    x0, _, x1, _ = frame(d, oid)
    assert x1 <= PAGE.width + 0.01 and x1 == pytest.approx(right * SCALE + PAD_X, abs=0.01)
    styles = [jobj(r, "updateParagraphStyle", "style") for r in d.part_requests("b2s_s000", "foot")
              if "updateParagraphStyle" in r]
    assert styles and all(s.get("alignment") == "END" for s in styles if "alignment" in s)
    assert x1 - x0 - 2 * PAD_X >= w + LINE_MARGIN - 0.01, "room for the words"
    assert not off_page(d, OFF_PAGE_TOLERANCE)
    assert_same_lines(lines, free_lines, oid, 1.0)


def test_left_aligned_lines_give_up_the_room_past_their_words(monkeypatch: pytest.MonkeyPatch) -> None:
    """Two single lines whose box reached 40 pt past the page (a single line's slack): the box
    ends at the page's edge, its left edge and every word where they were."""
    deck = deck_of(text_ir("A line of words\nAnd another one", [20.0, 60.0, 330.0, 84.0], "body", role="body"))
    d, free, lines, free_lines = both_ways(deck, monkeypatch)
    oid = "b2s_s000_t0"
    before, after = frame(free, oid), frame(d, oid)
    assert before[2] > PAGE.width + 10
    assert after[0] == pytest.approx(before[0], abs=0.01) and after[2] <= PAGE.width + 0.01
    assert not off_page(d, OFF_PAGE_TOLERANCE)
    assert_same_lines(lines, free_lines, oid, 0.01)


def test_a_title_keeps_to_the_page(monkeypatch: pytest.MonkeyPatch) -> None:
    """A frame title in the layout's TITLE placeholder (moved and sized by emit) as wide as the
    page's text: its box ends at the page, its words where they were."""
    deck = deck_of(text_ir("A Title That Runs Wide", [20.0, 10.0, 340.0, 30.0], "title", role="title"))
    d, free, lines, free_lines = both_ways(deck, monkeypatch)
    oid = "b2s_s000_t0"
    assert frame(free, oid)[2] > PAGE.width + 1
    assert not off_page(d, OFF_PAGE_TOLERANCE)
    assert_same_lines(lines, free_lines, oid, 0.01)


def test_a_box_near_the_bottom_ends_at_the_page() -> None:
    """A line whose baseline is 3 pt above the page's bottom: emit's 4 pt under the descent ran past
    it; the page takes that room back, never the words'."""
    deck = deck_of(text_ir("Bottom line", [20.0, 259.0, 120.0, 269.0], "low", role="body"))
    d = Emitted("bottom", deck)
    _, y0, _, y1 = frame(d, "b2s_s000_t0")
    assert y1 <= PAGE.height + 0.01
    baseline = (259.0 + 8) * SCALE
    assert y1 >= baseline + emit.DESCENT_EM * 10 * SCALE - 0.01, "the words' descent stays in the box"
    assert not off_page(d, OFF_PAGE_TOLERANCE)


def test_a_tilted_text_stands_on_its_words() -> None:
    """A text adopt wrote tilted 10 degrees (marked.unturned) is laid out upright and turned about
    the middle of its words. Turned a quarter turn as `classify.rotated_texts` are, it stood
    hundreds of points off the page."""
    el = text_ir("Tilted words", [100.0, 100.0, 200.0, 112.0], "tilt", role="body")
    el["rotation"] = 10.0
    d = Emitted("tilted", deck_of(el))
    x0, y0, x1, y1 = frame(d, "b2s_s000_t0")
    assert 0 <= x0 and 0 <= y0 and x1 <= PAGE.width and y1 <= PAGE.height
    assert x0 < 100.0 * SCALE < 200.0 * SCALE < x1 and y0 < 100.0 * SCALE < 112.0 * SCALE < y1, "about its words"
    assert x1 - x0 < 300 and y1 - y0 < 120


def test_a_number_ball_at_the_left_edge_keeps_its_box_on_the_page() -> None:
    """A list number's box is centred on its ball, as wide as the ball and an em per digit with the
    insets: a ball 20 pt from the edge put it 6 pt past. It narrows about its middle (the number
    stays centred) down to the number and LINE_MARGIN between the insets."""
    number: JsonObject = {**run("7", size=8.0), "center": [10.0, 120.0], "height": 9.5, "baseline": 123.4, "x0": 8.0}
    got, center, height = number_box_of(number)

    def box(page: Page | None) -> Box4:
        reqs = [slides_json(r) for r in number_box_requests_of(got, center, height, "s", "s_n", SCALE, FONTS, page)]
        props = jobj(reqs[0], "createShape", "elementProperties")
        return box_of(jobj(props, "transform"), pt_of(jat(props, "size", "width")), pt_of(jat(props, "size", "height")))

    before, after = box(None), box(PAGE)
    assert before[0] < -1
    assert after[0] >= 0 and (after[0] + after[2]) / 2 == pytest.approx(20.0, abs=0.01)
    digit = emit.slides_width([run("7", size=8.0)], SCALE, FONTS)
    assert digit is not None and after[2] - after[0] >= 2 * PAD_X + LINE_MARGIN + digit - 0.01


def test_a_diagram_label_box_narrows_about_its_middle() -> None:
    """A label's own box centred 30 pt from the page's edge and 200 pt wide: narrowed to 60 pt
    about the same middle, which still holds its words, LABEL_ROOM wider."""
    label = ir_types.run(run("ab", size=10.0), ir_types.At(where="a test label", path=""))
    assert isinstance(label, ir_types.Run)
    w = label_width_on(200.0, 690.0, [[label]], SCALE, FONTS, PAGE)
    assert w == pytest.approx(60.0)
    words = emit.slides_width([run("ab", size=10.0)], SCALE, FONTS)
    assert words is not None and w >= words * LABEL_ROOM + LINE_MARGIN + 2 * PAD_X
    assert label_width_on(200.0, 360.0, [[label]], SCALE, FONTS, PAGE) == 200.0, "one on the page keeps its box"


def test_a_hanging_label_tab_takes_the_words_to_indent_start() -> None:
    """`text_layout` sets a hanging label (`label<TAB>text`, indentFirstLine before indentStart) as
    emit writes it and Slides draws it: the label in its hang, the words from indentStart. Read
    as words starting at indentStart after the label and a 2 em tab, a tabbed line 4 pt inside
    its box wrapped (27_text_fit s7), and a box brought in to its words rewrapped."""
    style: JsonObject = {"fontFamily": "Lato", "fontSize": 18.0}
    rest = "the words after it"
    width = sum(text_layout.advance(ch, style, 18.0) for ch in rest)
    right = PAD_X + 120.0 + width + 4.0 + PAD_X
    def held(words: str) -> JsonObject:
        rb = readback([0.0, 0.0, right, 60.0], words)
        rb["run_spans"] = [[0, len(words), style]]
        rb["paragraph_styles"] = [{"alignment": "START", "indentStart": 120.0, "indentFirstLine": 10.0}]
        return rb

    lay = text_layout.layout(held(f"Label\t{rest}\n"))
    assert lay is not None and len(lay.lines) == 1
    (line,) = lay.lines
    assert line.x == pytest.approx(PAD_X + 10.0) and line.box[2] == pytest.approx(PAD_X + 120.0 + width)
    words = text_layout.span_box(lay, len("Label\t"), len("Label\t") + 3)
    assert words is not None and words.box[0] == pytest.approx(PAD_X + 120.0)
    spilled = text_layout.layout(held(f"A label longer than its hang\t{rest}\n"))
    assert spilled is not None and spilled.lines[0].tab_to is None, "a label past its hang: a plain tab"


TEMPLATE = Template(id="tpl", w=236.22, h=236.22)  # (an arc preset as Slides stores it: 3,000,000 EMU)


def arc_json(start: tuple[float, float], end: tuple[float, float], sweep: float, heads: bool) -> JsonObject:
    """A diagram of one arc in deck.json's words, with or without heads."""
    line: JsonObject = {"from": [start[0], start[1]], "to": [end[0], end[1]], "sweep": sweep, "stroke": "#336699",
                        "width": 0.8, "arrow_from": "OPEN_ARROW" if heads else None,
                        "arrow_to": "STEALTH_ARROW" if heads else None}
    return {"id": "d", "kind": "diagram", "role": "figure", "bbox": [0.0, 0.0, 360.0, 270.0], "nodes": [],
            "spans": [], "lines": [line]}


def parsed(el: JsonObject) -> DiagramElement:
    typed = ir_types.parse_element(el, "diagram")
    assert isinstance(typed, DiagramElement)
    return typed


def arc_diagram(x0: float, x1: float, sweep: float) -> DiagramElement:
    return parsed(arc_json((x0, 200.0), (x1, 200.0), sweep, True))


Copy = tuple[str, TemplateKey, JsonObject]  # a template copy: its object id, its template, its transform


def copies(el: DiagramElement, page: Page) -> tuple[list[Copy], list[JsonObject]]:
    """The arc copies (in order) and every request a diagram's plan makes on `page`; each template
    is a 3,000,000 EMU square, as Slides stores one, and its id names its key."""
    keys: list[TemplateKey] = []

    def template(key: TemplateKey) -> Template:
        keys.append(key)
        return Template(id=f"tpl{len(keys) - 1}", w=TEMPLATE.w, h=TEMPLATE.h)
    reqs = [slides_json(r) for r in diagram_requests_of(el, "s", "d", SCALE, FONTS, template, page, ())]
    out: list[Copy] = []
    for r, move in zip(reqs, reqs[1:]):
        if "duplicateObject" in r:
            src = jstr(r, "duplicateObject", "objectId")
            out.append((jstr(move, "updatePageElementTransform", "objectId"), keys[int(src[3:])],
                        jobj(move, "updatePageElementTransform", "transform")))
    return out, reqs


def preset_points(t: JsonObject, key: TemplateKey, samples: int) -> list[tuple[float, float]]:
    """Where a copy of the arc preset (`key`: from its start clockwise through its sweep) draws,
    in PDF pt: OOXML's `arc` evaluated at the size the copy shows (its ends where rays from its
    centre at adj1 and adj2 meet its ellipse) and put on the page by the copy's transform."""
    _, sweep, start = key
    assert sweep is not None
    a, b, c, d = jnum(t, "scaleX"), jnum(t, "shearX"), jnum(t, "shearY"), jnum(t, "scaleY")
    tx, ty = jnum(t, "translateX") / EMU_PER_PT, jnum(t, "translateY") / EMU_PER_PT
    w, h = TEMPLATE.w, TEMPLATE.h
    wd2, hd2 = w * math.hypot(a, c) / 2, h * math.hypot(b, d) / 2
    out: list[tuple[float, float]] = []
    for i in range(samples + 1):
        seen = math.radians((0.0 if start is None else start) + sweep * i / samples)
        t_ = math.atan2(wd2 * math.sin(seen), hd2 * math.cos(seen))
        u, v = w / 2 * (1 + math.cos(t_)), h / 2 * (1 + math.sin(t_))
        out.append(((a * u + b * v + tx) / SCALE, (c * u + d * v + ty) / SCALE))
    return out


def frame_of(t: JsonObject) -> Box4:
    a, b, c, d = jnum(t, "scaleX"), jnum(t, "shearX"), jnum(t, "shearY"), jnum(t, "scaleY")
    tx, ty = jnum(t, "translateX") / EMU_PER_PT, jnum(t, "translateY") / EMU_PER_PT
    corners = [(a * u + b * v + tx, c * u + d * v + ty) for u in (0.0, TEMPLATE.w) for v in (0.0, TEMPLATE.h)]
    return (min(x for x, _ in corners), min(y for _, y in corners), max(x for x, _ in corners), max(y for _, y in corners))


def on(box: Box4, page: Page) -> bool:
    return box[0] >= -0.02 and box[1] >= -0.02 and box[2] <= page.width + 0.02 and box[3] <= page.height + 0.02


def heading(points: list[tuple[float, float]], end: int) -> float:
    """The direction (degrees) a drawn curve leaves by at its first (0) or last (-1) point."""
    (x0, y0), (x1, y1) = (points[0], points[1]) if end == 0 else (points[-2], points[-1])
    return math.degrees(math.atan2(y1 - y0, x1 - x0))


def test_a_gentle_arc_whose_circle_leaves_the_page_is_smooth_elliptical_pieces() -> None:
    """A gentle 30 degree arc across the page: the arc preset's shape is its whole circle's box
    (radius 618 pt), far past the page, turned or upright. It is a few pieces of the same preset
    stretched to flat ellipses: each lands on its ends, leaves them in the arc's own directions
    (no kink where pieces meet, the heads pointing as the arc's), keeps within
    `ELLIPSE_TOLERANCE` of the arc and its box on the page; the heads are the end pieces'
    templates', at the arc's weight. (Before: straight pieces with corners at every join.)"""
    as_json = arc_json((20.0, 200.0), (340.0, 200.0), 30.0, True)
    el = parsed(as_json)
    (ln,) = el.lines
    circle = curves.arc_circle(ln.from_, ln.to, 30.0)
    plan = arc_plan(ln.from_, ln.to, 30.0, SCALE, PAGE)
    assert plan.form == "ellipses" and 1 <= len(plan.pieces) <= MAX_ELLIPSES
    made, reqs = copies(el, PAGE)
    assert not [r for r in reqs if "createLine" in r], "no straight pieces"
    n = len(plan.pieces)
    assert [oid for oid, _, _ in made] == [f"d_l0e{k}" for k in range(n)]
    assert [arc_heads(key[0]) for _, key, _ in made] == \
        [("OPEN_ARROW" if k == 0 else None, "STEALTH_ARROW" if k == n - 1 else None) for k in range(n)]
    ends = [ln.from_, *(curves.arc_point(circle, 30.0 * k / n) for k in range(1, n)), ln.to]
    for k, (_, key, t) in enumerate(made):
        assert on(frame_of(t), PAGE)
        points = preset_points(t, key, 400)
        assert math.dist(points[0], ends[k]) < 0.01 and math.dist(points[-1], ends[k + 1]) < 0.01
        assert max(abs(math.dist(p, circle.centre) - circle.radius) for p in points) <= ELLIPSE_TOLERANCE + 1e-6
        # leaving its ends along the circle's tangents (to the sampling's own step)
        for end, at in ((0, 30.0 * k / n), (-1, 30.0 * (k + 1) / n)):
            tangent = math.degrees(math.atan2(math.cos(math.radians(circle.start + at)),
                                              -math.sin(math.radians(circle.start + at))))
            assert abs((heading(points, end) - tangent + 180) % 360 - 180) < 0.2
    weights = [jnum(r, "updateShapeProperties", "shapeProperties", "outline", "weight", "magnitude")
               for r in reqs if "updateShapeProperties" in r]
    assert weights == [round(0.8 * SCALE, 2)] * n, "the arc's weight: its heads the arc's size"
    # DeckPlan's templates for the page are the ones asked for; sync's (any page) hold them
    assert element_template_keys_on(as_json, SCALE, PAGE) == [key for _, key, _ in made]
    assert set(element_template_keys_on(as_json, SCALE, PAGE)) <= set(element_template_keys(as_json, SCALE))


def test_an_arc_whose_turned_box_leaves_the_page_stands_upright() -> None:
    """29_tikz_diagrams page 3's `bend left` (66 degrees): its circle's square turned to where the
    arc starts reaches past the bottom, the same square upright does not. Its template starts
    where the arc starts (`upright_at`), the copy turned by under half a degree: one exact arc."""
    as_json = arc_json((78.41, 119.03), (160.66, 119.19), 65.68, False)
    jobjs(as_json, "lines")[0]["arrow_to"] = "STEALTH_ARROW"
    typed = parsed(as_json)
    page = Page(width=720.0, height=540.0)
    assert arc_plan(typed.lines[0].from_, typed.lines[0].to, 65.68, SCALE, page).form == "upright"
    (made,), reqs = copies(typed, page)
    oid, key, t = made
    assert oid == "d_l0u" and key == ("ARC>STEALTH_ARROW", 66.0, upright_at((78.41, 119.03), (160.66, 119.19), 65.68))
    assert jnum(t, "shearX") == pytest.approx(0.0, abs=0.01 * jnum(t, "scaleX"))
    assert on(frame_of(t), page)
    points = preset_points(t, key, 200)
    assert math.dist(points[0], (78.41, 119.03)) < 0.01 and math.dist(points[-1], (160.66, 119.19)) < 0.01


def test_an_arc_no_curve_fits_on_the_page_is_fine_chords() -> None:
    """An arc hugging the page's top edge: no copy of the preset - circle or ellipse, whose box
    runs past its ends - lies on the page. Only then is it straight pieces, turning no more than
    `KINK` degrees at a join (no corner to see at slide size), its heads on the end pieces."""
    el = parsed(arc_json((0.2, 3.0), (120.0, 3.0), -20.0, True))  # (bulging down, into the page)
    (line,) = el.lines
    page = Page(width=720.0, height=540.0)
    assert arc_plan(line.from_, line.to, -20.0, SCALE, page).form == "chords"
    made, reqs = copies(el, page)
    assert not made
    created = [jobj(r, "createLine") for r in reqs if "createLine" in r]
    step = fine_step(line, -20.0)
    assert 0 < step <= KINK and len(created) == math.ceil(20.0 / step)
    for c in created:
        props = jobj(c, "elementProperties")
        x0, y0, x1, y1 = box_of(jobj(props, "transform"), pt_of(jat(props, "size", "width")),
                                pt_of(jat(props, "size", "height")))
        assert 0 <= x0 and 0 <= y0 and x1 <= page.width and y1 <= page.height
    heads = [(jstr(jobj(r, "updateLineProperties"), "objectId"), jobj(r, "updateLineProperties", "lineProperties"))
             for r in reqs if "updateLineProperties" in r]
    assert [o for o, p in heads if "startArrow" in p and p["startArrow"] != "NONE"] == ["d_l0s"]
    assert [o for o, p in heads if "endArrow" in p and p["endArrow"] != "NONE"] == [f"d_l0sc{len(created) - 1}"]


def test_an_arc_on_the_page_stays_one_turned_arc() -> None:
    """A tight arc whose circle lies on the page is the template's copy under the line's own id,
    as before (sync finds an arc's template there)."""
    el = arc_diagram(150.0, 210.0, 200.0)
    (ln,) = el.lines
    assert arc_plan(ln.from_, ln.to, 200.0, SCALE, PAGE).form == "turned"
    made, reqs = copies(el, PAGE)
    assert [(oid, key) for oid, key, _ in made] == [("d_l0", ("ARC<OPEN_ARROW>STEALTH_ARROW", 200.0, None))]
    assert not [r for r in reqs if "createLine" in r]


def test_an_elliptical_template_is_the_arc_preset_about_its_top() -> None:
    """A template key's third part is where its arc starts (None: 0, the turned copy's)."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_template_shapes(slide, [("ARC>STEALTH_ARROW", 145.0, 270 - 145.0 / 2), ("ARC", 66.0, 237.0)])
    geometries = slide.shapes._spTree.findall(f".//{{{NS_A}}}prstGeom")
    adjust = [{gd.get("name"): gd.get("fmla") for gd in g.iter(f"{{{NS_A}}}gd")} for g in geometries]
    assert adjust == [{"adj1": f"val {round(197.5 * 60000)}", "adj2": f"val {round(342.5 * 60000)}"},
                      {"adj1": f"val {237 * 60000}", "adj2": f"val {303 * 60000}"}]
    with pytest.raises(ValueError):
        _add_template_shapes(slide, [("ARC", 30.0, 360.0)])


@pytest.mark.parametrize(("radius", "sweep"), [(150.0, 18.8), (178.5, 14.5), (75.8, 65.7), (600.0, 12.0), (40.0, 150.0)])
def test_elliptical_pieces_keep_to_their_arc(radius: float, sweep: float) -> None:
    """`ellipse_pieces` over the arcs 29_tikz_diagrams draws and beyond: each piece within
    `ELLIPSE_TOLERANCE` of the circle, flat enough to take `ELLIPSE_SPAN` of its ellipse."""
    half = math.radians(sweep) / 2
    start, end = (100.0, 100.0), (100.0 + 2 * radius * math.sin(half), 100.0)
    pieces = ellipse_pieces(start, end, sweep)
    assert pieces is not None
    for p in pieces:
        chord = math.dist(p.start, p.end)
        a, b, phi = ellipse_axes(chord, abs(p.sweep), p.visual)
        assert (a, b) == pytest.approx((p.a, p.b)) and math.degrees(phi) >= ELLIPSE_SPAN
        assert p.deviation <= ELLIPSE_TOLERANCE and a * math.sin(phi) == pytest.approx(chord / 2)
