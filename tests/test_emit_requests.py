"""Offline checks of the requests emit sends to Google (emit.plan_offline) over every built deck,
replayed on a model of the presentation, and unit tests of emit's placement helpers (no Google API).

The PDFs come from `python tests/decks/build.py` and `python tests/themes/sweep.py --build`."""

import re
import tempfile
from collections import defaultdict
from functools import lru_cache
from pathlib import Path

import numpy as np
import pytest

from beamer2slides import emit
from beamer2slides.classify import HOLE_PAD, classify
from beamer2slides.emit import (EMU_PER_PT, HOLE_FONT, HOLE_SPACE_EM, SLIDE_W, FontMapper, find_marks, fit_holes,
                                formula_shifts, hole_offset, hole_run, mark_alpha, number_box_requests,
                                measure_jobs, overlay_boxes, pick_gap, slide_holes, space_shift,
                                table_requests)
from beamer2slides.extract import extract, select_overlays
from beamer2slides.notes import prepare

TESTS = Path(__file__).resolve().parent
NBSP = "\xa0"
MAX_SHIFT = 15.0  # slide pt a predicted picture may move from its PDF place


def pdfs() -> list[Path]:
    out = sorted((TESTS / "decks" / "out").glob("*.pdf")) + sorted((TESTS / "decks" / "out" / "notes").glob("*.pdf"))
    out += sorted((TESTS / "themes" / "out").glob("*/talk.pdf"))
    demo = TESTS.parent / "examples" / "demo" / "demo.pdf"
    return out + ([demo] if demo.exists() else [])


# ---------------------------------------------------------------- replaying a deck's requests

def box_of(transform: dict, w: float, h: float) -> tuple[float, float, float, float]:
    """Page box (pt) of an element of size w × h under an affine transform."""
    t = {"scaleX": 0.0, "scaleY": 0.0, "shearX": 0.0, "shearY": 0.0, "translateX": 0.0, "translateY": 0.0, **transform}
    k = 1 / EMU_PER_PT if t.get("unit") == "EMU" else 1.0
    xs, ys = [], []
    for x, y in ((0, 0), (w, 0), (0, h), (w, h)):
        xs.append(t["scaleX"] * x + t["shearX"] * y + t["translateX"] * k)
        ys.append(t["shearY"] * x + t["scaleY"] * y + t["translateY"] * k)
    return min(xs), min(ys), max(xs), max(ys)


def pt_of(dim: dict) -> float:
    return dim["magnitude"] / EMU_PER_PT if dim["unit"] == "EMU" else dim["magnitude"]


def collapse(text: str) -> str:
    return re.sub(NBSP + "+", NBSP, text)


def run_text(runs: list[dict]) -> str:
    return "".join(r["text"] for r in runs)


class Emitted:
    """One deck's requests in the order emit sends them, replayed on a model of the presentation:
    objects and their pages, groups, text boxes' text. Violations go to `problems[invariant]`."""

    def __init__(self, name: str, deck: dict):
        self.name = name
        self.result = emit.plan_offline(deck)
        self.plan = self.result["plan"]
        self.slides = {s["page"]: s for s in self.plan.deck["slides"]}
        w, h = self.plan.deck["slides"][0]["size"]
        self.page_h = SLIDE_W * h / w
        self.problems: dict[str, list[str]] = defaultdict(list)
        self.page_of: dict[str, str | None] = {}  # object id -> page id (None: a page)
        self.parent: dict[str, str] = {}  # object id -> group id
        self.groups: dict[str, list[str]] = {}
        self.texts: dict[str, str] = {}
        self.font_max: dict[str, float] = {}
        self.sizes: dict[str, tuple[float, float]] = {}
        self.boxes: dict[str, tuple] = {}  # object id -> page box (pt) where its creation gives one
        self.replay()

    def flag(self, invariant: str, where: str, message: str) -> None:
        self.problems[invariant].append(f"{self.name} {where}: {message}")

    def create(self, oid: str, page: str | None, where: str) -> None:
        if not 5 <= len(oid) <= 50:
            self.flag("ids", where, f"object id {oid!r} is {len(oid)} characters")
        if oid in self.page_of:
            self.flag("ids", where, f"object id {oid} created twice")
        if page is not None and self.page_of.get(page, "") is not None:
            self.flag("ids", where, f"{oid} created on {page}, which is no page")
        self.page_of[oid] = page

    def need(self, oid: str, where: str, what: str = "") -> bool:
        if oid not in self.page_of:
            self.flag("ids", where, f"{what or 'request'} refers to {oid}, which does not exist (yet)")
            return False
        return True

    def replay(self) -> None:
        for slide_id, elements in self.result["page_elements"].items():
            self.create(slide_id, None, "import")
            for e in elements:
                self.create(e["objectId"], slide_id, "import")
                self.sizes[e["objectId"]] = (pt_of(e["size"]["width"]), pt_of(e["size"]["height"]))
                self.texts[e["objectId"]] = ""
            self.create(self.result["speaker_notes"][slide_id], slide_id, "import")
        for r in self.result["measure"]:
            self.apply(r, "measure_places")
        for slide_id, page, parts, _ in self.result["slides"]:
            for el, reqs in parts:
                where = f"{slide_id} {el['id'] if el else ''}".strip()
                for r in reqs:
                    self.apply(r, where)

    def apply(self, request: dict, where: str) -> None:
        (kind, body), = request.items()
        self.check_values(body, where)
        if kind == "createSlide":
            self.create(body["objectId"], None, where)
        elif kind in ("createShape", "createTable", "createLine"):
            props = body["elementProperties"]
            self.need(props["pageObjectId"], where, kind)
            self.create(body["objectId"], props["pageObjectId"], where)
            w, h = pt_of(props["size"]["width"]), pt_of(props["size"]["height"])
            self.sizes[body["objectId"]] = (w, h)
            self.boxes[body["objectId"]] = box_of(props["transform"], w, h)
            if kind == "createShape":
                self.texts[body["objectId"]] = ""
        elif kind == "duplicateObject":
            if self.need(body["objectId"], where, kind):
                for new in body["objectIds"].values():
                    self.create(new, self.page_of[body["objectId"]], where)
                    self.sizes[new] = self.sizes.get(body["objectId"], (0.0, 0.0))
                    self.texts[new] = ""
        elif kind == "deleteObject":
            if self.need(body["objectId"], where, kind):
                del self.page_of[body["objectId"]]
        elif kind == "groupObjects":
            children = body["childrenObjectIds"]
            if len(children) < 2 or len(set(children)) != len(children):
                self.flag("groups", where, f"group {body['groupObjectId']} of {children}")
            pages = {self.page_of.get(c) for c in children if self.need(c, where, kind)}
            for c in children:
                if c in self.parent:
                    self.flag("groups", where, f"{c} is in {self.parent[c]} and {body['groupObjectId']}")
                self.parent[c] = body["groupObjectId"]
            if len(pages) > 1:
                self.flag("groups", where, f"group {body['groupObjectId']} spans pages {pages}")
            self.create(body["groupObjectId"], pages.pop() if pages else None, where)
            self.groups[body["groupObjectId"]] = children
        elif kind == "updatePageElementsZOrder":
            for oid in body["pageElementObjectIds"]:
                self.need(oid, where, kind)
        elif kind == "updatePageElementTransform":
            if self.need(body["objectId"], where, kind) and body["applyMode"] == "ABSOLUTE":
                self.boxes[body["objectId"]] = box_of(body["transform"], *self.sizes.get(body["objectId"], (0.0, 0.0)))
        else:
            oid = body["objectId"]
            if not self.need(oid, where, kind):
                return
            if kind == "updateLineProperties":
                for end in ("startConnection", "endConnection"):
                    if end in body["lineProperties"]:
                        self.need(body["lineProperties"][end]["connectedObjectId"], where, end)
            link = body.get("style", {}).get("link", {})
            if "pageObjectId" in link:
                self.need(link["pageObjectId"], where, "link")
            if "fontSize" in body.get("style", {}):
                self.font_max[oid] = max(self.font_max.get(oid, 0.0), body["style"]["fontSize"]["magnitude"])
            if "cellLocation" in body or oid not in self.texts:
                return
            self.edit_text(kind, body, where)

    def edit_text(self, kind: str, body: dict, where: str) -> None:
        oid, text = body["objectId"], self.texts[body["objectId"]]
        rng = body.get("textRange", {})
        if rng.get("type") == "FIXED_RANGE":
            start, end = rng["startIndex"], rng["endIndex"]
            if not 0 <= start < end <= len(text) + 1:  # (+1: the implicit newline ending the text)
                self.flag("ranges", where, f"{kind} range {start}..{end} outside {oid}'s {len(text)} characters")
        if kind == "insertText":
            i = body.get("insertionIndex", 0)
            self.texts[oid] = text[:i] + body["text"] + text[i:]
        elif kind == "deleteText":
            self.texts[oid] = text[:rng["startIndex"]] + text[rng["endIndex"]:]
        elif kind == "createParagraphBullets":
            out, pos = [], 0
            for para in text.split("\n"):  # (paragraphs in the range lose their leading tabs)
                inside = pos < rng["endIndex"] and pos + len(para) + 1 > rng["startIndex"]
                out.append(para.lstrip("\t") if inside else para)
                pos += len(para) + 1
            self.texts[oid] = "\n".join(out)

    def check_values(self, body, where: str) -> None:
        """Font sizes, weights and element sizes positive; colours and alphas in [0, 1]."""
        stack = [body]
        while stack:
            node = stack.pop()
            if isinstance(node, list):
                stack += node
            elif isinstance(node, dict):
                for key, value in node.items():
                    if key in ("fontSize", "weight") and isinstance(value, dict) and not value["magnitude"] > 0:
                        self.flag("values", where, f"{key} {value}")
                    elif key == "rgbColor" and (set(value) - {"red", "green", "blue"} or
                                                not all(0 <= v <= 1 for v in value.values())):
                        self.flag("values", where, f"colour {value}")
                    elif key == "alpha" and not 0 <= value <= 1:
                        self.flag("values", where, f"alpha {value}")
                    elif key == "size" and isinstance(value, dict) and "width" in value and \
                            (value["width"]["magnitude"] < 0 or value["height"]["magnitude"] < 0):
                        self.flag("values", where, f"size {value}")
                    elif isinstance(value, (dict, list)):
                        stack.append(value)

    # helpers for the invariants

    def top(self, oid: str) -> str | None:
        return self.parent.get(oid)

    def part_requests(self, slide_id: str, element_id: str) -> list[dict]:
        _, _, parts, _ = next(s for s in self.result["slides"] if s[0] == slide_id)
        return [r for el, reqs in parts if el and el["id"] == element_id for r in reqs]

    def title_oids(self, slide: dict) -> set[str]:
        slide_id = f"b2s_s{slide['page']:03}"
        title = emit.title_element(slide)
        if title is None:
            return set()
        sub = emit.subtitle_element(slide, title)
        return {f"{slide_id}_t{title}"} | ({f"{slide_id}_t{sub}"} if sub is not None else set())


@lru_cache(maxsize=None)
def emitted() -> tuple[Emitted, ...]:
    out = []
    for pdf in pdfs():
        with tempfile.TemporaryDirectory() as tmp:
            prepared = prepare(pdf, Path(tmp))
            raw = extract(prepared.pdf, prepared.labels)
            for page in raw["pages"]:
                page["notes"] = prepared.notes.get(page["index"])
            deck = classify(select_overlays(raw, "last"))
        out.append(Emitted(str(pdf.relative_to(TESTS.parent)).replace("\\", "/"), deck))
    return tuple(out)


@pytest.fixture(scope="module")
def decks() -> tuple[Emitted, ...]:
    found = emitted()
    if not found:
        pytest.skip("no PDFs built")
    return found


def problems(decks, invariant: str) -> list[str]:
    return [p for d in decks for p in d.problems[invariant]]


def report(found: list[str]) -> str:
    return f"{len(found)} problem(s):\n" + "\n".join(found[:40])


# ---------------------------------------------------------------- invariants over every deck

def test_object_ids_exist_are_unique_and_well_formed(decks):
    assert not (found := problems(decks, "ids")), report(found)


def test_text_ranges_lie_within_the_text(decks):
    assert not (found := problems(decks, "ranges")), report(found)


def test_sizes_and_colours_are_valid(decks):
    assert not (found := problems(decks, "values")), report(found)


def test_groups_hold_existing_objects_once(decks):
    assert not (found := problems(decks, "groups")), report(found)


def test_every_element_is_emitted_once(decks):
    found = []
    for d in decks:
        for slide_id, page, parts, element_ids in d.result["slides"]:
            slide = d.slides[page]
            ids = [el["id"] for el, _ in parts if el]
            if ids != [e["id"] for e in slide["elements"]] or len(set(element_ids)) != len(element_ids):
                found.append(f"{d.name} {slide_id}: elements {ids} emitted as {element_ids}")
            spans: dict[str, str] = {}
            for e in slide["elements"]:
                if e["kind"] in ("text", "table", "diagram") or e.get("overlay"):
                    for s in e.get("spans", []):
                        if s in spans:
                            found.append(f"{d.name} {slide_id}: span {s} in both {spans[s]} and {e['id']}")
                        spans[s] = e["id"]
    assert not found, report(found)


def test_every_text_reaches_its_text_box(decks):
    found = []
    for d in decks:
        for slide_id, page, parts, element_ids in d.result["slides"]:
            for el, oid in zip(d.slides[page]["elements"], element_ids):
                if el["kind"] == "text":
                    want = "\n".join("".join(NBSP if r.get("hole") else r["text"] for r in p["runs"]) for p in el["paragraphs"])
                    got = d.texts.get(oid)
                    if got is None or collapse(got) != collapse(want):
                        found.append(f"{d.name} {slide_id} {el['id']}: {got!r} != {want!r}")
                elif el["kind"] == "table":
                    inserted = {(r["insertText"]["cellLocation"]["rowIndex"], r["insertText"]["cellLocation"]["columnIndex"]):
                                r["insertText"]["text"] for r in d.part_requests(slide_id, el["id"]) if "insertText" in r}
                    for i, row in enumerate(el["cells"]):
                        for j, runs in enumerate(row):
                            if run_text(runs).strip() and inserted.get((i, j), "").strip() != run_text(runs).strip():
                                found.append(f"{d.name} {slide_id} {el['id']} cell {i},{j}: {run_text(runs)!r}")
                elif el["kind"] == "diagram":
                    inserted = "\n".join(r["insertText"]["text"] for r in d.part_requests(slide_id, el["id"]) if "insertText" in r)
                    for node in el["nodes"]:
                        for runs in node["paragraphs"]:
                            if run_text(runs).strip() not in inserted:
                                found.append(f"{d.name} {slide_id} {el['id']}: label {run_text(runs)!r}")
    assert not found, report(found)


def test_anchored_pictures_share_a_group_with_their_text(decks):
    found = []
    for d in decks:
        for slide_id, page, parts, element_ids in d.result["slides"]:
            slide = d.slides[page]
            oids = {e["id"]: oid for e, oid in zip(slide["elements"], element_ids)}
            for e, oid in zip(slide["elements"], element_ids):
                if not e.get("anchor"):
                    continue
                if e["anchor"] not in oids:
                    found.append(f"{d.name} {slide_id} {e['id']}: anchor {e['anchor']} is not on the slide")
                    continue
                text = oids[e["anchor"]]
                if text in d.title_oids(slide):
                    continue  # (placeholders can't be grouped)
                for member in [oid] + ([f"{oid}n"] if e.get("number") else []):
                    if d.top(member) is None or d.top(member) != d.top(text):
                        found.append(f"{d.name} {slide_id} {e['id']}: {member} in group {d.top(member)}, "
                                     f"its text {text} in {d.top(text)}")
    assert not found, report(found)


def test_numbers_are_centred_on_their_ball(decks):
    found = []
    for d in decks:
        for slide_id, page, parts, element_ids in d.result["slides"]:
            pictures = {e["id"]: box for e, box in d.result["pictures"][page]}
            for e, oid in zip(d.slides[page]["elements"], element_ids):
                if not e.get("number"):
                    continue
                reqs = [r for r in d.part_requests(slide_id, e["id"]) if next(iter(r.values())).get("objectId") == f"{oid}n"]
                x0, y0, x1, y1 = d.boxes[f"{oid}n"]
                bx0, by0, bx1, by1 = pictures[e["id"]]
                if abs((x0 + x1 - bx0 - bx1) / 2) > 0.1 or abs((y0 + y1 - by0 - by1) / 2) > 0.1:
                    found.append(f"{d.name} {slide_id} {e['id']}: number box centre "
                                 f"({(x0 + x1) / 2:.2f}, {(y0 + y1) / 2:.2f}) != ball's ({(bx0 + bx1) / 2:.2f}, {(by0 + by1) / 2:.2f})")
                shape = [r["updateShapeProperties"]["shapeProperties"] for r in reqs if "updateShapeProperties" in r]
                para = [r["updateParagraphStyle"]["style"] for r in reqs if "updateParagraphStyle" in r]
                if [s.get("contentAlignment") for s in shape] != ["MIDDLE"] or [s.get("alignment") for s in para] != ["CENTER"]:
                    found.append(f"{d.name} {slide_id} {e['id']}: number not centred ({shape}, {para})")
                if d.texts.get(f"{oid}n") != e["number"]["text"]:
                    found.append(f"{d.name} {slide_id} {e['id']}: number text {d.texts.get(f'{oid}n')!r} != {e['number']['text']!r}")
    assert not found, report(found)


def test_holes_are_no_break_spaces_as_wide_as_the_hole(decks):
    found = []
    fonts = FontMapper()
    for d in decks:
        scale = d.plan.scale
        for slide_id, page, parts, element_ids in d.result["slides"]:
            slide = d.slides[page]
            pictures = {id(run): pic for _, _, run, pic in slide_holes(slide)}
            for (el, reqs), oid in zip(parts[1:], element_ids):
                if el["kind"] != "text":
                    continue
                holes = [r for p in el["paragraphs"] for r in p["runs"] if r.get("hole")]
                text = d.texts[oid]
                styled = [r["updateTextStyle"] for r in reqs if "updateTextStyle" in r
                          and r["updateTextStyle"]["style"].get("fontFamily") == HOLE_FONT
                          and r["updateTextStyle"]["textRange"]["type"] == "FIXED_RANGE"]
                gaps = [(s["textRange"]["endIndex"] - s["textRange"]["startIndex"], s["style"]["fontSize"]["magnitude"])
                        for s in styled if set(text[s["textRange"]["startIndex"]:s["textRange"]["endIndex"]]) == {NBSP}]
                if len(gaps) != len(holes):
                    found.append(f"{d.name} {slide_id} {el['id']}: {len(holes)} holes, {len(gaps)} runs of no-break spaces")
                    continue
                for run, (n, size) in zip(holes, gaps):
                    width, z = run["hole"] * scale, fonts(run, scale)[1]
                    if abs(n * HOLE_SPACE_EM * size - width) > 0.005 * n + 0.01 or size > z + 0.01:
                        found.append(f"{d.name} {slide_id} {el['id']}: {n} spaces of {size} pt for a hole of "
                                     f"{width:.2f} pt at {z} pt")
                    pic = pictures.get(id(run))
                    if pic and run["hole"] < pic["bbox"][2] - pic["bbox"][0] - 0.01:
                        found.append(f"{d.name} {slide_id} {pic['id']}: hole {run['hole']:.2f} narrower than its picture "
                                     f"{pic['bbox'][2] - pic['bbox'][0]:.2f}")
    assert not found, report(found)


def test_predicted_pictures_stay_near_their_pdf_place(decks):
    found = []
    for d in decks:
        scale = d.plan.scale
        for page, slide in d.slides.items():
            for pid, dx in d.plan.shifts[page].items():
                if abs(dx * scale) >= MAX_SHIFT:
                    found.append(f"{d.name} page {page + 1} {pid}: formula shift {dx * scale:.1f} pt")
            boxes = {e["id"]: e["bbox"] for e in slide["elements"]}
            for pid, (x0, x1) in d.plan.overlays[page].items():
                bx0, _, bx1, _ = boxes[pid]
                if abs((x0 - bx0) * scale) >= MAX_SHIFT or abs((x1 - x0) / (bx1 - bx0) - 1) > emit.OVERLAY_STRETCH + 1e-6:
                    found.append(f"{d.name} page {page + 1} {pid}: overlay {bx0:.1f}..{bx1:.1f} -> {x0:.1f}..{x1:.1f}")
    assert not found, report(found)


def test_overlay_marks_highlight_their_words_on_the_scratch_slide(decks):
    """Every mark of an overlay lies on a word found in its anchor's text box (measure_places
    highlights it): a right edge's word is the last word before it."""
    found = []
    for d in decks:
        plan = d.plan
        _, jobs = measure_jobs(plan.deck, plan.scale, plan.fonts, plan.placed, plan.page_slide)
        for sid, page, _, overlays in jobs:
            texts = {e["id"]: e for e in d.slides[page]["elements"] if e["kind"] == "text"}
            for o in overlays:
                pic = o["pic"]
                text = "\n".join(emit.slides_texts(texts[pic["anchor"]], plan.scale, plan.fonts))
                where = f"{d.name} page {page + 1} {pic['id']}"
                marks = [i for w in o["words"] for i, _ in w["marks"]]
                if sorted(marks) != list(range(len(pic["marks"]))):
                    found.append(f"{where}: marks {sorted(marks)} of {len(pic['marks'])} on words")
                for w in o["words"]:
                    for i, side in w["marks"]:
                        word = pic["marks"][i]["before"][-1][5] if side else None
                        got = text[w["range"][0]:w["range"][1]]
                        if side and " ".join(word.split()) != " ".join(got.split()):
                            found.append(f"{where}: mark {i} ends {word!r}, highlighted {got!r}")
    assert not found, report(found)


def test_bullets_are_styled_before_they_are_created(decks):
    """A bullet keeps the style its paragraph had when it was created (CLAUDE.md pitfall): family,
    size and colour first, bullets next, then runs in parts that never cover a whole item."""
    found = []
    for d in decks:
        for slide_id, page, parts, element_ids in d.result["slides"]:
            for (el, reqs), oid in zip(parts[1:], element_ids):
                if el["kind"] != "text" or not any(p["bullet"] for p in el["paragraphs"]):
                    continue
                where = f"{d.name} {slide_id} {el['id']}"
                inserted = next(r["insertText"]["text"] for r in reqs if "insertText" in r)
                dummies = {r["deleteText"]["textRange"]["startIndex"] for r in reqs if "deleteText" in r}
                spans, pos = [], 0  # (start, end) of each paragraph in the inserted text, dummies left out
                for para in inserted.split("\n"):
                    if not (pos in dummies and para == "-"):
                        spans.append((pos, pos + len(para), para))
                    pos += len(para) + 1
                if len(spans) != len(el["paragraphs"]):
                    found.append(f"{where}: {len(spans)} paragraphs inserted for {len(el['paragraphs'])}")
                    continue
                kinds = [next(iter(r)) for r in reqs]
                covered = set()
                for k, r in enumerate(reqs):
                    if "createParagraphBullets" not in r:
                        continue
                    rng = r["createParagraphBullets"]["textRange"]
                    for (start, end, para), p in zip(spans, el["paragraphs"]):
                        if not (start < rng["endIndex"] and end + 1 > rng["startIndex"]):
                            continue
                        covered.add(start)
                        if not p["bullet"]:
                            found.append(f"{where}: bullets over the plain paragraph {para!r}")
                            continue
                        if r["createParagraphBullets"]["bulletPreset"] != emit.bullet_preset(p["bullet"]) or \
                                len(para) - len(para.lstrip("\t")) != emit.bullet_level(p["bullet"], p["level"]):
                            found.append(f"{where}: {para!r} got {r['createParagraphBullets']['bulletPreset']} at "
                                         f"{len(para) - len(para.lstrip(chr(9)))} tabs")
                        before = [q["updateTextStyle"] for q in reqs[:k] if "updateTextStyle" in q
                                  and q["updateTextStyle"]["textRange"]["startIndex"] <= start
                                  and q["updateTextStyle"]["textRange"]["endIndex"] >= end]
                        need = {"fontFamily", "fontSize"} | ({"foregroundColor"} if p["bullet"].get("color") else set())
                        if not any(need <= set(q["fields"].split(",")) for q in before):
                            found.append(f"{where}: {para!r} has no {sorted(need)} before its bullet")
                for (start, _, para), p in zip(spans, el["paragraphs"]):
                    if p["bullet"] and start not in covered:
                        found.append(f"{where}: {para!r} has no bullet")
                # After the bullets exist (final indices, no tabs): no style request covers a whole item.
                last = max(i for i, kind in enumerate(kinds) if kind in ("createParagraphBullets", "deleteText"))
                pos, items = 0, []
                for p in el["paragraphs"]:
                    n = len(run_text([hole_run(r, d.plan.scale, d.plan.fonts) if r.get("hole") else r for r in p["runs"]]))
                    if p["bullet"] and n > 1:
                        items.append((pos, pos + n))
                    pos += n + 1
                for q in reqs[last + 1:]:
                    if "updateTextStyle" in q:
                        t = q["updateTextStyle"]["textRange"]
                        if (t["startIndex"], t["endIndex"]) in items:
                            found.append(f"{where}: one style request over the whole item {t}")
    assert not found, report(found)


def test_elements_lie_within_the_page(decks):
    """Pictures, shapes, tables and lines within 1 pt. A text box may reach past the right edge
    by its slack (a single line gets 15% or two ems of room, so a wider font never wraps it) and
    past the bottom by its 4 pt of room below the last line. A number box centred on a ball near
    the edge only needs its number inside."""
    found = []
    for d in decks:
        boxes = [(f"b2s_s{page:03}", e["id"], box, None) for page, pictures in d.result["pictures"].items()
                 for e, box in pictures]
        boxes += [(d.page_of[oid], oid, box, d.font_max.get(oid, 0.0) if oid in d.texts else None)
                  for oid, box in d.boxes.items() if str(d.page_of.get(oid)).startswith("b2s_s")]
        for slide_id, name, (x0, y0, x1, y1), text in boxes:
            if re.search(r"_f\d+n$", name):
                pad = (x1 - x0 - len(d.texts[name]) * text) / 2
                x0, x1 = x0 + pad, x1 - pad
            right, bottom = (1.0, 1.0) if text is None else (max(0.15 * (x1 - x0), 2 * text) + 1, 4.0)
            if x0 < -1 or y0 < -1 or x1 > SLIDE_W + right or y1 > d.page_h + bottom:
                found.append(f"{d.name} {slide_id} {name}: ({x0:.1f}, {y0:.1f}, {x1:.1f}, {y1:.1f}) "
                             f"outside {SLIDE_W:.0f} x {d.page_h:.0f}")
    assert not found, report(found)


# ---------------------------------------------------------------- the imported .pptx

def placeholders(prs) -> list[tuple]:
    return [(page.name, s.name, s.left, s.top, s.width, s.height) for page in [prs.slide_master, *prs.slide_layouts]
            for s in page.placeholders]


def test_wide_pptx_rescales_every_placeholder_once():
    """16:9 decks: layout placeholders that inherit the master's position used to be scaled twice,
    with x and width written as 0."""
    from pptx import Presentation
    square = placeholders(Presentation(emit.build_pptx(720.0, 540.0, [], [], {"color": "#ffffff"})))
    wide = placeholders(Presentation(emit.build_pptx(720.0, 405.0, [], [], {"color": "#ffffff"})))
    for (layout, name, x, y, w, h), (_, _, wx, wy, ww, wh) in zip(square, wide):
        assert (wx, ww) == (x, w) and ww > 0, (layout, name)
        assert abs(wy - 0.75 * y) <= 1 and abs(wh - 0.75 * h) <= 1, (layout, name)


def test_pptx_pictures_sit_at_their_boxes(tmp_path):
    from PIL import Image
    from pptx import Presentation
    Image.new("RGB", (4, 3), "red").save(tmp_path / "f.png")
    box = [101.37, 57.2, 180.05, 96.93]
    page = {"layout": "BLANK", "fill": None, "templates": False,
            "pictures": [{"file": tmp_path / "f.png", "bbox": box, "alt": "x^2", "title": "Formula"}]}
    pic, = Presentation(emit.build_pptx(453.54, 255.12, [], [page], {"color": "#ffffff"})).slides[0].shapes
    got = [pic.left, pic.top, pic.left + pic.width, pic.top + pic.height]
    assert all(abs(g - v * EMU_PER_PT) <= 1 for g, v in zip(got, box))
    assert pic._element.nvPicPr.cNvPr.get("descr") == "x^2"


def test_theme_decoration_keeps_what_backgrounds_share():
    """A header bar on every frame (one frame with a highlighted word in it, one without the bar):
    the decoration is the bar without the pixels where frames differ; the bare frame doesn't show it."""
    from beamer2slides.render import page_ground, theme_decoration
    white = np.array([255, 255, 255])
    frame = np.full((60, 80, 3), 255, np.uint8)
    frame[:10] = (51, 51, 179)
    highlight = frame.copy()
    highlight[4:6, 10:20] = (255, 200, 0)
    bare = np.full((60, 80, 3), 255, np.uint8)
    assert (page_ground(frame) == white).all()
    picture, inside, exact = theme_decoration([frame, highlight, bare], white)
    assert inside == [True, True, False] and not exact
    assert (picture[:10, :, 3] == 255).sum() == 800 - 20 and (picture[10:, :, 3] == 0).all()
    assert (picture[4:6, 10:20, 3] == 0).all() and tuple(picture[0, 0]) == (51, 51, 179, 255)
    assert theme_decoration([frame, frame], white)[2], "a frame that is ground plus decoration can inherit"
    assert theme_decoration([bare, bare], white)[0] is None


def test_pptx_layouts_carry_the_theme(tmp_path):
    """Decorations at the bottom of the layouts; a _V1 page layout is a copy with its own decoration."""
    from PIL import Image
    from pptx import Presentation
    for name in ("main", "title", "variant"):
        Image.new("RGBA", (8, 6), (0, 0, 255, 255)).save(tmp_path / f"{name}.png")
    pages = [{"layout": "TITLE_ONLY", "fill": None, "templates": False, "pictures": []},
             {"layout": "TITLE_ONLY_V1", "fill": None, "templates": False, "pictures": []},
             {"layout": "BLANK_V1", "fill": None, "templates": False, "pictures": []}]
    decorations = {"*": tmp_path / "main.png", "TITLE": tmp_path / "title.png", "*_V1": tmp_path / "variant.png"}
    prs = Presentation(emit.build_pptx(720.0, 405.0, [], pages, {"color": "#ffffff"}, decorations))
    pictures = {l.name: [s for s in l.shapes if s.shape_type == 13] for l in prs.slide_layouts}
    assert len(prs.slide_layouts) == 13 and all(len(p) == 1 for p in pictures.values())
    assert [s.slide_layout.name for s in prs.slides] == ["Title Only", "Title Only (theme 2)", "Blank (theme 2)"]
    layout = prs.slides[1].slide_layout
    assert layout.shapes[0].shape_type == 13 and any("Title" in p.name for p in layout.placeholders), "decoration first, title placeholder kept"
    assert (layout.shapes[0].width, layout.shapes[0].height) == (prs.slide_width, prs.slide_height)


# ---------------------------------------------------------------- numeric helpers

SCALE = SLIDE_W / 453.54
FONTS = FontMapper()


def run_of(text: str = "", size: float = 10.91, **extra) -> dict:
    return {"text": text, "font": "CMSS10", "family": "sans", "size": size, "bold": False, "italic": False,
            "smallcaps": False, "color": "#000000", "link": None, **extra}


@pytest.mark.parametrize("text", ["7", "12", "123"])
def test_number_box_is_centred_on_the_ball(text):
    number = {**run_of(text, 8.0), "center": [40.3, 120.7], "height": 9.5, "baseline": 123.4, "x0": 37.0}
    reqs = number_box_requests(number, "b2s_s001", "b2s_s001_f3n", SCALE, FONTS)
    props = reqs[0]["createShape"]["elementProperties"]
    w, h = pt_of(props["size"]["width"]), pt_of(props["size"]["height"])
    x0, y0, x1, y1 = box_of(props["transform"], w, h)
    assert abs((x0 + x1) / 2 - 40.3 * SCALE) < 0.01 and abs((y0 + y1) / 2 - 120.7 * SCALE) < 0.01
    assert w - 2 * emit.PAD_X >= len(text) * FONTS(number, SCALE)[1], "room for every digit: no wrap"
    assert reqs[1]["updateShapeProperties"]["shapeProperties"]["contentAlignment"] == "MIDDLE"
    assert reqs[2]["insertText"]["text"] == text and reqs[4]["updateParagraphStyle"]["style"]["alignment"] == "CENTER"


@pytest.mark.parametrize("hole", [0.4, 3.0, 17.67, 40.71, 120.0])
def test_hole_run_is_exactly_as_wide_as_the_hole(hole):
    run = run_of(NBSP, hole=hole)
    z = FONTS(run, SCALE)[1]
    out = hole_run(run, SCALE, FONTS)
    n, width = len(out["text"]), hole * SCALE
    assert set(out["text"]) == {NBSP} and out["hole"] == hole
    assert abs(n * HOLE_SPACE_EM * out["hole_size"] - width) <= 0.005 * n
    assert out["hole_size"] <= z + 0.005, "never taller than the line"
    assert n == 1 or (n - 1) * HOLE_SPACE_EM * z < width, "as few spaces as fit"


def three_holes() -> dict:
    """A line 'A [f] and [g] or [h] end.' with TeX's word spaces (3.63 pt) around three formulas."""
    words = [("A", 10.91, 7.26), ("and", 40.43, 17.0), ("or", 93.69, 10.3), ("end.", 132.25, 20.0)]
    holes = [(21.8, 16.0), (61.06, 30.0), (107.62, 22.0)]  # (hole_x0, picture width)
    runs, elements = [run_of("A ")], []
    for k, (x, w) in enumerate(holes):
        before = [[ww, "CMSS10", "sans", False, False, t, x0] for t, x0, ww in words[:k + 1]]
        runs += [run_of(NBSP, hole=w, hole_x0=x, before=before, next_x0=words[k + 1][1]), run_of(words[k + 1][0] + " ")]
        elements.append({"id": f"p0h{k}", "kind": "image", "anchor": "p0t0", "bbox": [x - HOLE_PAD, 92.0, x - HOLE_PAD + w, 106.0]})
    runs[-1]["text"] = "end."
    para = {"align": "left", "runs": runs, "lines": [{"baseline": 104.0, "x0": 10.91, "x1": 152.25}]}
    return {"page": 0, "size": [453.54, 255.12], "elements": elements + [{"id": "p0t0", "kind": "text", "paragraphs": [para]}]}


def test_several_holes_on_a_line_keep_their_pictures_in_place():
    slide = fit_holes(three_holes(), SCALE, FONTS)
    shifts = formula_shifts(slide, SCALE, FONTS)
    assert all(abs(shifts.get(f"p0h{k}", 0.0)) < 3 for k in range(3)), shifts
    run = slide["elements"][-1]["paragraphs"][0]["runs"][5]
    em = FONTS(run, SCALE)[1] / SCALE
    assert space_shift(run, em) < -30, "without the earlier holes both gaps look like stretched spaces"


def test_fit_holes_reaches_the_next_word_but_never_below_the_picture():
    slide = three_holes()
    tight = slide["elements"][-1]["paragraphs"][0]["runs"][3]
    tight["next_x0"] = tight["hole_x0"] + 10  # the next word starts inside the picture
    fitted = fit_holes(slide, SCALE, FONTS)
    holes = [r["hole"] for r in fitted["elements"][-1]["paragraphs"][0]["runs"] if r.get("hole")]
    space = emit.SYMBOL_ADVANCE_EM[" "] * FONTS(tight, SCALE)[1] / SCALE
    assert holes[0] == pytest.approx(40.43 - (10.91 + 7.26 + space), abs=0.01)
    assert holes[1] == 30.0, "as wide as the picture"
    followed = three_holes()
    followed["elements"][-1]["paragraphs"][0]["runs"][2]["text"] = " and "  # a space after the hole: no reach
    assert [r["hole"] for r in fit_holes(followed, SCALE, FONTS)["elements"][-1]["paragraphs"][0]["runs"] if r.get("hole")][0] == 16.0


def test_hole_offset_shares_the_spaces_like_the_pdf():
    slide = fit_holes(three_holes(), SCALE, FONTS)
    _, p, run, pic = slide_holes(slide)[1]
    z = FONTS(run, SCALE)[1]
    o = hole_offset(p, run, pic, SCALE, z)
    space = emit.SYMBOL_ADVANCE_EM[" "] * z
    left, right = space + o, (run["hole"] - (pic["bbox"][2] - pic["bbox"][0])) * SCALE - o
    pdf_left, pdf_right = pic["bbox"][0] - 57.43, 93.69 - pic["bbox"][2]
    assert left / right == pytest.approx(pdf_left / pdf_right, rel=0.01)
    glued = three_holes()
    glued["elements"][-1]["paragraphs"][0]["runs"][0]["text"] = "A"  # no word space before the formula
    _, p, run, pic = slide_holes(glued)[0]
    assert hole_offset(p, run, pic, SCALE, z) == 0.0


def overlay(marks_x: list[float], bbox=(100.0, 50.0, 200.0, 60.0)) -> dict:
    marks = [{"x": x, "hole_x0": x, "font": "CMSS10", "family": "sans", "size": 10.91, "bold": False, "italic": False,
              "before": []} for x in marks_x]
    return {"elements": [{"id": "p0o0", "kind": "image", "overlay": True, "anchor": "p0t0", "bbox": list(bbox), "marks": marks}]}


@pytest.mark.parametrize("drifts, expected", [
    ({100.0: 5.0, 200.0: 5.0}, (105.0, 205.0)),     # both ends drift alike: moved, not stretched
    ({100.0: 0.0, 200.0: 30.0}, (110.0, 220.0)),    # slope 0.3 clamped to OVERLAY_STRETCH
    ({100.0: 0.0, 200.0: -4.0}, (100.0, 196.0)),    # narrower by the slope (-0.04)
    ({150.0: 2.0, 155.0: 12.0}, (107.0, 207.0)),    # marks closer than 10 pt: no stretch
])
def test_overlay_boxes_drift_and_clamped_stretch(monkeypatch, drifts, expected):
    def fake(probe, scale, fonts):  # the drift of the mark at the probe's gap
        return {"gap": HOLE_PAD + drifts[probe["elements"][0]["paragraphs"][0]["runs"][0]["hole_x0"]]}
    monkeypatch.setattr(emit, "formula_shifts", fake)
    x0, x1 = overlay_boxes(overlay(list(drifts)), SCALE, FONTS)["p0o0"]
    assert (x0, x1) == pytest.approx(expected, abs=1e-6)


def test_coloured_text_is_no_mark():
    img = np.full((900, 1600, 3), 255, dtype=np.uint8)
    img[100:140, 200:300] = (255, 0, 0)  # a red link
    img[100:140, 400:500] = (0, 0, 255)  # a blue one
    img[200:240, 200:300] = (255, 0, 255)  # the mark
    marks = find_marks(mark_alpha(img, "#ff00ff"), 1600 / SLIDE_W)
    assert len(marks) == 1 and marks[0][1] == pytest.approx(200 * SLIDE_W / 1600, abs=0.5)


@pytest.mark.parametrize("label", ["below", "above"])
def test_a_braced_formula_is_measured_on_its_text_line(label):
    """An \\underbrace label below (or \\overbrace above) makes the picture taller than its line: the
    gap is looked for on the text line, not at the picture's middle, which lies a line away."""
    size = 10.91
    pic_box = [60.0, 88.0, 100.0, 124.0] if label == "below" else [60.0, 70.0, 100.0, 104.0]
    baselines = [100.0, 128.0, 141.0, 154.0] if label == "below" else [72.0, 100.0, 113.0, 126.0]
    line_at = 100.0
    hole = run_of(NBSP, size, hole=40.0, hole_x0=61.0, before=[[7.26, "CMSS10", "sans", False, False, "A", 10.91]],
                  next_x0=104.0)
    runs = [run_of("some text " * 3), run_of("A "), hole, run_of("words")]
    para = {"align": "left", "bullet": None, "level": 0, "size": size, "text_x0": 10.91, "tab_x0": None, "wrap_limit": 300.0,
            "runs": runs, "lines": [{"baseline": b, "x0": 10.91, "x1": 200.0} for b in baselines]}
    slide = {"page": 0, "size": [453.54, 255.12], "elements": [
        {"id": "p0h0", "kind": "image", "anchor": "p0t0", "bbox": pic_box},
        {"id": "p0t0", "kind": "text", "role": "body", "bbox": [10.91, 60.0, 200.0, 160.0], "paragraphs": [para]}]}
    _, jobs = measure_jobs({"slides": [slide]}, SCALE, FONTS, lambda el, n: el, {0: "b2s_s000"})
    (found,), = [job[2] for job in jobs]
    assert found["cy"] == pytest.approx((line_at - 0.35 * size) * SCALE)
    middle = (pic_box[1] + pic_box[3]) / 2 * SCALE
    assert abs(round((middle - found["cy"]) / found["pitch"])) == 1, "the picture's middle is a line away"
    mark = (found["x0"] + 2.0, found["cy"] - 6.0, found["x0"] + 2.0 + found["width"], found["cy"] + 6.0)
    assert pick_gap([mark], found["x0"], found["cy"], found["width"], found["pitch"]) == pytest.approx((2.0, 0.0))


def braced_phrase() -> dict:
    """'Gradient descent updates every parameter after each batch.' with a brace under
    'updates every parameter' (marks at both edges of 'updates' and 'parameter', as classify.mark writes them)."""
    words = [("Gradient", 10.91, 38.9), ("descent", 53.4, 33.91), ("updates", 91.03, 35.39), ("every", 130.05, 23.47),
             ("parameter", 157.16, 45.52), ("after", 206.3, 22.0)]
    marks = []
    for k in (2, 4):
        before = [[w, "CMSS10", "sans", False, False, t, x0] for t, x0, w in words[:k + 1]]
        for x, b in ((words[k][1], before[:-1]), (words[k][1] + words[k][2], before)):
            marks.append({"x": x, "hole_x0": x, "pads": 0.0, "font": "CMSS10", "family": "sans", "size": 10.91,
                          "bold": False, "italic": False, "before": b})
    para = {"align": "left", "bullet": None, "level": 0, "size": 10.91, "text_x0": 10.91, "tab_x0": None, "wrap_limit": 300.0,
            "runs": [run_of(" ".join(t for t, _, _ in words) + " each batch.")],
            "lines": [{"baseline": 106.13, "x0": 10.91, "x1": 283.62}]}
    return {"page": 0, "size": [453.54, 255.12], "elements": [
        {"id": "p0t0", "kind": "text", "role": "body", "bbox": [10.91, 98.0, 283.62, 109.0], "paragraphs": [para]},
        {"id": "p0f0", "kind": "image", "overlay": True, "anchor": "p0t0", "bbox": [89.83, 110.44, 203.99, 129.32], "marks": marks}]}


def test_overlay_words_are_highlighted_and_the_brace_fits_them():
    slide = braced_phrase()
    reqs, jobs = measure_jobs({"slides": [slide]}, SCALE, FONTS, lambda el, n: el, {0: "b2s_s000"})
    (_, _, _, (o,)), = jobs
    text = emit.slides_texts(slide["elements"][0], SCALE, FONTS)[0]
    assert [text[slice(*w["range"])] for w in o["words"]] == ["updates", "parameter"]
    assert len({w["colour"] for w in o["words"]}) == 2, "words on one line get different colours"
    highlights = [r["updateTextStyle"] for r in reqs if "backgroundColor" in r.get("updateTextStyle", {}).get("style", {})]
    assert [(h["textRange"]["startIndex"], h["textRange"]["endIndex"]) for h in highlights] == [w["range"] for w in o["words"]]
    # Slides sets 'updates' 3 pt (PDF) further right and 3 pt wider, 'parameter' 16 pt further: stretched 12%.
    drift = {91.03: 3.0, 126.42: 6.0, 157.16: 12.0, 202.68: 16.0}
    x = lambda v: (v + drift[v]) * SCALE
    marks = {w["colour"]: [(x(a), 160.0, x(b), 172.0), (x(a) + 200, 160.0, x(b) + 200, 172.0)]
             for w, (a, b) in zip(o["words"], [(91.03, 126.42), (157.16, 202.68)])}
    (dx, dy, sx), table = emit.overlay_move(o, marks, SCALE)
    b0, b1 = emit.fit_overlay(slide["elements"][1]["bbox"], list(drift.items()), emit.OVERLAY_STRETCH_MEASURED)
    x0, _, x1, _ = slide["elements"][1]["bbox"]
    assert dx == pytest.approx((b0 - x0) * SCALE) and dy == 0.0 and sx == pytest.approx((b1 - b0) / (x1 - x0))
    assert sx > 1 + emit.OVERLAY_STRETCH, "measured words stretch a brace further than predicted ones may"
    assert [m[2] for m in table] == pytest.approx(list(drift.values()), abs=0.01)
    assert emit.overlay_move(o, {c: [] for c in marks}, SCALE)[0] is None, "no word found: the prediction stays"


def test_measured_overlay_move_keeps_the_left_edge_under_a_relative_scale(decks):
    """measure_places' (dx, dy, scaleX): the RELATIVE transform scales about the page origin, so
    its translation makes the picture's left edge move by dx and its width scale by scaleX."""
    d = next((d for d in decks if any(e.get("marks") for s in d.slides.values() for e in s["elements"])), None)
    if d is None:
        pytest.skip("no deck with overlays built")
    slide = next(s for s in d.plan.deck["slides"] if any(e.get("marks") for e in s["elements"]))
    i, pic = next((i, e) for i, e in enumerate(slide["elements"]) if e.get("marks"))
    x0, _, x1, _ = next(box for e, box in d.result["pictures"][slide["page"]] if e["id"] == pic["id"])
    templates = [(100.0, 100.0)] * len(d.plan.keys)
    parts, _ = d.plan.slide_parts(slide, d.result["page_elements"], d.result["speaker_notes"],
                                  {pic["id"]: (4.0, 0.0, 1.2)}, templates)
    oid = f"b2s_s{slide['page']:03}_f{i}"
    t, = [r["updatePageElementTransform"] for _, reqs in parts for r in reqs
          if r.get("updatePageElementTransform", {}).get("objectId") == oid]
    assert t["applyMode"] == "RELATIVE"
    moved = [t["transform"]["scaleX"] * v + t["transform"]["translateX"] / EMU_PER_PT for v in (x0, x1)]
    assert moved == pytest.approx([x0 + 4.0, x0 + 4.0 + 1.2 * (x1 - x0)], abs=0.01)

# ---------------------------------------------------------------- right to left

# A Hebrew paragraph is a left-to-right paragraph to Slides unless it is told otherwise, and
# then its full stop lands at the wrong end and the cursor walks the wrong way. Synthetic, as
# `tests/test_bidi.py` is: the characters, not a font.
ALEF, BET, GIMEL = "א", "ב", "ג"


def hebrew_element(align: str = "left", lines=((100.0, 200.0),), bullet=None,
                   direction: str | None = "rtl", text: str | None = None) -> dict:
    para = {"align": align, "level": 0, "bullet": bullet, "size": 10.91, "text_x0": lines[0][0],
            "tab_x0": None, "wrap_limit": None,
            "runs": [run_of(f"{ALEF}{BET} {GIMEL}" if text is None else text)],
            "lines": [{"baseline": 60.0 + 12 * i, "x0": x0, "x1": x1} for i, (x0, x1) in enumerate(lines)]}
    if direction:
        para["direction"] = direction
    return {"id": "p0t0", "kind": "text", "role": "body", "code": False,
            "bbox": [min(l[0] for l in lines), 50.0, max(l[1] for l in lines), 50.0 + 12 * len(lines)],
            "paragraphs": [para]}


def paragraph_style(el: dict) -> dict:
    reqs = emit.text_box_requests(el, "b2s_s001", "b2s_s001_t0", SCALE, FONTS)
    return next(r["updateParagraphStyle"] for r in reqs if "updateParagraphStyle" in r)


def test_a_hebrew_paragraph_is_told_which_way_it_reads():
    style = paragraph_style(hebrew_element())
    assert style["style"]["direction"] == "RIGHT_TO_LEFT" and "direction" in style["fields"].split(",")


def test_a_left_to_right_paragraph_is_told_nothing_it_was_not_told_before():
    style = paragraph_style(hebrew_element(direction=None, text="one two"))
    assert "direction" not in style["style"] and "direction" not in style["fields"]
    assert style["fields"] == "alignment,lineSpacing,spaceAbove,spaceBelow,indentStart,indentFirstLine"


@pytest.mark.parametrize("align, lines, alignment", [
    ("right", ((120.0, 200.0), (100.0, 200.0)), "START"),   # flush right: where it starts
    ("left", ((100.0, 200.0),), "START"),                   # one line: nothing was measured
    ("left", ((100.0, 200.0), (100.0, 200.0)), "START"),    # justified: it starts at the right too
    ("left", ((100.0, 200.0), (100.0, 170.0)), "END"),      # really ragged right: it ends there
    ("center", ((110.0, 190.0), (100.0, 200.0)), "CENTER"),
])
def test_a_hebrew_paragraph_keeps_the_edge_it_is_drawn_against(align, lines, alignment):
    # START and END are the reading direction's own ends, so the page's left and right have to
    # be mirrored into them (emit.hugs), or every Hebrew paragraph moves across its box.
    assert paragraph_style(hebrew_element(align, lines))["style"]["alignment"] == alignment


def test_a_hebrew_paragraph_is_indented_from_the_right():
    # The second paragraph starts 10 pt short of the right margin the first one reaches.
    el = hebrew_element(lines=((100.0, 200.0),))
    short = hebrew_element(lines=((100.0, 190.0),))["paragraphs"][0]
    el["paragraphs"].append(short)
    styles = [r["updateParagraphStyle"]["style"]
              for r in emit.text_box_requests(el, "b2s_s001", "b2s_s001_t0", SCALE, FONTS)
              if "updateParagraphStyle" in r]
    assert pt_of(styles[0]["indentStart"]) == 0
    assert pt_of(styles[1]["indentStart"]) == pytest.approx(10 * SCALE, abs=0.01)


def test_a_hebrew_bullet_hangs_on_the_right_of_its_item():
    bullet = {"kind": "glyph", "text": "●", "bbox": [204.0, 52.0, 210.0, 58.0], "color": "#000000"}
    style = paragraph_style(hebrew_element(bullet=bullet))["style"]
    # indentFirstLine is measured from the paragraph's own edge, which is the bullet's right.
    assert pt_of(style["indentFirstLine"]) == pytest.approx((210.0 - 204.0) * SCALE + emit.BULLET_GAP, abs=0.01)
    assert style["direction"] == "RIGHT_TO_LEFT"


def hebrew_table() -> dict:
    return {"id": "p0b0", "kind": "table", "frame": [100.0, 50.0, 200.0, 70.0], "size": 10.91,
            "columns": [{"x0": 100.0, "x1": 140.0, "align": "left"},
                        {"x0": 160.0, "x1": 200.0, "align": "left"}],
            "cells": [[[run_of(f"{ALEF}{BET}")], [run_of("one")]]],
            "row_baselines": [60.0], "row_heights": [20.0], "rules": []}


def test_a_hebrew_table_cell_reads_right_to_left_and_stays_where_it_is_drawn():
    reqs = table_requests(hebrew_table(), "b2s_s001", "b2s_s001_b0", SCALE, FONTS)
    styles = {r["updateParagraphStyle"]["cellLocation"]["columnIndex"]: r["updateParagraphStyle"]["style"]
              for r in reqs if "updateParagraphStyle" in r}
    assert styles[0]["direction"] == "RIGHT_TO_LEFT" and styles[0]["alignment"] == "END"
    assert "direction" not in styles[1] and styles[1]["alignment"] == "START"
