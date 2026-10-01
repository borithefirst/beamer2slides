"""Producers and consumers of deck.json elements for `test_ir_matrix.py`.

A **producer** is a function `(case, home) -> Made`: a deck (the IR every stage passes on) and the
stage it is at. A **consumer** is a function `(Made, home) -> None` that runs the real code a
journey runs on such a deck, offline, and raises when an element crashes it or goes missing on the
way. Consumers work on a deep copy: several of them read one cached producer's deck, and render,
`fit_holes` and `pictured_shapes` write into what they are handed.

Nothing here talks to Google. What only Google could answer is stood in for the way the other
offline tests do: the imported .pptx is read back with python-pptx (its slides are what the Drive
import copies from), a deck's read-back is `slides_sim.presentation_of`, and an adopted deck's
`presentations.get` is made from the target's own object ids (`pres_of`).

The decks are JSON as the stages hand them on (`JsonObject`), read through `json_reads`; what is
built here for a package function is its typed form (a `google_types.Presentation`, the .pptx's
pages as `PageDict`s)."""

import copy
import io
import json
import os
import shutil
from collections import Counter
from collections.abc import Callable, Collection, Mapping, Sequence
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Literal, TypedDict, TypeVar

import pytest
from pptx.shapes.base import BaseShape

from beamer2slides.adopt_sync import Folds
from beamer2slides.emit import OfflinePlan, Part
from beamer2slides.emit_model import PptxTable
from beamer2slides.emit_theme import BgKey
from beamer2slides.google_types import Dimension, Page, PageElement, Presentation, SlidesRequest
from beamer2slides.inverse import Candidate, Workspace
from beamer2slides.json_types import Json, JsonObject, as_object, as_optional_str
from beamer2slides.raw_types import RawDoc

from .json_reads import jarr, jint, jnum, jnums, jobj, jobjs, jstr

TESTS = Path(__file__).parent
DECKS = TESTS / "decks" / "out"

Stage = Literal["classified", "rendered", "planned", "read"]
_T = TypeVar("_T")


@dataclass(frozen=True, kw_only=True)
class AdoptedSource:
    """A showcase deck's read (`target`, deck_ir's JSON), the source adopt wrote for it (`tex`), the
    inverse `Workspace` it compiles in and the PDF that compile made (a copy, `pdf`)."""
    target: JsonObject
    tex: Path
    workspace: Workspace
    pdf: Path


@dataclass(frozen=True, kw_only=True)
class Extra:
    """What else a consumer of a `Made` needs, None where its producer made none: extract's `raw`,
    the `adopted` source and the deck_ir read it came from (`target`), the `folds` a base keeps,
    the elements a fallback deck `refused`, the pull loop's `candidate`, and the `source` a read-back
    was read from."""
    raw: RawDoc | None
    adopted: AdoptedSource | None
    target: JsonObject | None
    folds: Folds | None
    refused: list[tuple[int, str]] | None
    candidate: Candidate | None
    source: "Made | None"


NOTHING = Extra(raw=None, adopted=None, target=None, folds=None, refused=None, candidate=None, source=None)


@dataclass(frozen=True)
class Made:
    """What a producer made: `deck` at `stage` ("classified", "rendered", "planned", "read"), the
    folder its files (backgrounds, figures) are relative to, the PDF it came from, the width in
    slide pt of the deck it is written into, and what else a consumer of it needs (`extra`).
    (Positional, and the last two defaulted, as test_sync_containment builds one.)"""
    deck: JsonObject
    stage: Stage
    out: Path | None
    pdf: Path | None
    page_width: float | None = None
    extra: Extra = NOTHING


def need(value: "_T | None", what: str) -> _T:
    """`value`, which the producer of the deck at hand makes."""
    assert value is not None, f"the producer made no {what}"
    return value


# ---------------------------------------------------------------- reading a deck

def slides_of(deck: JsonObject) -> list[JsonObject]:
    return jobjs(deck, "slides")


def elements_of(slide: JsonObject) -> list[JsonObject]:
    return jobjs(slide, "elements")


def page_of(slide: JsonObject) -> int:
    return jint(slide, "page")


def with_elements(slide: JsonObject, elements: Sequence[JsonObject]) -> JsonObject:
    listed: list[Json] = [*elements]
    return {**slide, "elements": listed}


def with_slides(deck: JsonObject, slides: Sequence[JsonObject]) -> JsonObject:
    listed: list[Json] = [*slides]
    return {**deck, "slides": listed}


def read_json(value: object, where: str) -> JsonObject:
    """`value` as the next process reads it back: JSON, tuples as lists and every key a string."""
    return as_object(json.loads(json.dumps(value, ensure_ascii=False)), where)


# ---------------------------------------------------------------- producers

def _convert_pdf(pdf: Path, out: Path) -> tuple[JsonObject, RawDoc, Path]:
    """convert's local half (`__main__.cmd_convert`, `agent.deck_tools._prepare`,
    `sync.build_ours`): notes taken out, extract, the last overlay steps, classify, render."""
    from beamer2slides.classify import classify
    from beamer2slides.extract import extract, select_overlays
    from beamer2slides.ir import deck_json
    from beamer2slides.notes import prepare
    from beamer2slides.render import render_backgrounds

    out.mkdir(parents=True, exist_ok=True)
    prepared = prepare(pdf, out)
    raw = extract(prepared.pdf, prepared.labels)
    for page in raw["pages"]:
        page["notes"] = prepared.notes.get(page["index"])
    raw = select_overlays(raw, "last")
    deck = deck_json(classify(raw))
    render_backgrounds(prepared.pdf, raw, deck, out, frozenset())
    return deck, raw, prepared.pdf


def built_pdf(name: str) -> Path:
    pdf = DECKS / f"{name}.pdf"
    if not pdf.exists():
        pytest.skip(f"{pdf.name} not built (python tests/decks/build.py)")
    return pdf


def converted(case: str, home: Path) -> Made:
    """A beamer PDF of the test decks, classified and rendered as `convert` does."""
    pdf = built_pdf(case)
    deck, raw, used = _convert_pdf(pdf, home / "out")
    return Made(deck, "rendered", home / "out", used, None, replace(NOTHING, raw=raw))


def adopted_source(case: str, home: Path, target: JsonObject | None) -> AdoptedSource:
    """A showcase deck (or `target`, any deck_ir read) adopted (`adopt.bootstrap`, fonts none as
    in test_adopt_compiles, nothing downloaded) and its source compiled in an inverse `Workspace`,
    as `adopt_sync.convert_source` and the pull loop compile it."""
    from beamer2slides import adopt
    from beamer2slides.inverse import tex_env

    from .test_adopt_compiles import engine

    read = showcase_target(case) if target is None else target
    tex = home / "tree" / "main.tex"
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("B2S_FONTS", str(home / "no-fonts-here"))
        mp.delenv("B2S_FONT_FETCH", raising=False)
        mp.setenv("B2S_NO_DOWNLOADS", "1")  # (offline: no font or picture fetch whatever the deck names)
        adopt._FAMILIES.clear()
        try:
            text = adopt.bootstrap(read, tex, False, None)
        finally:
            adopt._FAMILIES.clear()
    if shutil.which(engine(text), path=tex_env()["PATH"]) is None:
        if os.environ.get("B2S_REQUIRE_TEX"):
            pytest.fail(f"{engine(text)} not found")
        pytest.skip(f"{engine(text)} not found")
    ws = Workspace(tex, home / "work", handout=False, engine=None, fresh=True)
    pdf, err = ws.compile()
    assert pdf is not None, err
    # (a copy: the workspace compiles again into its build folder when the pull loop wants notes)
    kept = home / "compiled.pdf"
    shutil.copyfile(pdf, kept)
    return AdoptedSource(target=read, tex=tex, workspace=ws, pdf=kept)


def adopted(case: str, home: Path, source: AdoptedSource) -> Made:
    """The source adopt wrote for a showcase deck, compiled and converted (classify reads its
    marks: `marked.classify_marked`; render pictures what Slides cannot draw)."""
    deck, raw, used = _convert_pdf(source.pdf, home / "out")
    return Made(deck, "rendered", home / "out", used, deck_width(source.target),
                replace(NOTHING, raw=raw, adopted=source, target=source.target))


def deck_width(target: JsonObject) -> float:
    """The width in slide pt of the deck adopt read (its PDF pt times the read's scale)."""
    return jnum(target, "slides", 0, "size", 0) * jnum(target, "scale")


def json_round_trip(made: Made) -> Made:
    """deck.json as the next process reads it: `deck_prepare` writes it and `deck_upload` reads it
    back (agent/deck_tools.py), and a base's `ir` comes back from Drive the same way. Tuples are
    lists and every key a string."""
    deck = read_json(made.deck, "deck.json")
    return Made(deck, made.stage, made.out, made.pdf, made.page_width, made.extra)


def folded(made: Made) -> Made:
    """An adopted source's conversion folded against the deck's own boxes, as `adopt_sync.record`
    (via `convert_source`) and a later `sync.build_ours` fold it."""
    from beamer2slides.adopt_sync import deck_folds, fold_slides

    deck = copy.deepcopy(made.deck)
    folds = deck_folds(need(made.extra.target, "deck_ir read"))
    fold_slides(deck, folds)
    return Made(deck, made.stage, made.out, made.pdf, made.page_width, replace(made.extra, folds=folds))


def blocks_merged(deck: JsonObject) -> JsonObject:
    """The deck with every slide's blocks merged (`emit.merge_blocks`), as emit plans it."""
    from beamer2slides.emit import merge_blocks

    return with_slides(deck, [with_elements(s, merge_blocks(elements_of(s))) for s in slides_of(deck)])


def with_fallbacks(made: Made) -> Made:
    """What `emit.emit` builds the deck from again when the API refused elements: every element
    but the pictures refused, each replaced by a crop of its region (`emit.fallback_pictures`)."""
    from beamer2slides.emit import fallback_pictures

    deck = blocks_merged(copy.deepcopy(made.deck))
    refused = [(page_of(s), jstr(e, "id")) for s in slides_of(deck) for e in elements_of(s) if e["kind"] != "image"]
    return Made(fallback_pictures(deck, refused, need(made.out, "out folder")), made.stage, made.out, made.pdf,
                made.page_width, replace(made.extra, refused=refused))


def candidate(made: Made, home: Path) -> Made:
    """The pull loop's own reading of an adopted source (`inverse.Workspace.build`: classify and
    `keep_visible_shapes`, no render), from the compile the producer already made."""
    source = need(made.extra.adopted, "adopted source")
    cand = source.workspace.build(home / "classify", any(s.get("notes") for s in slides_of(source.target)),
                                  compiled=(source.pdf, ""))
    assert not isinstance(cand, str), cand
    return Made(cand.deck, "classified", None, cand.pdf, made.page_width, replace(made.extra, candidate=cand))


def target(case: str, home: Path) -> Made:
    """What `deck_ir` read from a deck a person made in Slides (the showcase fixtures)."""
    t = showcase_target(case)
    return Made(t, "read", None, None, deck_width(t), replace(NOTHING, target=t))


# ---------------------------------------------------------------- showcase decks and their variants

def blank_tables(t: JsonObject) -> None:
    """Every table's words gone: a grid of (coloured) empty cells, as a person lays one out before
    filling it in (a corpus deck has one)."""
    for s in slides_of(t):
        for e in elements_of(s):
            if e["kind"] == "table":
                rows: list[Json] = []
                for row in jarr(e, "rows"):
                    blank: list[Json] = ["" for _ in jarr(row)]
                    rows.append(blank)
                e["rows"] = rows
                for cell in jobjs(e, "table_cells"):
                    none: list[Json] = []
                    cell["paragraphs"] = none


def ellipsis_then_letter(t: JsonObject) -> None:
    """The last paragraph of the slide's last text box ends on an ellipsis after its last word,
    and a paragraph of one letter follows it (a corpus deck's labels under a quote)."""
    box = [e for e in elements_of(slides_of(t)[0]) if e["kind"] == "text"][-1]
    paragraphs = jarr(box, "paragraphs")
    last = jobj(paragraphs[-1])
    last_run = jobjs(last, "runs")[-1]
    last_run["text"] = jstr(last_run, "text").rstrip() + "…"
    letter = copy.deepcopy(last)
    letter["runs"] = [{**last_run, "text": "E"}]
    letter["lines"] = jarr(last, "lines")[-1:]
    paragraphs.append(letter)


def note_opening_on_a_break(t: JsonObject) -> None:
    """Speaker notes whose second paragraph opens on a soft break (Shift+Enter at its start)."""
    slides_of(t)[0]["notes"] = "Say why first.\n\x0bThen say how."


# case -> (showcase deck, the one slide kept, what is changed on it)
VARIANTS: dict[str, tuple[str, int, Callable[[JsonObject], None]]] = {
    "hashing-blank_table": ("hashing", 4, blank_tables),
    "hashing-ellipsis": ("hashing", 1, ellipsis_then_letter),
    "hashing-note_break": ("hashing", 1, note_opening_on_a_break),
}


def showcase_of(case: str) -> str:
    return VARIANTS[case][0] if case in VARIANTS else case


def showcase_target(case: str) -> JsonObject:
    """The deck_ir read of a showcase deck (`tests/decks/foreign/showcase`), or of a variant: one
    of its slides with one thing a person's deck can have that no showcase slide has."""
    from .test_adopt_compiles import SHOWCASE, target_of

    name = showcase_of(case)
    if not (SHOWCASE / name / "target.json").exists():
        pytest.skip(f"{name}: no showcase fixture (tools/showcase.py fixture)")
    t = as_object(target_of(name), f"{name}/target.json")
    if case in VARIANTS:
        _, keep, change = VARIANTS[case]
        t = copy.deepcopy(t)
        t = with_slides(t, [{**slides_of(t)[keep], "page": 0}])
        change(t)
    return t


def read_back(made: Made) -> Made:
    """`deck_ir` of the deck convert would make of `made` (`slides_sim.presentation_of`): what
    `pull` reads from a converted deck nobody edited."""
    from beamer2slides.deck_ir import deck_ir

    from .slides_sim import presentation_of

    size = jnums(made.deck, "slides", 0, "size")
    return Made(deck_ir(presentation_of(copy.deepcopy(made.deck)), size, None, None, None, False, None), "read",
                None, None, None, replace(NOTHING, source=made))


# ---------------------------------------------------------------- what every consumer checks

def ids(deck: JsonObject) -> list[list[str]]:
    return [[jstr(e, "id") for e in elements_of(s)] for s in slides_of(deck)]


def same_elements(before: JsonObject, after: JsonObject, what: str) -> None:
    """No element came or went: `merge_blocks`, `fit_holes` and the folds reshape, they never drop."""
    assert [len(elements_of(s)) for s in slides_of(before)] == [len(elements_of(s)) for s in slides_of(after)] \
        and ids(before) == ids(after), f"{what} changed the elements: {ids(before)} -> {ids(after)}"


def merged(deck: JsonObject) -> JsonObject:
    out = blocks_merged(deck)
    assert [sorted(i) for i in ids(deck)] == [sorted(i) for i in ids(out)], "merge_blocks dropped an element"
    return out


def written(slide: JsonObject, parts: Sequence[Part], element_ids: Sequence[str], carried: Collection[str],
            what: str) -> None:
    """Every element of `slide` has an object: created by its own requests, or carried by what the
    slide was copied from (a picture or a table the .pptx brought, a layout placeholder).
    `parts`: `DeckPlan.slide_parts`' (the first part is the slide's cleanup, then one per element)."""
    from beamer2slides.emit import created_ids

    elements = elements_of(slide)
    assert len(element_ids) == len(elements) == len(parts[1:1 + len(element_ids)]), \
        f"{what}: {len(elements)} elements, {len(element_ids)} object ids"
    missing: list[str] = []
    for el, oid, (part_el, reqs) in zip(elements, element_ids, parts[1:]):
        assert part_el is not None and part_el["id"] == el["id"], f"{what}: parts out of step with the elements"
        made = created_ids(reqs)
        # (a diagram of one object gets no group, Slides grouping two or more: its object is the
        # diagram, `snapshot.attach_readback`)
        if oid not in carried and oid not in made and not (el["kind"] == "diagram" and made):
            missing.append(f"{el['kind']} {el['id']} ({oid})")
        elif not says(el, oid, reqs):
            missing.append(f"{el['kind']} {el['id']} ({oid}) without its words")
    assert not missing, f"{what}: elements with no object: {missing}"


def copied_ids(request: SlidesRequest) -> dict[str, str]:
    """The object ids a slide's duplicateObject gives its copies (`emit.copy_request`)."""
    dup = request.get("duplicateObject")
    ids = None if dup is None else dup.get("objectIds")
    assert ids is not None, f"a copy that is no duplicateObject with objectIds: {request}"
    return dict(ids)


def says(el: JsonObject, oid: str, reqs: Sequence[Mapping[str, object]]) -> bool:
    """A text element with words has them inserted into its object (a placeholder the slide was
    copied with exists whether or not anything is written into it)."""
    if el["kind"] != "text" or not any(jstr(r, "text").strip() for p in jobjs(el.get("paragraphs", []))
                                       for r in jobjs(p, "runs")):
        return True
    for r in reqs:
        inserted = r.get("insertText")   # (emit's typed requests, or sync's JSON)
        if isinstance(inserted, Mapping) and inserted.get("objectId") == oid:
            words = inserted.get("text")
            if isinstance(words, str) and words.strip():
                return True
    return False


# ---------------------------------------------------------------- convert (emit.build_deck, offline)

PLACEHOLDER_TYPES = {"TITLE": "TITLE", "CENTER_TITLE": "CENTERED_TITLE", "SUBTITLE": "SUBTITLE"}


class ColorFill(TypedDict):
    color: str


class PictureFill(TypedDict):
    picture: Path


Fill = ColorFill | PictureFill


class PictureDict(TypedDict):
    """A picture of a `build_pptx` page (`emit_pptx.pptx_page` reads it)."""
    file: Path | str
    bbox: list[float]
    alt: str | None
    title: str


class PageDict(TypedDict):
    """A source slide of the .pptx as `build_pptx` takes it (`emit_pptx.pptx_page` reads it)."""
    layout: str
    fill: Fill | None
    pictures: list[PictureDict]
    tables: list[PptxTable]
    templates: bool


def source_slides(pptx: io.BytesIO) -> list[JsonObject]:
    """The slides of the .pptx as the Drive import brings them (`presentations.get`'s
    pageElements, as far as `DeckPlan.copy_request` reads them), in the order python-pptx wrote
    them: layout placeholders, pictures, tables, template shapes."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    def size(shape: BaseShape) -> JsonObject:
        return {"width": {"magnitude": shape.width or 0, "unit": "EMU"},
                "height": {"magnitude": shape.height or 0, "unit": "EMU"}}

    out: list[JsonObject] = []
    for k, slide in enumerate(Presentation(pptx).slides):
        els: list[Json] = []
        for j, shape in enumerate(slide.shapes):
            oid = f"src{k:03}_{j}"
            if shape.is_placeholder:
                kind = shape.placeholder_format.type.name
                els.append({"objectId": oid, "size": size(shape),
                            "shape": {"placeholder": {"type": PLACEHOLDER_TYPES.get(kind, kind)}}})
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                els.append({"objectId": oid, "size": size(shape), "image": {}})
            elif getattr(shape, "has_table", False):  # (a graphic frame holding a table)
                els.append({"objectId": oid, "size": size(shape), "table": {}})
            else:
                els.append({"objectId": oid, "size": size(shape), "shape": {}})
        out.append({"objectId": f"src{k:03}", "pageElements": els})
    return out


def picture_title(e: JsonObject) -> str:
    """The title a picture of the .pptx gets (`emit.PICTURE_TITLES` by its role)."""
    from beamer2slides.emit import PICTURE_TITLES

    role = e.get("role")
    return PICTURE_TITLES.get(role, "Figure") if isinstance(role, str) else "Figure"


WHITE: BgKey = ("color", "#ffffff")


def convert(made: Made, home: Path) -> None:
    """`emit.build_deck` up to what only Google can answer: the plan, the backgrounds and theme,
    the .pptx with every page's real pictures, tables and template shapes, phase 1's copy of each
    imported slide (against that .pptx read back), the layouts' styles and texts, and phase 2's
    requests for every slide, with the table margins the base records."""
    from beamer2slides import emit

    deck = copy.deepcopy(made.deck)
    out = need(made.out, "out folder")
    page_w, page_h = jnums(deck, "slides", 0, "size")
    plan = emit.DeckPlan(merged(deck), emit.SLIDE_W, pptx_tables=True, contain=False)
    same_elements(merged(deck), plan.deck, "DeckPlan")
    deck, scale, fonts = plan.deck, plan.scale, plan.fonts
    bg_key = {page_of(s): emit.background_key(s, out) for s in slides_of(deck)}
    bg_file = {bg_key[page_of(s)]: out / jstr(s, "background") for s in slides_of(deck)
               if not s.get("background_color")}
    counts = Counter(bg_key.values())
    shared = counts.most_common(1)[0][0] if counts and counts.most_common(1)[0][1] >= 2 else None

    def fill(key: BgKey) -> Fill:
        return {"color": key[1]} if key[0] == "color" else {"picture": bg_file[key]}

    def group(s: JsonObject) -> str:
        return "TITLE" if emit.slide_layout(s)[0] == "TITLE" else "*"

    theme = emit.plan_theme(deck, out, bg_key)
    master_fill = fill(WHITE if shared is None else shared)
    if theme:
        if shared is None or all(theme["exact"].get(group(s)) == shared for s in slides_of(deck)
                                 if bg_key[page_of(s)] == shared):
            master_fill = {"color": theme["ground"]}
    pages: list[PageDict] = [{
        "layout": theme["layouts"][page_of(s)] if theme else emit.slide_layout(s)[0],
        "fill": None if bg_key[page_of(s)] == shared else fill(bg_key[page_of(s)]),
        "pictures": [{"file": out / jstr(e, "file"), "bbox": bbox, "alt": as_optional_str(e.get("alt"), "alt"),
                      "title": picture_title(e)} for e, bbox in plan.pictures(s)],
        "tables": plan.tables(s),
        "templates": plan.uses_templates[page_of(s)],
    } for s in slides_of(deck)]
    pictures = {page_of(s): {jstr(e, "id") for e, _ in plan.pictures(s)} for s in slides_of(deck)}
    for s in slides_of(deck):
        for e in elements_of(s):
            if e["kind"] == "image":
                assert jstr(e, "id") in pictures[page_of(s)] and (out / jstr(e, "file")).is_file(), \
                    f"slide {page_of(s) + 1}: picture {e['id']} has no file the .pptx could carry"
    pptx = emit.build_pptx(page_w, page_h, plan.keys, pages, master_fill,
                           None if theme is None else theme["decorations"])
    sources = source_slides(pptx)
    assert len(sources) == len(slides_of(deck)), "a source slide per deck slide"

    template_sizes: list[tuple[float, float]] = []
    page_elements: dict[str, list[JsonObject]] = {}
    carried: dict[str, set[str]] = {}
    for slide, source in zip(slides_of(deck), sources):
        request, sizes = plan.copy_request(slide, source)
        template_sizes = template_sizes or sizes
        new = copied_ids(request)
        sid = new[jstr(source, "objectId")]
        page_elements[sid] = [{"objectId": new.get(jstr(e, "objectId"), f"{jstr(e, 'objectId')}_copy"),
                               "size": e["size"]} for e in jobjs(source, "pageElements")]
        carried[sid] = set(new.values())
    speaker_notes = {sid: f"{sid}_notes" for sid in page_elements}

    ground = emit.master_ground(shared, bg_file, page_w)
    emit.layout_style_spec(deck, scale, fonts, emit.PPTX_TITLE_DY, ground)
    for ti, el in enumerate(jobjs(deck.get("layout_texts") or [])):
        assert emit.text_box_requests(el, "layout", f"{emit.LAYOUT_TEXT_PREFIX}0_{ti}", scale, fonts)

    for slide in slides_of(deck):
        sid = f"b2s_s{page_of(slide):03}"
        parts, element_ids = plan.slide_parts(slide, page_elements, speaker_notes, {}, template_sizes, None)
        written(slide, parts, element_ids, carried[sid], f"slide {page_of(slide) + 1}")
        emit.element_objects(parts, element_ids)
        for el in elements_of(slide):
            if el["kind"] == "table":
                emit.pptx_table(el, scale, fonts, emit.SLIDE_W / scale)["margins"]


# ---------------------------------------------------------------- what sync writes, element by element

def emission(made: Made, home: Path) -> None:
    """`sync.mark_emitted`'s two sides: a slide's emission worked out from that slide alone
    (`sync.emitted_elements` over `emit.slide_emission`), for every slide of the plan a sync makes
    at the deck's own width (`DeckPlan(page_width)`, `sync.build_ours`, `adopt_sync.convert_source`)
    - and for a base's IR, which is exactly such a slide read back from JSON."""
    from beamer2slides.emit import SLIDE_W, DeckPlan
    from beamer2slides.sync import emitted_elements

    deck = copy.deepcopy(made.deck)
    plan = DeckPlan(merged(deck), made.page_width or SLIDE_W, pptx_tables=False, contain=False)
    same_elements(merged(deck), plan.deck, "DeckPlan")
    for slide in slides_of(plan.deck):
        names = [f"e{i}" for i in range(len(elements_of(slide)))]
        got = emitted_elements(slide, names, plan.scale, plan.fonts)
        assert len(got) == len(elements_of(slide)), f"slide {page_of(slide) + 1}: an element emitted nothing"
        again = read_json(slide, "slide")
        assert emitted_elements(again, names, plan.scale, plan.fonts) == got, \
            f"slide {page_of(slide) + 1}: its IR read back from a base emits differently"


# ---------------------------------------------------------------- the base convert records

def grid_json(rows: Sequence[Sequence[float]]) -> list[Json]:
    out: list[Json] = []
    for r in rows:
        row: list[Json] = [*r]
        out.append(row)
    return out


def strings_json(rows: Sequence[Sequence[str]]) -> list[Json]:
    out: list[Json] = []
    for r in rows:
        row: list[Json] = [*r]
        out.append(row)
    return out


def convert_state(deck: JsonObject) -> tuple[OfflinePlan, JsonObject]:
    """(plan_offline, emit.json's state) for a deck, as `build_deck` records it (element objects,
    groups, the .pptx tables' margins)."""
    from beamer2slides.emit import SLIDE_W, element_objects, plan_offline, pptx_table

    off = plan_offline(deck)
    plan = off["plan"]
    slides: list[Json] = []
    for n, (slide, (sid, page, parts, element_ids)) in enumerate(zip(slides_of(plan.deck), off["slides"])):
        copied = copied_ids(off["copies"][n])
        written(slide, parts, element_ids, set(copied.values()), f"slide {page + 1}")
        objects, groups = element_objects(parts, element_ids)
        margins: JsonObject = {str(i): grid_json(pptx_table(el, plan.scale, plan.fonts, SLIDE_W / plan.scale)["margins"])
                               for i, el in enumerate(elements_of(slide)) if el["kind"] == "table"}
        slides.append({"page": page, "objectId": sid, "elements": [*element_ids], "objects": strings_json(objects),
                       "groups": [*groups], "table_margins": margins})
    state: JsonObject = {"presentationId": "simulated", "scale": plan.scale, "slides": slides}
    return off, state


def base_of(made: Made) -> tuple[JsonObject, JsonObject]:
    """The base `convert` records (`snapshot.build_base`), over the simulated read-back, as a
    later sync loads it (JSON); and that read-back."""
    from beamer2slides import snapshot
    from beamer2slides.google_types import presentation

    from .slides_sim import presentation_of

    deck = copy.deepcopy(made.deck)
    off, state = convert_state(deck)
    pres = presentation_of(copy.deepcopy(deck))
    base = snapshot.build_base(off["plan"].deck, need(made.out, "out folder"),
                               presentation(pres, "the simulated deck"), state, made.pdf, 0, False, "last", None)
    for entry, slide in zip(slides_of(base), slides_of(off["plan"].deck)):
        assert [jstr(e, "id") for e in elements_of(entry)] == [jstr(e, "id") for e in elements_of(slide)], \
            f"slide {page_of(slide) + 1}: the base lost an element"
        assert len({jstr(e, "key") for e in elements_of(entry)}) == len(elements_of(entry)), \
            f"slide {page_of(slide) + 1}: two elements share a key"
    return read_json(base, "base.json"), pres


def convert_base(made: Made, home: Path) -> None:
    base_of(made)


def pres_of(target: JsonObject) -> Presentation:
    """A `presentations.get` of the deck adopt read, as far as `adopt_sync.build_base` reads it:
    its slides and the objects on them, by the ids the read gave (no test holds a recorded one)."""
    scale = jnum(target, "scale")

    def emu(v: float) -> Dimension:
        return {"magnitude": v * scale * 12700, "unit": "EMU"}

    slides: list[Page] = []
    for s in slides_of(target):
        seen: set[str] = set()
        els: list[PageElement] = []
        for e in elements_of(s):
            oid = as_optional_str(e.get("object"), "object")
            if not oid or oid in seen:
                continue
            seen.add(oid)
            x0, y0, x1, y1 = jnums(e, "bbox")
            els.append({"objectId": oid, "size": {"width": emu(x1 - x0), "height": emu(y1 - y0)},
                        "transform": {"scaleX": 1, "scaleY": 1, "translateX": x0 * scale * 12700,
                                      "translateY": y0 * scale * 12700, "unit": "EMU"},
                        "shape": {"shapeType": "TEXT_BOX"}})
        slides.append({"objectId": jstr(s, "objectId"), "pageElements": els})
    w, h = jnums(target, "slides", 0, "size")
    return {"presentationId": "adopted", "revisionId": "r1",
            "pageSize": {"width": emu(w), "height": emu(h)}, "slides": slides, "layouts": [], "masters": []}


def adopt_base_of(made: Made) -> tuple[JsonObject, Presentation]:
    """`adopt_sync.record`'s second half on a folded conversion: the plan at the deck's width,
    the labels checked, the conversion paired with the deck's objects (`pair_elements`,
    `explained_by_layout`, `inside_tables`, `drawn_from`) and the base recorded."""
    from beamer2slides import adopt_sync
    from beamer2slides.emit import DeckPlan
    from beamer2slides.google_types import as_json

    target = need(made.extra.target, "deck_ir read")
    deck = copy.deepcopy(made.deck)
    plan = DeckPlan(merged(deck), need(made.page_width, "deck width"), pptx_tables=False, contain=False)
    assert adopt_sync.labels_match(plan.deck, target) is None, adopt_sync.labels_match(plan.deck, target)
    pres = pres_of(target)
    base = adopt_sync.build_base(plan.deck, need(made.out, "out folder"), target, as_json(pres, "the adopted deck"),
                                 need(made.pdf, "PDF"), "last", made.extra.folds)
    for entry, slide in zip(slides_of(base), slides_of(plan.deck)):
        assert [jstr(e, "id") for e in elements_of(entry)] == [jstr(e, "id") for e in elements_of(slide)], \
            f"slide {page_of(slide) + 1}: the base lost an element"
    return read_json(base, "base.json"), pres


def adopt_base(made: Made, home: Path) -> None:
    adopt_base_of(made)


# ---------------------------------------------------------------- sync (build_ours and the writes)

def duplicated(r: JsonObject) -> set[str]:
    """The object ids a duplicateObject request names for its copies."""
    dup = r.get("duplicateObject")
    ids = dup.get("objectIds") if isinstance(dup, dict) else None
    return {jstr(v) for v in ids.values()} if isinstance(ids, dict) else set()


def resync(made: Made, home: Path) -> None:
    """A sync of the same PDF into the deck the base describes: `sync.build_ours` (which folds
    against an adopted base's boxes, plans at the deck's width, keys and records the new side),
    `mark_emitted` on every paired slide (fast off: both emissions are worked out), the merge
    against the deck as the base read it (`merge.plan_merge`), theme sync's fresh side, and
    `Sync.slide_requests` recreating every element of every slide - a title into its
    placeholder, a table refilled in place as `table_refill` does - with the staging .pptx that
    brings the pictures."""
    from beamer2slides import merge, snapshot, sync, theme_sync
    from beamer2slides.emit import SLIDE_W, build_pptx, title_element
    from beamer2slides.google_types import presentation

    from .fake_google import NoDrive, NoSlides

    if made.extra.folds is not None:
        base, pres = adopt_base_of(made)
    else:
        base, simulated = base_of(made)
        pres = presentation(simulated, "the simulated deck")
    # (the width went in as `overlays` while build_ours had defaults: planned at SLIDE_W whatever the deck)
    base_pictures = snapshot.find_base_pictures(base, snapshot.PictureFolders(kept=(made.out,), rendered=None,
                                                                              held=None)) \
        if made.out else snapshot.NO_PICTURES
    ours = sync.build_ours_of(need(made.pdf, "PDF"), home / "ours", base, "last", made.page_width or SLIDE_W,
                              base_pictures)
    # a base this converter just wrote is in today's form: nothing to rewrite, nothing it cannot read
    assert snapshot.base_form_json(ours.base_forms) == []
    for s, entry in zip(slides_of(ours.deck), ours.slides):
        assert [jstr(e, "id") for e in elements_of(entry)] == [jstr(e, "id") for e in elements_of(s)]
    sync.mark_emitted(base, ours.slides, ours.deck, ours.pairs, ours.plan.scale, ours.plan.fonts,
                      fast=False, unread=[])
    merge.plan_merge(base, sync.ours_json(ours), snapshot.read_presentation(pres))
    theme_sync.ours_side(sync.theme_ours_of(ours))

    s = sync.Sync(NoSlides(), NoDrive(), "offline", base, ours, home / "ours", dry_run=True, measure=False,
                  trust_generation=True, check_plan=None, follow_labels=False, take_source=(), facts=None,
                  way_back=None)
    pictures: dict[str, str] = {}
    for j, slide in enumerate(slides_of(ours.deck)):
        title = title_element(slide)
        in_place: dict[int, sync.Refilled] = {}
        if title is not None:  # (a placeholder holding words, emptied first)
            in_place[title] = sync.InPlace(id=f"live{j}_title", size=(sync.STAND_IN, sync.STAND_IN), text="x")
        for i, el in enumerate(elements_of(slide)):
            if el["kind"] == "table":
                in_place[i] = sync.TableFill(id=f"live{j}_tab{i}", cells=[], steps=[], shift=(0.0, 0.0), margins=[])
            if el["kind"] == "image":
                pictures[str(ours.out / jstr(el, "file"))] = "picture"
        if slide.get("background") and not slide.get("background_color"):
            pictures[str(ours.out / jstr(slide, "background"))] = "background"
        units = list(range(len(elements_of(slide))))
        reqs, objects, new_oid, _ = s.slide_requests(j, units, f"live{j}", in_place, {}, {}, True, frozenset())
        created = {jstr(r, k, "objectId") for r in reqs
                   for k in ("createShape", "createLine", "createTable", "createImage") if k in r} \
            | {v for r in reqs for v in duplicated(r)} \
            | {jstr(r, "groupObjects", "groupObjectId") for r in reqs if "groupObjects" in r}
        lost = [f"{el['kind']} {el['id']}" for i, el in enumerate(elements_of(slide))
                if i not in objects or (new_oid[i] not in created and i not in in_place
                                        and not (el["kind"] == "diagram" and created & set(objects[i])))
                or not says(el, new_oid[i], reqs)]
        assert not lost, f"slide {page_of(slide) + 1}: sync would write nothing for {lost}"
    if pictures:
        page_w, page_h = jnums(ours.deck, "slides", 0, "size")
        still = [f for f, what in pictures.items() if what != "background"]
        pages: list[PageDict] = [{"layout": "BLANK", "fill": None, "templates": False, "tables": [], "pictures": [
            {"file": f, "bbox": [0, 0, *s._fit(f)], "alt": f"b2s-stage:{k + n}", "title": "stage"}
            for n, f in enumerate(still[k:k + 40])]} for k in range(0, len(still), 40)]
        pages += [{"layout": "BLANK", "fill": {"picture": Path(f)}, "templates": False, "tables": [], "pictures": []}
                  for f, what in pictures.items() if what == "background"]
        white: ColorFill = {"color": "#ffffff"}
        assert len(source_slides(build_pptx(page_w, page_h, [], pages, white, None))) == len(pages)


# ---------------------------------------------------------------- pull and adopt (deck_ir's side)

def compare_read_back(made: Made, home: Path) -> None:
    """`compare` of a converted deck with what `deck_ir` reads from it (the pull of a converted,
    unedited deck): it runs, and every element of the source is compared."""
    from beamer2slides.compare import TOL, compare, without_keys
    from beamer2slides.inverse import typed_target

    source = need(made.extra.source, "converted deck")
    comp = compare(without_keys(copy.deepcopy(source.deck)), typed_target(copy.deepcopy(made.deck)), TOL, {})
    comp.open()
    comp.summary()


def ignore(line: str) -> None:
    """A log that keeps nothing."""


def round0(made: Made, home: Path) -> None:
    """The pull loop's first round on an adopted source (`inverse.converge`, `adopt_replay`):
    picture hashes, `compare` with the deck adopt read, the frame guard's look (no thumbnails:
    residuals and words) and the planner's edits."""
    from beamer2slides.compare import TOL, compare
    from beamer2slides.frame_guard import FrameGuard
    from beamer2slides.inverse import Planner, class_pt_option, fresh_context, picture_hashes, typed_target

    cand = need(made.extra.candidate, "pull loop candidate")
    target = need(made.extra.target, "deck_ir read")
    ws = need(made.extra.adopted, "adopted source").workspace
    typed = typed_target(target)
    hashes = picture_hashes(cand, typed, home / "work")
    comp = compare(cand.current(), typed, TOL, hashes)
    comp.summary()
    FrameGuard(typed, None, ignore).observe(0, cand, comp)
    Planner(cand, comp, target, fresh_context(class_pt_option(ws.source)), ws, set(), {}, hashes).plan()


def adopt_read(made: Made, home: Path) -> None:
    """What `adopt._adopt` asks of the deck it read besides the bootstrap (which every adopted
    producer runs): the pictures it has to leave out or crop, and the fold records a base keeps."""
    from beamer2slides import adopt
    from beamer2slides.adopt_sync import deck_folds

    t = copy.deepcopy(made.deck)
    adopt.pictures_missing(t)
    adopt.pictures_from_thumbnail(t)
    folds = deck_folds(t)
    assert len(folds) == len([s for s in slides_of(t)]), "a fold record per slide label"
