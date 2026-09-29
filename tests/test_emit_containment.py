"""One element emit cannot plan does not fail the conversion (`emit.DeckPlan.contain`).

Reported (688ebf4): `deck_upload` of a source `deck_adopt` wrote died with KeyError 'custom', then
'flip' - `marked.py`'s read of an adopted source made shape elements emit's planner had never
seen. Now such an element is the picture of its region, as an element Slides refused is
(`emit.fallback_element`), with a warning; the other elements are written as they would have
been. Under `emit.STRICT_ENV` (the offline suite's setting) the failure is raised instead.
No Google API.
"""

from collections.abc import Callable
from pathlib import Path

import pytest

from beamer2slides import emit
from beamer2slides.emit_model import Template, TemplateKey
from beamer2slides.emit_pptx import shape_element_requests as real_shape_element_requests
from beamer2slides.emit_pptx import shape_requests as real_shape_requests
from beamer2slides.ir_types import IRError, MarkedShape, ShapeElement
from beamer2slides.json_types import Json, JsonObject

from .json_reads import jint, jnums, jobjs, jstr
from .test_classify import deck

ShapePlanner = Callable[[ShapeElement | MarkedShape, str, str, float, Callable[[TemplateKey], Template]], list[JsonObject]]


@pytest.fixture
def lenient(monkeypatch: pytest.MonkeyPatch) -> None:
    """Containment on, whatever the suite's default."""
    monkeypatch.setenv(emit.STRICT_ENV, "0")


def raising_for(target: str) -> ShapePlanner:
    """`shape_element_requests` failing for the element `target` as a missing field would. (emit
    plans a parsed shape: `el` is its ir_types record.)"""
    def shape_element_requests(el: ShapeElement | MarkedShape, slide_id: str, object_id: str, scale: float,
                               template_for: Callable[[TemplateKey], Template]) -> list[JsonObject]:
        if el.id == target:
            raise KeyError("flip")
        return real_shape_element_requests(el, slide_id, object_id, scale, template_for)
    return shape_element_requests


def shape(eid: str, x0: float, **fields: Json) -> JsonObject:
    el: JsonObject = {"kind": "shape", "id": eid, "role": "panel", "shape": "RECTANGLE", "flip": False, "radius": 0.0,
                      "bbox": [x0, 30.0, x0 + 100.0, 90.0], "fill": "#3366cc", "spans": [], "drawing": f"d-{eid}"}
    el.update(fields)
    return el


def adopted_custom() -> JsonObject:
    """688ebf4's element: a marked shape of a kind no Slides preset has, with no `flip`."""
    return {"kind": "shape", "id": "p0-custom", "role": "panel", "shape": "custom", "bbox": [20.0, 30.0, 120.0, 90.0],
            "fill": "#ff0000", "outline": None, "radius": 0.0, "spans": [], "drawings": ["d0"], "mark": "p0-custom"}


# What parsing says of adopted_custom (emit parses each element before planning it: ir_types).
NO_FLIP = "IRError: slide page 0, element p0-custom: lacks required key 'flip'"


def synthetic(*elements: JsonObject) -> JsonObject:
    return {"source": {"pdf": "nowhere.pdf"},
            "slides": [{"page": 0, "size": [360.0, 270.0], "elements": [e for e in elements]}]}


def element_parts(planned: emit.OfflinePlan) -> dict[tuple[int, str], tuple[str, list[JsonObject]]]:
    """(page, element id) -> (object id, the element's requests) of a plan_offline result."""
    out: dict[tuple[int, str], tuple[str, list[JsonObject]]] = {}
    for _, page, parts, ids in planned["slides"]:
        elements = [(el, reqs) for el, reqs in parts if el is not None]
        for (el, reqs), oid in zip(elements, ids):
            out[(page, jstr(el, "id"))] = (oid, reqs)
    return out


def part_element(part: emit.Part) -> JsonObject:
    """The element of an element's part (a slide's own part has none)."""
    el, _ = part
    assert el is not None, "the slide's own part"
    return el


@pytest.mark.needs_decks("out/04_theme_blocks-handout.pdf")
def test_an_element_emit_cannot_plan_is_the_picture_of_its_region(lenient: None, monkeypatch: pytest.MonkeyPatch,
                                                                    tmp_path: Path, capsys: pytest.CaptureFixture[str]
                                                                    ) -> None:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    from beamer2slides.emit_theme import slide_layout

    classified = deck("04_theme_blocks")
    before = emit.plan_offline(classified)
    page, target = next((jint(s, "page"), jstr(e, "id")) for s in before["plan"].slides() for e in jobjs(s, "elements")
                        if e["kind"] == "shape" and e.get("block") is None)
    monkeypatch.setattr(emit, "shape_element_requests", raising_for(target))
    after = emit.plan_offline(classified)

    assert after["contained"] == [{"page": page, "id": target, "kind": "shape", "error": "KeyError: 'flip'"}]
    slide = next(s for s in after["plan"].slides() if s["page"] == page)
    old_slide = next(s for s in before["plan"].slides() if s["page"] == page)
    elements = jobjs(slide, "elements")
    i = next(k for k, e in enumerate(elements) if e["id"] == target)
    was = jobjs(old_slide, "elements")[i]
    assert was["id"] == target  # (in its place: z-order and every other object id kept)
    x0, y0, x1, y1 = jnums(was, "bbox")
    assert elements[i] == {"kind": "image", "id": target, "role": "fallback",
                           "bbox": [x0 - 2, y0 - 2, x1 + 2, y1 + 2], "file": f"figures/fallback-{target}.png"}
    old, new = element_parts(before), element_parts(after)
    assert old.keys() == new.keys()
    assert new[(page, target)] == (f"b2s_s{page:03}_f{i}", [{"updatePageElementsZOrder": {
        "pageElementObjectIds": [f"b2s_s{page:03}_f{i}"], "operation": "BRING_TO_FRONT"}}])
    assert {k: v for k, v in new.items() if k != (page, target)} == {k: v for k, v in old.items() if k != (page, target)}

    # build_deck's plan says so, and cuts the picture out of the page for the .pptx
    capsys.readouterr()
    plan = emit.upload_plan(classified, tmp_path)
    assert f"warning: slide {page + 1}: shape {target} could not be planned (KeyError: 'flip'); using a picture" \
        in capsys.readouterr().out
    assert (tmp_path / "figures" / f"fallback-{target}.png").stat().st_size > 0
    page_w, page_h = jnums(plan.slides()[0], "size")
    pages: list[dict[str, object]] = [
        {"layout": slide_layout(s)[0], "fill": None, "templates": plan.uses_templates[jint(s, "page")],
         "pictures": [{"file": tmp_path / jstr(e, "file"), "bbox": bbox, "alt": None, "title": "Picture"}
                      for e, bbox in plan.pictures(s) if e.get("role") == "fallback"],  # (render made none)
         "tables": plan.tables(s)} for s in plan.slides()]
    prs = Presentation(emit.build_pptx(page_w, page_h, plan.keys, pages, {"color": "#ffffff"}, None))
    k = [jint(s, "page") for s in plan.slides()].index(page)
    assert sum(sh.shape_type == MSO_SHAPE_TYPE.PICTURE for sh in prs.slides[k].shapes) == 1


@pytest.mark.needs_decks("out/04_theme_blocks-handout.pdf")
def test_the_refused_rebuild_starts_from_the_merged_deck(lenient: None) -> None:
    """`emit` rebuilds with refused elements as pictures from `DeckPlan.merged`: the deck with its
    blocks merged, as emit used to merge them before planning (a refused body's picture reaches up
    under its title bar), and a contained element a picture already."""
    classified = deck("04_theme_blocks")
    merged = emit.DeckPlan(classified, pptx_tables=True, contain=True).merged
    assert merged["slides"] == [{**s, "elements": emit.merge_blocks(s["elements"])} for s in classified["slides"]]
    merged = emit.DeckPlan(synthetic(adopted_custom(), shape("p0-panel", 150.0)), pptx_tables=True, contain=True).merged
    assert [e["kind"] for e in jobjs(merged, "slides", 0, "elements")] == ["image", "shape"]


@pytest.mark.needs_decks("out/04_theme_blocks-handout.pdf")
def test_strict_raises_what_would_be_contained(monkeypatch: pytest.MonkeyPatch) -> None:
    classified = deck("04_theme_blocks")
    target = next(e["id"] for s in classified["slides"] for e in s["elements"] if e["kind"] == "shape")
    monkeypatch.setattr(emit, "shape_element_requests", raising_for(target))
    monkeypatch.setenv(emit.STRICT_ENV, "1")
    with pytest.raises(KeyError, match="flip"):
        emit.plan_offline(classified)


def test_an_adopted_custom_shape_without_flip_is_contained(lenient: None) -> None:
    """688ebf4's element: a shape kind no Slides preset has, no `flip`, reaching emit. It fails
    where emit parses it, and is contained as a planning failure was."""
    planned = emit.plan_offline(synthetic(adopted_custom(), shape("p0-panel", 150.0)))
    assert planned["contained"] == [{"page": 0, "id": "p0-custom", "kind": "shape", "error": NO_FLIP}]
    (_, _, parts, ids), = planned["slides"]
    assert ids == ["b2s_s000_f0", "b2s_s000_s1"]
    assert [part_element(p)["kind"] for p in parts[1:3]] == ["image", "shape"]
    assert planned["pictures"][0][0][0]["role"] == "fallback"
    assert parts[2][1] == shape_requests_of(shape("p0-panel", 150.0), "b2s_s000_s1")


def test_the_same_element_strict_is_the_same_error(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv(emit.STRICT_ENV, "1")
    with pytest.raises(IRError, match="lacks required key 'flip'"):
        emit.plan_offline(synthetic(adopted_custom()))


def shape_requests_of(el: JsonObject, oid: str) -> list[JsonObject]:
    return real_shape_requests(el, "b2s_s000", oid, emit.SLIDE_W / 360.0, None)


def test_a_step_over_the_whole_slide_blames_the_element_it_tripped_on(lenient: None) -> None:
    """merge_blocks reads every block shape's `flip`: a body without one fails the slide's merge,
    not an element's own planning. The body is the picture; its title bar, which the failing code
    held too, is not (a KeyError's key is what the culprit lacks)."""
    bar = shape("p0-bar", 20.0, shape="ROUND_2_SAME_RECTANGLE", block=0, radius=4.0)
    body = shape("p0-body", 20.0, shape="ROUND_2_SAME_RECTANGLE", block=0, radius=4.0, title_bar=[20.0, 10.0, 120.0, 30.0])
    del body["flip"]
    planned = emit.plan_offline(synthetic(bar, body, shape("p0-other", 200.0)))
    assert [(c["id"], c["error"]) for c in planned["contained"]] == [("p0-body", "KeyError: 'flip'")]
    kinds = {jstr(el, "id"): el["kind"] for el in jobjs(planned["plan"].slides()[0], "elements")}
    assert kinds == {"p0-bar": "shape", "p0-body": "image", "p0-other": "shape"}


def test_two_elements_missing_the_same_field_are_both_pictures(lenient: None) -> None:
    """Neither lets the slide through alone; each is to blame, and nothing else is."""
    heads = [shape(f"p0-bar{k}", 20.0 + 150 * k, shape="ROUND_2_SAME_RECTANGLE", block=k, radius=4.0) for k in (0, 1)]
    bodies = [shape(f"p0-body{k}", 20.0 + 150 * k, shape="ROUND_2_SAME_RECTANGLE", block=k, radius=4.0,
                    title_bar=[20.0 + 150 * k, 10.0, 120.0 + 150 * k, 30.0]) for k in (0, 1)]
    for b in bodies:
        del b["flip"]
    planned = emit.plan_offline(synthetic(*heads, *bodies))
    assert sorted(c["id"] for c in planned["contained"]) == ["p0-body0", "p0-body1"]


def test_containment_is_deterministic(lenient: None) -> None:
    """sync diffs what emit would write (`sync.mark_emitted`): the same deck, the same pictures."""
    one, two = (emit.plan_offline(synthetic(adopted_custom(), shape("p0-panel", 150.0))) for _ in range(2))
    assert one["contained"] == two["contained"] and one["slides"] == two["slides"]
    assert emit.slide_emission(one["plan"].slides()[0], one["plan"].scale, one["plan"].fonts)["parts"] == \
        emit.slide_emission(two["plan"].slides()[0], two["plan"].scale, two["plan"].fonts)["parts"]


def test_an_element_that_trips_only_on_googles_sizes_goes_the_refused_way(lenient: None,
                                                                           monkeypatch: pytest.MonkeyPatch) -> None:
    """build_deck plans each slide again with the sizes the import brought: an element failing only
    then gets no requests and (its index, the error) - build_deck turns it into a refused one, and
    the rebuild with pictures follows (`emit`). Without that list, the error is raised."""
    planned = emit.plan_offline(synthetic(shape("p0-a", 20.0), shape("p0-b", 150.0)))
    plan = planned["plan"]
    slide = plan.slides()[0]
    monkeypatch.setattr(emit, "shape_element_requests", raising_for("p0-a"))
    late: list[tuple[int, Exception]] = []
    parts, ids = plan.slide_parts(slide, planned["page_elements"], planned["speaker_notes"], {}, [], late)
    assert [(i, type(e)) for i, e in late] == [(0, KeyError)]
    assert ids == ["b2s_s000_s0", "b2s_s000_s1"] and parts[1][1] == [] and parts[2][1]
    with pytest.raises(KeyError):
        plan.slide_parts(slide, planned["page_elements"], planned["speaker_notes"], {}, [])
    monkeypatch.setenv(emit.STRICT_ENV, "1")
    with pytest.raises(KeyError):
        plan.slide_parts(slide, planned["page_elements"], planned["speaker_notes"], {}, [], [])


def test_an_element_with_no_box_is_never_lost_silently(lenient: None) -> None:
    """No region, no picture: the planning error is raised rather than the element dropped."""
    custom = adopted_custom()
    del custom["bbox"]
    with pytest.raises(IRError, match="lacks required key 'bbox'"):
        emit.plan_offline(synthetic(custom, shape("p0-panel", 150.0)))
