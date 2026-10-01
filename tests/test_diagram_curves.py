"""TikZ curves as Slides arcs (`curves`): the fit on synthetic cubics, a curved edge through
classify, its template (`emit_pptx` arc preset) and the requests that put a copy of it in place."""

import math
from collections.abc import Callable

import pytest
from pptx import Presentation

from beamer2slides import curves, ir_types
from beamer2slides.emit_diagrams import arc_chords, arc_template_key, arc_template_key_of, diagram_requests
from beamer2slides.emit_metrics import FontMapper
from beamer2slides.emit_model import JsonMap, TemplateKey
from beamer2slides.emit_pptx import NS_A, _add_template_shapes, arc_heads, arc_kind
from beamer2slides.google_types import slides_json
from beamer2slides.gslides import EMU_PER_PT
from beamer2slides.ir import Arrow, DiagramElement, element_json
from beamer2slides.json_types import JsonObject
from beamer2slides.raw_types import PathItem

from .json_reads import jnum, jobj, jstr, jstrs
from .test_charts_diagrams import SANS, Page, body_text, lines, rect
from .test_diagrams_overlays import diagram_of

Point = tuple[float, float]
Cubic = tuple[Point, Point, Point, Point]
LOOSENESS = 0.3915  # TikZ's control distance per unit of chord for a 'bend' (looseness 1)


def bend(start: Point, end: Point, angle: float) -> Cubic:
    """TikZ's `bend left=angle` cubic from `start` to `end` on the page (y down)."""
    (x0, y0), (x1, y1) = start, end
    d = LOOSENESS * math.dist(start, end)
    heading = math.atan2(y1 - y0, x1 - x0)
    out, back = heading - math.radians(angle), heading + math.pi + math.radians(angle)
    return start, (x0 + d * math.cos(out), y0 + d * math.sin(out)), (x1 + d * math.cos(back), y1 + d * math.sin(back)), end


def path_of(*cubics: Cubic) -> list[PathItem]:
    return [("c", [[x, y] for x, y in c]) for c in cubics]


def at(c: Cubic, t: float) -> Point:
    u = 1 - t
    return (u ** 3 * c[0][0] + 3 * u * u * t * c[1][0] + 3 * u * t * t * c[2][0] + t ** 3 * c[3][0],
            u ** 3 * c[0][1] + 3 * u * u * t * c[1][1] + 3 * u * t * t * c[2][1] + t ** 3 * c[3][1])


def piece_distance(p: Point, piece: curves.Piece) -> float:
    if piece.sweep is None:
        (ax, ay), (bx, by) = piece.start, piece.end
        dx, dy = bx - ax, by - ay
        t = max(0.0, min(1.0, ((p[0] - ax) * dx + (p[1] - ay) * dy) / (dx * dx + dy * dy)))
        return math.hypot(p[0] - ax - t * dx, p[1] - ay - t * dy)
    return curves.arc_deviation([p], piece.start, piece.end, curves.drawn_sweep(piece.sweep))


def assert_follows(cubics: list[Cubic], pieces: list[curves.Piece]) -> None:
    """The pieces join end to end from the path's start to its end, and every point of the path
    lies within the tolerance of one of them (the arcs as drawn, their sweeps rounded)."""
    assert pieces[0].start == cubics[0][0] and pieces[-1].end == cubics[-1][3]
    assert all(a.end == b.start for a, b in zip(pieces, pieces[1:]))
    for c in cubics:
        for k in range(41):
            p = at(c, k / 40)
            assert min(piece_distance(p, q) for q in pieces) <= curves.ARC_TOLERANCE + 1e-6, p


# ---------------------------------------------------------------- the fit

def test_a_bend_left_edge_is_one_arc_turning_clockwise_on_the_page() -> None:
    """bend left=30 between two nodes on a level: one arc of about 2 x 30 degrees, bulging up,
    which on the page (y down) turns clockwise: a positive sweep."""
    c = bend((100.0, 100.0), (200.0, 100.0), 30)
    pieces = curves.path_pieces(path_of(c), curves.ARC_TOLERANCE)
    assert pieces is not None and len(pieces) == 1
    (arc,) = pieces
    assert arc.sweep is not None and 58 < arc.sweep < 66
    assert_follows([c], pieces)
    circle = curves.arc_circle(arc.start, arc.end, arc.sweep)
    assert circle.centre[1] > 100  # (its centre below the chord: the arc bulges up)
    right = curves.path_pieces(path_of(bend((100.0, 100.0), (200.0, 100.0), -30)), curves.ARC_TOLERANCE)
    assert right is not None and len(right) == 1 and right[0].sweep == pytest.approx(-arc.sweep)


def test_a_loop_is_a_chain_of_arcs() -> None:
    """A self loop (two cubics leaving a node's top and coming back): a few arcs, all turning one way."""
    loop: list[Cubic] = [((110.0, 127.0), (106.0, 110.0), (110.0, 105.0), (113.0, 105.0)),
                         ((113.0, 105.0), (116.0, 105.0), (120.0, 110.0), (116.0, 127.0))]
    pieces = curves.path_pieces(path_of(*loop), curves.ARC_TOLERANCE)
    assert pieces is not None and 2 <= len(pieces) <= curves.MAX_PIECES
    assert all(p.sweep is not None and p.sweep > 0 for p in pieces)
    assert_follows(loop, pieces)


def test_a_cubic_that_is_no_arc_is_arcs_turning_both_ways() -> None:
    """An S (out=60, in=60: the edge leaves up and comes in from below): no one circle, so arcs
    that turn one way and then the other."""
    s: Cubic = ((0.0, 0.0), (30.0, -50.0), (70.0, 50.0), (100.0, 0.0))
    pieces = curves.path_pieces(path_of(s), curves.ARC_TOLERANCE)
    assert pieces is not None and len(pieces) >= 2
    sweeps = [p.sweep for p in pieces if p.sweep is not None]
    assert min(sweeps) < 0 < max(sweeps)
    assert_follows([s], pieces)


def test_a_straight_cubic_is_a_straight_piece() -> None:
    line: Cubic = ((0.0, 0.0), (10.0, 5.0), (40.0, 20.0), (60.0, 30.0))
    assert curves.path_pieces(path_of(line), curves.ARC_TOLERANCE) == [
        curves.Piece(start=(0.0, 0.0), end=(60.0, 30.0), sweep=None)]


def test_a_coil_is_no_chain_of_arcs() -> None:
    """An inductor's coil: more pieces than `MAX_PIECES`, so the drawing stays a picture."""
    coil: list[Cubic] = []
    for x in (4.0 * k for k in range(10)):
        coil += [((x, 0.0), (x, -6.0), (x + 4.0, -6.0), (x + 4.0, 0.0)), ((x + 4.0, 0.0), (x + 4.0, 3.0), (x, 3.0), (x, 0.0))]
    assert curves.path_pieces(path_of(*coil), curves.ARC_TOLERANCE) is None


def test_a_nearly_straight_curve_is_no_huge_arc() -> None:
    """A long gentle curve: its circle would be metres wide (the arc's shape is its circle's box),
    so it is straight pieces within the tolerance instead."""
    gentle: Cubic = ((0.0, 0.0), (100.0, -1.5), (200.0, -1.5), (300.0, 0.0))
    pieces = curves.path_pieces(path_of(gentle), curves.ARC_TOLERANCE)
    assert pieces is not None
    assert all(p.sweep is None or curves.arc_circle(p.start, p.end, p.sweep).radius <= curves.MAX_RADIUS for p in pieces)
    assert_follows([gentle], pieces)


# ---------------------------------------------------------------- arc geometry

def test_an_arc_circle_lies_right_of_the_way_for_a_clockwise_arc() -> None:
    c = curves.arc_circle((0.0, 0.0), (100.0, 0.0), 90.0)
    assert c.centre == pytest.approx((50.0, 50.0)) and c.radius == pytest.approx(50 * math.sqrt(2))
    assert c.start == pytest.approx(-135.0)
    assert curves.arc_point(c, 90.0) == pytest.approx((100.0, 0.0))
    assert curves.arc_point(c, 45.0) == pytest.approx((50.0, 50.0 - 50 * math.sqrt(2)))  # (the top: it bulges up)
    half = curves.arc_circle((0.0, 0.0), (100.0, 0.0), -180.0)
    assert half.centre == pytest.approx((50.0, 0.0)) and curves.arc_point(half, -90.0) == pytest.approx((50.0, 50.0))


def test_an_arc_grows_along_its_circle_to_its_arrow_tip() -> None:
    """A filled head's tip: the arc runs on along its circle, so the head points along the curve."""
    ux, uy = curves.end_direction((0.0, 0.0), (100.0, 0.0), 90.0, "to")
    assert (ux, uy) == pytest.approx((math.sqrt(0.5), math.sqrt(0.5)))  # leaving down and right
    assert curves.end_direction((0.0, 0.0), (100.0, 0.0), 90.0, "from") == pytest.approx((-math.sqrt(0.5), math.sqrt(0.5)))
    r = 50 * math.sqrt(2)
    start, end, sweep = curves.extended((0.0, 0.0), (100.0, 0.0), 90.0, "to", r * math.radians(10))
    assert start == (0.0, 0.0) and sweep == pytest.approx(100.0)
    assert math.dist(end, (50.0, 50.0)) == pytest.approx(r, abs=0.01) and end[1] > 0
    start, end, sweep = curves.extended((0.0, 0.0), (100.0, 0.0), -90.0, "from", r * math.radians(10))
    assert end == (100.0, 0.0) and sweep == pytest.approx(-100.0) and start[0] < 0 and start[1] < 0  # (up and left: it bulges down)


# ---------------------------------------------------------------- the template

ARC_LINE: JsonObject = {"from": [0.0, 0.0], "to": [100.0, 0.0], "sweep": 62.37, "stroke": "#336699", "width": 0.8,
                        "arrow_from": None, "arrow_to": "STEALTH_ARROW"}


def line_of(o: JsonObject) -> ir_types.DiagramLine:
    return ir_types.diagram_line(o, ir_types.At(where="a test line", path=""))


def test_an_arcs_template_is_its_heads_and_its_rounded_sweep() -> None:
    """Rounded to `ARC_STEP`; the start angle and the way it turns are the copy's transform, so an
    arc mirrored or turned shares its template."""
    assert arc_template_key(ARC_LINE) == ("ARC>STEALTH_ARROW", 62.0, None)
    mirrored = line_of({**ARC_LINE, "sweep": -61.8, "to": [0.0, 100.0]})
    assert arc_template_key_of(mirrored) == ("ARC>STEALTH_ARROW", 62.0, None)
    assert arc_template_key_of(line_of({**ARC_LINE, "arrow_from": "FILL_ARROW", "arrow_to": None})) == \
        ("ARC<FILL_ARROW", 62.0, None)
    assert arc_template_key({**ARC_LINE, "sweep": None}) is None
    assert arc_template_key_of(line_of({k: v for k, v in ARC_LINE.items() if k != "sweep"})) is None


@pytest.mark.parametrize("start", [None, "OPEN_ARROW", "STEALTH_ARROW", "FILL_ARROW"])
@pytest.mark.parametrize("end", [None, "OPEN_ARROW", "STEALTH_ARROW", "FILL_ARROW"])
def test_an_arc_kind_reads_back_its_heads(start: Arrow | None, end: Arrow | None) -> None:
    assert arc_heads(arc_kind(start, end)) == (start, end)


def test_other_kinds_are_no_arcs() -> None:
    assert [arc_heads(k) for k in ("ROUND_RECTANGLE", "BENT_CONNECTOR", "ARCH", "ARC<NOPE", "ARC>", "ARC<")] == \
        [None] * 6


def test_the_pptx_template_is_the_arc_preset_with_its_heads() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    keys: list[TemplateKey] = [("ROUND_RECTANGLE", 0.1, None), ("ARC<OPEN_ARROW>FILL_ARROW", 62.0, None), ("ARC", 180.0, None)]
    _add_template_shapes(slide, keys)
    geometries = slide.shapes._spTree.findall(f".//{{{NS_A}}}prstGeom")
    assert [g.get("prst") for g in geometries] == ["roundRect", "arc", "arc"]
    adjust = {gd.get("name"): gd.get("fmla") for gd in geometries[1].iter(f"{{{NS_A}}}gd")}
    assert adjust == {"adj1": "val 0", "adj2": f"val {62 * 60000}"}
    outline = slide.shapes[1]._element.spPr.find(f"{{{NS_A}}}ln")
    assert outline is not None
    head, tail = outline.find(f"{{{NS_A}}}headEnd"), outline.find(f"{{{NS_A}}}tailEnd")
    assert head is not None and tail is not None and (head.get("type"), tail.get("type")) == ("arrow", "triangle")
    bare = slide.shapes[2]._element.spPr.find(f"{{{NS_A}}}ln")
    assert bare is not None and bare.find(f"{{{NS_A}}}headEnd") is None and bare.find(f"{{{NS_A}}}tailEnd") is None
    with pytest.raises(ValueError):
        _add_template_shapes(slide, [("ARC", None, None)])


# ---------------------------------------------------------------- a curved edge, classified and emitted

def stealth_tip(base: Point, direction: Point) -> list[PathItem]:
    """A small filled head whose base is at `base`, pointing along `direction` (a unit vector)."""
    (bx, by), (ux, uy) = base, direction
    nx, ny = -uy, ux
    tip = (bx + 4 * ux, by + 4 * uy)
    return lines((bx + 1.6 * nx, by + 1.6 * ny), tip, (bx - 1.6 * nx, by - 1.6 * ny), closed=True)


def curved_edge_page() -> tuple[Page, Point]:
    """Two boxes on a level and a `bend left` edge from the first's top to the second's, its head
    drawn at its end; returns the page and where the head's tip is."""
    p = Page()
    for x0, name in ((60.0, "Lexer"), (160.0, "Parser")):
        p.draw(rect(x0, 100, x0 + 50, 120), "s", None, "#000000", 0.4)
        p.text(name, x0 + 8, 113, 7.97, SANS, None, "#000000")
    c = bend((85.0, 100.0), (185.0, 100.0), 30)
    p.draw(path_of(c), "s", None, "#336699", 0.8)
    ux, uy = c[3][0] - c[2][0], c[3][1] - c[2][1]
    norm = math.hypot(ux, uy)
    p.draw(stealth_tip(c[3], (ux / norm, uy / norm)), "f", "#336699", None, None)
    body_text(p, 225)
    return p, (c[3][0] + 4 * ux / norm, c[3][1] + 4 * uy / norm)


def the_diagram() -> tuple[DiagramElement, Point]:
    p, tip = curved_edge_page()
    return diagram_of(p), tip


def test_a_curved_edge_is_an_arc_reaching_its_arrow_tip() -> None:
    el, tip = the_diagram()
    (ln,) = el["lines"]
    sweep = ln.get("sweep")
    assert sweep is not None and 60 < sweep < 72
    assert ln["from"] == [85.0, 100.0] and ln["arrow_to"] == "FILL_ARROW" and ln["arrow_from"] is None
    # grown along its circle to the tip (which TikZ set on the cubic's own end tangent, a few
    # degrees off the arc's)
    assert math.dist((ln["to"][0], ln["to"][1]), tip) < 0.4


def template_for(asked: list[TemplateKey]) -> Callable[[TemplateKey], JsonMap]:
    """Every template is one 236.22 pt square (3,000,000 EMU, as Slides stores a shape); the keys
    asked for go to `asked`."""
    def template(key: TemplateKey) -> JsonMap:
        asked.append(key)
        return {"id": "tpl", "w": 236.22, "h": 236.22}
    return template


def placed(transform: JsonObject, local: Point) -> Point:
    """Where the template's point `local` (pt in its own box) goes on the page (pt)."""
    x, y = local
    return ((jnum(transform, "scaleX") * x + jnum(transform, "shearX") * y) + jnum(transform, "translateX") / EMU_PER_PT,
            (jnum(transform, "shearY") * x + jnum(transform, "scaleY") * y) + jnum(transform, "translateY") / EMU_PER_PT)


def template_point(angle: float) -> Point:
    """The arc preset's point at `angle` degrees (clockwise from its right) in a 236.22 pt square."""
    r = 236.22 / 2
    return r + r * math.cos(math.radians(angle)), r + r * math.sin(math.radians(angle))


@pytest.mark.parametrize("sweep", [62.37, -62.37, 200.0, -95.5])
def test_an_arc_copy_is_turned_onto_its_ends(sweep: float) -> None:
    """The template runs 0 -> |drawn sweep| clockwise; one ABSOLUTE transform (turned, mirrored
    for an anticlockwise arc) puts its ends on the line's, its middle on the arc."""
    start, end = (80.0, 100.0), (180.0, 130.0)
    el: JsonObject = {"id": "d", "kind": "diagram", "role": "figure", "bbox": [0.0, 0.0, 300.0, 300.0], "nodes": [],
                      "spans": [], "lines": [{**ARC_LINE, "from": [start[0], start[1]], "to": [end[0], end[1]], "sweep": sweep}]}
    asked: list[TemplateKey] = []
    reqs = [slides_json(r) for r in diagram_requests(el, "s", "d", 1.5, FontMapper(), template_for(asked))]
    drawn = curves.drawn_sweep(sweep)
    assert asked == [("ARC>STEALTH_ARROW", abs(drawn), None)]
    assert jstr(reqs[0], "duplicateObject", "objectId") == "tpl"
    assert jobj(reqs[0], "duplicateObject", "objectIds") == {"tpl": "d_l0"}
    move = jobj(reqs[1], "updatePageElementTransform")
    assert move["applyMode"] == "ABSOLUTE"
    t = jobj(move, "transform")
    assert placed(t, template_point(0)) == pytest.approx((1.5 * start[0], 1.5 * start[1]), abs=0.01)
    assert placed(t, template_point(abs(drawn))) == pytest.approx((1.5 * end[0], 1.5 * end[1]), abs=0.01)
    circle = curves.arc_circle(start, end, drawn)
    mx, my = curves.arc_point(circle, drawn / 2)
    assert placed(t, template_point(abs(drawn) / 2)) == pytest.approx((1.5 * mx, 1.5 * my), abs=0.01)
    look = jobj(reqs[3], "updateShapeProperties", "shapeProperties")
    assert jobj(look, "outline", "outlineFill", "solidFill", "color", "rgbColor") == \
        pytest.approx({"red": 0x33 / 255, "green": 0x66 / 255, "blue": 0x99 / 255})
    assert jnum(look, "outline", "weight", "magnitude") == 1.2
    assert jstr(look, "shapeBackgroundFill", "propertyState") == "NOT_RENDERED"
    # An arc is a shape: no line properties, no connection to a node.
    assert not [r for r in reqs if "updateLineProperties" in r or "createLine" in r]
    assert "dashStyle" not in jstr(reqs[3], "updateShapeProperties", "fields")  # (a solid arc writes none)


def test_a_dashed_arc_has_a_dashed_outline() -> None:
    el: JsonObject = {"id": "d", "kind": "diagram", "role": "figure", "bbox": [0.0, 0.0, 300.0, 300.0], "nodes": [],
                      "spans": [], "lines": [{**ARC_LINE, "from": [80.0, 100.0], "to": [180.0, 130.0], "sweep": 62.37,
                                              "dash": "DASH"}]}
    reqs = [slides_json(r) for r in diagram_requests(el, "s", "d", 1.0, FontMapper(), template_for([]))]
    look = jobj(reqs[3], "updateShapeProperties")
    assert jstr(look, "fields").endswith(",outline.dashStyle")
    assert jstr(look, "shapeProperties", "outline", "dashStyle") == "DASH"


def test_a_curved_edge_through_emit_is_one_arc_copy_in_its_group() -> None:
    el, _ = the_diagram()
    asked: list[TemplateKey] = []
    reqs = [slides_json(r) for r in diagram_requests(element_json(el), "s", "d", 1.0, FontMapper(), template_for(asked))]
    arc = [k for k in asked if arc_heads(k[0]) is not None]
    assert arc == [arc_template_key(jobj(element_json(el), "lines", 0))]
    groups = [jstrs(r, "groupObjects", "childrenObjectIds") for r in reqs if "groupObjects" in r]
    assert "d_l0" in groups[-1]


def test_without_a_template_an_arc_is_straight_pieces() -> None:
    """(diagram_requests with no template: the tests' and a deck with no template shapes)"""
    ln = line_of({**ARC_LINE, "arrow_from": "OPEN_ARROW"})
    pieces = arc_chords("d_l0", ln, 62.37)
    assert len(pieces) == math.ceil(62.37 / 15) and pieces[0][0] == "d_l0" and pieces[1][0] == "d_l0c1"
    assert pieces[0][1] == ln.from_ and pieces[-1][2] == ln.to
    assert [(q.arrow_from, q.arrow_to, q.sweep) for _, _, _, q in pieces] == \
        [("OPEN_ARROW", None, None)] + [(None, None, None)] * (len(pieces) - 2) + [(None, "STEALTH_ARROW", None)]
    el: JsonObject = {"id": "d", "kind": "diagram", "role": "figure", "bbox": [0.0, 0.0, 300.0, 300.0], "nodes": [],
                      "spans": [], "lines": [ARC_LINE]}
    reqs = [slides_json(r) for r in diagram_requests(el, "s", "d", 1.0, FontMapper(), None)]
    assert len([r for r in reqs if "createLine" in r]) == len(pieces)


# ---------------------------------------------------------------- the IR

def test_an_old_line_reads_as_straight_and_writes_back_as_it_was() -> None:
    old: JsonObject = {"from": [1.0, 2.0], "to": [3.0, 4.0], "arrow_from": None, "arrow_to": "OPEN_ARROW",
                       "stroke": "#000000", "width": 0.4}
    ln = line_of(old)
    assert ln.sweep is None and ir_types.diagram_line_json(ln) == old
    arc = line_of(ARC_LINE)
    assert arc.sweep == 62.37 and ir_types.diagram_line_json(arc) == ARC_LINE


@pytest.mark.parametrize("bad", [
    {**ARC_LINE, "sweep": 0.0}, {**ARC_LINE, "sweep": 360.0}, {**ARC_LINE, "sweep": -400.0},
    {**ARC_LINE, "via": [50.0, 0.0], "bend": "vh"}, {**ARC_LINE, "to": [0.0, 0.0]}, {**ARC_LINE, "sweep": "62"},
])
def test_an_arc_that_cannot_be_is_refused(bad: JsonObject) -> None:
    with pytest.raises(ir_types.IRError):
        line_of(bad)
