"""Beamer's ▶ bullets as themselves: a box whose bullets no Slides preset draws comes in the .pptx
as a text shell carrying `a:buChar` ► (the same shape, larger in Slides; emit_text.text_shell_of, emit_pptx._add_text_shell), and the
API pass writes its words in (`text_shell_requests`) instead of creating the box and giving it
createParagraphBullets' ➢. Sync, which creates every box through the API, keeps ➢."""

from copy import deepcopy

from beamer2slides import emit
from beamer2slides.emit import EMU_PER_PT, SLIDE_W
from beamer2slides.emit_metrics import CHAR_BULLETS
from beamer2slides.emit_model import PptxText, Shell, text_of
from beamer2slides.emit_pptx import INSET_PT, SHELL_CHAR
from beamer2slides.emit_text import shell_route, text_shell_of, text_shell_requests_of
from beamer2slides.google_types import slides_json
from beamer2slides.json_types import Json, JsonObject

from .json_reads import jint, jnum, jobj, jobjs, jstr
from .test_emit_hunt import column_list_of
from .test_emit_requests import FONTS, pt_of, text_requests

PDF_W = 362.83  # column_list_of's page (Goettingen, 4:3)
SCALE = SLIDE_W / PDF_W
TRIANGLE = CHAR_BULLETS["triangle"][0]
NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def nested_list() -> JsonObject:
    """column_list_of with its second item a subitem."""
    el = deepcopy(column_list_of())
    jobj(el, "paragraphs", 1)["level"] = 1
    return el


def shell_of(el: JsonObject) -> PptxText:
    x0, y0, x1, y1 = (jnum(el, "bbox", k) * SCALE for k in range(4))
    shell = text_shell_of(text_of(el), (x0, y0, x1, y1), SCALE, FONTS)
    assert shell is not None
    return shell


def test_only_boxes_of_triangle_bullets_are_shells() -> None:
    el = column_list_of()
    assert shell_route(text_of(el))
    # a box mixing ▶ with a preset's bullet keeps the API's way (its presets are what was measured)
    mixed = deepcopy(el)
    jobj(mixed, "paragraphs", 1, "bullet")["text"] = "●"
    assert not shell_route(text_of(mixed))
    # numbers, no bullets at all, and an empty paragraph (it would lose its bullet at the import)
    numbered = deepcopy(el)
    jobj(numbered, "paragraphs", 1)["bullet"] = {"kind": "number", "text": "2.", "color": "#000000",
                                                  "bbox": [21.03, 113.7, 29.51, 124.61]}
    assert not shell_route(text_of(numbered))
    plain = deepcopy(el)
    for p in jobjs(plain, "paragraphs"):
        p["bullet"] = None
    assert not shell_route(text_of(plain))
    empty = deepcopy(el)
    jobj(empty, "paragraphs", 1)["runs"] = []
    assert not shell_route(text_of(empty))


def test_the_pptx_carries_a_shell_with_its_bullets_as_characters() -> None:
    from pptx import Presentation
    from pptx.shapes.autoshape import Shape
    shell = shell_of(nested_list())
    assert [(p.level, p.char, p.color) for p in shell.paragraphs] == [(0, TRIANGLE, "#3333b3"), (1, TRIANGLE, "#3333b3")]
    page: dict[str, object] = {"layout": "BLANK", "fill": None, "pictures": [], "tables": [], "shells": [shell],
                               "templates": False}
    prs = Presentation(emit.build_pptx(PDF_W, 272.13, [], [page], {"color": "#ffffff"}, None))
    box, = [s for s in prs.slides[0].shapes if isinstance(s, Shape)]
    x0, y0, x1, y1 = shell.box
    assert (box.left, box.top, box.width, box.height) == tuple(round(v * EMU_PER_PT) for v in (x0, y0, x1 - x0, y1 - y0))
    body = box.text_frame._txBody
    pr = body.find(f"{NS}bodyPr")
    assert pr is not None
    assert {k: pr.get(k) for k in ("lIns", "tIns", "rIns", "bIns", "wrap", "anchor")} == {
        **{k: str(round(INSET_PT * EMU_PER_PT)) for k in ("lIns", "tIns", "rIns", "bIns")}, "wrap": "square", "anchor": "t"}
    assert len(pr) == 0  # no spAutoFit: the box keeps the size emit gives it
    paragraphs = body.findall(f"{NS}p")
    assert len(paragraphs) == 2
    for p, want in zip(paragraphs, shell.paragraphs):
        ppr = p.find(f"{NS}pPr")
        assert ppr is not None and ppr.get("lvl") == str(want.level)
        char = ppr.find(f"{NS}buChar")
        assert char is not None and char.get("char") == TRIANGLE == "►"
        size = ppr.find(f"{NS}buSzPct")
        assert size is not None and size.get("val") == "100000"
        colour = ppr.find(f"{NS}buClr/{NS}srgbClr")
        assert colour is not None and colour.get("val") == "3333B3"
        run, = p.findall(f"{NS}r")
        rpr = run.find(f"{NS}rPr")
        assert rpr is not None and rpr.get("sz") == str(round(want.size * 100))
        assert run.findtext(f"{NS}t") == SHELL_CHAR  # (an empty bulleted paragraph loses its bullet)


def test_a_paragraph_without_a_bullet_is_said_so_in_the_shell() -> None:
    from pptx import Presentation
    from pptx.shapes.autoshape import Shape
    el = deepcopy(column_list_of())
    jobj(el, "paragraphs", 0)["bullet"] = None
    shell = shell_of(el)
    assert [p.char for p in shell.paragraphs] == [None, TRIANGLE]
    page: dict[str, object] = {"layout": "BLANK", "fill": None, "pictures": [], "shells": [shell], "templates": False}
    prs = Presentation(emit.build_pptx(PDF_W, 272.13, [], [page], {"color": "#ffffff"}, None))
    box, = [s for s in prs.slides[0].shapes if isinstance(s, Shape)]
    first = box.text_frame._txBody.findall(f"{NS}p")[0].find(f"{NS}pPr")
    assert first is not None and first.find(f"{NS}buNone") is not None and first.find(f"{NS}buChar") is None


def test_the_shell_is_filled_where_a_created_box_would_stand() -> None:
    el = nested_list()
    shell = shell_of(el)
    base_w, base_h = shell.box[2] - shell.box[0], shell.box[3] - shell.box[1]
    reqs = [slides_json(r) for r in text_shell_requests_of(text_of(el), "b2s_s003", "b2s_s003_t1", SCALE, FONTS,
                                                         Shell(base_w=base_w, base_h=base_h), None, None, None, None)]
    kinds = [next(iter(r)) for r in reqs]
    assert "createShape" not in kinds and "createParagraphBullets" not in kinds
    created = text_requests(el, SCALE)
    # the same box: the shell's size times its scale is the created box's size, at the same place
    made = jobj(next(r for r in created if "createShape" in r), "createShape", "elementProperties")
    moved = jobj(reqs[0], "updatePageElementTransform")
    assert jstr(moved, "applyMode") == "ABSOLUTE"
    t = jobj(moved, "transform")
    assert abs(jnum(t, "scaleX") * base_w - pt_of(jobj(made, "size", "width"))) < 1e-3  # (EMU)
    assert abs(jnum(t, "scaleY") * base_h - pt_of(jobj(made, "size", "height"))) < 1e-3
    assert (jnum(t, "translateX"), jnum(t, "translateY")) == \
        (jnum(made, "transform", "translateX"), jnum(made, "transform", "translateY"))
    assert jobj(reqs[1], "updatePageElementsZOrder")["operation"] == "BRING_TO_FRONT"  # where a new box lands
    # the words: in front of each placeholder character, the last paragraph first, the character out after
    text = "\n".join(SHELL_CHAR for _ in shell.paragraphs)
    for r in reqs:
        if "insertText" in r:
            i = jint(r, "insertText", "insertionIndex")
            text = text[:i] + jstr(r, "insertText", "text") + text[i:]
        elif "deleteText" in r:
            rng = jobj(r, "deleteText", "textRange")
            assert jstr(rng, "type") == "FIXED_RANGE"
            assert text[jint(rng, "startIndex"):jint(rng, "endIndex")] == SHELL_CHAR
            text = text[:jint(rng, "startIndex")] + text[jint(rng, "endIndex"):]
    want = "\n".join(jstr(p, "runs", 0, "text") for p in jobjs(el, "paragraphs"))
    assert text == want
    # the pitch and the text's own indents as the created box has them (only the bullet is another glyph)
    def paragraph_styles(rs: list[JsonObject]) -> list[dict[str, Json]]:
        keep = ("alignment", "lineSpacing", "spaceAbove", "spaceBelow", "indentStart")
        return [{k: v for k, v in jobj(r, "updateParagraphStyle", "style").items() if k in keep}
                for r in rs if "updateParagraphStyle" in r]
    assert paragraph_styles(reqs) == paragraph_styles(created)
    # nothing styles a whole item (it would restyle its bullet)
    items = {(0, len(want.split("\n")[0])), (len(want.split("\n")[0]) + 1, len(want))}
    for r in reqs:
        if "updateTextStyle" in r:
            rng = jobj(r, "updateTextStyle", "textRange")
            assert (jint(rng, "startIndex"), jint(rng, "endIndex")) not in items


def test_a_created_triangle_box_keeps_the_preset() -> None:
    """What sync creates (every box through the API: DeckPlan(pptx_tables=False), slide_emission;
    the corpus check is test_emit_requests.test_triangle_bullets_come_with_the_pptx): ➢ as before."""
    created = text_requests(column_list_of(), SCALE)
    presets = [jstr(r, "createParagraphBullets", "bulletPreset") for r in created if "createParagraphBullets" in r]
    assert presets == ["BULLET_ARROW3D_CIRCLE_SQUARE"]
