"""What Drive keeps when a presentation's content is replaced by `files.update` (what `convert`
does when it rebuilds a deck in place), and what of it can be restored through the API.

    python tools/probe_revision_history.py [--wait 30] [--keep]

The probe works on decks it creates itself and deletes again (nothing of yours is touched):

  1. a .pptx is uploaded as a Google Slides deck        (marker ALPHA)
  2. a text box is added through the Slides API         (marker EDIT - "the person edits the deck")
  3. the file's content is replaced with another .pptx  (marker BETA - "convert rebuilds the deck")
  4. `revisions.list` is read and every revision exported as .pptx: which markers does it hold?
  5. `revisions.update(keepForever=True)` is tried on the revision that holds the edit
  6. the revision holding the edit is uploaded again as a new presentation: is the edit back?

Findings go to docs/sync.md. Run it twice with different --wait to see how Drive groups changes
into revisions.
"""

import argparse
import io
import json
import sys
import time
import uuid
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from googleapiclient.errors import HttpError  # noqa: E402
from googleapiclient.http import MediaIoBaseUpload  # noqa: E402

from beamer2slides.google_auth import drive_service, slides_service  # noqa: E402
from beamer2slides.gslides import execute  # noqa: E402
from deck_backup import download, revisions  # noqa: E402

PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
SLIDES_MIME = "application/vnd.google-apps.presentation"


def pptx_with(text: str) -> io.BytesIO:
    from pptx import Presentation
    from pptx.util import Pt
    pres = Presentation()
    slide = pres.slides.add_slide(pres.slide_layouts[6])
    box = slide.shapes.add_textbox(Pt(50), Pt(50), Pt(400), Pt(60))
    box.text_frame.text = text
    buf = io.BytesIO()
    pres.save(buf)
    buf.seek(0)
    return buf


def markers_in(data: bytes, markers: dict[str, str]) -> list[str]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        blob = b"".join(z.read(n) for n in z.namelist() if n.endswith(".xml"))
    return [name for name, text in markers.items() if text.encode() in blob]


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--wait", type=float, default=30, help="seconds between the steps")
    ap.add_argument("--keep", action="store_true", help="leave the probe decks in Drive")
    args = ap.parse_args()
    drive, slides = drive_service(), slides_service()
    tag = uuid.uuid4().hex[:8].upper()
    markers = {"ALPHA": f"ALPHA-{tag}", "EDIT": f"EDIT-{tag}", "BETA": f"BETA-{tag}"}
    report: dict = {"tag": tag, "wait": args.wait, "steps": []}
    made = []

    pid = execute(drive.files().create(body={"name": f"b2s revision probe {tag}", "mimeType": SLIDES_MIME},
                                       media_body=MediaIoBaseUpload(pptx_with(markers["ALPHA"]), mimetype=PPTX_MIME),
                                       fields="id"))["id"]
    made.append(pid)
    report["presentationId"] = pid
    report["steps"].append({"step": "upload", "at": time.strftime("%H:%M:%S"), "marker": "ALPHA"})
    try:
        time.sleep(args.wait)
        page = execute(slides.presentations().get(presentationId=pid))["slides"][0]["objectId"]
        execute(slides.presentations().batchUpdate(presentationId=pid, body={"requests": [
            {"createShape": {"objectId": f"probe{tag}", "shapeType": "TEXT_BOX", "elementProperties": {
                "pageObjectId": page, "size": {"width": {"magnitude": 300, "unit": "PT"},
                                               "height": {"magnitude": 40, "unit": "PT"}},
                "transform": {"scaleX": 1, "scaleY": 1, "translateX": 60, "translateY": 200, "unit": "PT"}}}},
            {"insertText": {"objectId": f"probe{tag}", "text": markers["EDIT"]}}]}))
        edited_at = execute(drive.files().get(fileId=pid, fields="modifiedTime"))["modifiedTime"]
        report["steps"].append({"step": "api edit", "at": time.strftime("%H:%M:%S"), "marker": "EDIT",
                                "modifiedTime": edited_at})
        time.sleep(args.wait)

        execute(drive.files().update(fileId=pid, media_body=MediaIoBaseUpload(pptx_with(markers["BETA"]),
                                                                             mimetype=PPTX_MIME), fields="id"))
        report["steps"].append({"step": "files.update (the rebuild)", "at": time.strftime("%H:%M:%S"), "marker": "BETA"})
        time.sleep(args.wait)

        revs = revisions(drive, pid)
        report["revisions"] = []
        for r in revs:
            links = r.get("exportLinks") or {}
            entry = {"id": r["id"], "modifiedTime": r["modifiedTime"], "keepForever": r.get("keepForever"),
                     "pptx_export": PPTX_MIME in links}
            if PPTX_MIME in links:
                data = download(links[PPTX_MIME])
                entry["markers"] = markers_in(data, markers)
                entry["bytes"] = len(data)
            report["revisions"].append(entry)
            time.sleep(1.5)  # the export endpoint rate-limits

        holds_edit = [r for r in report["revisions"] if "EDIT" in r.get("markers", [])]
        report["edit_survives_the_rebuild"] = bool(holds_edit)
        if holds_edit:
            rid = holds_edit[-1]["id"]
            try:
                kept = execute(drive.revisions().update(fileId=pid, revisionId=rid, body={"keepForever": True},
                                                        fields="id,keepForever"))
                report["keepForever"] = {"accepted": True, "result": kept}
            except HttpError as e:
                report["keepForever"] = {"accepted": False, "error": str(e)[:200]}
            data = download((execute(drive.revisions().get(fileId=pid, revisionId=rid, fields="exportLinks"))
                             .get("exportLinks") or {})[PPTX_MIME])
            back = execute(drive.files().create(body={"name": f"b2s revision probe {tag} restored", "mimeType": SLIDES_MIME},
                                                media_body=MediaIoBaseUpload(io.BytesIO(data), mimetype=PPTX_MIME),
                                                fields="id"))["id"]
            made.append(back)
            text = execute(slides.presentations().get(presentationId=back))
            report["restored"] = {"presentationId": back, "revision": rid,
                                  "holds_the_edit": markers["EDIT"] in json.dumps(text)}
    finally:
        if not args.keep:
            for fid in made:
                try:
                    execute(drive.files().delete(fileId=fid))
                except HttpError as e:
                    print(f"could not delete the probe deck {fid}: {e}")
            report["cleaned_up"] = made
    print(json.dumps(report, indent=1))
    return 0


if __name__ == "__main__":
    sys.exit(main())
