"""Probe: native drop shadows on shapes. The Slides API can't set a shadow, but a .pptx
imported through Drive can carry one. Does it survive the import, duplicateObject (of the
shape and of its slide), a fill change and a transform change?

Creates a deck from a python-pptx file with a template slide holding a shadowed rounded
rectangle, duplicates the slide and the shape, restyles the copy, prints the shadow
properties the API reports, and saves a thumbnail to out/probe_shadow.png.

Usage: python tools/probe_shadow.py
"""

import copy
import io
import json
from pathlib import Path

from googleapiclient.http import MediaIoBaseUpload
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

from beamer2slides.google_auth import drive_service, slides_service
from beamer2slides.gslides import EMU_PER_PT, emu, execute, save_thumbnail

OUT = Path(__file__).resolve().parents[1] / "out"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def shadow_xml(blur: float, dist: float, alpha: int) -> etree._Element:
    return etree.fromstring(
        f'<a:effectLst xmlns:a="{A}"><a:outerShdw blurRad="{round(blur * EMU_PER_PT)}" '
        f'dist="{round(dist * EMU_PER_PT)}" dir="2700000" algn="tl" rotWithShape="0">'
        f'<a:srgbClr val="000000"><a:alpha val="{alpha}"/></a:srgbClr></a:outerShdw></a:effectLst>')


def main() -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(720 * EMU_PER_PT), Emu(405 * EMU_PER_PT)
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(40), Pt(120), Pt(300), Pt(120))
    shape.fill.solid()
    shape.line.fill.background()
    shape.adjustments[0] = 0.05
    shape.element.spPr.append(shadow_xml(8, 6, 40000))
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    drive, slides = drive_service(), slides_service()
    f = execute(drive.files().create(
        body={"name": "b2s probe shadow", "mimeType": "application/vnd.google-apps.presentation"},
        media_body=MediaIoBaseUpload(buf, mimetype="application/vnd.openxmlformats-officedocument.presentationml.presentation"),
        fields="id"))
    pid = f["id"]
    pres = execute(slides.presentations().get(presentationId=pid))
    tpl = pres["slides"][0]
    for pe in tpl["pageElements"]:
        print(pe["objectId"], pe["shape"].get("shapeType"), pe["shape"].get("placeholder"),
              json.dumps(pe["shape"]["shapeProperties"].get("shadow")))
    box = next(pe for pe in tpl["pageElements"] if "placeholder" not in pe["shape"])
    title = next(pe for pe in tpl["pageElements"] if "placeholder" in pe["shape"])
    reqs = [
        {"duplicateObject": {"objectId": tpl["objectId"], "objectIds": {
            tpl["objectId"]: "probe_slide2", box["objectId"]: "probe_tpl2", title["objectId"]: "probe_title2"}}},
        {"duplicateObject": {"objectId": "probe_tpl2", "objectIds": {"probe_tpl2": "probe_copy"}}},
        {"updateShapeProperties": {"objectId": "probe_copy", "fields": "shapeBackgroundFill.solidFill.color",
                                   "shapeProperties": {"shapeBackgroundFill": {"solidFill": {"color": {"rgbColor": {"red": 0.9, "green": 0.94, "blue": 0.9}}}}}}},
        {"updatePageElementTransform": {"objectId": "probe_copy", "applyMode": "ABSOLUTE", "transform": {
            "scaleX": 1.5, "scaleY": 0.6, "unit": "EMU", "translateX": 360 * EMU_PER_PT, "translateY": 60 * EMU_PER_PT}}},
        {"insertText": {"objectId": "probe_title2", "text": "Duplicated slide"}},
        {"updateSlideProperties": {"objectId": tpl["objectId"], "fields": "isSkipped", "slideProperties": {"isSkipped": True}}},
    ]
    try:
        execute(slides.presentations().batchUpdate(presentationId=pid, body={"requests": reqs}))
    except Exception as e:
        print("batch failed:", e)
        reqs = reqs[:-1]
        execute(slides.presentations().batchUpdate(presentationId=pid, body={"requests": reqs}))
    pres = execute(slides.presentations().get(presentationId=pid))
    for s in pres["slides"]:
        print("slide", s["objectId"], s.get("slideProperties", {}).get("isSkipped"),
              s["slideProperties"].get("layoutObjectId"))
        for pe in s["pageElements"]:
            print("  ", pe["objectId"], pe["shape"].get("shapeType"),
                  json.dumps(pe["shape"]["shapeProperties"].get("shadow")))
    save_thumbnail(slides, pid, "probe_slide2", OUT / "probe_shadow.png")
    print(f"https://docs.google.com/presentation/d/{pid}/edit")


if __name__ == "__main__":
    main()
