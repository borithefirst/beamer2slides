"""Probe: bullets a text shell carries as `a:buChar` against the presets' own, and characters for
bullets no preset draws (docs/project-notes.md "Triangle bullets through the .pptx").

A text box holding beamer's ▶ comes in the .pptx as a shell whose bullets are characters
(emit_text.shell_route, emit_pptx._add_text_shell); its other bullets are written as their presets'
glyphs (emit_metrics.PRESET_CHARS: ● ○ ■ ❏ ★ ◆ ◇) and sized and placed with the presets' measures
(BULLET_SHAPES, from tools/probe_bullets.py), on the assumption that Slides draws a buChar glyph in
the run's face as createParagraphBullets draws its preset's. This measures, on Google's renderer,
at 24 pt in blue beside black Lato "HH" (indentFirstLine FIRST, indentStart START, as emit writes):

- A: each preset glyph as a shell's buChar and as createParagraphBullets' bullet in a box the API
  made: ink height (em), bottom against the baseline (em) and gap from the ink's right edge to
  indentFirstLine (em). Equal rows: the assumption holds. Else the shell rows are the measures a
  shell's preset glyphs need (a table beside CHAR_BULLETS).
- A mixed shell (► over ●, ● over ►): both glyphs drawn, each line's as in its own box (an
  imported list could hand one paragraph's glyph to the next).
- B: characters for Zapf Dingbats' thin eight-spoked ✴ (pifont \\ding{86}..., real_africa-remote-
  sens-30 s39: written as the star preset, a solid ★), and ⬤ for beamer's ball (a disc of 0.41 em
  held at its text's size, smaller than the PDF's ball): ink height and width (em), fill (ink
  pixels over its box: ✴'s is 0.21) and pixels of another colour than the bullet's (a colour emoji).
- C: ▶ and ► as words (no bullet: an empty item's label kept as text, real_c-error-handling s27)
  in Lato and in Noto Sans Symbols 2: ink height (em). MSAM10's ▶ is 0.58 em.

The probe deck is a .pptx imported by Drive (two slides), moved to the trash at the end (--keep
keeps it). Results: out/probe_shell_bullets/ (slide-N.png, results.json, a line per row printed).

--dry-run: no Google. Writes the .pptx and every request (ids checked: unique, 5-50 characters)
under out/probe_shell_bullets/dry-run/, and measures a stand-in thumbnail drawn here with local
faces at the planned places (proving the regions, baselines and colours are read right).

Usage: python tools/probe_shell_bullets.py [--dry-run] [--keep]
"""

import io
import json
import sys
from dataclasses import asdict, dataclass
from pathlib import Path

import numpy as np
from PIL import Image, ImageDraw, ImageFont

from beamer2slides import gapi
from beamer2slides.emit_metrics import PAD_X, PRESET_CHARS
from beamer2slides.emit_model import PptxText, ShellParagraph
from beamer2slides.emit_pptx import _add_text_shell
from beamer2slides.google_types import (BulletPreset, SlidesOptionalColor, SlidesRange, SlidesRequest, SlidesService,
                                        file_id, object_id, slides_json)
from beamer2slides.gslides import EMU_PER_PT, execute, pt, save_thumbnail, text_box
from beamer2slides.ir_types import Color

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "out" / "probe_shell_bullets"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
SLIDE_W, SLIDE_H = 720.0, 405.0
SIZE, FIRST, START = 24.0, 60.0, 90.0
BOX_W, BOX_H, PITCH = 230.0, 40.0, 44.0
BLUE_HEX = Color("#0000ff")
BLUE: SlidesOptionalColor = {"opaqueColor": {"rgbColor": {"blue": 1}}}
BLACK: SlidesOptionalColor = {"opaqueColor": {"rgbColor": {}}}
WORDS = "HH"
# preset glyph -> (preset, nesting level showing it), as tools/probe_bullets.py writes them
PRESETS: dict[str, tuple[BulletPreset, int]] = {
    PRESET_CHARS["disc"]: ("BULLET_DISC_CIRCLE_SQUARE", 0), PRESET_CHARS["circle"]: ("BULLET_DISC_CIRCLE_SQUARE", 1),
    PRESET_CHARS["square"]: ("BULLET_DISC_CIRCLE_SQUARE", 2), PRESET_CHARS["open_square"]: ("BULLET_CHECKBOX", 0),
    PRESET_CHARS["star"]: ("BULLET_STAR_CIRCLE_SQUARE", 0), PRESET_CHARS["diamond"]: ("BULLET_DIAMOND_CIRCLE_SQUARE", 0),
    PRESET_CHARS["open_diamond"]: ("BULLET_DIAMONDX_HOLLOWDIAMOND_SQUARE", 1)}
ASTERISKS = ["✴", "✳", "✻", "❊", "❋", "✱", "✲", "∗", "*", "⬤", "●"]
WORD_GLYPHS = [("▶", "Lato"), ("►", "Lato"), ("▶", "Noto Sans Symbols 2"), ("►", "Noto Sans Symbols 2")]


@dataclass(frozen=True, kw_only=True)
class Row:
    """One measured box: what it holds and where (slide index, top-left pt)."""
    name: str
    char: str
    way: str  # "shell", "api", "mixed" or "words"
    font: str
    slide: int
    x: float
    y: float
    lines: int


@dataclass(frozen=True, kw_only=True)
class Ink:
    """A row's blue ink, per em of SIZE."""
    height: float
    width: float
    bottom: float  # above the baseline (+) or under it (-)
    gap: float     # ink's right edge before indentFirstLine (bullets)
    fill: float    # ink pixels over the ink's box
    other: int     # pixels of a colour neither the bullet's nor the words'


def layout() -> list[Row]:
    rows: list[Row] = []
    for i, ch in enumerate(PRESETS):
        for c, way in enumerate(("shell", "api")):
            rows.append(Row(name=f"{way}_{i}", char=ch, way=way, font="Lato", slide=0, x=10 + c * 240,
                            y=8 + i * PITCH, lines=1))
    for j, pair in enumerate(("►●", "●►")):
        rows.append(Row(name=f"mixed_{j}", char=pair, way="mixed", font="Lato", slide=0, x=490, y=8 + j * 2 * PITCH,
                        lines=2))
    for i, ch in enumerate(ASTERISKS):
        rows.append(Row(name=f"char_{i}", char=ch, way="shell", font="Lato", slide=1, x=10 + (i // 8) * 240,
                        y=8 + (i % 8) * PITCH, lines=1))
    for i, (ch, font) in enumerate(WORD_GLYPHS):
        rows.append(Row(name=f"words_{i}", char=ch, way="words", font=font, slide=1, x=490, y=8 + i * PITCH, lines=1))
    return rows


def shell_of(row: Row) -> PptxText:
    chars = list(row.char) if row.way == "mixed" else [row.char]
    box = (row.x, row.y, row.x + BOX_W, row.y + BOX_H * row.lines)
    return PptxText(box=box, paragraphs=tuple(ShellParagraph(level=0, char=c, color=BLUE_HEX, size=SIZE, font="Lato",
                                                             text_size=SIZE) for c in chars))


def build() -> bytes:
    from pptx import Presentation
    from pptx.util import Emu
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(round(SLIDE_W * EMU_PER_PT)), Emu(round(SLIDE_H * EMU_PER_PT))
    slides = [prs.slides.add_slide(prs.slide_layouts[6]) for _ in range(2)]
    for row in layout():
        if row.way in ("shell", "mixed"):
            _add_text_shell(slides[row.slide], shell_of(row))
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def fixed(s: int, e: int) -> SlidesRange:
    return {"type": "FIXED_RANGE", "startIndex": s, "endIndex": e}


def words_style(oid: str, at: int) -> list[SlidesRequest]:
    """The words black Lato, a request per letter (one over a whole paragraph restyles its bullet)."""
    return [{"updateTextStyle": {"objectId": oid, "textRange": fixed(k, k + 1), "fields": "fontFamily,fontSize,foregroundColor",
                                 "style": {"fontFamily": "Lato", "fontSize": pt(SIZE), "foregroundColor": BLACK}}}
            for k in range(at, at + len(WORDS))]


def indents(oid: str, start: int, end: int) -> SlidesRequest:
    return {"updateParagraphStyle": {"objectId": oid, "textRange": fixed(start, end),
                                     "fields": "indentFirstLine,indentStart,spaceAbove,spaceBelow,lineSpacing",
                                     "style": {"indentFirstLine": pt(FIRST), "indentStart": pt(START), "spaceAbove": pt(0),
                                               "spaceBelow": pt(0), "lineSpacing": 100}}}


def requests(rows: list[Row], shells: dict[str, str], pages: list[str]) -> list[SlidesRequest]:
    """Every request after the import: shells filled as emit_text.text_shell_requests_of fills one,
    the API's boxes made as tools/probe_bullets.py makes them, the words' boxes."""
    reqs: list[SlidesRequest] = []
    for row in rows:
        if row.way in ("shell", "mixed"):
            oid = shells[row.name]
            for i in reversed(range(row.lines)):
                at = i * 2  # (each paragraph: its placeholder character and a newline)
                fill: list[SlidesRequest] = [
                    {"insertText": {"objectId": oid, "text": WORDS, "insertionIndex": at}},
                    {"deleteText": {"objectId": oid, "textRange": fixed(at + len(WORDS), at + len(WORDS) + 1)}}]
                reqs += fill
            for i in range(row.lines):
                at = i * (len(WORDS) + 1)
                reqs += words_style(oid, at) + [indents(oid, at, at + len(WORDS))]
        elif row.way == "api":
            oid = f"probe_{row.name}"
            preset, level = PRESETS[row.char]
            text = "x\n" + "\t" * level + WORDS  # (the dummy "x" paragraph makes the level absolute)
            made: list[SlidesRequest] = [
                text_box(oid, pages[row.slide], row.x, row.y, BOX_W, BOX_H),
                {"insertText": {"objectId": oid, "text": text}},
                {"updateTextStyle": {"objectId": oid, "textRange": {"type": "ALL"}, "fields": "fontFamily,fontSize,foregroundColor",
                                     "style": {"fontFamily": "Lato", "fontSize": pt(SIZE), "foregroundColor": BLUE}}},
                {"createParagraphBullets": {"objectId": oid, "bulletPreset": preset, "textRange": fixed(0, len(text))}},
                {"deleteText": {"objectId": oid, "textRange": fixed(0, 2)}}]
            reqs += made
            reqs += words_style(oid, 0) + [indents(oid, 0, len(WORDS))]
        else:
            oid = f"probe_{row.name}"
            text = f"H{row.char}H"
            words: list[SlidesRequest] = [
                text_box(oid, pages[row.slide], row.x, row.y, BOX_W, BOX_H),
                {"insertText": {"objectId": oid, "text": text}},
                {"updateTextStyle": {"objectId": oid, "textRange": {"type": "ALL"}, "fields": "fontFamily,fontSize,foregroundColor",
                                     "style": {"fontFamily": "Lato", "fontSize": pt(SIZE), "foregroundColor": BLACK}}},
                {"updateTextStyle": {"objectId": oid, "textRange": fixed(1, 2), "fields": "fontFamily,foregroundColor",
                                     "style": {"fontFamily": row.font, "foregroundColor": BLUE}}}]
            reqs += words
    return reqs


def checked(reqs: list[SlidesRequest]) -> None:
    """Every id a request creates is unique and 5-50 characters long (Slides refuses others)."""
    made: list[str] = []
    for r in reqs:
        shape = r.get("createShape")
        oid = None if shape is None else shape.get("objectId")
        if oid is not None:
            made.append(oid)
    bad = [i for i in made if not 5 <= len(i) <= 50]
    if bad or len(set(made)) != len(made):
        raise SystemExit(f"object ids refused by Slides: {bad or made}")


def measure(img: Image.Image, row: Row, line: int) -> Ink | None:
    """The blue ink of one line of `row` (0-based) in a slide's thumbnail, per em of SIZE."""
    im = np.asarray(img.convert("RGB")).astype(int)
    k = im.shape[1] / SLIDE_W
    x0, y0 = round(row.x * k), round(row.y * k)
    x1, y1 = round((row.x + BOX_W + 10) * k), round((row.y + BOX_H * row.lines) * k)
    r, g, b = (im[y0:y1, x0:x1, c] for c in range(3))
    blue = (b > 150) & (r < 120) & (g < 120)
    black = (r + g + b) < 200
    # (a colour of its own: saturated, and not the bullet's blue blended into white at its edges)
    other = (np.maximum(np.maximum(r, g), b) - np.minimum(np.minimum(r, g), b) > 60) & (b < np.maximum(r, g))
    rows_black = np.nonzero(black.any(axis=1))[0]
    if not len(rows_black):
        return None
    # the lines' baselines: the foot of each run of rows holding black ink
    feet = [int(a) for a, nxt in zip(rows_black, list(rows_black[1:]) + [10 ** 9]) if nxt - a > 1]
    if line >= len(feet):
        return None
    lo = feet[line - 1] + round(0.4 * SIZE * k) if line else 0  # (past the line above's descent)
    base = feet[line] + 1
    band = blue[lo:base + round(0.4 * SIZE * k)]
    ys, xs = np.nonzero(band)
    if not len(ys):
        return None
    top, bottom = (ys.min() + lo) / k, (ys.max() + 1 + lo) / k
    left, right = xs.min() / k, (xs.max() + 1) / k
    area = (ys.max() - ys.min() + 1) * (xs.max() - xs.min() + 1)
    return Ink(height=round((bottom - top) / SIZE, 3), width=round((right - left) / SIZE, 3),
               bottom=round((base / k - bottom) / SIZE, 3), gap=round((PAD_X + FIRST - right) / SIZE, 3),
               fill=round(len(ys) / area, 2), other=int(other[lo:base + round(0.4 * SIZE * k)].sum()))


def stand_in(rows: list[Row], slide: int) -> Image.Image:
    """A thumbnail drawn here: each row's words in black at its baseline and its glyph in blue where
    Slides would end it, in local faces (only the measuring is under test)."""
    k = 1600 / SLIDE_W
    img = Image.new("RGB", (1600, round(SLIDE_H * k)), "white")
    draw = ImageDraw.Draw(img)
    try:
        face: ImageFont.FreeTypeFont | ImageFont.ImageFont = ImageFont.truetype("seguisym.ttf", round(SIZE * k))
    except OSError:
        face = ImageFont.load_default()
    for row in rows:
        if row.slide != slide:
            continue
        for i, ch in enumerate(row.char if row.way == "mixed" else [row.char]):
            base = (row.y + 7.2 + 0.95 * SIZE + i * 1.2 * SIZE) * k
            draw.text(((row.x + PAD_X + START) * k, base), WORDS, fill="black", font=face, anchor="ls")
            draw.text(((row.x + PAD_X + FIRST - 0.1 * SIZE) * k, base), ch, fill=(0, 0, 255), font=face, anchor="rs")
    return img


def report(rows: list[Row], images: list[Image.Image]) -> dict[str, object]:
    results: dict[str, object] = {}
    for row in rows:
        for line in range(row.lines):
            ink = measure(images[row.slide], row, line)
            ch = row.char[line] if row.way == "mixed" else row.char
            key = f"{row.name}/{line}" if row.lines > 1 else row.name
            results[key] = {"char": ch, "way": row.way, "font": row.font, "ink": None if ink is None else asdict(ink)}
            print(f"{key:10s} {row.way:6s} {ch} {row.font:20s} " + ("no ink" if ink is None else
                  f"height {ink.height:.3f} em, width {ink.width:.3f}, bottom {ink.bottom:+.3f}, gap {ink.gap:.3f}, "
                  f"fill {ink.fill:.2f}, other colour px {ink.other}"))
    return results


def main() -> None:
    dry, keep = "--dry-run" in sys.argv, "--keep" in sys.argv
    rows = layout()
    out = OUT / "dry-run" if dry else OUT
    out.mkdir(parents=True, exist_ok=True)
    data = build()
    (out / "probe.pptx").write_bytes(data)
    if dry:
        shells = {r.name: f"imported_{r.name}" for r in rows if r.way in ("shell", "mixed")}
        reqs = requests(rows, shells, ["page_one", "page_two"])
        checked(reqs)
        (out / "requests.json").write_text(json.dumps([slides_json(r) for r in reqs], indent=1, ensure_ascii=False),
                                           encoding="utf-8")
        images = [stand_in(rows, s) for s in (0, 1)]
        for s, img in enumerate(images):
            img.save(out / f"slide-{s + 1}.png")
        (out / "results.json").write_text(json.dumps(report(rows, images), indent=1, ensure_ascii=False), encoding="utf-8")
        print(f"dry run: {len(reqs)} requests, {out}")
        return
    from beamer2slides.google_auth import drive_service, slides_service
    drive, slides = drive_service(None), slides_service(None)
    pid = file_id(execute(drive.files().create(
        body={"name": "b2s probe shell bullets", "mimeType": "application/vnd.google-apps.presentation"},
        media_body=gapi.media_upload(io.BytesIO(data), PPTX_MIME), fields="id")), "the uploaded probe deck")
    try:
        run(slides, pid, rows, out)
    finally:
        if keep:
            print(f"kept https://docs.google.com/presentation/d/{pid}/edit")
        else:
            execute(drive.files().update(fileId=pid, body={"trashed": True}))
            print("probe deck moved to the trash")


def run(slides: SlidesService, pid: str, rows: list[Row], out: Path) -> None:
    pres = execute(slides.presentations().get(presentationId=pid))
    pages = pres.get("slides", [])
    page_ids = [object_id(p) for p in pages]
    shells: dict[str, str] = {}
    for s, page in enumerate(pages):
        # (the importer drops the shapes' names; the shells come back in the order they were made)
        mine = [r for r in rows if r.slide == s and r.way in ("shell", "mixed")]
        boxes = [object_id(pe) for pe in page.get("pageElements", []) if "shape" in pe]
        if len(boxes) != len(mine):
            raise SystemExit(f"slide {s + 1}: {len(boxes)} shapes imported for {len(mine)} shells")
        shells.update({r.name: oid for r, oid in zip(mine, boxes)})
    reqs = requests(rows, shells, page_ids)
    checked(reqs)
    execute(slides.presentations().batchUpdate(presentationId=pid, body={"requests": reqs}))
    images: list[Image.Image] = []
    for s, page_id in enumerate(page_ids):
        png = out / f"slide-{s + 1}.png"
        save_thumbnail(slides, pid, page_id, png, None)
        images.append(Image.open(png))
    got = execute(slides.presentations().get(presentationId=pid))
    (out / "after.json").write_text(json.dumps(got.get("slides", []), indent=1, ensure_ascii=False), encoding="utf-8")
    (out / "results.json").write_text(json.dumps(report(rows, images), indent=1, ensure_ascii=False), encoding="utf-8")
    print(out)


if __name__ == "__main__":
    main()
