"""Probe: pictures and backgrounds carried by the imported .pptx instead of public URLs.

The Slides API only inserts images from publicly fetchable URLs, which protected Google
Workspace domains forbid. A .pptx imported through Drive can carry the pictures itself. Does
the import keep: a master background picture (inherited by layouts and slides), a slide's own
background picture or colour, pictures with their alt text, empty title placeholders? And does
duplicateObject of such a slide keep all of that under the object IDs we choose?

Usage: python tools/probe_pptx_pictures.py [presentationId-to-reuse]
"""

import io
import json
import sys
from pathlib import Path

from googleapiclient.http import MediaIoBaseUpload
from lxml import etree
from pptx import Presentation
from pptx.util import Emu, Pt

from beamer2slides.google_auth import drive_service, slides_service
from beamer2slides.gslides import EMU_PER_PT, execute, save_thumbnail

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "out" / "probe_pptx_pictures"
SRC = ROOT / "out" / "gdg-talk"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def picture_background(part, c_sld, png: Path) -> None:
    _, rid = part.get_or_add_image_part(str(png))
    bg = etree.fromstring(
        f'<p:bg xmlns:p="{P}" xmlns:a="{A}" xmlns:r="{R}"><p:bgPr><a:blipFill dpi="0" rotWithShape="1">'
        f'<a:blip r:embed="{rid}"/><a:srcRect/><a:stretch><a:fillRect/></a:stretch></a:blipFill>'
        f'<a:effectLst/></p:bgPr></p:bg>')
    old = c_sld.find(f"{{{P}}}bg")
    if old is not None:
        c_sld.remove(old)
    c_sld.insert(0, bg)


def colour_background(c_sld, hex_colour: str) -> None:
    bg = etree.fromstring(
        f'<p:bg xmlns:p="{P}" xmlns:a="{A}"><p:bgPr><a:solidFill><a:srgbClr val="{hex_colour}"/></a:solidFill>'
        f'<a:effectLst/></p:bgPr></p:bg>')
    c_sld.insert(0, bg)


def build() -> io.BytesIO:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(720 * EMU_PER_PT), Emu(405 * EMU_PER_PT)
    master = prs.slide_master
    picture_background(master.part, master.element.find(f"{{{P}}}cSld"), SRC / "backgrounds" / "bg-005.png")

    s1 = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only, inherits the master background
    pic = s1.shapes.add_picture(str(SRC / "figures" / "p4h1.png"), Pt(100), Pt(150), Pt(200), Pt(100))
    pic._element.nvPicPr.cNvPr.set("descr", "b2s_f3")
    pic._element.nvPicPr.cNvPr.set("title", "Figure")

    s2 = prs.slides.add_slide(prs.slide_layouts[6])  # Blank, own background picture
    picture_background(s2.part, s2.element.find(f"{{{P}}}cSld"), SRC / "backgrounds" / "bg-009.png")

    s3 = prs.slides.add_slide(prs.slide_layouts[0])  # Title, own colour
    colour_background(s3.element.find(f"{{{P}}}cSld"), "FDE293")
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf


def main() -> None:
    drive, slides = drive_service(), slides_service()
    media = MediaIoBaseUpload(build(), mimetype=PPTX_MIME)
    if len(sys.argv) > 1:
        pid = sys.argv[1]
        execute(drive.files().update(fileId=pid, media_body=media, fields="id"))
    else:
        pid = execute(drive.files().create(
            body={"name": "b2s probe pptx pictures", "mimeType": "application/vnd.google-apps.presentation"},
            media_body=media, fields="id"))["id"]
    pres = execute(slides.presentations().get(presentationId=pid))

    def fill(page):
        f = page.get("pageProperties", {}).get("pageBackgroundFill", {})
        return {k: (v if k != "stretchedPictureFill" else {"contentUrl": bool(v.get("contentUrl")), "size": v.get("size")})
                for k, v in f.items()}

    for m in pres["masters"]:
        print("master", m["objectId"], json.dumps(fill(m)))
    for l in pres["layouts"][:3]:
        print("layout", l["objectId"], json.dumps(fill(l)))
    for s in pres["slides"]:
        print("slide", s["objectId"], json.dumps(fill(s)))
        for pe in s.get("pageElements", []):
            kind = next(k for k in ("shape", "image", "table", "line", "elementGroup") if k in pe)
            print("   ", pe["objectId"], kind, pe.get("title"), pe.get("description"),
                  pe.get("shape", {}).get("placeholder"), json.dumps(pe.get("transform")), json.dumps(pe.get("size")))

    first = pres["slides"][0]
    ids = {first["objectId"]: "b2s_s900", **{pe["objectId"]: f"b2s_s900_e{j}" for j, pe in enumerate(first["pageElements"])}}
    second = pres["slides"][1]
    ids2 = {second["objectId"]: "b2s_s901"}
    execute(slides.presentations().batchUpdate(presentationId=pid, body={"requests": [
        {"duplicateObject": {"objectId": first["objectId"], "objectIds": ids}},
        {"duplicateObject": {"objectId": second["objectId"], "objectIds": ids2}},
        {"updatePageElementsZOrder": {"pageElementObjectIds": ["b2s_s900_e1"], "operation": "BRING_TO_FRONT"}},
        {"insertText": {"objectId": "b2s_s900_e0", "text": "Duplicated"}},
    ]}))
    pres = execute(slides.presentations().get(presentationId=pid))
    for s in pres["slides"]:
        if s["objectId"].startswith("b2s_"):
            print("dup", s["objectId"], json.dumps(fill(s)), [(pe["objectId"], pe.get("description")) for pe in s.get("pageElements", [])])
    save_thumbnail(slides, pid, "b2s_s900", OUT / "dup-900.png")
    save_thumbnail(slides, pid, "b2s_s901", OUT / "dup-901.png")
    print(f"https://docs.google.com/presentation/d/{pid}/edit")


if __name__ == "__main__":
    main()
