"""Probe: a beamer block whose body grows with its text. The body shape carries the body
text with "resize shape to fit text" (spAutoFit) and a top inset clearing the title bar;
the title bar shape carries the title. Neither autofit nor insets can be set through the
API, so they come from an imported .pptx. Checks what the API reports after import, after
duplicateObject and after replacing the text through the API.

Usage: python tools/probe_autofit_block.py
"""

import io
import json
from pathlib import Path

from googleapiclient.http import MediaIoBaseUpload
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

from beamer2slides.google_auth import drive_service, slides_service
from beamer2slides.gslides import EMU_PER_PT, execute, save_thumbnail

OUT = Path(__file__).resolve().parents[1] / "out"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def main() -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(720 * EMU_PER_PT), Emu(540 * EMU_PER_PT)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    body = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(40), Pt(100), Pt(640), Pt(60))
    body.adjustments[0] = 0.12
    body.fill.solid()
    body.fill.fore_color.rgb = RGBColor(0xE9, 0xE9, 0xF3)
    body.line.fill.background()
    body.element.spPr.append(etree.fromstring(
        f'<a:effectLst xmlns:a="{A}"><a:outerShdw blurRad="{8 * EMU_PER_PT}" dist="{6 * EMU_PER_PT}" dir="2700000" '
        f'algn="tl" rotWithShape="0"><a:srgbClr val="000000"><a:alpha val="50000"/></a:srgbClr></a:outerShdw></a:effectLst>'))
    tf = body.text_frame
    tf.text = "Body text of a standard block that is long enough to wrap when the block gets narrower."
    tf.paragraphs[0].runs[0].font.size = Pt(22)
    tf.paragraphs[0].runs[0].font.color.rgb = RGBColor(0, 0, 0)
    body_pr = tf._txBody.find(f"{{{A}}}bodyPr")
    body_pr.set("tIns", str(36 * EMU_PER_PT))
    body_pr.set("lIns", str(7 * EMU_PER_PT))
    body_pr.set("anchor", "t")
    for child in list(body_pr):
        body_pr.remove(child)
    body_pr.append(etree.fromstring(f'<a:spAutoFit xmlns:a="{A}"/>'))

    head = slide.shapes.add_shape(MSO_SHAPE.ROUND_2_SAME_RECTANGLE, Pt(40), Pt(100), Pt(640), Pt(30))
    head.adjustments[0] = 0.25
    head.adjustments[1] = 0.0
    head.fill.solid()
    head.fill.fore_color.rgb = RGBColor(0x26, 0x26, 0x86)
    head.line.fill.background()
    head.element.spPr.append(etree.fromstring(f'<a:effectLst xmlns:a="{A}"/>'))
    head.text_frame.text = "Standard block"
    head.text_frame.paragraphs[0].runs[0].font.size = Pt(22)
    head.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    head_pr = head.text_frame._txBody.find(f"{{{A}}}bodyPr")
    head_pr.set("anchor", "ctr")
    head_pr.set("tIns", "0")
    head_pr.set("bIns", "0")

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    drive, slides = drive_service(), slides_service()
    pid = execute(drive.files().create(
        body={"name": "b2s probe autofit block", "mimeType": "application/vnd.google-apps.presentation"},
        media_body=MediaIoBaseUpload(buf, mimetype=PPTX_MIME), fields="id"))["id"]
    pres = execute(slides.presentations().get(presentationId=pid))
    page = pres["slides"][0]
    for pe in page["pageElements"]:
        props = pe["shape"]["shapeProperties"]
        print(pe["objectId"], pe["shape"]["shapeType"], json.dumps(props.get("autofit")), props.get("contentAlignment"),
              pe["size"], pe["transform"])
    body_id = page["pageElements"][0]["objectId"]
    head_id = page["pageElements"][1]["objectId"]
    execute(slides.presentations().batchUpdate(presentationId=pid, body={"requests": [
        {"duplicateObject": {"objectId": body_id, "objectIds": {body_id: "probe_body2"}}},
        {"duplicateObject": {"objectId": head_id, "objectIds": {head_id: "probe_head2"}}},
        {"updatePageElementTransform": {"objectId": "probe_body2", "applyMode": "RELATIVE", "transform": {
            "scaleX": 1, "scaleY": 1, "translateX": 0, "translateY": 200 * EMU_PER_PT, "unit": "EMU"}}},
        {"updatePageElementTransform": {"objectId": "probe_head2", "applyMode": "RELATIVE", "transform": {
            "scaleX": 1, "scaleY": 1, "translateX": 0, "translateY": 200 * EMU_PER_PT, "unit": "EMU"}}},
        {"deleteText": {"objectId": "probe_body2", "textRange": {"type": "ALL"}}},
        {"insertText": {"objectId": "probe_body2", "text": "Text replaced through the API, on two lines\nsecond paragraph"}},
        {"groupObjects": {"groupObjectId": "probe_block2", "childrenObjectIds": ["probe_body2", "probe_head2"]}},
    ]}))
    pres = execute(slides.presentations().get(presentationId=pid))
    for pe in pres["slides"][0]["pageElements"]:
        for c in pe.get("elementGroup", {}).get("children", [pe]):
            props = c["shape"]["shapeProperties"]
            print(c["objectId"], json.dumps(props.get("autofit")), c["size"], c["transform"])
    save_thumbnail(slides, pid, page["objectId"], OUT / "probe_autofit_block.png")
    print(f"https://docs.google.com/presentation/d/{pid}/edit")


if __name__ == "__main__":
    main()
