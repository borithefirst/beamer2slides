"""Probe: diagram nodes that keep their label inside when moved, and arrows that follow.

- A template rectangle from an imported .pptx with zero text insets (the API can't set
  padding): duplicated, restyled and given a label through the API. Does the label stay on one
  line in a node as tight as a TikZ node?
- A straight line between two nodes with startConnection/endConnection on connection sites.
  Where does the API put the line, and what does it report?

Saves out/probe_diagram.png and prints the line and shape properties.

Usage: python tools/probe_diagram.py
"""

import io
import json
from pathlib import Path

from googleapiclient.http import MediaIoBaseUpload
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

from beamer2slides.google_auth import drive_service, slides_service
from beamer2slides.gslides import EMU_PER_PT, execute, pt, save_thumbnail

OUT = Path(__file__).resolve().parents[1] / "out"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def main() -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(720 * EMU_PER_PT), Emu(405 * EMU_PER_PT)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for kind in (MSO_SHAPE.RECTANGLE, MSO_SHAPE.OVAL):
        shape = slide.shapes.add_shape(kind, Pt(10), Pt(10), Pt(100), Pt(100))
        shape.fill.solid()
        shape.element.spPr.append(etree.fromstring(f'<a:effectLst xmlns:a="{A}"/>'))
        body_pr = shape.text_frame._txBody.find(f"{{{A}}}bodyPr")
        for side in ("lIns", "tIns", "rIns", "bIns"):
            body_pr.set(side, "0")
        body_pr.set("anchor", "ctr")
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    drive, slides = drive_service(), slides_service()
    pid = execute(drive.files().create(
        body={"name": "b2s probe diagram", "mimeType": "application/vnd.google-apps.presentation"},
        media_body=MediaIoBaseUpload(buf, mimetype=PPTX_MIME), fields="id"))["id"]
    pres = execute(slides.presentations().get(presentationId=pid))
    page = pres["slides"][0]
    rect_tpl, oval_tpl = [e["objectId"] for e in page["pageElements"]]

    def node(tpl: str, oid: str, x: float, y: float, w: float, h: float, text: str, size: float) -> list[dict]:
        return [
            {"duplicateObject": {"objectId": tpl, "objectIds": {tpl: oid}}},
            {"updatePageElementTransform": {"objectId": oid, "applyMode": "ABSOLUTE", "transform": {
                "scaleX": w / 236.22, "scaleY": h / 236.22, "translateX": x * EMU_PER_PT, "translateY": y * EMU_PER_PT, "unit": "EMU"}}},
            {"updateShapeProperties": {"objectId": oid, "fields": "shapeBackgroundFill.solidFill.color,outline.outlineFill.solidFill.color,outline.weight",
                                       "shapeProperties": {"shapeBackgroundFill": {"solidFill": {"color": {"rgbColor": {"red": 0.85, "green": 0.9, "blue": 1}}}},
                                                           "outline": {"outlineFill": {"solidFill": {"color": {"rgbColor": {}}}}, "weight": pt(0.8)}}}},
            {"insertText": {"objectId": oid, "text": text}},
            {"updateTextStyle": {"objectId": oid, "textRange": {"type": "ALL"}, "fields": "fontFamily,fontSize,foregroundColor",
                                 "style": {"fontFamily": "Lato", "fontSize": pt(size), "foregroundColor": {"opaqueColor": {"rgbColor": {}}}}}},
            {"updateParagraphStyle": {"objectId": oid, "textRange": {"type": "ALL"}, "fields": "alignment",
                                      "style": {"alignment": "CENTER"}}},
        ]

    # "Result" at 22 pt Lato is ~62 pt wide; TikZ inner sep adds ~13 pt: a 75 pt node.
    reqs = node(rect_tpl, "node_a", 60, 80, 75, 34, "Result", 22) + node(oval_tpl, "node_b", 300, 150, 60, 60, "B", 22) \
        + node(rect_tpl, "node_c", 60, 250, 140, 34, "Convert all", 22)
    reqs += [
        {"createLine": {"objectId": "line_ab", "lineCategory": "STRAIGHT", "elementProperties": {
            "pageObjectId": page["objectId"], "size": {"width": {"magnitude": 165 * EMU_PER_PT, "unit": "EMU"},
                                                         "height": {"magnitude": 83 * EMU_PER_PT, "unit": "EMU"}},
            "transform": {"scaleX": 1, "scaleY": 1, "translateX": 135 * EMU_PER_PT, "translateY": 97 * EMU_PER_PT, "unit": "EMU"}}}},
        {"updateLineProperties": {"objectId": "line_ab", "fields": "endArrow,startConnection,endConnection",
                                  "lineProperties": {"endArrow": "OPEN_ARROW",
                                                     "startConnection": {"connectedObjectId": "node_a", "connectionSiteIndex": 3},
                                                     "endConnection": {"connectedObjectId": "node_b", "connectionSiteIndex": 2}}}},
        {"deleteObject": {"objectId": rect_tpl}}, {"deleteObject": {"objectId": oval_tpl}},
    ]
    try:
        execute(slides.presentations().batchUpdate(presentationId=pid, body={"requests": reqs}))
    except Exception as e:
        print("batch failed:", str(e)[:400])
    pres = execute(slides.presentations().get(presentationId=pid))
    for e in pres["slides"][0]["pageElements"]:
        kind = "line" if "line" in e else e["shape"]["shapeType"]
        extra = json.dumps({k: v for k, v in e["line"]["lineProperties"].items() if "onnection" in k}) if "line" in e else ""
        print(e["objectId"], kind, e.get("size"), e["transform"], extra)
    save_thumbnail(slides, pid, pres["slides"][0]["objectId"], OUT / "probe_diagram.png")
    print(f"https://docs.google.com/presentation/d/{pid}/edit")


if __name__ == "__main__":
    main()
