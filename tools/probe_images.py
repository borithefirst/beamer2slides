"""Probe: what a live deck gives back of a picture a person inserted and edited.

Pull must recover pictures at their best resolution, with the edits a person made in Slides
(crop, transparency, recolour, brightness/contrast, rotation, outline) as properties it can write
as LaTeX options rather than pixels baked into a file. Which source keeps what?

  (a) `image.contentUrl` bytes: pixel size, format, what is baked into them;
  (b) the deck exported as .pptx (Drive `files.export`): ppt/media and the picture XML
      (srcRect, alphaModFix, lum, grayscl/duotone, xfrm rot, ln);
  (c) `image.sourceUrl`.

Pictures reach the deck two ways: a staging .pptx import carrying every edit as DrawingML (what a
person's uploaded .pptx gives, and the only way to set crop/transparency/recolour: the API can't),
and `createImage` from the staging contentUrls into the converted probe deck (what an inserted
picture gives), where `updateImageProperties` and a rotating transform are tried field by field.

Usage: python tools/probe_images.py [--folder out/agent-pull-images/probe] [--keep]
The probe deck is the convert output in --folder (reused every run: its probe slide is replaced);
the staging file is deleted at the end. Results: <folder>/probe-images/findings.json and a table.
"""

import argparse
import hashlib
import io
import json
import math
import re
import struct
import sys
import time
import zipfile
from pathlib import Path

import numpy as np
from googleapiclient.errors import HttpError
from googleapiclient.http import MediaIoBaseUpload
from lxml import etree
from PIL import Image, ImageDraw

from beamer2slides.deck_ir import fetch_url
from beamer2slides.google_auth import credentials, drive_service, slides_service
from beamer2slides.gslides import EMU_PER_PT, execute

ROOT = Path(__file__).resolve().parents[1]
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
SLIDE_ID = "b2s_probe_images"
CROP = {"l": 0.10, "t": 0.20, "r": 0.30, "b": 0.05}


# ---------------------------------------------------------------- test pictures

def photo(w: int = 3000, h: int = 2000) -> Image.Image:
    """Photo-like: smooth sky and hills, a sun, fine 1 px stripes and small text, mild noise."""
    rng = np.random.default_rng(7)
    y, x = np.mgrid[0:h, 0:w].astype(np.float32)
    img = np.zeros((h, w, 3), np.float32)
    img[..., 0] = 90 + 80 * y / h
    img[..., 1] = 140 + 60 * y / h
    img[..., 2] = 230 - 40 * y / h
    hill = h * (0.62 + 0.08 * np.sin(x / w * 7) + 0.03 * np.sin(x / w * 31))
    ground = y > hill
    img[ground] = np.stack([60 + 40 * np.sin(x[ground] / 37), 120 + 30 * np.cos(y[ground] / 23), 50 + 0 * x[ground]], -1)
    sun = (x - 0.78 * w) ** 2 + (y - 0.22 * h) ** 2 < (0.08 * h) ** 2
    img[sun] = [250, 220, 90]
    img += rng.normal(0, 4, img.shape)
    out = Image.fromarray(np.clip(img, 0, 255).astype(np.uint8))
    d = ImageDraw.Draw(out)
    for k in range(0, 400, 2):  # 1 px stripes: only survive at full resolution
        d.line([(100 + k, 100), (100 + k, 500)], fill=(0, 0, 0))
    for row in range(40):
        d.text((700, 100 + row * 12), "fine print 0123456789 " * 6, fill=(20, 20, 20))
    return out


def frames_gif(path: Path) -> None:
    frames = [Image.new("RGB", (240, 160), c) for c in ((220, 30, 30), (30, 160, 30), (30, 30, 220))]
    frames[0].save(path, save_all=True, append_images=frames[1:], duration=300, loop=0)


def emf_file(path: Path) -> None:
    """A minimal EMF: header, a red rectangle, EOF."""
    def rec(kind: int, payload: bytes) -> bytes:
        return struct.pack("<II", kind, 8 + len(payload)) + payload
    bounds = struct.pack("<4i", 0, 0, 199, 99)
    frame = struct.pack("<4i", 0, 0, 5291, 2645)  # 0.01 mm
    body = [
        rec(39, struct.pack("<II", 1, 0) + struct.pack("<BBBB", 220, 30, 30, 0) + struct.pack("<I", 0)),  # CREATEBRUSHINDIRECT
        rec(37, struct.pack("<I", 1)),  # SELECTOBJECT
        rec(43, struct.pack("<4i", 10, 10, 190, 90)),  # RECTANGLE
        rec(14, struct.pack("<III", 0, 16, 20)),  # EOF
    ]
    n_records = 1 + len(body)
    header_payload = (bounds + frame + struct.pack("<I", 0x464D4520) + struct.pack("<I", 0x10000) + b"SIZE"
                      + struct.pack("<I", n_records) + struct.pack("<HH", 2, 0) + struct.pack("<I", 0)
                      + struct.pack("<I", 0) + struct.pack("<I", 0) + struct.pack("<ii", 1920, 1080)
                      + struct.pack("<ii", 508, 286))
    header = rec(1, header_payload)
    data = header + b"".join(body)
    data = data[:48] + struct.pack("<I", len(data)) + data[52:]
    path.write_bytes(data)


# ---------------------------------------------------------------- staging .pptx

VARIANTS = [  # name, file, DrawingML edits
    ("plain", "photo.png", {}),
    ("crop", "photo.png", {"srcRect": CROP}),
    ("alpha50", "photo.png", {"alpha": 50000}),
    ("bright", "photo.png", {"lum": (30000, -20000)}),
    ("gray", "photo.png", {"grayscl": True}),
    ("duotone", "photo.png", {"duotone": True}),
    ("rot30", "photo.png", {"rot": 30}),
    ("outline", "photo.png", {"ln": ("FF0000", 3)}),
    ("crop_rot", "photo.png", {"srcRect": CROP, "rot": -15}),
    ("jpeg", "photo.jpg", {}),
    ("png1600", "photo1600.png", {}),
    ("jpeg1600", "photo1600.jpg", {}),
    ("gif", "anim.gif", {}),
    ("emf", "rect.emf", {}),
]


def staging_pptx(folder: Path) -> io.BytesIO:
    from pptx import Presentation
    from pptx.opc.packuri import PackURI
    from pptx.util import Emu, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(720 * EMU_PER_PT), Emu(405 * EMU_PER_PT)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for k, (name, file, edits) in enumerate(VARIANTS):
        x, y = 10 + (k % 5) * 142, 10 + (k // 5) * 130
        w, h = 130, 87
        if edits.get("srcRect"):
            c = edits["srcRect"]
            h = w * (1 - c["t"] - c["b"]) * 2000 / ((1 - c["l"] - c["r"]) * 3000)
        pic = slide.shapes.add_picture(str(folder / file), Pt(x), Pt(y), Pt(w), Pt(h))
        pic._element.nvPicPr.cNvPr.set("descr", f"probe:{name}")
        pic._element.nvPicPr.cNvPr.set("title", name)
        if file.endswith(".emf"):  # python-pptx files EMF as WMF
            part = pic.part.related_part(pic._element.blipFill.blip.rEmbed)
            part.partname = PackURI(str(part.partname).rsplit(".", 1)[0] + ".emf")
            part._content_type = "image/x-emf"
        blip = pic._element.blipFill.find(f"{{{A}}}blip")
        if edits.get("srcRect"):
            c = edits["srcRect"]
            pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = c["l"], c["t"], c["r"], c["b"]
        if edits.get("alpha"):
            etree.SubElement(blip, f"{{{A}}}alphaModFix", amt=str(edits["alpha"]))
        if edits.get("lum"):
            etree.SubElement(blip, f"{{{A}}}lum", bright=str(edits["lum"][0]), contrast=str(edits["lum"][1]))
        if edits.get("grayscl"):
            etree.SubElement(blip, f"{{{A}}}grayscl")
        if edits.get("duotone"):
            duo = etree.SubElement(blip, f"{{{A}}}duotone")
            etree.SubElement(duo, f"{{{A}}}prstClr", val="black")
            etree.SubElement(etree.SubElement(duo, f"{{{A}}}srgbClr", val="4285F4"), f"{{{A}}}tint", val="45000")
        if edits.get("rot"):
            pic.rotation = edits["rot"]
        if edits.get("ln"):
            pic.line.color.rgb = __import__("pptx.dml.color", fromlist=["RGBColor"]).RGBColor.from_string(edits["ln"][0])
            pic.line.width = Pt(edits["ln"][1])
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------- analysis helpers

def describe(data: bytes) -> dict:
    fmt = "png" if data[:4] == b"\x89PNG" else "jpeg" if data[:2] == b"\xff\xd8" else "gif" if data[:3] == b"GIF" \
        else "webp" if data[8:12] == b"WEBP" else "emf" if data[40:44] == b" EMF" else "wmf" if data[:4] == b"\xd7\xcd\xc6\x9a" \
        else "svg" if b"<svg" in data[:500] else "?"
    out = {"format": fmt, "bytes": len(data), "sha1": hashlib.sha1(data).hexdigest()[:12]}
    try:
        img = Image.open(io.BytesIO(data))
        out.update(size=list(img.size), mode=img.mode, frames=getattr(img, "n_frames", 1))
        if fmt == "jpeg":
            q = img.quantization or {}
            out["jpeg_q0"] = sum(q.get(0, [])) // max(1, len(q.get(0, [])))
    except Exception as e:  # EMF without a renderer, SVG
        out["open_error"] = type(e).__name__
    return out


def grey(img: Image.Image, w: int = 240) -> np.ndarray:
    if img.mode in ("RGBA", "LA", "P"):
        img = img.convert("RGBA")
        ground = Image.new("RGBA", img.size, (255, 255, 255, 255))
        ground.alpha_composite(img)
        img = ground
    h = max(1, round(w * img.size[1] / img.size[0]))
    return np.asarray(img.convert("L").resize((w, h), Image.BILINEAR), np.float32)


def distance(a: Image.Image, b: Image.Image) -> float:
    ga, gb = grey(a), grey(b)
    if ga.shape != gb.shape:
        gb = np.asarray(Image.fromarray(gb.astype(np.uint8)).resize(ga.shape[::-1]), np.float32)
    return round(float(np.abs(ga - gb).mean() / 255), 4)


def stripes_visible(img: Image.Image, ref_size: tuple[int, int]) -> bool:
    """The 1 px stripes at x 100..500 of the 3000 px original still alternate pixel by pixel."""
    if img.size[0] < ref_size[0]:
        return False
    row = np.asarray(img.convert("L"))[300, 110:390].astype(int)
    return bool(np.abs(np.diff(row)).mean() > 60)


def baked(name: str, img: Image.Image, original: Image.Image, plain: Image.Image | None) -> dict:
    """Which of the variant's edits the pixels show."""
    out: dict = {}
    if name.startswith("crop"):
        w, h = original.size
        c = CROP
        cropped = original.crop((round(c["l"] * w), round(c["t"] * h), round((1 - c["r"]) * w), round((1 - c["b"]) * h)))
        out["crop"] = "baked" if distance(img, cropped) < distance(img, original) else "kept"
    if name == "alpha50":
        alpha = np.asarray(img.convert("RGBA"))[..., 3] if img.mode in ("RGBA", "LA", "P") else None
        out["transparency"] = "baked" if alpha is not None and alpha.mean() < 200 else "kept"
    if name in ("bright", "gray", "duotone") and plain is not None:
        a = np.asarray(img.convert("RGB").resize((200, 133)), np.float32)
        b = np.asarray(plain.convert("RGB").resize((200, 133)), np.float32)
        out[name] = "baked" if np.abs(a - b).mean() > 8 else "kept"
    if name.startswith("rot") or name == "crop_rot":
        out["rotation"] = "baked" if abs(img.size[0] / img.size[1] - (original.size[0] / original.size[1])) > 0.05 \
            and not name.startswith("crop") else "kept"
    if name == "outline":
        arr = np.asarray(img.convert("RGB"), np.float32)
        edge = np.concatenate([arr[:3].reshape(-1, 3), arr[-3:].reshape(-1, 3)])
        out["outline"] = "baked" if (edge[:, 0] > 200).mean() > 0.8 and (edge[:, 1] < 60).mean() > 0.8 else "kept"
    return out


def url_variants(url: str) -> list[str]:
    """Size parameters googleusercontent URLs accept (=s0 is the original on photos URLs)."""
    m = re.match(r"(?P<base>.*?)=s\d+(?P<query>\?.*)?$", url)
    if not m:
        return []
    q = m.group("query") or ""
    return [(v or "(none)", m.group("base") + v + q) for v in ("=s0", "=s3000", "=s4096", "=w3000-h2000", "=s1024", "")]


def props(pe: dict) -> dict:
    ip = pe.get("image", {}).get("imageProperties", {})
    t = pe.get("transform", {})
    rot = math.degrees(math.atan2(t.get("shearY", 0.0), t.get("scaleX", 1.0))) if t else 0.0
    return {"imageProperties": ip, "rotation": round(rot, 2), "sourceUrl": pe.get("image", {}).get("sourceUrl"),
            "size": pe.get("size"), "transform": t}


# ---------------------------------------------------------------- export

def export_pptx(drive, creds, fid: str, dest: Path) -> dict:
    try:
        data = execute(drive.files().export(fileId=fid, mimeType=PPTX_MIME))
        how = "files.export"
    except HttpError as e:
        from google.auth.transport.requests import AuthorizedSession
        link = execute(drive.files().get(fileId=fid, fields="exportLinks"))["exportLinks"][PPTX_MIME]
        r = AuthorizedSession(creds).get(link, timeout=300)
        r.raise_for_status()
        data = r.content
        how = f"exportLinks (files.export: {e.resp.status} {str(e)[:80]})"
    dest.write_bytes(data)
    return {"how": how, "bytes": len(data)}


def pptx_pictures(path: Path) -> list[dict]:
    out = []
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        for slide in sorted(n for n in names if re.match(r"ppt/slides/slide\d+\.xml$", n)):
            rels_name = slide.replace("slides/", "slides/_rels/") + ".rels"
            rels = {}
            if rels_name in names:
                for rel in etree.fromstring(z.read(rels_name)):
                    rels[rel.get("Id")] = rel.get("Target")
            root = etree.fromstring(z.read(slide))
            for pic in root.iter(f"{{{P}}}pic"):
                c_nv = pic.find(f"{{{P}}}nvPicPr/{{{P}}}cNvPr")
                blip = pic.find(f"{{{P}}}blipFill/{{{A}}}blip")
                target = rels.get(blip.get(f"{{{R}}}embed")) if blip is not None else None
                media = "ppt/" + target.replace("../", "") if target else None
                src = pic.find(f"{{{P}}}blipFill/{{{A}}}srcRect")
                xfrm = pic.find(f"{{{P}}}spPr/{{{A}}}xfrm")
                ln = pic.find(f"{{{P}}}spPr/{{{A}}}ln")
                entry = {"slide": slide, "descr": c_nv.get("descr"), "name": c_nv.get("name"), "media": media,
                         "srcRect": dict(src.attrib) if src is not None else None,
                         "rot": int(xfrm.get("rot", 0)) / 60000 if xfrm is not None else None,
                         "blip_children": [etree.QName(ch).localname + json.dumps(dict(ch.attrib)) for ch in blip] if blip is not None else [],
                         "ln": etree.tostring(ln).decode()[:200] if ln is not None and len(ln) else None}
                if media and media in names:
                    entry["media_info"] = describe(z.read(media))
                out.append(entry)
    return out


# ---------------------------------------------------------------- main

def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--folder", type=Path, default=ROOT / "out" / "agent-pull-images" / "probe")
    ap.add_argument("--keep", action="store_true", help="keep the probe slide and the staging file")
    args = ap.parse_args()
    folder = args.folder.resolve()
    work = folder / "probe-images"
    work.mkdir(parents=True, exist_ok=True)
    pid = json.loads((folder / "emit.json").read_text(encoding="utf-8"))["presentationId"]
    creds = credentials()
    slides, drive = slides_service(creds), drive_service(creds)

    original = photo()
    original.save(work / "photo.png", optimize=False)
    original.save(work / "photo.jpg", quality=92)
    small = original.resize((1600, 1067), Image.LANCZOS)
    small.save(work / "photo1600.png")
    small.save(work / "photo1600.jpg", quality=92)
    frames_gif(work / "anim.gif")
    emf_file(work / "rect.emf")
    findings: dict = {"files": {n: describe((work / n).read_bytes()) for n in
                                ("photo.png", "photo.jpg", "photo1600.png", "photo1600.jpg", "anim.gif", "rect.emf")}}
    print("test files:", json.dumps(findings["files"], indent=1))

    t = time.time()
    fid = execute(drive.files().create(body={"name": "beamer2slides image probe staging (temporary)",
                                             "mimeType": "application/vnd.google-apps.presentation"},
                                       media_body=MediaIoBaseUpload(staging_pptx(work), mimetype=PPTX_MIME, resumable=True),
                                       fields="id"))["id"]
    findings["staging_upload_s"] = round(time.time() - t, 1)
    try:
        staged = execute(slides.presentations().get(presentationId=fid))
        imported = {}
        for pe in staged["slides"][0].get("pageElements", []):
            d = pe.get("description") or ""
            if d.startswith("probe:"):
                imported[d[6:]] = pe
        findings["import_missing"] = [n for n, _, _ in VARIANTS if n not in imported]
        plain_img = None
        rows = {}
        for name, file, edits in VARIANTS:
            pe = imported.get(name)
            if pe is None:
                continue
            url = pe["image"]["contentUrl"]
            data = fetch_url(url)
            (work / f"import-{name}.bin").write_bytes(data)
            info = describe(data)
            row = {"edits": edits, "contentUrl_host": url.split("/")[2], "contentUrl_tail": url[-40:],
                   "content": info, "props": props(pe),
                   "same_bytes_as_file": info["sha1"] == findings["files"][file]["sha1"]}
            try:
                img = Image.open(io.BytesIO(data))
                img.load()
                if name == "plain":
                    plain_img = img
                    row["stripes_full_res"] = stripes_visible(img, original.size)
                row["baked"] = baked(name, img, original, plain_img)
                row["distance_to_original"] = distance(img, original)
            except Exception as e:
                row["baked_error"] = type(e).__name__
            if name == "plain":
                row["url_variants"] = {}
                for label, v in url_variants(url):
                    try:
                        row["url_variants"][label] = describe(fetch_url(v))
                    except Exception as e:
                        row["url_variants"][label] = str(e)[:80]
            rows[name] = row
        findings["import"] = rows

        export = work / "staging-export.pptx"
        findings["staging_export"] = export_pptx(drive, creds, fid, export)
        findings["staging_export_pictures"] = pptx_pictures(export)
        for p in findings["staging_export_pictures"]:
            name = (p["descr"] or "")[6:]
            if name in rows and p.get("media_info"):
                p["same_bytes_as_contentUrl"] = p["media_info"]["sha1"] == rows[name]["content"]["sha1"]

        # ---- the converted probe deck: createImage from the staging URLs, then API edits
        pres = execute(slides.presentations().get(presentationId=pid, fields="slides(objectId)"))
        reqs = []
        if any(s["objectId"] == SLIDE_ID for s in pres.get("slides", [])):
            reqs.append({"deleteObject": {"objectId": SLIDE_ID}})
        reqs.append({"createSlide": {"objectId": SLIDE_ID, "slideLayoutReference": {"predefinedLayout": "BLANK"}}})
        created = {}
        for k, name in enumerate(["plain", "crop", "alpha50", "gray", "rot30", "gif", "emf"]):
            if name not in imported:
                continue
            oid = f"b2s_probe_img_{name}"
            created[name] = oid
            reqs.append({"createImage": {"objectId": oid, "url": imported[name]["image"]["contentUrl"], "elementProperties": {
                "pageObjectId": SLIDE_ID, "size": {"width": {"magnitude": 150 * EMU_PER_PT, "unit": "EMU"},
                                                   "height": {"magnitude": 100 * EMU_PER_PT, "unit": "EMU"}},
                "transform": {"scaleX": 1, "scaleY": 1, "translateX": (20 + (k % 4) * 175) * EMU_PER_PT,
                              "translateY": (20 + (k // 4) * 125) * EMU_PER_PT, "unit": "EMU"}}}})
        edit_targets = {}
        for name in ("crop", "transparency", "brightness", "contrast", "recolor", "outline", "rotation"):
            oid = f"b2s_probe_edit_{name}"
            edit_targets[name] = oid
            reqs.append({"createImage": {"objectId": oid, "url": imported["plain"]["image"]["contentUrl"], "elementProperties": {
                "pageObjectId": SLIDE_ID, "size": {"width": {"magnitude": 150 * EMU_PER_PT, "unit": "EMU"},
                                                   "height": {"magnitude": 100 * EMU_PER_PT, "unit": "EMU"}},
                "transform": {"scaleX": 1, "scaleY": 1, "translateX": 20 * EMU_PER_PT, "translateY": 270 * EMU_PER_PT,
                              "unit": "EMU"}}}})
        execute(slides.presentations().batchUpdate(presentationId=pid, body={"requests": reqs}))
        api_edits = {
            "crop": {"cropProperties": {"leftOffset": CROP["l"], "rightOffset": CROP["r"], "topOffset": CROP["t"],
                                        "bottomOffset": CROP["b"]}},
            "transparency": {"transparency": 0.5},
            "brightness": {"brightness": 0.3},
            "contrast": {"contrast": -0.2},
            "recolor": {"recolor": {"name": "GRAYSCALE"}},
            "outline": {"outline": {"outlineFill": {"solidFill": {"color": {"rgbColor": {"red": 1.0}}}},
                                    "weight": {"magnitude": 3, "unit": "PT"}}},
        }
        api = {}
        for name, value in api_edits.items():
            field = next(iter(value))
            try:
                execute(slides.presentations().batchUpdate(presentationId=pid, body={"requests": [
                    {"updateImageProperties": {"objectId": edit_targets[name], "imageProperties": value, "fields": field}}]}))
                api[name] = {"request": "accepted"}
            except HttpError as e:
                reason = " ".join(str(e.reason).split())[:160]
                api[name] = {"request": f"rejected {e.resp.status}: {reason}"}
        # rotation: 30 degrees about the centre
        page = execute(slides.presentations().pages().get(presentationId=pid, pageObjectId=SLIDE_ID))
        rot_pe = next(e for e in page["pageElements"] if e["objectId"] == edit_targets["rotation"])
        tr = rot_pe["transform"]
        W, H = rot_pe["size"]["width"]["magnitude"], rot_pe["size"]["height"]["magnitude"]
        sx, sy = tr.get("scaleX", 1), tr.get("scaleY", 1)
        cx = tr.get("translateX", 0) + sx * W / 2
        cy = tr.get("translateY", 0) + sy * H / 2
        th = math.radians(30)
        a, b, d, e_ = sx * math.cos(th), -sy * math.sin(th), sx * math.sin(th), sy * math.cos(th)
        execute(slides.presentations().batchUpdate(presentationId=pid, body={"requests": [
            {"updatePageElementTransform": {"objectId": edit_targets["rotation"], "applyMode": "ABSOLUTE", "transform": {
                "scaleX": a, "shearX": b, "shearY": d, "scaleY": e_, "unit": "EMU",
                "translateX": cx - (a * W / 2 + b * H / 2), "translateY": cy - (d * W / 2 + e_ * H / 2)}}}]}))
        api["rotation"] = {"request": "accepted"}
        page = execute(slides.presentations().pages().get(presentationId=pid, pageObjectId=SLIDE_ID))
        by_id = {e["objectId"]: e for e in page["pageElements"]}
        plain_live = None
        live = {}
        for name, oid in list(created.items()) + [(f"edit_{n}", o) for n, o in edit_targets.items()]:
            pe = by_id.get(oid)
            if pe is None:
                continue
            data = fetch_url(pe["image"]["contentUrl"])
            info = describe(data)
            row = {"content": info, "props": props(pe),
                   "same_bytes_as_import": info["sha1"] == rows.get(name, rows["plain"])["content"]["sha1"]}
            try:
                img = Image.open(io.BytesIO(data))
                img.load()
                if name == "plain":
                    plain_live = img
                    row["stripes_full_res"] = stripes_visible(img, original.size)
                vname = {"edit_crop": "crop", "edit_transparency": "alpha50", "edit_brightness": "bright",
                         "edit_recolor": "gray", "edit_outline": "outline", "edit_rotation": "rot30"}.get(name, name)
                row["baked"] = baked(vname, img, original, plain_live or plain_img)
            except Exception as ex:
                row["baked_error"] = type(ex).__name__
            if name.startswith("edit_"):
                row.update(api[name[5:]])
            live[name] = row
        findings["live"] = live

        export = work / "deck-export.pptx"
        findings["deck_export"] = export_pptx(drive, creds, pid, export)
        findings["deck_export_pictures"] = [p for p in pptx_pictures(export)
                                            if p["descr"] is None or not str(p["descr"]).startswith("b2s")]
        if not args.keep:
            execute(slides.presentations().batchUpdate(presentationId=pid, body={"requests": [
                {"deleteObject": {"objectId": SLIDE_ID}}]}))
    finally:
        if not args.keep:
            execute(drive.files().delete(fileId=fid))
        else:
            print(f"staging kept: https://docs.google.com/presentation/d/{fid}/edit")
    (work / "findings.json").write_text(json.dumps(findings, indent=1, default=str), encoding="utf-8")
    print_table(findings)
    print(f"findings: {work / 'findings.json'}")


def print_table(f: dict) -> None:
    print("\n| picture | source | format | pixels | baked into contentUrl | kept as property |")
    print("|---|---|---|---|---|---|")
    for side in ("import", "live"):
        for name, row in f.get(side, {}).items():
            c = row["content"]
            ip = row["props"]["imageProperties"]
            kept = ", ".join(k for k in ip if k not in ("outline", "shadow") or ip[k].get("propertyState") != "NOT_RENDERED")
            if row["props"]["rotation"]:
                kept += f", rotation {row['props']['rotation']}"
            print(f"| {name} | {side} | {c['format']} | {c.get('size')} | {row.get('baked')} | {kept} {row.get('request', '')} |")
    for key in ("staging_export_pictures", "deck_export_pictures"):
        print(f"\n{key}:")
        for p in f.get(key, []):
            print(f"  {p['descr']}: {p.get('media_info', {}).get('format')} {p.get('media_info', {}).get('size')} "
                  f"srcRect={p['srcRect']} rot={p['rot']} blip={p['blip_children']} ln={bool(p['ln'])}")


if __name__ == "__main__":
    sys.exit(main())
