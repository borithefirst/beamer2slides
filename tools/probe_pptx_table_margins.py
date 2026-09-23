"""Probe: can a table imported from a .pptx have rows shorter than the API's 14.4 pt + one line?

The Slides API cannot set a table cell's padding (TableCellProperties has only a fill and the
content alignment), and an API-made row is at least 14.4 + 1.195·z·lineSpacing pt
(tools/probe_table_rows.py). A .pptx table cell carries its own margins
(`<a:tcPr marL= marR= marT= marB=>`). This probe uploads a .pptx with one-column tables of
several margins and sizes through Drive's conversion, then

- reads them back (row heights, cell properties: is the padding visible anywhere?);
- duplicates an empty template table with `duplicateObject` and fills it through the API the way
  emit.table_requests does (insertText, text and paragraph style, minRowHeight, borders);
- fills an imported empty table in place, and edits an imported filled one (replace text, insert
  a row, fill and align a cell);
- duplicates the whole slide (emit's phase 1 path);
- makes the same table with createTable for comparison;

and measures the row pitch on Google's renderer: every table gets black borders and the
thumbnail is scanned down a column where no text is, for the border lines.

The presentation is moved to the trash at the end (keep it with --keep).

Usage: python tools/probe_pptx_table_margins.py [--keep]
"""

import io
import json
import sys
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Pt

from beamer2slides import gapi
from beamer2slides.google_auth import drive_service, slides_service
from beamer2slides.gslides import EMU_PER_PT, emu, execute, pt, save_thumbnail

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "out" / "probe_pptx_table_margins"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
NO_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # "No Style, No Grid"
ROWS = 4
WIDTH = 64.0
TEXT = "Hxg"

# name: (margins l, t, r, b in pt or None for the .pptx default, size, lnSpc %, text in the .pptx, (x, y))
IMPORTED = {
    "imp_def_10": (None, 10, 100, True, (10, 30)),
    "imp_m0_10": ((0, 0, 0, 0), 10, 100, True, (84, 30)),
    "imp_m0_8": ((0, 0, 0, 0), 8, 100, True, (158, 30)),
    "imp_m0_14": ((0, 0, 0, 0), 14, 100, True, (232, 30)),
    "imp_t0_10": ((7.2, 0, 7.2, 0), 10, 100, True, (306, 30)),
    "imp_t1_10": ((0, 1, 0, 1), 10, 100, True, (380, 30)),
    "imp_m0_10_ls70": ((0, 0, 0, 0), 10, 70, True, (454, 30)),
    "tpl_m0": ((0, 0, 0, 0), 10, 100, False, (640, 300)),   # the template, duplicated to (10, 220)
    "fill_m0": ((0, 0, 0, 0), 10, 100, False, (84, 220)),   # filled in place through the API
    "mod_m0_10": ((0, 0, 0, 0), 10, 100, True, (232, 220)),  # imported, then edited through the API
    "tpl_t0_2x2": ((7.2, 0, 7.2, 0), 10, 100, False, (640, 200)),  # template grown by insertTable*
}
DUP = ("dup_m0", (10, 220))
GROWN = ("grown_t0", (380, 220))
API = ("api_10", (158, 220))


def _table(slide, name, margins, size, spacing, text, xy, rows=ROWS, cols=1):
    x, y = xy
    frame = slide.shapes.add_table(rows, cols, Pt(x), Pt(y), Pt(WIDTH * cols), Pt(rows))
    frame.name = name
    tbl = frame._element.graphic.graphicData.tbl
    pr = tbl.tblPr
    for flag in ("firstRow", "bandRow"):
        pr.attrib.pop(flag, None)
    style = pr.find(f"{{{A}}}tableStyleId")
    if style is None:
        style = etree.SubElement(pr, f"{{{A}}}tableStyleId")
    style.text = NO_STYLE
    for row in frame.table.rows:
        row.height = Emu(EMU_PER_PT)  # 1 pt: the importer grows it to what the text needs
    for r in range(rows):
        for c in range(cols):
            cell = frame.table.cell(r, c)
            if margins is not None:
                cell.margin_left, cell.margin_top, cell.margin_right, cell.margin_bottom = (Emu(round(m * EMU_PER_PT)) for m in margins)
            if text:
                para = cell.text_frame.paragraphs[0]
                run = para.add_run()
                run.text = TEXT
                run.font.size = Pt(size)
                run.font.name = "Lato"
                if spacing != 100:
                    para.line_spacing = spacing / 100
                para.space_before = Pt(0)
                para.space_after = Pt(0)
    return frame


def build() -> io.BytesIO:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(720 * EMU_PER_PT), Emu(405 * EMU_PER_PT)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for name, (margins, size, spacing, text, xy) in IMPORTED.items():
        rows, cols = (2, 2) if name == "tpl_t0_2x2" else (ROWS, 1)
        _table(slide, name, margins, size, spacing, text, xy, rows, cols)
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf


def fill_requests(oid: str, rows: int, size: float, cols: int = 1) -> list[dict]:
    """What emit.table_requests does to a cell with text."""
    reqs = []
    for r in range(rows):
        for c in range(cols):
            loc = {"rowIndex": r, "columnIndex": c}
            reqs += [
                {"insertText": {"objectId": oid, "cellLocation": loc, "text": TEXT}},
                {"updateTextStyle": {"objectId": oid, "cellLocation": loc, "textRange": {"type": "ALL"},
                                     "style": {"fontFamily": "Lato", "fontSize": pt(size)}, "fields": "fontFamily,fontSize"}},
                {"updateParagraphStyle": {"objectId": oid, "cellLocation": loc, "textRange": {"type": "ALL"},
                                          "style": {"alignment": "START", "lineSpacing": 100, "spaceAbove": pt(0),
                                                    "spaceBelow": pt(0), "indentStart": pt(0), "indentFirstLine": pt(0)},
                                          "fields": "alignment,lineSpacing,spaceAbove,spaceBelow,indentStart,indentFirstLine"}},
            ]
    reqs.append({"updateTableRowProperties": {"objectId": oid, "rowIndices": list(range(rows)),
                                              "tableRowProperties": {"minRowHeight": emu(1)}, "fields": "minRowHeight"}})
    return reqs


def borders(oid: str, rows: int, cols: int = 1) -> dict:
    return {"updateTableBorderProperties": {
        "objectId": oid, "borderPosition": "ALL",
        "tableRange": {"location": {"rowIndex": 0, "columnIndex": 0}, "rowSpan": rows, "columnSpan": cols},
        "tableBorderProperties": {"tableBorderFill": {"solidFill": {"color": {"rgbColor": {}}, "alpha": 1}},
                                  "weight": pt(0.75)},
        "fields": "tableBorderFill.solidFill.color,tableBorderFill.solidFill.alpha,weight"}}


def move(oid: str, xy) -> dict:
    return {"updatePageElementTransform": {"objectId": oid, "applyMode": "ABSOLUTE", "transform": {
        "scaleX": 1, "scaleY": 1, "unit": "EMU",
        "translateX": round(xy[0] * EMU_PER_PT), "translateY": round(xy[1] * EMU_PER_PT)}}}


def lines_in(img: Image.Image, x_pt: float, y0_pt: float, y1_pt: float) -> list[float]:
    """y (pt) of the dark horizontal lines crossing x between y0 and y1."""
    k = img.width / 720
    x = round(x_pt * k)
    ys = []
    run = []
    for y in range(max(0, round(y0_pt * k)), min(img.height, round(y1_pt * k))):
        dark = sum(img.getpixel((x, y))[:3]) < 3 * 110
        if dark:
            run.append(y)
        elif run:
            ys.append(sum(run) / len(run) / k)
            run = []
    if run:
        ys.append(sum(run) / len(run) / k)
    return [round(v, 2) for v in ys]


def row_heights(el: dict) -> list[float]:
    return [round(r["rowHeight"].get("magnitude", 0) / EMU_PER_PT, 2) for r in el["table"]["tableRows"]]


def main() -> None:
    keep = "--keep" in sys.argv
    OUT.mkdir(parents=True, exist_ok=True)
    drive, slides = drive_service(), slides_service()
    media = gapi.media_upload(build(), PPTX_MIME)
    pid = execute(drive.files().create(
        body={"name": "b2s probe pptx table margins", "mimeType": "application/vnd.google-apps.presentation"},
        media_body=media, fields="id"))["id"]
    try:
        run(slides, pid)
    finally:
        if keep:
            print(f"kept https://docs.google.com/presentation/d/{pid}/edit")
        else:
            execute(drive.files().update(fileId=pid, body={"trashed": True}))
            print("presentation moved to the trash")


def run(slides, pid: str) -> None:
    pres = execute(slides.presentations().get(presentationId=pid))
    (OUT / "imported.json").write_text(json.dumps(pres["slides"], indent=1), encoding="utf-8")
    page = pres["slides"][0]
    by_name = {}
    for pe in page.get("pageElements", []):
        if "table" in pe:
            # The importer drops the shape names; tables come back in creation order.
            by_name[list(IMPORTED)[len(by_name)]] = pe
    print("imported row heights (pt), before any API request:")
    for name, pe in by_name.items():
        print(f"  {name:16s} {row_heights(pe)}  cell props {json.dumps(pe['table']['tableRows'][0]['tableCells'][0].get('tableCellProperties'))}")
    oid = {name: pe["objectId"] for name, pe in by_name.items()}
    dup_id, grown_id, api_id = "probe_dup_m0", "probe_grown_t0", "probe_api_10"
    sid = page["objectId"]
    reqs: list[dict] = [
        {"duplicateObject": {"objectId": oid["tpl_m0"], "objectIds": {oid["tpl_m0"]: dup_id}}},
        move(dup_id, DUP[1]),
        *fill_requests(dup_id, ROWS, 10),
        *fill_requests(oid["fill_m0"], ROWS, 10),
        # A 2x2 template grown to 4x2 by insertTableRows: do new cells keep the margins?
        {"duplicateObject": {"objectId": oid["tpl_t0_2x2"], "objectIds": {oid["tpl_t0_2x2"]: grown_id}}},
        move(grown_id, GROWN[1]),
        {"insertTableRows": {"tableObjectId": grown_id, "cellLocation": {"rowIndex": 1, "columnIndex": 0},
                             "insertBelow": True, "number": 2}},
        {"insertTableColumns": {"tableObjectId": grown_id, "cellLocation": {"rowIndex": 0, "columnIndex": 1},
                                "insertRight": True, "number": 1}},
        {"deleteTableColumn": {"tableObjectId": grown_id, "cellLocation": {"rowIndex": 0, "columnIndex": 2}}},
        *fill_requests(grown_id, ROWS, 10, cols=2),
        {"updateTableColumnProperties": {"objectId": grown_id, "columnIndices": [0, 1],
                                         "tableColumnProperties": {"columnWidth": emu(WIDTH / 2)}, "fields": "columnWidth"}},
        {"createTable": {"objectId": api_id, "rows": ROWS, "columns": 1, "elementProperties": {
            "pageObjectId": sid, "size": {"width": emu(WIDTH), "height": emu(ROWS)},
            "transform": {"scaleX": 1, "scaleY": 1, "unit": "EMU",
                          "translateX": round(API[1][0] * EMU_PER_PT), "translateY": round(API[1][1] * EMU_PER_PT)}}}},
        *fill_requests(api_id, ROWS, 10),
        # An imported, filled table edited through the API.
        {"deleteText": {"objectId": oid["mod_m0_10"], "cellLocation": {"rowIndex": 0, "columnIndex": 0},
                        "textRange": {"type": "ALL"}}},
        {"insertText": {"objectId": oid["mod_m0_10"], "cellLocation": {"rowIndex": 0, "columnIndex": 0}, "text": "New"}},
        {"updateTextStyle": {"objectId": oid["mod_m0_10"], "cellLocation": {"rowIndex": 0, "columnIndex": 0},
                             "textRange": {"type": "ALL"}, "style": {"fontFamily": "Lato", "fontSize": pt(10)},
                             "fields": "fontFamily,fontSize"}},
        {"updateTableCellProperties": {"objectId": oid["mod_m0_10"],
                                       "tableRange": {"location": {"rowIndex": 2, "columnIndex": 0}, "rowSpan": 1, "columnSpan": 1},
                                       "tableCellProperties": {"contentAlignment": "MIDDLE",
                                                               "tableCellBackgroundFill": {"solidFill": {"color": {"rgbColor": {"red": 1, "green": 0.9, "blue": 0.6}}}}},
                                       "fields": "contentAlignment,tableCellBackgroundFill.solidFill.color"}},
        {"insertTableRows": {"tableObjectId": oid["mod_m0_10"], "cellLocation": {"rowIndex": 3, "columnIndex": 0},
                             "insertBelow": True, "number": 1}},
        {"insertText": {"objectId": oid["mod_m0_10"], "cellLocation": {"rowIndex": 4, "columnIndex": 0}, "text": TEXT}},
        {"updateTextStyle": {"objectId": oid["mod_m0_10"], "cellLocation": {"rowIndex": 4, "columnIndex": 0},
                             "textRange": {"type": "ALL"}, "style": {"fontFamily": "Lato", "fontSize": pt(10)},
                             "fields": "fontFamily,fontSize"}},
        {"updateTableRowProperties": {"objectId": oid["mod_m0_10"], "rowIndices": [0, 1, 2, 3, 4],
                                      "tableRowProperties": {"minRowHeight": emu(1)}, "fields": "minRowHeight"}},
    ]
    for name in IMPORTED:
        if name not in ("tpl_m0", "tpl_t0_2x2"):
            reqs.append(borders(oid[name], ROWS + (name == "mod_m0_10")))
    reqs += [borders(dup_id, ROWS), borders(grown_id, ROWS, 2), borders(api_id, ROWS)]
    execute(slides.presentations().batchUpdate(presentationId=pid, body={"requests": reqs}))
    # emit's phase 1: the whole source slide duplicated under our ids.
    page_ids = {sid: "probe_slide2", **{pe["objectId"]: f"probe_s2_{i}" for i, pe in enumerate(page["pageElements"])}}
    execute(slides.presentations().batchUpdate(presentationId=pid, body={"requests": [
        {"duplicateObject": {"objectId": sid, "objectIds": page_ids}}]}))
    got = execute(slides.presentations().get(presentationId=pid))
    (OUT / "after.json").write_text(json.dumps(got["slides"], indent=1), encoding="utf-8")
    save_thumbnail(slides, pid, sid, OUT / "slide.png")
    save_thumbnail(slides, pid, "probe_slide2", OUT / "slide-dup.png")
    img = Image.open(OUT / "slide.png").convert("RGB")
    img2 = Image.open(OUT / "slide-dup.png").convert("RGB")
    els = {pe["objectId"]: pe for pe in got["slides"][0]["pageElements"]}
    dup_els = {pe["objectId"]: pe for pe in got["slides"][1]["pageElements"]}
    named = {**{n: oid[n] for n in IMPORTED}, DUP[0]: dup_id, GROWN[0]: grown_id, API[0]: api_id}
    where = {**{n: v[4] for n, v in IMPORTED.items()}, DUP[0]: DUP[1], GROWN[0]: GROWN[1], API[0]: API[1]}
    results = {}
    print("\nafter the API requests: API row heights / rendered row pitch (pt)")
    for name, o in named.items():
        if name in ("tpl_m0", "tpl_t0_2x2"):
            continue
        x, y = where[name]
        scan_x = x + WIDTH - 6  # right of the text (a 2-column table: right of column 1's text)
        found = lines_in(img, scan_x, y - 3, y + 160)
        pitch = [round(b - a, 2) for a, b in zip(found, found[1:])]
        dup_name = page_ids.get(o)
        found2 = lines_in(img2, scan_x, y - 3, y + 160) if dup_name in dup_els else []
        results[name] = {"api": row_heights(els[o]), "rendered": pitch,
                         "slide_dup_api": row_heights(dup_els[dup_name]) if dup_name in dup_els else None,
                         "slide_dup_rendered": [round(b - a, 2) for a, b in zip(found2, found2[1:])]}
        print(f"  {name:16s} api {results[name]['api']}  rendered {pitch}  "
              f"| slide dup api {results[name]['slide_dup_api']} rendered {results[name]['slide_dup_rendered']}")
    (OUT / "results.json").write_text(json.dumps(results, indent=1), encoding="utf-8")
    print(OUT / "slide.png")


if __name__ == "__main__":
    main()
